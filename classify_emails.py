"""
Email classification and invoice extraction test system for a multi-location
pharmacy group. Reads every .txt, .eml, or .msg file in test_emails/ (.msg is
Outlook's native export format), sends each to the Claude API for
classification/extraction, and writes results to output/results.xlsx,
output/invoice_history.xlsx, and output/qb_ready.xlsx (all real Excel
workbooks, one invoice/payment per row).

test_emails/ is a clean inbox: successfully processed emails are MOVED out
into processed_emails/ (not just copied), so future runs only ever scan
what's actually pending — no re-reading/re-skipping an ever-growing pile of
old files. Emails that error out (API/parse failure) are left in place so
they're retried automatically next run. As a safety net, output/results.xlsx
is also checked by filename in case a file ever ends up back in test_emails/
after already succeeding.

qb_ready.xlsx accumulates every QB-ready invoice ever found (not just this
run's) and has two columns that are never recomputed once set: "Entered in
QB" (fill this in by hand once you've copy/pasted a row into QuickBooks) and
"Date Added" (when the row first appeared). Every sheet also has "Date
Added"/"date_added" recording when that row was first written.

"File Path" is saved for QB-bound records (non-skip invoices and non-skip
credit memos) AND for payments (kept as proof of payment) — one document PER
record, saved under invoice_files/ and named
"Store - Vendor - Invoice # - Amount - Date". It's the real PDF/image
attachment the record was read from, or a PDF rendering of the email body
when the data came from the email text. Skip-QB records, statements/ship
notices, and flagged items get NO file (for flagged items, open the email in
review_pending/ instead). This is separate from processed_emails/, which
archives the original raw source email file itself.

HUMAN REVIEW LOOP (review_queue.xlsx + review_pending/ + learnings.xlsx):
  - An email the AI flags for review does NOT go to results.xlsx; it lands in
    review_queue.xlsx with read-only context (filename, vendor, store,
    invoice #, and the AI's question) plus two blank columns for you:
    "Answer" (your decision for this item) and "Notes to Learn" (a general
    rule you want the AI to remember). The source email itself is moved to
    review_pending/ so you can open it while answering; the row's File Path
    points there.
  - On the next run, each answered row is reprocessed — your Answer and all
    accumulated learnings are injected into the AI — so the item resolves and
    lands in results.xlsx, and that row is cleared from the queue. Rows you
    haven't answered stay. Answers that error on reprocess stay for retry.
  - Each "Notes to Learn" is appended to learnings.xlsx, a durable rule set
    injected into the AI on EVERY future run so new similar emails classify
    correctly without being re-flagged. You can hand-edit/prune learnings.xlsx.

Usage:
    python classify_emails.py

Requires the ANTHROPIC_API_KEY environment variable to be set.
"""

import atexit
import base64
import email
import hashlib
import html as html_module
import io
import json
import mimetypes
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime, timedelta, timezone
from email.policy import default as email_policy
from email.utils import parseaddr, parsedate_to_datetime
from pathlib import Path

import extract_msg
import httpx
from anthropic import Anthropic
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from pypdf import PdfReader, PdfWriter

# Fuzzy vendor matching for standardization (C-accelerated; handles word reordering + minor
# spelling drift far better than exact string compare). Defensive import so the feature
# degrades to the stdlib difflib fallback instead of crashing if rapidfuzz isn't installed.
try:
    from rapidfuzz import fuzz as _rapidfuzz
except ImportError:
    _rapidfuzz = None
    import difflib

# Attachment types Claude can read natively (PDF via document blocks, common
# raster formats via image blocks) vs. plain-text-ish types we can just inline.
SUPPORTED_DOCUMENT_TYPES = {"application/pdf"}
SUPPORTED_IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp"}
INLINABLE_TEXT_TYPES = {"text/plain", "text/csv"}
# .xlsx / .xlsm only (openpyxl, already a dependency) — legacy binary .xls is NOT
# covered (different format, would need a separate library); an .xls attachment
# will still come through as unreadable for now.
SUPPORTED_SPREADSHEET_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",  # .xlsx
    "application/vnd.ms-excel.sheet.macroEnabled.12",                     # .xlsm
}

# Some .msg attachments report a generic/missing MIME type (e.g. a real PDF
# tagged application/octet-stream). When that happens we fall back to the file
# extension so a genuine PDF/image isn't treated as unreadable.
EXT_TO_MEDIA_TYPE = {
    ".pdf": "application/pdf",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".txt": "text/plain",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
}
MEDIA_TYPE_TO_EXT = {
    "application/pdf": ".pdf",
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/gif": ".gif",
    "image/webp": ".webp",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.ms-excel.sheet.macroEnabled.12": ".xlsm",
}

MODEL = "claude-sonnet-5"
# Output budget. A single email can hold many invoices (a combined multi-page
# scan), and each invoice/credit-memo/payment is a sizable JSON object, so this
# needs headroom — too low and the JSON is truncated (or comes back empty on a
# big scanned batch: stop_reason=max_tokens with zero usable output).
# claude-sonnet-5's real output ceiling is 128K, BUT this pipeline calls the API
# non-streaming, and the anthropic SDK hard-raises ("Streaming is required...")
# once max_tokens would risk the 10-min non-streaming timeout — i.e. above
# ~21,333 (128000 * 600s/3600s). So ~21,000 is the max SAFE non-streaming value;
# we sit just under the ceiling for maximum first-pass headroom without a
# streaming rework or beta header. (Reaching the full 128K would require
# switching this call to client.messages.stream(...).get_final_message() — a
# deliberate non-goal here.) Scans/images too dense even for this land in the
# page-chunking path below (CHUNK_THRESHOLD_PAGES) or fail into review.
MAX_TOKENS = 21000

# Auto-split oversized single-PDF scans before classification. A stack of separate
# paper invoices scanned at the office MFP arrives as ONE big PDF; asking the model
# to enumerate every invoice across all pages into one JSON response blows the
# output-token budget (empty response, stop_reason=max_tokens). When the ONLY
# attachment is a PDF with more than CHUNK_THRESHOLD_PAGES pages, it's split into
# sequential CHUNK_SIZE_PAGES-page chunks, each classified on its own, then merged.
# Narrow to exactly that case — every other email is one classify call as before.
CHUNK_THRESHOLD_PAGES = 8   # only chunk a lone PDF with MORE pages than this
CHUNK_SIZE_PAGES = 6        # pages per chunk when splitting

# --- Storage mode (backward compatible) --------------------------------------
# No DATABASE_URL  -> Excel only, exactly as before.
# DATABASE_URL set -> write to PostgreSQL; pass --excel to ALSO write the Excel
#                     workbooks. --no-db forces the old Excel-only behavior.
# Drive upload: each saved invoice/payment file goes to Google Drive and its share link becomes the
#                     record's file path. ON by default in DB mode since 2026-10-08 — runs started by
#                     hand without --drive had left 298 records' documents only on this PC. --no-drive
#                     (or UPLOAD_TO_DRIVE=0) turns it off for a run; Excel-only mode stays off unless
#                     --drive / UPLOAD_TO_DRIVE=1. --drive is still accepted (Fetch New Invoices.bat).
_ARGS = set(sys.argv[1:])
USE_DB = bool(os.environ.get("DATABASE_URL")) and "--no-db" not in _ARGS
WRITE_EXCEL = (not USE_DB) or ("--excel" in _ARGS)


def _drive_enabled(args, env_value, use_db):
    """Whether this run uploads to Drive: --no-drive / UPLOAD_TO_DRIVE=0 always win; otherwise on with
    --drive, UPLOAD_TO_DRIVE=1, or DB mode."""
    env_value = (env_value or "").strip()
    if "--no-drive" in args or env_value == "0":
        return False
    return "--drive" in args or env_value == "1" or bool(use_db)


USE_DRIVE = _drive_enabled(_ARGS, os.environ.get("UPLOAD_TO_DRIVE"), USE_DB)
# One-off retries of ALREADY-queued items (opt-in; default runs are unaffected):
#   --retry-unreadable : re-run items flagged for an unreadable attachment/zip (run
#                        once after the auto-unzip improvement to clear the backlog).
#   --retry-open-all   : re-run EVERY open item once (broad — after any improvement).
RETRY_UNREADABLE = "--retry-unreadable" in _ARGS
RETRY_OPEN_ALL = "--retry-open-all" in _ARGS
# Reprocess ONE specific open review item by its displayed id (from the Review page), without
# touching any other open item. Ids are kept across runs (save_review_queue keeps each row's id).
RETRY_ID = next((a.split("=", 1)[1] for a in _ARGS if a.startswith("--retry-id=")), None)
# One-off maintenance: remove redundant re-record rows (the churned-filename dup bug).
DEDUPE_RECORDS = "--dedupe-records" in _ARGS
# Sales orders / order acknowledgements are HELD this many days before recording, so a
# formal invoice for the same vendor+order# can supersede them (no double entry).
try:
    SALES_ORDER_HOLD_DAYS = float(os.environ.get("SALES_ORDER_HOLD_DAYS", "2"))
except ValueError:
    SALES_ORDER_HOLD_DAYS = 2.0

# filename -> gmail deep-link, loaded once per run (see load_gmail_links()).
GMAIL_LINKS = {}

_DB = None
_DRIVE = None


def _database():
    """Lazy PostgreSQL handle (tables created on first use)."""
    global _DB
    if _DB is None:
        import database
        database.init_db()
        _DB = database
    return _DB


def _drive():
    global _DRIVE
    if _DRIVE is None:
        from drive_uploader import DriveUploader
        _DRIVE = DriveUploader()
    return _DRIVE


def load_gmail_links():
    """filename -> gmail deep-link, from gmail_fetcher's sidecar and/or the DB."""
    links = {}
    sidecar = OUTPUT_DIR / "gmail_links.json"
    if sidecar.exists():
        try:
            for fn, meta in json.loads(sidecar.read_text(encoding="utf-8")).items():
                if meta.get("gmail_link"):
                    links[fn] = meta["gmail_link"]
        except Exception as e:
            print(f"(could not read gmail_links.json: {e})")
    if USE_DB:
        try:
            links.update(_database().get_gmail_links())
        except Exception as e:
            print(f"(could not read gmail links from DB: {e})")
    return links


# Why Drive couldn't start this run (e.g. the Google login expired), so the rest of the run skips it
# instead of retrying for every file — and the documents that stayed local because of it or a failed
# upload, for the end-of-run reminder (upload_local_files_to_drive.py catches them up later).
_DRIVE_UNAVAILABLE = None
DRIVE_FALLBACKS = []


def upload_record_files(classification):
    """When Drive upload is on, upload each clean record's saved file to Google Drive
    and replace its _file_path with the shareable link. A file that still fails after the
    uploader's retries keeps its local path (noted in DRIVE_FALLBACKS); if Drive can't start at
    all, it's switched off for the rest of the run rather than failing the run."""
    global _DRIVE_UNAVAILABLE
    if not USE_DRIVE:
        return
    for rtype, rec in _all_records(classification):
        local = rec.get("_file_path")
        if local and os.path.exists(local):
            subfolder = "payment_files" if rtype == "payment" else "invoice_files"
            # Remember the local copy: if a record-time check flags this record after all, its review
            # item reuses this document (_late_review_document) instead of having none.
            rec["_local_file_path"] = local
            if _DRIVE_UNAVAILABLE:
                DRIVE_FALLBACKS.append(local)
                continue
            try:
                uploader = _drive()
            except (Exception, SystemExit) as e:      # google_auth raises SystemExit without credentials.json
                _DRIVE_UNAVAILABLE = str(e) or e.__class__.__name__
                print(f"  Drive upload unavailable this run ({_DRIVE_UNAVAILABLE}) — documents stay on this PC.")
                DRIVE_FALLBACKS.append(local)
                continue
            try:
                rec["_file_path"] = uploader.upload(subfolder, local)
            except Exception as e:
                DRIVE_FALLBACKS.append(local)
                print(f"  Drive upload failed for {os.path.basename(local)} after retries: {e} "
                      f"— kept the local copy.")

EMAILS_DIR = Path(__file__).parent / "test_emails"
OUTPUT_DIR = Path(__file__).parent / "output"
RESULTS_PATH = OUTPUT_DIR / "results.xlsx"
QB_READY_PATH = OUTPUT_DIR / "qb_ready.xlsx"
HISTORY_PATH = OUTPUT_DIR / "invoice_history.xlsx"
PAYMENT_HISTORY_PATH = OUTPUT_DIR / "payment_history.xlsx"
REVIEW_QUEUE_PATH = OUTPUT_DIR / "review_queue.xlsx"
LEARNINGS_PATH = OUTPUT_DIR / "learnings.xlsx"

TODAY = date.today().isoformat()

# review_queue.xlsx: items the AI flagged for human review land here INSTEAD of
# results.xlsx. The context columns are read-only (so you know what you're
# answering); "Answer" and "Notes to Learn" are the two columns you fill in.
# On the next run, answered rows are reprocessed (answer + learnings injected)
# and moved into results, then cleared from this queue. Unanswered rows stay.
# "Filename" is the standalone split file for THIS flagged record (in
# review_pending/); "Source Email" is the original email it came from. Because
# each flagged record is its own split file, answering + reprocessing touches
# only that record — the clean records in the same batch are untouched.
REVIEW_QUEUE_CONTEXT_FIELDS = [
    "Filename", "Source Email", "Date Flagged", "Email Type", "Vendor Matched",
    "Store Matched", "Invoice #", "Amount", "Question (AI review reason)",
    "Duplicate Group",
    "File Path", "Gmail Link",
]
REVIEW_QUEUE_ANSWER_FIELDS = ["Answer", "Notes to Learn"]
REVIEW_QUEUE_FIELDNAMES = REVIEW_QUEUE_CONTEXT_FIELDS + REVIEW_QUEUE_ANSWER_FIELDS

# learnings.xlsx: durable knowledge distilled from your "Notes to Learn". Its
# contents are injected into the system prompt on every run so the AI applies
# them to all future emails. You can hand-edit/prune this file.
LEARNINGS_FIELDNAMES = ["Learned Rule", "Date Added", "From Email"]

HISTORY_FIELDNAMES = [
    "vendor_matched", "store_matched", "invoice_number", "invoice_date",
    "invoice_amount", "source_file", "source_email_id", "date_added",
]

PAYMENT_HISTORY_FIELDNAMES = [
    "vendor_matched", "store_matched", "reference_number", "invoice_number", "payment_date",
    "amount", "source_file", "source_email_id", "date_added",
]

RESULTS_FIELDNAMES = [
    "Filename", "Date Added", "File Path", "Gmail Link", "Record Type", "Email Type", "Confidence", "Needs Review",
    "Review Reason", "Route To", "Routing Reason",
    "Vendor Raw", "Vendor Matched", "Vendor Confidence", "Skip QB Import",
    "Store Raw", "Store Matched", "Store Confidence",
    "Invoice #", "Invoice Date", "Amount", "Ship Charge", "Sales Tax", "Class",
    "Duplicate Suspect", "Duplicate Reason", "Duplicate Ignored", "Duplicate Of", "Notes",
    "Skipped Attachments", "Error", "Raw Response",
]

# "Entered in QB" is a manual column the user fills in after copy/pasting a row
# into QuickBooks; "Date Added" records when the row first appeared on this
# sheet. Both must survive every regeneration of qb_ready.xlsx (see
# load_qb_ready_preserved) rather than being recomputed from scratch.
QB_READY_FIELDNAMES = [
    "Filename", "Store", "Vendor", "Invoice #", "Invoice Date", "Amount",
    "Possible Duplicate", "Entered in QB", "Date Added", "File Path",
]

# Archived copy of every successfully-resolved source email, so test_emails/
# can be cleared out without losing the originals.
PROCESSED_DIR = Path(__file__).parent / "processed_emails"

# Emails flagged for review (or that couldn't be auto-read) live here — one
# place to open them while answering review_queue.xlsx. Each run reprocesses the
# ones you've answered; resolved emails move on to processed_emails/, the rest
# stay here.
REVIEW_PENDING_DIR = Path(__file__).parent / "review_pending"

# Per-record source documents (what "File Path" points to). Invoices and credit
# memos go in invoice_files/, payments (proof of payment) go in payment_files/.
# Each file is the attachment the record was read from — split to just that
# record's pages when one PDF holds several records — or a PDF render of the
# email body. Distinct from processed_emails/, which holds the raw email.
INVOICE_FILES_DIR = Path(__file__).parent / "invoice_files"
PAYMENT_FILES_DIR = Path(__file__).parent / "payment_files"


def _find_soffice():
    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ):
        if Path(candidate).exists():
            return candidate
    return shutil.which("soffice")


SOFFICE_PATH = _find_soffice()

# This machine's network runs traffic through a TLS-inspecting proxy (Techloq),
# whose root CA isn't in Python's default trust store. ca_bundle.pem is
# certifi's bundle plus that proxy's root cert; fall back to certifi defaults
# if it's not present (e.g. on a different machine).
CA_BUNDLE = Path(__file__).parent / "ca_bundle.pem"

# SYSTEM_PROMPT_TEMPLATE holds everything EXCEPT the two curated lists, which are
# now editable DB reference data (stores; the Exclude from QB flags on vendor_name_map). build_system_prompt()
# splices them in at run time via the %%STORE_LINES%% / %%EXCLUDED_VENDORS%% /
# %%CARD_STORE_MAP%% sentinels (plain .replace, not str.format — the prompt contains
# literal { } from the JSON schema). Seed defaults live in database.py (DEFAULT_STORES /
# DEFAULT_EXCLUDED_VENDORS / DEFAULT_CARD_STORE_MAP), which is also the Excel-mode fallback.
SYSTEM_PROMPT_TEMPLATE = """You are an invoice processing assistant for a multi-location pharmacy group. Your job is to analyze incoming emails and extract structured data.

Pharmacy entities, with known aliases/DBAs and physical addresses (use these addresses to resolve store identity when the store name on an invoice doesn't match one of the six names below):

%%STORE_LINES%%

Other known related entity (NOT one of the six operating pharmacies — if this address appears, flag needs_human_review rather than assuming it maps to a specific store): Lakewood Rx Inc, formerly Rx Fleet Inc — 10 Prag Blvd #203, Monroe, NY 10950.

An address match is a strong signal even when the name on the invoice doesn't match any entity name above — e.g. an invoice billed to "Lakewood Pharmacy" (not a real entity name) but shipped to 911 E County Line Rd, Lakewood NJ should be matched to Refuah Pharmacy by address, not left as unknown.

Vendors excluded from QuickBooks import (skip_qb_import: true): %%EXCLUDED_VENDORS%%

This list names the VENDOR that issued the document, never the store it's billed or shipped to. Some entries are our own pharmacies, because bills one of our stores sends another are excluded. A bill FROM any other vendor TO one of those pharmacies (the pharmacy appears as Bill To, Ship To or the customer) is a normal bill: set skip_qb_import false for it, unless the multi-store cost-split rule below applies. Decide skip_qb_import only by who issued the document.

IMPORTANT: this exclusion applies to INVOICES and CREDIT MEMOS only, NOT to payments. A payment is always entered in QuickBooks regardless of vendor — so for every payment record set skip_qb_import to false, even when the vendor is on the excluded list above.

Some vendors send ONE combined invoice (single invoice number/document) that splits a shared cost across MULTIPLE pharmacy store locations within that SAME document — each store gets its own line-item section and subtotal, rather than each store receiving its own separate invoice. Lakewood Rx enters this whole invoice as ONE combined bill on the Lakewood Rx entity, split across different QuickBooks accounts/classes manually — NOT as separate per-store QuickBooks entries. When an invoice or credit memo has this shape (multiple different store-name section headers, each with its own subtotal, all under ONE invoice number), set skip_qb_import: true for EVERY record produced from it (you should still produce one record per store section as usual, for the audit trail — just mark all of them skip_qb_import), and set that record's "notes" to something like "multi-store cost-split invoice — entered as one combined Lakewood Rx bill, not individually in QB." A current example: Sageteq's recurring monthly "IT Service Plan" invoice (vendor Sageteq, Bill To: Lakewood Rx, sections for Lakewood Rx / 3M Pharmacy / Carepax / Madison Pharmacy / Maple Medex Pharmacy / Refuah Pharmacy / Westgate Pharmacy, each with its own subtotal). Apply this rule by SHAPE, to any vendor's invoice that splits cost across multiple stores in one document this way — not only Sageteq. Do NOT apply it to an invoice that bills only ONE store, even from a vendor (like Sageteq) that also sometimes sends this kind of multi-store split invoice.

On ANY record (invoice, credit memo, or payment), if a card ending is shown or mentioned anywhere on the document or memo (e.g. "Visa ending 1234", "xxxx-1234", "card ****1234", "paid with card ending 1234"), ALWAYS capture its last 4 digits in that record's "card_last4" (digits only). Leave it null only when no card ending is shown.

Card last-4 to store: the card usually identifies which pharmacy the document belongs to. For any invoice, credit memo, or payment whose card ends in one of these last-4 digits, set that record's store_matched to the mapped store with HIGH confidence and do NOT flag it for review merely because the store would otherwise be uncertain:
%%CARD_STORE_MAP%%

Shipment/delivery notifications and account statements are informational, not actionable — they don't have an invoice or payment being submitted. Set needs_human_review to false and route_to to "none" for these UNLESS something is actually wrong (e.g. a disputed amount) — don't flag them for review just because the store or amount isn't fully certain.

A PROMOTIONAL/MARKETING email — a new-product announcement, a sale/discount blast, a newsletter (recognize by its own wording: "NEW from <vendor>", "meet our new product", "get X% off", "our sweetest launch yet", etc., and/or being sent to a mailing list rather than addressed to a specific order) — is also informational and should be routed to "none" with needs_human_review false, EVEN IF it has an attachment (a blank order form, catalog, price list, flyer) that this pipeline can't read. A blank order-form template is not an actual order or invoice, and the email's own promotional wording is enough by itself to know that, without needing to see the attachment. Reserve "an attachment that couldn't be read" as a review trigger for a document that otherwise reads like a REAL transaction — its own text says something like "invoice attached," "please remit," "balance due," or similar — where the unreadable attachment might genuinely be the bill itself.

Email date for sanity checks: do NOT trust the email's own "Date:" header as the email's date — some store scanners have mis-set internal clocks, so that header can be off by YEARS. If this message includes an "[EMAIL RECEIVED — authoritative ...]" line, treat THAT timestamp as the email's true date. Never flag an invoice for review merely because its invoice date is later than, or far from, the "Date:" header. You MAY still flag an invoice whose date is implausibly in the FUTURE relative to the authoritative received date (a real transmittal-date problem), or an invoice with no plausible date at all.

If an invoice is marked as REVISED, CORRECTED, AMENDED, or is a re-issue of an earlier invoice (often same or related invoice number, changed amount), set needs_human_review to true and explain in review_reason — these must be checked by a human before QB entry so a corrected invoice isn't entered alongside the original as a double entry.

Before citing a mismatch between the sum of visible line items and an invoice's stated total as a review reason, recount every line item carefully, one at a time, and re-add them — do not estimate. A miscounted sum produces a FALSE discrepancy that sends a perfectly correct invoice to a human for no reason. Only flag a line-item/total mismatch after you have verified the sum is actually correct and it still does not match the stated total.

A scanned or photographed multi-page document sometimes has its pages out of physical order (e.g. the page printed "Page 2" was scanned before the page printed "Page 1"), upside-down, or sideways. None of that is a reason to report a line-item/total mismatch. Before recounting (per the instruction above), first work out the TRUE reading order and orientation: use each page's own printed page number/label if present, or content continuity otherwise (a running item list, a "continued" marker, matching invoice number/header repeated on each page) — then read every page, including any upside-down or sideways one, in its correct upright orientation and true sequence, exactly as you would if it had been scanned right-side up and in order. Recount and total the line items using THIS corrected order/orientation, not the physical scan order. Only report a line-item/total mismatch as a review reason if it still doesn't reconcile after correcting for order and orientation — being out-of-sequence or upside-down is never by itself a valid reason to flag a mismatch.

Separately, before writing invoice_amount (and credit_amount / payment_amount), re-read the exact digits printed next to "Total" / "Amount Due" / "Balance" on the document one more time and confirm the number you are about to output matches them exactly, digit for digit — do not let a nearby number (a price, a quantity, an item code, a PO/order number) bleed into the total. When a single line item's Qty x Price already equals the printed total, that is a strong confirmation signal — use it to double-check your reading, not to substitute a different total than what's printed.

Set "total_source" on every record that has an amount. If you read the amount directly from a printed "Total" / "Amount Due" / "Balance" figure on the document, set total_source to "stated". If the document has NO printed Total / Amount Due / Balance anywhere — only itemized line items with no summary figure — set total_source to "derived" rather than inventing certainty; this is expected and fine, it is NOT an error to note in notes/review_reason by itself. Likewise, if a "Page X of Y" indicator (or any other evidence) shows this record's total-bearing page is absent from what you were given, set total_source to "derived" for whatever amount you could determine from what's visible. (A downstream code rule routes "derived" totals to human review, so you don't need to also flag them yourself.)

If a single email contains an invoice AND CONCRETE payment information (a check number, transaction/auth/confirmation number, card ending, ACH reference, or a distinct payment date/method — e.g. "paid via check #123"), produce TWO separate records: one in "invoices" for the bill and one in "payments" for the payment.

A vendor invoice that ALSO shows its own payment/remittance details (a "Payment against invoice #X", "Payment applied", or similar section on the SAME invoice showing it was settled, with a payment ID/date/amount) is NOT a payment-only document — it is a paid INVOICE. Create the INVOICE record for it exactly as you would for an unpaid one (full line items, vendor, store, total, etc.) AND a separate PAYMENT record for the settlement, with the payment's invoice_number field set to that invoice's number. When a single payment settles MULTIPLE invoices in one document (the same payment ID appears against more than one invoice), create one payment record per invoice it was applied to, each with its own invoice_number — do NOT collapse them into one payment record, and do NOT record only the payments while leaving the invoices out.

BUT: only create a payment record when there is ACTUAL payment information worth recording. If an invoice is merely marked "Paid" / "Balance $0.00" with NO payment identifier beyond the fact that it was paid and the same invoice amount, do NOT create a payment record — a payment row that only repeats the invoice amount adds no information. In that case record just the invoice (you may note "marked paid" in the invoice's notes). Never collapse a purchase into a single payment record.

When one email — or one scanned/combined PDF — contains MULTIPLE invoices, treat each invoice as an INDEPENDENT record. Set "needs_human_review" (and "review_reason") individually on each invoice/credit_memo/payment that actually needs it, and leave the others false so the clean ones can be processed normally. Do NOT flag an entire batch just because one invoice in it is questionable — the invoices are separate bills that merely happen to be scanned together. Use the TOP-LEVEL needs_human_review only for a whole-email problem (the email as a whole is ambiguous/spam/unreadable, or has no line-item records); when there are invoice/payment/credit-memo records, rely on each record's own needs_human_review.

A PAID order confirmation or receipt is still a PURCHASE, not merely a payment. When a document shows an order/order-summary/invoice that has already been paid or auto-charged (common with vendors like Shaklee that charge a card at order time), create an INVOICE record for the order (the bill). Add a separate PAYMENT record ONLY if concrete payment info is present (transaction/auth/confirmation number, card ending, check #, etc.) — not just because it says "Paid". Do NOT classify a purchase as only a payment just because it was already paid — the invoice/bill still needs to be recorded. This rule applies ONLY when the document shows the order was actually paid/charged — it does NOT apply to an unpaid sales order or order acknowledgement with no payment indication; that case is governed by is_sales_order above (set true, do not record it as an invoice).

Payment reminders / past-due notices / dunning emails: some emails are not a NEW bill at all — they are a reminder that an EARLIER, already-issued invoice is due or overdue ("payment reminder", "second notice", "past due", "your invoice #12345 is overdue", "friendly reminder your balance is due", a statement-style demand for one specific prior invoice). When an email is clearly this kind of reminder about a SPECIFIC prior invoice (it references that invoice's number and/or amount and is asking for payment of it, rather than presenting a fresh itemized bill for new goods/services), produce ONE invoice record for the referenced invoice as usual — fill in vendor_matched, invoice_number, and invoice_amount from what the reminder references (so the pipeline can match it against records on file) — and set "is_payment_reminder": true on that record. Do NOT set needs_human_review for a payment reminder, and do NOT treat it as a sales order. This flag only identifies the email as a reminder; it is NOT a duplicate judgment (you cannot see our records) — the pipeline itself decides whether the referenced invoice is already on file. Only use is_payment_reminder for a reminder that clearly points at a PRIOR invoice; a first-time bill that merely says "due upon receipt" or "due today" is a NORMAL invoice (is_payment_reminder false), not a reminder. If a reminder references SEVERAL prior invoices at once, that is a statement/aging report — classify it as email_type "statement" instead and use statement_invoices, not this flag.

route_to guidance: a real vendor bill that AP needs to handle should route to "accounts_payable" — this INCLUDES invoices from excluded (skip_qb_import) vendors, because those are still real bills AP pays, just not entered into QuickBooks. Do not route a real invoice to "none" merely because it's skip_qb_import. Use "none" only for purely informational emails (shipment notices, statements, recalls, etc.) that need no AP action.

For EACH invoice, payment, and credit memo, set "source_document" to indicate where that record's data came from: the exact attachment filename it was read from (the attachments are labeled "[Attachment N: filename]" in this message), or the literal string "email_body" if the data came from the email's own text rather than an attachment.

Also set "source_pages" for each record: when a single PDF attachment contains MULTIPLE records (e.g. a multi-page scan with several invoices), give the 1-based page range within that PDF that belongs to THIS record — e.g. "1", "2-3", or "1,4". If the record spans the whole document, or the source is a single page / an image / the email body, set source_pages to null. This lets each saved file contain only that record's own pages.

Also set "page_rotations" for each record: scans and photos (e.g. WhatsApp IMG-*.jpg) often arrive upside-down or sideways. For each page you attribute to this record, report the CLOCKWISE rotation in degrees — one of 0, 90, 180, or 270 — needed to make that page display UPRIGHT for a human reviewer. Report it as an object keyed by the 1-based page number WITHIN source_document (the SAME numbering as source_pages), e.g. {"1": 0, "2": 90} means page 2 must be turned 90° clockwise to be upright. If source_document is a single IMAGE (a photo/scan — source_pages is null), still report its rotation under key "1", e.g. {"1": 90} for a sideways photo. Use 0 (or omit the page) for pages that are already upright; set page_rotations to null if none need rotating. This ONLY rotates the saved file for review — it must NOT change any value you extract or which pages belong to which record. Correct ROTATION only (upside-down / sideways); do NOT report anything for mirrored or flipped pages (that's a reflection, not a rotation, and is out of scope).

Also OPTIONALLY set "blank_pages" and "page_order" for each record — both are CONSERVATIVE, opt-in cleanups of the saved file only (they never change any extracted value); MOST documents need NEITHER, so leave both null/empty unless clearly warranted. "blank_pages": a list of 1-based source_document page numbers (from among this record's source_pages) that are ENTIRELY blank — no visible content at all — which will be dropped from the saved file. Be strict: a mostly-white page that still shows a faint letterhead, footer, page number, or any stray marks is NOT blank; when unsure, OMIT it (default to keeping the page). Use null or an empty list if no page is blank. "page_order": normally null. Set it to a list that is an EXACT PERMUTATION of this record's source_pages, in the correct READING order, ONLY when you can clearly read a printed page indicator (e.g. "Page 2 of 4", a running header/footer) showing the physical scan order is wrong — e.g. source_pages "1-3" scanned out of order but printed "Page 3, Page 1, Page 2" would be page_order [2,3,1] to restore reading order. If the pages already read in correct physical order, OR you cannot confidently read the printed page numbers, leave page_order null — reordering on a guess is worse than leaving the scan order alone, so "no reordering" is the safe default.

Classify the email and return ONLY a valid JSON object — no commentary, no markdown:
{
  "email_type": "invoice | payment | statement | shipment_notice | correspondence | spam | unknown",
  "confidence": "high | medium | low",
  "needs_human_review": true/false,
  "review_reason": "string or null",
  "route_to": "accounts_payable | BMW | owner | none",
  "routing_reason": "string",
  "vendor_matched": "normalized vendor name if this email is about a specific vendor, else null — ALWAYS fill this in when identifiable, even for statements/shipment notices/correspondence that have no invoices/payments below",
  "store_matched": "Maple Pharmacy | Madison Pharmacy | Refuah Pharmacy | Westgate Pharmacy | 3M Pharmacy | BMW | Lakewood Rx | unknown, or null if not applicable — ALWAYS fill this in when identifiable, even for statements/shipment notices/correspondence",
  "reference_number": "order #, statement #, tracking #, or other identifying number mentioned, or null — used as a fallback identifier when there's no formal invoice/payment number. Do NOT use a persistent customer/account number (one that would appear identically on every document for this account); prefer a per-transaction identifier, else null.",
  "invoices": [
    {
      "vendor_name_raw": "as it appears in the email",
      "vendor_matched": "normalized vendor name",
      "vendor_confidence": "high | medium | low",
      "skip_qb_import": true/false,
      "store_raw": "as it appears in the email",
      "store_matched": "Maple Pharmacy | Madison Pharmacy | Refuah Pharmacy | Westgate Pharmacy | 3M Pharmacy | BMW | Lakewood Rx | unknown",
      "store_confidence": "high | medium | low",
      "invoice_number": "string or null",
      "invoice_date": "YYYY-MM-DD or null",
      "invoice_amount": number or null,
      "total_source": "'stated' if invoice_amount was read directly from a printed Total/Amount Due/Balance figure on the document, or 'derived' if NO such printed total exists anywhere on the document and this number was instead computed by summing the visible line items. Always set this — never leave it null when an amount is present.",
      "shipping_charge": "number or null — shipping/freight charge shown on the invoice, if any",
      "is_sales_order": "true if this document is a SALES ORDER / order acknowledgement / order confirmation rather than a formal invoice — e.g. titled 'Sales Order' / 'Order Acknowledgement' / 'Order Confirmation', shows an 'Order Total' (not an invoice amount / amount due), or the email says 'we have received your order'. Showing full, itemized pricing does NOT make it an invoice — sales order confirmations routinely show full pricing; that alone is never a reason to set this to false. If the document itself is titled/labeled 'Sales Order' or shows 'Order Number'/'Order Total' rather than an invoice number/amount due, set this true even if every line item and price is shown, even if the document is handwritten or informal, and even if there is no other reason to doubt the amount. Example: a small vendor slip headed 'Sales Order #' with a handwritten quantity/item/amount table is a sales order — is_sales_order: true — NOT an invoice, no matter how complete the pricing looks. Do not invent or assume a 'policy' for processing a sales order as an invoice — no such policy exists; the correct handling of a sales order is to set is_sales_order true and let it be held, not to record it as an invoice yourself. Otherwise false.",
      "sales_tax": "number or null — sales tax charged on the invoice, if any",
      "not_received_note": "string or null — if this document shows checkmarks, strikeouts, handwritten notes, or other markings indicating specific line items were NOT received / short-shipped / back-ordered, briefly describe what's marked (e.g. 'Item ABC-123 marked not received, qty 2'). Null if no such markup is present.",
      "not_received_amount": "number or null — if not_received_note is set AND the dollar impact of the marked items is determinable from the document (their line total, or an explicit adjustment shown), the amount by which those items reduce the invoice total. Null if not determinable or not_received_note is null.",
      "card_last4": "the last 4 digits of any card ending shown/mentioned on this document or memo (e.g. 'paid with card ending 1234'), digits only, or null if none is shown",
      "class": "the QuickBooks class/department if this document EXPLICITLY states one, else null (a default class is filled in later by rule)",
      "source_document": "attachment filename this invoice was read from, or 'email_body'",
      "source_pages": "page range within source_document for this record (e.g. '1', '2-3'), or null",
      "page_rotations": "object mapping 1-based source_document page -> clockwise degrees (0/90/180/270) to make it upright, e.g. {\"2\": 90}, or null",
      "blank_pages": "OPTIONAL list of 1-based source_document page numbers (within source_pages) that are ENTIRELY blank and should be dropped from the saved file; null/empty unless a page is truly blank (see instructions above — be strict)",
      "page_order": "OPTIONAL exact permutation of this record's source_pages giving correct reading order, ONLY when a printed 'Page X of Y' shows the scan order is wrong; null otherwise (the safe default)",
      "needs_human_review": true/false,
      "review_reason": "string or null — why THIS invoice needs review, if it does",
      "is_payment_reminder": "true ONLY if this is a payment-reminder / past-due / dunning notice about an EARLIER invoice, not a fresh bill (see the payment-reminder instruction above); false for a normal invoice — default false",
      "is_duplicate_suspect": false,
      "duplicate_reason": null,
      "notes": "anything unusual worth flagging"
    }
  ],
  "credit_memos": [
    {
      "vendor_name_raw": "as it appears in the email",
      "vendor_matched": "normalized vendor name",
      "vendor_confidence": "high | medium | low",
      "skip_qb_import": true/false,
      "store_raw": "as it appears in the email",
      "store_matched": "Maple Pharmacy | Madison Pharmacy | Refuah Pharmacy | Westgate Pharmacy | 3M Pharmacy | BMW | Lakewood Rx | unknown",
      "store_confidence": "high | medium | low",
      "credit_memo_number": "string or null",
      "credit_memo_date": "YYYY-MM-DD or null",
      "credit_amount": "number or null — the credit as a POSITIVE number (it will be entered in QB as a negative)",
      "total_source": "'stated' if credit_amount was read directly from a printed Total/Amount/Balance figure on the document, or 'derived' if NO such printed total exists and this number was instead computed by summing the visible line items. Always set this — never leave it null when an amount is present.",
      "card_last4": "the last 4 digits of any card ending shown/mentioned on this document or memo, digits only, or null if none is shown",
      "class": "the QuickBooks class/department if this document EXPLICITLY states one, else null (a default class is filled in later by rule)",
      "source_document": "attachment filename this credit memo was read from, or 'email_body'",
      "source_pages": "page range within source_document for this record (e.g. '1', '2-3'), or null",
      "page_rotations": "object mapping 1-based source_document page -> clockwise degrees (0/90/180/270) to make it upright, e.g. {\"2\": 90}, or null",
      "blank_pages": "OPTIONAL list of 1-based source_document page numbers (within source_pages) that are ENTIRELY blank and should be dropped from the saved file; null/empty unless a page is truly blank (see instructions above — be strict)",
      "page_order": "OPTIONAL exact permutation of this record's source_pages giving correct reading order, ONLY when a printed 'Page X of Y' shows the scan order is wrong; null otherwise (the safe default)",
      "needs_human_review": true/false,
      "review_reason": "string or null — why THIS credit memo needs review, if it does",
      "notes": "anything unusual worth flagging"
    }
  ],
  "payments": [
    {
      "vendor_name_raw": "string",
      "vendor_matched": "string",
      "skip_qb_import": false,
      "store_matched": "string",
      "payment_amount": number or null,
      "total_source": "'stated' if payment_amount was read directly from a printed total/amount-paid figure on the receipt, or 'derived' if NO such printed total exists and this number was instead computed by summing line items. Always set this — never leave it null when an amount is present.",
      "payment_date": "YYYY-MM-DD or null",
      "reference_number": "the UNIQUE identifier for THIS SPECIFIC transaction — a check number, confirmation/authorization code, or per-payment transaction/receipt ID — or null if none is shown. Do NOT use a persistent customer/account number here (e.g. a 'Customer #', 'Account #', or similar ID that would appear identically on EVERY payment receipt this vendor ever sends for this same account/store) — that is not a per-transaction identifier and using it as reference_number will cause completely unrelated payments to be wrongly flagged as duplicates of each other. If the only number shown is this kind of persistent account/customer number and there's no separate per-transaction reference, leave this null.",
      "invoice_number": "the specific invoice number this payment was applied against, if the document states one (e.g. a line reading 'Payment against invoice #TRINV-685027' or similar remittance detail tied to one invoice) — or null if this payment isn't tied to one specific invoice number. This is DIFFERENT from reference_number (the transaction/check/ACH id, which can legitimately repeat across several invoices settled by the same payment) — invoice_number is what disambiguates which invoice each application belongs to when that happens.",
      "card_last4": "the last 4 digits of the card used for this payment (digits only, string), or null if not a card payment or not shown",
      "class": "the QuickBooks class/department if this payment EXPLICITLY states one, else null (a default class is filled in later by rule)",
      "source_document": "attachment filename this payment was read from, or 'email_body'",
      "source_pages": "page range within source_document for this record (e.g. '1', '2-3'), or null",
      "page_rotations": "object mapping 1-based source_document page -> clockwise degrees (0/90/180/270) to make it upright, e.g. {\"2\": 90}, or null",
      "blank_pages": "OPTIONAL list of 1-based source_document page numbers (within source_pages) that are ENTIRELY blank and should be dropped from the saved file; null/empty unless a page is truly blank (see instructions above — be strict)",
      "page_order": "OPTIONAL exact permutation of this record's source_pages giving correct reading order, ONLY when a printed 'Page X of Y' shows the scan order is wrong; null otherwise (the safe default)",
      "needs_human_review": true/false,
      "review_reason": "string or null — why THIS payment needs review, if it does",
      "notes": "string"
    }
  ],
  "statement_invoices": [
    {
      "vendor": "vendor the statement is from (normalized), or the vendor named for this line",
      "invoice_number": "the invoice/document number as listed on the statement",
      "invoice_date": "YYYY-MM-DD or null",
      "amount": "number or null — the amount the statement lists for this invoice"
    }
  ]
}

ONLY when email_type is "statement" (a vendor account statement or aging report that itemizes open/paid invoices), populate "statement_invoices" with EVERY invoice the statement references — one entry per listed invoice — so we can reconcile them against our records. This is reconciliation metadata; the statement itself still records no invoice/payment/credit_memo. For any NON-statement email, omit "statement_invoices" or return [].

Always set is_duplicate_suspect to false and duplicate_reason to null. Do not infer duplication from wording or tone (e.g. "reminder", "adjusted amount") — you cannot see the pharmacy's invoice records, so you have no basis for that judgment. Duplicate detection is performed separately by comparing this invoice against previously recorded invoices."""


# --------------------------------------------------------------------------- #
# Editable reference data (stores / vendor exclusions / sender_store_map).
# Read from the DB when configured (curated in the web UI), else the seed
# defaults in database.py — so behavior is identical before any curation and in
# Excel-only mode. See PROJECT_STATUS.md "Reference tables".
# --------------------------------------------------------------------------- #

# sender email -> store, loaded once per run in main(); empty until then so the
# store signal is simply absent (never wrong) in code paths that don't set it.
SENDER_STORE_MAP = {}
# recipient (To/Cc) email -> store, same lifecycle: identifies the store on email
# FROM third-party vendors (the mirror of SENDER_STORE_MAP). Empty until main() loads it.
CONTACT_STORE_MAP = {}
# card last-4 -> store, loaded once per run in main() (same lifecycle as above);
# empty until then, so a payment's store resolution is unaffected before it loads.
CARD_STORE_MAP = {}
# QuickBooks class-default rules (list of {store_matched, vendor_matched, class}),
# loaded once per run in main(); empty until then, so class is left as the model set it.
CLASS_RULES = []
# Vendor/account standardization data, loaded once per run in main() (DB-only). Empty until
# then / in Excel mode, which makes apply_vendor_standardization() a no-op.
VENDOR_STD = {}
VENDOR_FUZZY_THRESHOLD = 88   # 0-100 scale; tune here if still too strict/loose
VENDOR_FUZZY_MARGIN = 4       # best match must beat the runner-up by this much
# filename -> authoritative received time (ISO string), loaded once per run.
RECEIVED_AT = {}
# split-file name -> ORIGINAL source-email filename, set during Phase-1 reprocess so
# a reprocessed record keeps the original email's identity for self-duplicate checks.
ORIGINAL_EMAIL = {}
_SYSTEM_PROMPT_CACHE = None


def load_received_dates():
    """filename -> authoritative received time (ISO string), from gmail_fetcher's
    sidecar and/or the emails table. Used so the invoice-date sanity check runs
    against the TRUE receive time, never the scanner-clock 'Date:' header."""
    received = {}
    sidecar = OUTPUT_DIR / "gmail_links.json"
    if sidecar.exists():
        try:
            for fn, meta in json.loads(sidecar.read_text(encoding="utf-8")).items():
                if meta.get("received_at"):
                    received[fn] = meta["received_at"]
        except Exception as e:
            print(f"(could not read received dates from gmail_links.json: {e})")
    if USE_DB:
        try:
            received.update(_database().get_received_dates())
        except Exception as e:
            print(f"(could not read received dates from DB: {e})")
    return received


def _received_from_eml(path):
    """Fallback true-receive time for a manually-dropped .eml: the TOP 'Received:'
    header (added last by the delivering server, so closest to real arrival).
    Returns a datetime or None."""
    if path.suffix.lower() != ".eml":
        return None
    try:
        with open(path, "rb") as f:
            msg = email.message_from_binary_file(f, policy=email_policy)
        for hdr in (msg.get_all("Received") or []):
            if ";" in hdr:
                try:
                    return parsedate_to_datetime(hdr.rsplit(";", 1)[1].strip())
                except (TypeError, ValueError, IndexError):
                    continue
    except Exception:
        pass
    return None


def _received_date_str(path):
    """The authoritative received date for this source file as 'YYYY-MM-DD HH:MM'
    (or None). Prefers the captured internalDate (RECEIVED_AT), falling back to the
    .eml's top Received header. NEVER uses the unreliable 'Date:' header."""
    iso = RECEIVED_AT.get(path.name)
    if iso:
        try:
            return datetime.fromisoformat(iso).strftime("%Y-%m-%d %H:%M %Z").strip()
        except ValueError:
            return iso
    dt = _received_from_eml(path)
    return dt.strftime("%Y-%m-%d %H:%M %Z").strip() if dt else None


def load_excluded_vendors():
    """Vendor names on the QB-exclusion list (DB table, or seed defaults)."""
    if USE_DB:
        try:
            return _database().load_excluded_vendors()
        except Exception as e:
            print(f"(could not read excluded_vendors from DB, using defaults: {e})")
    import database
    return list(database.DEFAULT_EXCLUDED_VENDORS)


def load_stores():
    """The pharmacy entities (name/aliases/address/notes) — DB table or defaults."""
    if USE_DB:
        try:
            return _database().load_stores()
        except Exception as e:
            print(f"(could not read stores from DB, using defaults: {e})")
    import database
    return [dict(s) for s in database.DEFAULT_STORES]


def load_sender_store_map():
    """{normalized sender_email: store} for self-scan mailboxes — DB or defaults."""
    if USE_DB:
        try:
            return _database().load_sender_store_map()
        except Exception as e:
            print(f"(could not read sender_store_map from DB, using defaults: {e})")
    import database
    return {e.strip().lower(): s for e, s in database.DEFAULT_SENDER_STORE_MAP}


def load_contact_store_map():
    """{normalized contact_email: store} for To/Cc addresses that identify a store on
    email FROM third-party vendors — DB or defaults."""
    if USE_DB:
        try:
            return _database().load_contact_store_map()
        except Exception as e:
            print(f"(could not read contact_store_map from DB, using defaults: {e})")
    import database
    return {e.strip().lower(): s for e, s in database.DEFAULT_CONTACT_STORE_MAP}


def _norm_last4(v):
    """Normalize a card last-4: strip non-digits, keep the trailing 4 digits."""
    digits = "".join(ch for ch in str(v or "") if ch.isdigit())
    return digits[-4:] if digits else ""


def load_card_store_map():
    """{normalized card_last4: store} for cards whose store is known — DB or defaults."""
    if USE_DB:
        try:
            return _database().load_card_store_map()
        except Exception as e:
            print(f"(could not read card_store_map from DB, using defaults: {e})")
    import database
    return {_norm_last4(c): s for c, s in database.DEFAULT_CARD_STORE_MAP}


def apply_card_store_map(rec):
    """Deterministic store override (safety net): if this record's card_last4 is on
    the card→store map, set its store_matched to the mapped store. One card = one
    store, so this is authoritative and overrides the model. Applies uniformly to
    invoice, credit_memo and payment records; mirrors the prompt's map instruction.
    Additive — a no-op when the record has no mapped card (CARD_STORE_MAP empty in
    Excel mode until curated), so the CLI is unaffected until cards are mapped."""
    last4 = _norm_last4(rec.get("card_last4"))
    if last4 and last4 in CARD_STORE_MAP:
        rec["store_matched"] = CARD_STORE_MAP[last4]


def load_class_rules():
    """The QuickBooks class-default rules — DB (curated) or the seed defaults. Each is
    {store_matched, vendor_matched, class}; blank store/vendor = 'any'."""
    if USE_DB:
        try:
            return _database().load_class_rules()
        except Exception as e:
            print(f"(could not read class_rules from DB, using defaults: {e})")
    import database
    return [dict(r) for r in database.DEFAULT_CLASS_RULES]


def apply_class_rules(rec):
    """Fill a record's QuickBooks class from the best-matching class_rule ONLY when the
    document didn't already state one (an explicit class always wins). Deterministic,
    post-classification. A rule matches when (rule.store blank OR == the record's store)
    AND (rule.vendor blank OR == the record's vendor), with at least one of store/vendor
    specified. Precedence: both-fields > vendor-only > store-only. No-op when CLASS_RULES
    is empty (Excel mode / uncurated), so the CLI is unaffected until rules are added."""
    if (rec.get("class") or "").strip():
        return                                    # explicit class from the doc wins
    store = _norm(rec.get("store_matched"))
    vendor = _norm(rec.get("vendor_matched"))
    best = None      # (precedence, class); higher precedence wins
    for r in CLASS_RULES:
        rs = _norm(r.get("store_matched"))
        rv = _norm(r.get("vendor_matched"))
        if not rs and not rv:
            continue                              # a rule must specify store and/or vendor
        if rs and rs != store:
            continue
        if rv and rv != vendor:
            continue
        prec = (2 if (rs and rv) else 1 if rv else 0)   # both > vendor-only > store-only
        cls = (r.get("class") or "").strip()
        if cls and (best is None or prec > best[0]):
            best = (prec, cls)
    if best:
        rec["class"] = best[1]


def apply_vendor_exclusion(rec):
    """Deterministic safety net: force skip_qb_import True whenever the vendor is on the
    excluded_vendors list, regardless of what the model set. The model is still told the list
    in the prompt (so it can reason about it correctly in the first place), but that alone
    occasionally misses — this makes the outcome certain instead of probabilistic.
    Invoices/credit memos only; payments are NEVER excluded from QB entry regardless of vendor
    (existing rule, unchanged). Only ever flips False -> True, never the reverse — no regression
    risk, it can only catch cases the model already should have caught."""
    if is_excluded_vendor(rec.get("vendor_matched")):
        rec["skip_qb_import"] = True


# The prompt asks the model to write a note like this on a multi-store cost-split invoice.
_COST_SPLIT_NOTE = re.compile(r"cost[\s-]*split|combined lakewood rx bill", re.I)


def _cost_split_numbers(classification):
    """Invoice / credit memo numbers this ONE email splits across several stores: the multi-store
    cost-split shape (one record per store section, all under one number), which the prompt marks
    skip_qb_import whatever the vendor."""
    stores = {}
    for rtype, rec in _all_records(classification or {}):
        if rtype == "payment":
            continue
        number = _record_fields(rtype, rec)[2]
        if number:
            stores.setdefault(_norm(str(number)), set()).add(_norm(rec.get("store_matched")))
    return {n for n, s in stores.items() if len(s) > 1}


def apply_unlisted_skip_gate(rec, rtype, classification):
    """Deterministic safety net (2026-10-12), the other direction from apply_vendor_exclusion: an
    invoice / credit memo the model marked skip_qb_import although its VENDOR isn't on the exclusion
    list is not excluded silently. The usual cause: our own pharmacies (3M, Refuah, Westgate, Madison)
    are on the list for bills between the stores, and the model applied it to the store being BILLED —
    24 real bills by 2026-10-12, its own notes saying e.g. "Refuah is the store, Easyfast is the vendor".
    Clears the flag and sends the record to review, so a person confirms it before it reaches QuickBooks.
    Leaves alone: excluded vendors, held sales orders, the multi-store cost-split shape (the same number
    for several stores in this email, or the model's cost-split note), and an email reprocessed with a
    reviewer's answer (the model followed the person, so this can't loop). Idempotent."""
    if rtype not in ("invoice", "credit_memo") or not rec.get("skip_qb_import") or rec.get("is_sales_order"):
        return
    vendor = (rec.get("vendor_matched") or "").strip()
    if not vendor or is_excluded_vendor(vendor) or (classification or {}).get("_reviewer_answered"):
        return
    number = _record_fields(rtype, rec)[2]
    if (number and _norm(str(number)) in _cost_split_numbers(classification)) \
            or _COST_SPLIT_NOTE.search(rec.get("notes") or ""):
        return
    store = rec.get("store_matched")
    rec["skip_qb_import"] = False
    rec["needs_human_review"] = True
    why = (f", apparently because the store it's billed to ({store}) is on the excluded-vendor list — but the "
           f"vendor ({vendor}) isn't" if store and is_excluded_vendor(store)
           else f", but {vendor} isn't on the excluded-vendor list")
    _append_text(rec, "review_reason",
                 f"The AI marked this bill not for QuickBooks{why}. Confirm it's a real bill to enter in "
                 f"QuickBooks (answer \"enter it in QuickBooks\"), or Ignore it if it shouldn't be entered.")


# Credit memos recorded so far, keyed (standard-vendor dedupe key, normalized number), plus the clean
# ones seen earlier in this run. Loaded on first use (one run = one process); None until then.
_CREDIT_MEMO_POOL = None


def _credit_memo_pool():
    global _CREDIT_MEMO_POOL
    if _CREDIT_MEMO_POOL is None:
        _CREDIT_MEMO_POOL = {}
        rows = []
        if USE_DB:
            try:
                rows = _database().load_credit_memo_keys()
            except Exception as e:
                print(f"(could not load credit memos for the duplicate check: {e})")
        for r in rows:
            key = (_dedupe_vendor_key(r.get("vendor_matched")), _norm(r.get("credit_memo_number")))
            if key[0] and key[1]:
                _CREDIT_MEMO_POOL.setdefault(key, []).append(r)
    return _CREDIT_MEMO_POOL


def apply_credit_memo_duplicate_gate(classification, filename):
    """Credit memos have no duplicate ledger, so before 2026-10-12 a credit sent twice (vendor email +
    store scan, or a resend) reached QB entry twice and was taken twice. This sends a credit memo to
    review when one with the same vendor (standard name, like DU-01) and number was already recorded
    from a different email, or seen clean earlier in this run: a person ignores the second copy or
    confirms it's a separate credit. It never auto-ignores (credit memos have no duplicate columns).
    Skips skip-QB and excluded-vendor credits (never entered), credits without a vendor or number, and an email reprocessed
    with a reviewer's answer (the person has seen this reason, so it can't loop). The same email (a
    reprocess, or the split file of the original email) never matches itself. Runs right after
    apply_review_gates, before documents are filed, so a flagged copy gets its own review document."""
    if not classification or classification.get("_reviewer_answered"):
        return
    source_id = _record_source_id(filename)
    own = {filename, source_id}
    pool = _credit_memo_pool()
    for cm in classification.get("credit_memos") or []:
        if cm.get("skip_qb_import") or is_excluded_vendor(cm.get("vendor_matched")):
            continue
        key = (_dedupe_vendor_key(cm.get("vendor_matched")), _norm(cm.get("credit_memo_number")))
        if not key[0] or not key[1]:
            continue
        priors = [p for p in pool.get(key, []) if p.get("email_filename") not in own]
        if not priors:
            if not record_flagged(cm):
                pool.setdefault(key, []).append({
                    "id": None, "vendor_matched": cm.get("vendor_matched"),
                    "credit_memo_number": cm.get("credit_memo_number"),
                    "amount": cm.get("credit_amount"), "email_filename": source_id,
                    "entered_in_qb": False})
            continue

        def _amt(v):
            try:
                return abs(float(v))
            except (TypeError, ValueError):
                return None
        this_amt = _amt(cm.get("credit_amount"))
        same = [p for p in priors if this_amt is not None and _amt(p.get("amount")) is not None
                and abs(_amt(p.get("amount")) - this_amt) < 0.005]
        p = (same or priors)[0]
        where = (f"credit memo record #{p['id']}" if p.get("id")
                 else f"earlier in this run, from {p.get('email_filename')}")
        p_amt = _amt(p.get("amount"))
        in_qb = ", already entered in QuickBooks" if p.get("entered_in_qb") else ""
        num, vendor = cm.get("credit_memo_number"), cm.get("vendor_matched")
        if same:
            note = (f"Credit memo #{num} from {vendor} was already recorded from another email ({where}"
                    f"{', $' + format(p_amt, ',.2f') if p_amt is not None else ''}{in_qb}) — this looks "
                    f"like a second copy of the same credit. Ignore this item if it is, or answer \"enter "
                    f"it in QuickBooks\" if it's a separate credit.")
        else:
            note = (f"Credit memo #{num} from {vendor} was already recorded from another email ({where}"
                    f"{', for $' + format(p_amt, ',.2f') if p_amt is not None else ''}{in_qb}); this copy "
                    f"is {'$' + format(this_amt, ',.2f') if this_amt is not None else 'missing an amount'} "
                    f"— possibly a revised credit memo. Ignore this item if the earlier one stands, or "
                    f"answer with what to enter (and ignore the earlier record on its page if this one "
                    f"replaces it).")
        cm["needs_human_review"] = True
        _append_text(cm, "review_reason", note)


def _append_text(rec, field, note):
    """Append `note` to rec[field] (space-separated) unless that exact sentence is already there.
    The review gates run twice on a record — at classification time (apply_review_gates) and again
    as a safety net at record time (build_result_rows) — and each used to append its sentence on
    BOTH passes, so the reviewer saw it doubled ("No invoice date found … No invoice date found …")."""
    current = rec.get(field) or ""
    if note in current:
        return
    rec[field] = (current + " " if current else "") + note


def apply_store_review_gate(rec, is_payment=False):
    """Deterministic safety net: a record with no identifiable store must not reach
    QB entry silently. Call AFTER apply_card_store_map() has had its chance to fill
    store_matched from the card->store map. If it's still null/blank/"unknown", force
    human review instead of letting an unresolved store through.

    EXCLUDED-VENDOR invoices/credit memos are skipped: they're kept out of QB via
    skip_qb_import regardless, so an unknown store there is harmless and reviewing it
    buys nothing. PAYMENTS are NEVER excluded (all vendors enter QB as payments), so a
    payment is ALWAYS gated on an unknown store regardless of vendor — pass is_payment=True.
    Only ever flips False -> True, same convention as the other apply_* safety nets."""
    store = (rec.get("store_matched") or "").strip().lower()
    if store and store != "unknown":
        return
    if not is_payment and is_excluded_vendor(rec.get("vendor_matched")):
        return  # excluded-vendor invoice/credit memo never reaches QB — don't force review
    rec["needs_human_review"] = True
    note = ("Store could not be identified — needs a human to assign the correct "
            "store before this goes to QuickBooks.")
    _append_text(rec, "review_reason", note)


def apply_date_review_gate(rec, date_field):
    """Deterministic safety net: a record with no date must not reach QB entry
    silently — QuickBooks needs a real transaction date. Pass the record's own date
    field name (invoice_date / credit_memo_date / payment_date). Only ever flips
    False -> True, same convention as the other apply_* safety nets."""
    if not (rec.get(date_field) or "").strip():
        rec["needs_human_review"] = True
        note = (f"No {date_field.replace('_', ' ')} found — needs a human to confirm "
                "or enter the date before this goes to QuickBooks.")
        _append_text(rec, "review_reason", note)


def _trusts_line_item_totals(rec):
    """True when the record's vendor is ticked "Trust line-item totals" in Setup › Vendor names.
    Matched EXACTLY on any of that vendor's spellings (standard name, aliases, per-store names — the
    VENDOR_STD lookup), never fuzzily, so a near-miss vendor is still checked. False in Excel mode or
    before the standardization tables load."""
    trusted = (VENDOR_STD or {}).get("trusted_totals")
    if not trusted:
        return False
    name = rec.get("vendor_matched") or rec.get("vendor_name_raw") or ""
    return VENDOR_STD["lookup"].get(_norm_vendor_std_key(name)) in trusted


def apply_amount_review_gate(rec, amount_field):
    """Deterministic safety net: a record with no usable amount, or one whose amount was
    COMPUTED from line items rather than read from a printed total, must not reach QB entry
    silently. Pass the record's own amount field name (invoice_amount / credit_amount /
    payment_amount). Unlike apply_store_review_gate, this is NOT skipped for excluded-vendor
    invoices — a missing/unverified dollar amount still matters for AP even when the bill isn't
    entered into QuickBooks (excluded vendors are still real bills AP pays). Only ever flips
    needs_human_review False -> True, same convention as the other apply_* safety nets.

    Two checks:
      1) amount missing / zero / negative -> flag (fully deterministic).
      2) total_source == 'derived' -> flag: the amount was summed from line items because NO
         printed Total/Amount Due/Balance was found (often a missing total-bearing page), so it
         needs a human to verify. This depends on the model self-reporting total_source; a
         missing/None total_source is a NO-OP (absence is NOT treated as derived OR stated).
         EXCEPT a vendor ticked "Trust line-item totals" in Setup (bills that only ever list line
         items): accepted with a note instead (_trusts_line_item_totals). Check 1 still applies."""
    amt = rec.get(amount_field)
    if amt is None or (isinstance(amt, (int, float)) and amt <= 0):
        rec["needs_human_review"] = True
        note = ("No usable amount was found (missing, zero, or negative) — needs a human to "
                "confirm or enter the correct amount before this goes to QuickBooks.")
        _append_text(rec, "review_reason", note)
    elif rec.get("total_source") == "derived" and _trusts_line_item_totals(rec):
        _append_text(rec, "notes", "[Amount added up from the line items (no printed total) — accepted "
                                   "because this vendor is set to trust line-item totals in Setup.]")
    elif rec.get("total_source") == "derived":
        rec["needs_human_review"] = True
        note = ("This amount was calculated from line items, not read from a printed total on the "
                "document (no stated Total/Amount Due/Balance was found, possibly because a page is "
                "missing) — needs a human to verify the correct total before this goes to QuickBooks.")
        _append_text(rec, "review_reason", note)


def apply_review_gates(classification):
    """Apply the deterministic review gates (card→store, store / scoped-learning / date / amount) to
    EVERY record, so `needs_human_review` is FINAL at classification time — before anything downstream
    (save_review_files, newly_flagged, apply_duplicate_checks, build_result_rows) decides the record's
    fate. These gates USED to run only inside build_result_rows, which executes AFTER
    apply_duplicate_checks has already added the record to invoice_history AND after newly_flagged was
    computed — so a gate-flagged record fell into a crack: in history, not recorded, not queued for
    review, and its emails-row status never updated. Running them here closes that. Idempotent (gates
    only flip needs_human_review False→True, and never repeat their sentence); build_result_rows
    re-runs them as a harmless safety net. The vendor/account check joins them here (2026-10-07) so a
    record it flags gets its own review_pending/ split file + preview — see
    apply_vendor_standardization_gate()."""
    if not classification:
        return
    for inv in classification.get("invoices") or []:
        if inv.get("is_sales_order"):
            continue                              # held, not recorded/reviewed here
        apply_card_store_map(inv)
        apply_store_review_gate(inv)
        apply_scoped_learning_gate(inv)
        apply_date_review_gate(inv, "invoice_date")
        apply_amount_review_gate(inv, "invoice_amount")
        apply_vendor_standardization_gate(inv)
        apply_unlisted_skip_gate(inv, "invoice", classification)   # before files are saved: a real bill gets one
    for cm in classification.get("credit_memos") or []:
        apply_card_store_map(cm)
        apply_store_review_gate(cm)
        apply_scoped_learning_gate(cm)
        apply_date_review_gate(cm, "credit_memo_date")
        apply_amount_review_gate(cm, "credit_amount")
        apply_vendor_standardization_gate(cm)
        apply_unlisted_skip_gate(cm, "credit_memo", classification)
    for pay in classification.get("payments") or []:
        apply_card_store_map(pay)
        apply_store_review_gate(pay, is_payment=True)
        apply_date_review_gate(pay, "payment_date")
        apply_amount_review_gate(pay, "payment_amount")


# Vendor-scoped store rules (Setup › Store rules / store_scope_rules), loaded once per run in main().
# None until then -> the seed defaults below (Excel mode, or a script that never loads them).
STORE_SCOPE_RULES = None
_STORE_RULE_ONLY, _STORE_RULE_EXCEPT = "only these vendors", "every vendor except these"
_DEFAULT_STORE_SCOPE_RULES = [
    {"reference": "Carepax", "store_matched": "BMW", "applies_to": _STORE_RULE_ONLY,
     "vendors": ["Sageteq", "Microsoft"]},
    {"reference": "Accuscript", "store_matched": "Lakewood Rx", "applies_to": _STORE_RULE_ONLY,
     "vendors": ["East Coast Building Maintenance"]},
    {"reference": "112 Hillside", "store_matched": "Westgate Pharmacy", "applies_to": _STORE_RULE_EXCEPT,
     "vendors": ["Braun"]},
]


def load_store_scope_rules():
    """The Setup › Store rules for this run (DB), else the seed defaults."""
    if USE_DB:
        try:
            return _database().load_store_scope_rules()
        except Exception as e:
            print(f"(could not read store rules from DB, using defaults: {e})")
    return [dict(r) for r in _DEFAULT_STORE_SCOPE_RULES]


def _squash(text):
    """Lowercase + collapse whitespace, so '112  Hillside' / '112 HILLSIDE' match '112 Hillside'."""
    return " ".join((text or "").lower().split())


def apply_scoped_learning_gate(rec):
    """Deterministic safety net for standing-learning SCOPE LEAKS. Standing learnings are injected
    into the prompt for EVERY email, so a store mapping meant for one vendor can leak onto another
    vendor's invoice — the Sageteq Bill-To learning's 'Carepax = BMW' did, on Abbott #1605/#1604
    ('Carepax' is not itself a store). Each Setup › Store rules row enforces one vendor scope in CODE
    instead of relying on the model to self-limit per the prompt wording (2026-10-08 — this used to be
    one hard-coded Carepax→BMW check, now the first seeded rule):
      'only these vendors'        — store resolved to <store> from a <reference> mention (store_raw)
                                    on any other vendor -> review. (Carepax→BMW: Sageteq, Microsoft;
                                    Accuscript→Lakewood Rx: East Coast Building Maintenance.)
      'every vendor except these' — a <reference> mention resolved to a DIFFERENT store on any other
                                    vendor -> review. (112 Hillside = Westgate, except Braun, which
                                    bills every location there.)
    A vendor matches when one of the rule's vendor names appears in vendor_matched (case-insensitive).
    Only ever flips needs_human_review False→True (same convention as the other apply_* safety nets);
    never fires when store_raw doesn't mention the reference (e.g. an invoice that literally names
    'BMW Technologies LLC')."""
    store_raw = _squash(rec.get("store_raw"))
    if not store_raw:
        return
    store = _norm(rec.get("store_matched"))
    vendor = _norm(rec.get("vendor_matched"))
    rules = STORE_SCOPE_RULES if STORE_SCOPE_RULES is not None else _DEFAULT_STORE_SCOPE_RULES
    for rule in rules:
        ref = _squash(rule.get("reference"))
        rule_store = _norm(rule.get("store_matched"))
        if not ref or not rule_store or ref not in store_raw:
            continue
        names = [v for v in (rule.get("vendors") or []) if _norm(v)]
        listed = any(_norm(v) in vendor for v in names) if vendor else False
        who = ", ".join(names) or "no vendor"
        label = rule.get("store_matched")
        if _norm(rule.get("applies_to")) == _STORE_RULE_EXCEPT:
            if store == rule_store or listed or store in ("", "unknown"):
                continue                          # (no store at all is the store check's job)
            note = (f"This record mentions '{rule.get('reference')}', which means {label} for every vendor "
                    f"except {who} (Setup › Store rules), but its store was set to "
                    f"{rec.get('store_matched') or 'none'} — a human must confirm the real store before this "
                    f"goes to QuickBooks.")
        else:
            if store != rule_store or listed:
                continue
            note = (f"Store was resolved to {label} from a '{rule.get('reference')}' reference on a "
                    f"{rec.get('vendor_matched') or 'unnamed-vendor'} record. '{rule.get('reference')}' = "
                    f"{label} is a rule for {who} only (Setup › Store rules), so a human must confirm the "
                    f"real store before this goes to QuickBooks.")
        rec["needs_human_review"] = True
        _append_text(rec, "review_reason", note)


def load_vendor_standardization():
    """Everything needed to standardize a vendor + resolve its store-specific QB names, built
    once per run. DB-only feature — returns {} if not USE_DB, which makes
    apply_vendor_standardization() a no-op until it's seeded. Re-keys the DB's bare lookup with
    THIS file's real normalization so all the fuzzy/normalization logic lives in one place."""
    if not USE_DB:
        return {}
    try:
        raw = _database().load_vendor_standardization()
    except Exception as e:
        print(f"(could not read vendor standardization tables from DB: {e})")
        return {}
    raw["lookup"] = {_norm_vendor_std_key(k): v for k, v in raw["lookup"].items()}
    return raw


_STORE_KEY = {
    "maple pharmacy": "maple", "madison pharmacy": "madison",
    "refuah pharmacy": "refuah", "westgate pharmacy": "westgate",
    "3m pharmacy": "3m", "bmw": "bmw", "lakewood rx": "lakewood_rx",
}


def _norm_vendor_std_key(value):
    """Normalization for vendor-standardization matching (separate from _norm_vendor_dedupe,
    which is duplicate-comparison-only, though it reuses the same legal-suffix stripping):
    lowercase, strip legal suffixes, '&'->'and', drop remaining punctuation, collapse
    whitespace."""
    name = _norm_vendor_dedupe(value)
    name = name.replace("&", " and ")
    name = re.sub(r"[.,'\-]", " ", name)
    name = re.sub(r"\s+", " ", name).strip()
    return name


def _best_fuzzy_vendor_match(key, lookup):
    """(standard_vendor_name, score 0-100) for the best fuzzy match of `key` against every
    normalized spelling in `lookup` (dict of normalized_key -> standard name), or
    (None, best_score) if nothing clears the bar. Requires BOTH a minimum similarity AND a
    clear margin over the runner-up (a DIFFERENT standard vendor), so two similarly-spelled but
    genuinely different vendors are never silently merged — an ambiguous case still falls
    through to a review flag."""
    if not key or not lookup:
        return None, 0
    scored = []
    for cand_key, std in lookup.items():
        if not cand_key:
            continue
        if _rapidfuzz:
            ratio = _rapidfuzz.token_sort_ratio(key, cand_key)
        else:
            ratio = difflib.SequenceMatcher(None, key, cand_key).ratio() * 100
        scored.append((ratio, std))
    if not scored:
        return None, 0
    scored.sort(key=lambda x: x[0], reverse=True)
    best_ratio, best_std = scored[0]
    if best_ratio < VENDOR_FUZZY_THRESHOLD:
        return None, best_ratio
    runner_up = next((r for r, s in scored[1:] if s != best_std), 0)
    if best_ratio - runner_up < VENDOR_FUZZY_MARGIN:
        return None, best_ratio          # too close to call — don't guess
    return best_std, best_ratio


def resolve_standard_vendor(raw_display, std_tables=None):
    """The standard vendor name for a raw vendor spelling, using the SAME exact-then-fuzzy
    lookup apply_vendor_standardization() uses (so 'these spellings are the same vendor' has a
    single source of truth — the vendor_name_map standard/alias/store spellings in VENDOR_STD,
    matched via _norm_vendor_std_key + _best_fuzzy_vendor_match). Returns the standard name, or
    None if nothing matches (or the tables aren't seeded). `std_tables` lets a caller pass an
    already-loaded load_vendor_standardization() dict instead of the module VENDOR_STD."""
    std = std_tables if std_tables is not None else VENDOR_STD
    if not std:
        return None
    name = (raw_display or "").strip()
    if not name:
        return None
    key = _norm_vendor_std_key(name)
    standard = std["lookup"].get(key)
    if not standard:
        standard, _score = _best_fuzzy_vendor_match(key, std["lookup"])
    return standard


def apply_vendor_standardization(rec):
    """Deterministic, post-classification (same slot as apply_class_rules). Matches rec's
    vendor to Joel's standardized vendor list — exact match first (after normalization), fuzzy
    match second — looks up its QB account, and resolves both to the store-specific spelling for
    QB entry. A vendor with no match at all (exact or fuzzy), OR one that matches but has no
    account assigned, is flagged for human review rather than guessed: the record goes to the review
    queue (its flag is set early by apply_vendor_standardization_gate, so it gets its own split file),
    except an excluded-vendor invoice/credit memo, which is recorded with review suppressed. A FUZZY
    (non-exact) match is accepted automatically but noted in the record's notes for audit. No-op if
    VENDOR_STD is empty (not seeded / Excel mode)."""
    if not VENDOR_STD:
        return
    raw_display = rec.get("vendor_matched") or rec.get("vendor_name_raw") or ""
    if not raw_display.strip():
        return
    key = _norm_vendor_std_key(raw_display)
    standard = VENDOR_STD["lookup"].get(key)
    if not standard:
        standard, score = _best_fuzzy_vendor_match(key, VENDOR_STD["lookup"])
        if standard:
            note = f"Vendor standardized via fuzzy match ({score:.0f}%): '{raw_display}' -> '{standard}'."
            _append_text(rec, "notes", note)
    if not standard:
        rec["needs_human_review"] = True
        note = (f"No standardized vendor mapping found for '{raw_display}'. Add it to "
                "the Vendor Names reference list (Setup), then reprocess.")
        _append_text(rec, "review_reason", note)
        return
    rec["standard_vendor"] = standard
    account = VENDOR_STD["vendor_account"].get(standard)
    if not account:
        rec["needs_human_review"] = True
        note = (f"'{standard}' has no QuickBooks account assigned yet. Add it to the "
                "Vendor Accounts reference list (Setup), then reprocess.")
        _append_text(rec, "review_reason", note)
        return
    rec["standard_account"] = account
    store_key = _STORE_KEY.get(_norm(rec.get("store_matched")))
    v_names = VENDOR_STD["vendor_names"].get(standard) or {}
    a_names = VENDOR_STD["account_names"].get(account) or {}
    rec["qb_vendor"] = (v_names.get(store_key) if store_key else None) or standard
    rec["qb_account"] = (a_names.get(store_key) if store_key else None) or account


def apply_vendor_standardization_gate(rec):
    """Classification-time pass of the vendor/account check, so its review flag is FINAL before the
    record's documents are saved. apply_vendor_standardization() itself runs only at record time
    (build_result_rows), AFTER assign_invoice_files had already filed the document under
    invoice_files/ as a clean record — so a record it flagged reached the queue with NO review_pending/
    split file: no preview, and a review row keyed to the WHOLE source email (answering one item
    reprocessed every record in that email). Flagging here instead means save_review_files gives the
    record its own split file like any other flagged record.

    Only decides the flag. A vendor that resolves cleanly is left to the record-time call, which fills
    standard/QB names as before. Excluded-vendor invoices/credit memos are also left to the record-time
    call: they're never queued, and flagging them this early would additionally skip their duplicate
    check — so their behavior stays exactly as it was. is_excluded_vendor() (a DB read) is consulted
    only when the check would flag. The record-time re-run can't double the sentence (_append_text)."""
    if not VENDOR_STD:
        return
    probe = {"vendor_matched": rec.get("vendor_matched"), "vendor_name_raw": rec.get("vendor_name_raw"),
             "store_matched": rec.get("store_matched")}
    apply_vendor_standardization(probe)
    if not probe.get("needs_human_review"):
        return
    if is_excluded_vendor(rec.get("vendor_matched")):
        return
    apply_vendor_standardization(rec)


# The two exact review-reason sentences apply_vendor_standardization() emits (see above).
# reapply_vendor_standardization_backfill() strips ONLY these when a mapping is later added in
# Setup, leaving every other reason (store/date/dup/...) intact. The quoted name is matched with
# [^']*, so a vendor name containing an apostrophe simply won't strip (safe: the row keeps its
# reason and stays in review) rather than over-matching.
_VENDOR_STD_REASON_PATTERNS = [
    re.compile(r"No standardized vendor mapping found for '[^']*'\. Add it to the Vendor Names "
               r"reference list \(Setup\), then reprocess\."),
    re.compile(r"'[^']*' has no QuickBooks account assigned yet\. Add it to the Vendor Accounts "
               r"reference list \(Setup\), then reprocess\."),
]


def _strip_vendor_std_reasons(reason):
    """Remove only the vendor/account sentences this module generates from a review_reason,
    collapsing the leftover whitespace. Returns the remaining reason (possibly '')."""
    text = reason or ""
    for pat in _VENDOR_STD_REASON_PATTERNS:
        text = pat.sub("", text)
    return re.sub(r"\s+", " ", text).strip()


def reapply_vendor_standardization_backfill():
    """Additive backfill: re-resolve vendor standardization on UNENTERED invoices/credit_memos
    whose mapping was incomplete, using the CURRENT Setup tables — so adding a Vendor Names alias
    or a Vendor -> account row updates the existing recorded rows without a full email reprocess.

    Scope (deliberately narrow): entered_in_qb = FALSE rows in invoices + credit_memos where
    standard_vendor / standard_account / qb_account is still NULL. Payments have no
    standardization columns and are skipped. Does NOT touch entered_in_qb = TRUE rows, store
    matching, class rules, dedup, or apply_vendor_standardization() itself.

    For each candidate it reloads VENDOR_STD fresh, re-runs apply_vendor_standardization() against
    the row's existing vendor_matched / store_matched, and — only if it NOW resolves (qb_account
    set) — writes standard_vendor / standard_account / qb_vendor / qb_account back, strips just the
    vendor/account review sentences from review_reason, and clears needs_review ONLY if no other
    reason remains. Each change is logged to edit_log for an audit trail.

    Returns (checked_count, resolved_count)."""
    global VENDOR_STD
    VENDOR_STD = load_vendor_standardization()
    if not VENDOR_STD:
        return (0, 0)

    db = _database()
    conn = db.connect()
    conn.autocommit = False
    checked = 0
    resolved = 0
    try:
        with conn.cursor() as cur:
            for table in ("invoices", "credit_memos"):
                cur.execute(
                    f"SELECT id, vendor_matched, store_matched, review_reason, needs_review, notes, "
                    f"       standard_vendor, standard_account, qb_vendor, qb_account "
                    f"  FROM {table} "
                    f" WHERE entered_in_qb IS NOT TRUE "
                    # Excluded (skip_qb) vendors never enter QB, so they don't need a standardized
                    # vendor/account and must not be flagged 'needs account' — leave them alone.
                    f"   AND skip_qb_import IS NOT TRUE "
                    f"   AND (standard_vendor IS NULL OR standard_account IS NULL OR qb_account IS NULL)"
                )
                cols = [d[0] for d in cur.description]
                rows = [dict(zip(cols, r)) for r in cur.fetchall()]
                for row in rows:
                    checked += 1
                    rec = {"vendor_matched": row.get("vendor_matched"),
                           "store_matched": row.get("store_matched")}
                    apply_vendor_standardization(rec)
                    # Persist as long as the VENDOR now resolves — even if it has no account yet.
                    # A vendor-resolved-but-account-less record must move from 'add vendor' to
                    # 'assign account' (standard_vendor set, qb_account still NULL); the old
                    # "only write when qb_account resolves" gate silently discarded that progress,
                    # so the record kept showing 'add vendor' and re-inviting the same click.
                    if not rec.get("standard_vendor"):
                        continue  # vendor still doesn't resolve (no mapping/alias yet) — leave as-is
                    # Rebuild the review reason: strip the two vendor/account sentences we generate,
                    # then re-add whichever one still applies. apply_vendor_standardization() leaves
                    # its "no account assigned yet" sentence on rec.review_reason when the vendor
                    # resolved but has no account; when fully resolved it leaves none.
                    base_reason = _strip_vendor_std_reasons(row.get("review_reason"))
                    apply_reason = rec.get("review_reason")
                    new_reason = " ".join(x for x in (base_reason, apply_reason) if x).strip() or None
                    new_needs = bool(new_reason)
                    # Preserve any fuzzy-match audit note apply() produced (rec['notes'] holds
                    # only the fresh note, since rec started empty) by appending to the row's.
                    fuzzy_note = rec.get("notes")
                    new_notes = row.get("notes")
                    if fuzzy_note:
                        new_notes = ((new_notes + " ") if new_notes else "") + fuzzy_note
                    updates = {
                        "standard_vendor": rec.get("standard_vendor"),
                        "standard_account": rec.get("standard_account"),
                        "qb_vendor": rec.get("qb_vendor"),
                        "qb_account": rec.get("qb_account"),
                        "review_reason": new_reason,
                        "needs_review": new_needs,
                        "notes": new_notes,
                    }
                    changed = any(str(row.get(f)) != str(updates[f]) for f in
                                  ("standard_vendor", "standard_account", "qb_vendor",
                                   "qb_account", "review_reason", "needs_review"))
                    if not changed:
                        continue  # already in this state (e.g. re-run) — nothing to write
                    # Count as "resolved" when the record materially progressed: a vendor newly
                    # standardized (add vendor -> assign account) OR an account newly assigned.
                    if ((not row.get("standard_vendor")) and rec.get("standard_vendor")) or \
                       ((not row.get("qb_account")) and rec.get("qb_account")):
                        resolved += 1
                    cur.execute(
                        f"UPDATE {table} SET standard_vendor=%(standard_vendor)s, "
                        f"standard_account=%(standard_account)s, qb_vendor=%(qb_vendor)s, "
                        f"qb_account=%(qb_account)s, review_reason=%(review_reason)s, "
                        f"needs_review=%(needs_review)s, notes=%(notes)s, edited_at=now() "
                        f"WHERE id=%(id)s",
                        {**updates, "id": row["id"]})
                    # Audit trail: log each field that actually changed.
                    for field in ("standard_vendor", "standard_account", "qb_vendor",
                                  "qb_account", "review_reason", "needs_review"):
                        old_v = row.get(field)
                        new_v = updates[field]
                        if str(old_v) != str(new_v):
                            cur.execute(
                                "INSERT INTO edit_log (table_name, record_id, field, old_value, "
                                "new_value, changed_by) VALUES (%s,%s,%s,%s,%s,%s)",
                                (table, row["id"], field,
                                 str(old_v) if old_v is not None else None,
                                 str(new_v) if new_v is not None else None,
                                 "auto-resolved by vendor/account backfill"))
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return (checked, resolved)


def reapply_excluded_vendor_backfill():
    """Additive backfill mirroring reapply_vendor_standardization_backfill(): when a vendor is newly
    added to excluded_vendors, flip skip_qb_import -> TRUE on the ALREADY-recorded, not-yet-entered
    invoices/credit_memos that now match the exclusion list. is_excluded_vendor() is read fresh each
    call and is correct for NEW records, but nothing re-evaluated existing rows — so an
    already-recorded invoice from a freshly-excluded vendor kept showing on /qb (needing an account
    it will never get). This closes that gap.

    Scope (deliberately narrow, only ever flips FALSE -> TRUE — never un-excludes): entered_in_qb IS
    NOT TRUE AND skip_qb_import IS NOT TRUE rows in invoices + credit_memos. Matching reuses
    is_excluded_vendor() directly (the SAME normalized-substring-either-direction logic + the same
    live excluded_vendors list) — ONE source of truth, no reimplementation. Payments are never
    excluded from QB, so they're untouched. Each flip is logged to edit_log.

    Returns the number of records newly skipped."""
    db = _database()
    conn = db.connect()
    conn.autocommit = False
    flipped = 0
    try:
        with conn.cursor() as cur:
            for table in ("invoices", "credit_memos"):
                cur.execute(
                    f"SELECT id, vendor_matched FROM {table} "
                    f" WHERE entered_in_qb IS NOT TRUE AND skip_qb_import IS NOT TRUE "
                    f"   AND vendor_matched IS NOT NULL AND vendor_matched <> ''")
                rows = cur.fetchall()
                for rec_id, vendor in rows:
                    if not is_excluded_vendor(vendor):
                        continue
                    cur.execute(
                        f"UPDATE {table} SET skip_qb_import = TRUE, edited_at = now() WHERE id = %s",
                        (rec_id,))
                    cur.execute(
                        "INSERT INTO edit_log (table_name, record_id, field, old_value, new_value, "
                        "changed_by) VALUES (%s,%s,%s,%s,%s,%s)",
                        (table, rec_id, "skip_qb_import", "False", "True",
                         "auto-skipped by excluded-vendor backfill"))
                    flipped += 1
        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
    return flipped


def _store_prompt_line(s):
    """One bullet for the prompt, rebuilt from a store's structured fields."""
    name = (s.get("store_name") or "").strip()
    aliases = (s.get("aliases") or "").strip()
    address = (s.get("address") or "").strip()
    notes = (s.get("notes") or "").strip()
    line = f"- {name}"
    if aliases:
        line += f" ({aliases})"
    if address:
        line += f" — {address}"
    if notes:
        if not line.rstrip().endswith((".", "!", "?")):
            line += "."
        line += f" {notes}"
    return line


def build_system_prompt(stores=None, excluded_vendors=None, card_store_map=None):
    """Splice the current stores + excluded-vendor list + card→store map into the
    prompt template."""
    stores = stores if stores is not None else load_stores()
    excluded_vendors = excluded_vendors if excluded_vendors is not None else load_excluded_vendors()
    card_store_map = card_store_map if card_store_map is not None else load_card_store_map()
    store_lines = "\n".join(_store_prompt_line(s) for s in stores) or "(none configured)"
    vendor_list = ", ".join(v for v in excluded_vendors if v) or "(none configured)"
    card_lines = "\n".join(f"- {last4} -> {store}"
                           for last4, store in sorted(card_store_map.items())
                           if last4 and store) or "(none configured)"
    return (SYSTEM_PROMPT_TEMPLATE
            .replace("%%STORE_LINES%%", store_lines)
            .replace("%%EXCLUDED_VENDORS%%", vendor_list)
            .replace("%%CARD_STORE_MAP%%", card_lines))


def get_system_prompt():
    """The run's system prompt, built once from the reference tables and cached."""
    global _SYSTEM_PROMPT_CACHE
    if _SYSTEM_PROMPT_CACHE is None:
        _SYSTEM_PROMPT_CACHE = build_system_prompt()
    return _SYSTEM_PROMPT_CACHE


def _extract_sender(text):
    """The sender's email address from an email's 'From:' header line, lowercased."""
    for line in (text or "").splitlines():
        if line.lower().startswith("from:"):
            return (parseaddr(line[5:])[1] or "").strip().lower()
    return ""


def _extract_recipients(text):
    """All To + Cc addresses from the header lines, lowercased."""
    from email.utils import getaddresses
    addrs = []
    for line in (text or "").splitlines():
        low = line.lower()
        if low.startswith("to:") or low.startswith("cc:"):
            addrs += [a.lower() for _, a in getaddresses([line.split(":", 1)[1]]) if a]
    return addrs


def _sender_store_hint(text):
    """A decisive store instruction when the email came from a known self-scan
    mailbox (sender_store_map). Empty string when the sender isn't mapped."""
    sender = _extract_sender(text)
    store = SENDER_STORE_MAP.get(sender) if sender else None
    if not store:
        return ""
    return (
        f"\n\n[STORE SIGNAL — this email was sent from {sender}, a known self-scan "
        f"mailbox for {store}. Treat the store as {store} (store_matched = \"{store}\") "
        f"with HIGH confidence; this is a self-scan. Only choose a different store if the "
        f"invoice document ITSELF clearly and explicitly states a different pharmacy, and "
        f"do NOT flag for review merely because the store would otherwise be uncertain.]"
    )


def _contact_store_hint(text):
    """Decisive store instruction when a To/Cc address is a known contact for a specific
    store (contact_store_map). Distinct from the sender-based hint: this covers an external
    biller/vendor emailing US where a recipient address (not the sender) is what identifies
    the store. Empty string when no recipient matches, or when matched recipients point to
    CONFLICTING stores (ambiguous — leave it to normal review in that case)."""
    stores = {CONTACT_STORE_MAP[a] for a in _extract_recipients(text) if a in CONTACT_STORE_MAP}
    if len(stores) != 1:
        return ""
    store = next(iter(stores))
    return (
        f"\n\n[STORE SIGNAL — this email is addressed to/cc'd to a known "
        f"contact for {store}. Treat the store as {store} (store_matched = "
        f"\"{store}\") with HIGH confidence, and do NOT flag for review merely "
        f"because the store would otherwise be uncertain. Only choose a "
        f"different store if the invoice document ITSELF clearly and "
        f"explicitly states a different pharmacy.]"
    )


def load_emails():
    if not EMAILS_DIR.exists():
        print(f"No {EMAILS_DIR} directory found.")
        return []
    files = sorted(
        p for p in EMAILS_DIR.iterdir()
        if p.is_file() and p.suffix.lower() in (".txt", ".eml", ".msg")
    )
    return files


def _strip_html(markup):
    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", markup)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = html_module.unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n\n", text)
    return text.strip()


def _iter_leaf_attachments(part):
    """Yield only LEAF (non-multipart) attachment parts, descending into any multipart
    container an email nests as an "attachment". Some senders wrap the real invoice like
    multipart/mixed -> [multipart/alternative body, multipart/mixed -> the PDF]; there
    iter_attachments() yields the inner multipart/mixed container itself, and calling
    get_content() on a multipart part raises KeyError(content_type) — which previously
    failed the WHOLE email and silently dropped the invoice inside. Recursing to the
    leaves reaches the actual document."""
    for sub in part.iter_attachments():
        if sub.is_multipart():
            yield from _iter_leaf_attachments(sub)
        else:
            yield sub


def _read_eml(path):
    with open(path, "rb") as f:
        msg = email.message_from_binary_file(f, policy=email_policy)

    body_part = msg.get_body(preferencelist=("plain", "html"))
    if body_part is not None:
        body = body_part.get_content()
        if body_part.get_content_type() == "text/html":
            body = _strip_html(body)
    else:
        body = ""

    headers = "\n".join(
        f"{h}: {msg.get(h)}" for h in ("From", "To", "Cc", "Subject", "Date") if msg.get(h)
    )
    text = f"{headers}\n\n{body}".strip()

    # Raw (unstripped) HTML body for high-fidelity fallback rendering — kept
    # separate from `text` (which stays plaintext for classification). None
    # when the email has no HTML part.
    html_part = msg.get_body(preferencelist=("html",))
    html_body = html_part.get_content() if html_part is not None else None

    attachments = []
    for part in _iter_leaf_attachments(msg):
        filename = part.get_filename() or "attachment"
        mimetype = part.get_content_type()
        # Guard each attachment so one unreadable part can't fail the whole email (and
        # silently drop a real invoice sitting in another part): skip it with a note and
        # keep going, same as the email-body/other-attachment fallback everywhere else.
        try:
            data = part.get_content()
        except Exception as e:
            print(f"  skipping unreadable attachment '{filename}' ({mimetype}): {e}")
            continue
        if isinstance(data, str):
            data = data.encode("utf-8", errors="replace")
        cid = part.get("Content-ID")
        if cid:
            cid = cid.strip().strip("<>")
        attachments.append({"filename": filename, "mimetype": mimetype, "data": data, "content_id": cid})

    return text, attachments, html_body


def _read_msg(path):
    m = extract_msg.Message(str(path))
    try:
        cc = getattr(m, "cc", None)          # extract_msg exposes .cc; None if unavailable
        headers = "\n".join(filter(None, [
            f"From: {m.sender}" if m.sender else None,
            f"To: {m.to}" if m.to else None,
            f"Cc: {cc}" if cc else None,
            f"Subject: {m.subject}" if m.subject else None,
            f"Date: {m.date}" if m.date else None,
        ]))
        body = m.body
        # Capture the raw HTML body unconditionally (for high-fidelity fallback
        # rendering), separately from the plaintext `text` used for classification.
        html_raw = m.htmlBody
        if isinstance(html_raw, bytes):
            html_raw = html_raw.decode("utf-8", errors="replace")
        if not body and html_raw:
            body = _strip_html(html_raw)
        text = f"{headers}\n\n{body or ''}".strip()
        html_body = html_raw or None

        attachments = []
        for att in m.attachments:
            filename = att.getFilename() or "attachment"
            mimetype = att.mimetype or mimetypes.guess_type(filename)[0]
            if isinstance(att.data, bytes):
                cid = getattr(att, "cid", None) or getattr(att, "contentId", None)
                if cid:
                    cid = str(cid).strip().strip("<>")
                attachments.append({"filename": filename, "mimetype": mimetype, "data": att.data, "content_id": cid})

        return text, attachments, html_body
    finally:
        m.close()


def read_email_content(path):
    suffix = path.suffix.lower()
    if suffix == ".msg":
        return _read_msg(path)
    if suffix == ".eml":
        return _read_eml(path)
    if suffix in EXT_TO_MEDIA_TYPE and suffix != ".txt" and suffix != ".csv":
        # A standalone document (a split per-record PDF/image pulled from a
        # batch for review) — treat the file itself as the sole attachment.
        return "", [{
            "filename": path.name,
            "mimetype": EXT_TO_MEDIA_TYPE[suffix],
            "data": path.read_bytes(),
        }], None
    return path.read_text(encoding="utf-8", errors="replace"), [], None


def _read_headers_only(path):
    """Just the From/To/Cc/Subject/Date header block of an original .eml/.msg source (no
    body, no attachments) — used to recover header context for a retry/reprocess whose own
    source is a split PDF/PNG with no headers of its own. Best-effort: returns "" on any
    failure (missing file, unsupported type, parse error)."""
    try:
        suffix = path.suffix.lower()
        if suffix == ".eml":
            text, _, _ = _read_eml(path)
        elif suffix == ".msg":
            text, _, _ = _read_msg(path)
        else:
            return ""
        # keep only the header block (everything before the blank-line body separator)
        return text.split("\n\n", 1)[0]
    except Exception:
        return ""


def render_text_to_pdf(text, out_path):
    """Render plain text to a PDF via LibreOffice headless conversion — used
    as a stand-in invoice document when an email has no usable PDF/image
    attachment, so there's still something to open/print for QB entry."""
    if not SOFFICE_PATH:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        tmp_txt = Path(tmp) / "email.txt"
        tmp_txt.write_text(text, encoding="utf-8")
        try:
            subprocess.run(
                [SOFFICE_PATH, "--headless", "--convert-to", "pdf", "--outdir", tmp, str(tmp_txt)],
                check=True, timeout=60, capture_output=True,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired):
            return False
        produced = Path(tmp) / "email.pdf"
        if not produced.exists():
            return False
        shutil.copy2(produced, out_path)
    return True


# --- HTML→PDF fallback rendering via headless system Edge (Playwright) --------
# A lazily-launched, process-wide singleton Edge browser reused for every record
# in a batch run (relaunching per record would be far slower) and closed at
# process exit. This is only ever driven from the CLI batch path
# (classify_emails.py / run_all.py); the always-on FastAPI review UI
# (web_db.py / review_app.py) does not call these functions, so a plain atexit
# close is safe. If Playwright or Edge isn't available, _get_edge_browser()
# returns None and callers fall back to the plaintext render_text_to_pdf.
_PLAYWRIGHT = None
_EDGE_BROWSER = None
_EDGE_UNAVAILABLE = False

_CID_REF_RE = re.compile(r"cid:([^\"')\s>]+)", re.IGNORECASE)


def _close_edge_browser():
    global _EDGE_BROWSER, _PLAYWRIGHT
    try:
        if _EDGE_BROWSER is not None:
            _EDGE_BROWSER.close()
    except Exception:
        pass
    try:
        if _PLAYWRIGHT is not None:
            _PLAYWRIGHT.stop()
    except Exception:
        pass
    _EDGE_BROWSER = None
    _PLAYWRIGHT = None


def _get_edge_browser():
    """Return the shared headless system-Edge browser, launching it once on
    first use. Returns None (and remembers the failure) if Playwright isn't
    installed or Edge can't be launched, so callers fall back to text rendering."""
    global _PLAYWRIGHT, _EDGE_BROWSER, _EDGE_UNAVAILABLE
    if _EDGE_UNAVAILABLE:
        return None
    if _EDGE_BROWSER is not None:
        return _EDGE_BROWSER
    try:
        from playwright.sync_api import sync_playwright
        _PLAYWRIGHT = sync_playwright().start()
        _EDGE_BROWSER = _PLAYWRIGHT.chromium.launch(channel="msedge", headless=True)
        atexit.register(_close_edge_browser)
        return _EDGE_BROWSER
    except Exception as e:
        print(f"  HTML-to-PDF rendering unavailable ({e}); using plaintext fallback")
        _EDGE_UNAVAILABLE = True
        _close_edge_browser()
        return None


def _inline_cid_images(html, cid_images):
    """Replace cid:XXXX references in the HTML with data: URIs built from the
    email's inline-image attachments, so embedded logos/icons render even
    though there's no network fetch for them. Remote http(s) images are left
    as-is for the browser to fetch live."""
    if not cid_images:
        return html

    def repl(m):
        key = m.group(1).strip().strip("<>")
        entry = cid_images.get(key)
        if entry is None:
            return m.group(0)
        data, mimetype = entry
        b64 = base64.b64encode(data).decode("ascii")
        return f"data:{mimetype or 'application/octet-stream'};base64,{b64}"

    return _CID_REF_RE.sub(repl, html)


def render_html_to_pdf(html, cid_images, out_path):
    """Render an email's real HTML body to PDF via headless system Edge
    (Playwright), so the fallback document looks like the actual email
    instead of a plaintext dump. Returns False on any failure (missing
    playwright/Edge, launch error, timeout) — caller falls back to
    render_text_to_pdf, never raises."""
    if not html:
        return False
    try:
        browser = _get_edge_browser()
        if browser is None:
            return False
        html = _inline_cid_images(html, cid_images or {})
        with tempfile.TemporaryDirectory() as tmp:
            tmp_html = Path(tmp) / "email.html"
            tmp_html.write_text(html, encoding="utf-8")
            page = browser.new_page()
            try:
                # file:// so relative/data URIs resolve; remote http(s) images
                # are fetched live over the network (via the OS trust store).
                try:
                    page.goto(tmp_html.resolve().as_uri(), wait_until="load", timeout=30000)
                except Exception as nav_e:
                    # Forwarded / marketing HTML (e.g. "your order is confirmed!") often embeds
                    # many remote images + tracking pixels; wait_until="load" blocks on ALL of
                    # them and can hit the 30s timeout intermittently (slow/blocked image hosts,
                    # batch-run contention) — which used to fail the whole render and leave the
                    # record with NO document. Instead, render whatever DID load: a partial page
                    # is far more useful to a reviewer than nothing.
                    print(f"  HTML render: page load did not fully settle ({nav_e}); "
                          f"rendering what loaded")
                page.pdf(path=str(out_path), print_background=True)
            finally:
                page.close()
        return Path(out_path).exists() and Path(out_path).stat().st_size > 0
    except Exception as e:
        print(f"  HTML-to-PDF render failed ({e}); using plaintext fallback")
        return False


# Vendor invoice-portal domains whose public (no-login) view shows enough to record the
# invoice — a link-only email (no attachment) linking here has its portal PAGE fetched and
# rendered to PDF for classification, instead of only the email body. A LIST on purpose:
# add a domain ONLY after verifying that vendor's portal is publicly viewable the same way
# (implementations vary; this fetches a real external site). Verified: BillerGenie.
LINK_INVOICE_DOMAINS = ["customer.billergenie.com"]


def _find_invoice_link(html):
    """First href in the HTML whose host is (or is a subdomain of) a LINK_INVOICE_DOMAINS
    entry, else None. Used to detect a link-only invoice email from a verified portal."""
    if not html:
        return None
    from urllib.parse import urlparse
    for m in re.finditer(r'href=["\']([^"\']+)["\']', html, re.IGNORECASE):
        raw = html_module.unescape(m.group(1))
        try:
            host = (urlparse(raw).hostname or "").lower()
        except ValueError:
            continue
        if host and any(host == d or host.endswith("." + d) for d in LINK_INVOICE_DOMAINS):
            return raw
    return None


def _has_usable_document(attachments):
    """True if any attachment is a real document to classify from — i.e. NOT just an inline
    image embedded in the HTML (a logo/banner referenced by cid, like BillerGenie's logo.png).
    Those decorative inline images aren't the invoice, so an email carrying only them is treated
    as attachment-less for the link-fetch decision."""
    for a in attachments:
        inline_image = bool(a.get("content_id")) and _effective_media_type(a) in SUPPORTED_IMAGE_TYPES
        if not inline_image:
            return True
    return False


def render_url_to_pdf(url, out_path):
    """Fetch an external invoice-portal URL in headless system Edge (Playwright) and save the
    real invoice PDF, so a link-only invoice email is classified from the actual invoice — not
    the email body.

    Two-step, best-effort:
      1. The landed portal page is often just a thin balance/"Pay By Card" widget (vendor
         address and line items live one click deeper). If it exposes a "View Invoice" /
         ...ExportToPdf link (same already-trusted domain), fetch THAT directly — it returns a
         genuine PDF document, which we save byte-for-byte (real layout + line items, far better
         extraction than a screenshot of the summary).
      2. If no such link is present (older/other BillerGenie tenants, or a portal that renders
         the invoice inline), fall back to today's behavior: render the loaded page to PDF.

    Returns False on any failure (no playwright/Edge, timeout, load error) — the caller falls
    back to the email body, so this is never a hard dependency on the site being up."""
    try:
        browser = _get_edge_browser()
        if browser is None:
            return False
        page = browser.new_page(viewport={"width": 1000, "height": 1200})
        try:
            page.goto(url, wait_until="load", timeout=30000)
            # These portals are client-side apps that fetch the invoice data AFTER load —
            # wait for the network to settle so the numbers are on the page (best-effort;
            # proceed anyway if it doesn't idle in time).
            try:
                page.wait_for_load_state("networkidle", timeout=15000)
            except Exception:
                pass
            # Bail to the email-body fallback if the page rendered essentially blank (site
            # down / structure changed / needs login with nothing public) — don't save an
            # empty document as if it were the invoice.
            try:
                if len((page.inner_text("body") or "").strip()) < 40:
                    return False
            except Exception:
                pass
            # Step 1: follow the "View Invoice" / ExportToPdf link to the real PDF, if present.
            # Prefer the href match (the actual PDF endpoint) over the visible text, and reuse
            # this page's request context so it carries the same session/cookies as the portal.
            try:
                link = (page.query_selector('a[href*="ExportToPdf" i]')
                        or page.query_selector('a:has-text("View Invoice")'))
                if link:
                    href = link.get_attribute("href")
                    if href:
                        from urllib.parse import urljoin
                        pdf_url = urljoin(page.url, href)
                        resp = page.request.get(pdf_url, timeout=30000)
                        ctype = (resp.headers.get("content-type") or "").lower()
                        if resp.ok and "application/pdf" in ctype:
                            body = resp.body()
                            if body:
                                Path(out_path).write_bytes(body)
                                return Path(out_path).stat().st_size > 0
                        # Not a PDF (or fetch not OK) — fall through to page-render below.
            except Exception:
                pass  # any trouble following the link -> fall back to rendering the page
            # Step 2 (fallback): render the loaded page itself. The portal hides its content
            # under print CSS (page.pdf() defaults to print media) — emulate screen so the
            # visible invoice actually renders.
            page.emulate_media(media="screen")
            page.pdf(path=str(out_path), print_background=True, prefer_css_page_size=True)
        finally:
            page.close()
        return Path(out_path).exists() and Path(out_path).stat().st_size > 0
    except Exception as e:
        print(f"  invoice-link fetch failed ({e}); using email body")
        return False


def _unique_path(directory, stem, suffix):
    candidate = directory / f"{stem}{suffix}"
    if not candidate.exists():
        return candidate
    i = 2
    while True:
        candidate = directory / f"{stem}__{i}{suffix}"
        if not candidate.exists():
            return candidate
        i += 1


def _effective_media_type(att):
    """The attachment's real media type. Trusts the declared MIME type when
    it's one we recognize; otherwise (generic application/octet-stream, or
    missing) falls back to the file extension."""
    mt = (att.get("mimetype") or "").lower()
    if (mt in SUPPORTED_DOCUMENT_TYPES or mt in SUPPORTED_IMAGE_TYPES
            or mt in INLINABLE_TEXT_TYPES or mt in SUPPORTED_SPREADSHEET_TYPES):
        return mt
    ext = Path(att.get("filename") or "").suffix.lower()
    return EXT_TO_MEDIA_TYPE.get(ext, mt)


# --- Zip auto-expansion --------------------------------------------------------
# Emails whose invoices arrive zipped are expanded in memory so the contained
# PDFs/images are classified automatically. A zip that can't be safely expanded
# (locked/corrupt/oversized/nested/no supported files) is LEFT AS-IS — exactly
# today's behavior: the record still flags and the manual-upload path handles it.
# Caps use ZipInfo.file_size (UNCOMPRESSED) so a zip bomb is rejected BEFORE inflating.
ZIP_MAX_ENTRIES = 100
ZIP_MAX_INNER_BYTES = 25 * 1024 * 1024        # 25 MB, single inner file
ZIP_MAX_TOTAL_BYTES = 100 * 1024 * 1024       # 100 MB, total extracted


def _is_zip_attachment(att):
    mt = (att.get("mimetype") or "").lower()
    if mt in ("application/zip", "application/x-zip-compressed", "multipart/x-zip"):
        return True
    return Path(att.get("filename") or "").suffix.lower() == ".zip"


def _expand_one_zip(data, stem):
    """(files, skip_reason). skip_reason is a short string when the WHOLE zip should
    be left as-is (printed by the caller); None on a handled zip (files may be empty =
    no supported members found). Nested zips are skipped (no recursion / zip-bomb)."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data or b""))
    except Exception:
        return [], "corrupt or not a zip"
    try:
        infos = [zi for zi in zf.infolist() if not zi.is_dir()]
        if len(infos) > ZIP_MAX_ENTRIES:
            return [], f"too many entries ({len(infos)})"
        total = 0
        for zi in infos:                          # size guards BEFORE any extraction
            if zi.file_size > ZIP_MAX_INNER_BYTES:
                return [], f"an inner file exceeds {ZIP_MAX_INNER_BYTES // (1024 * 1024)}MB"
            total += zi.file_size
        if total > ZIP_MAX_TOTAL_BYTES:
            return [], f"extracted size exceeds {ZIP_MAX_TOTAL_BYTES // (1024 * 1024)}MB"

        files = []
        for zi in infos:
            inner = (zi.filename or "").rsplit("/", 1)[-1]
            if Path(inner).suffix.lower() == ".zip":
                continue                          # do NOT recurse into nested zips
            media = _effective_media_type({"filename": inner, "mimetype": None})
            if (media not in SUPPORTED_DOCUMENT_TYPES and media not in SUPPORTED_IMAGE_TYPES
                    and media not in INLINABLE_TEXT_TYPES
                    and media not in SUPPORTED_SPREADSHEET_TYPES):
                continue                          # unsupported member -> drop
            try:
                b = zf.read(zi)                   # RuntimeError if password-protected
            except (RuntimeError, zipfile.BadZipFile):
                return [], "password-protected or unreadable"
            # Keep a "<zip> :: <inner>" label so source_document citation stays traceable.
            files.append({"filename": f"{stem} :: {inner}", "mimetype": media, "data": b})
        return files, None
    finally:
        zf.close()


def expand_zip_attachments(attachments):
    """Replace each ZIP attachment with its contained supported files (PDF/image/text),
    so ALL downstream use (the model AND source-document saving) sees the real documents
    instead of an unreadable zip. Non-zip attachments pass through unchanged; overall
    order is preserved. A zip that can't be safely expanded is left in place (today's
    behavior). Same {filename, mimetype, data} dict shape, so nothing downstream changes."""
    out = []
    for att in attachments:
        if not _is_zip_attachment(att):
            out.append(att)
            continue
        name = att.get("filename") or "archive.zip"
        stem = Path(name).stem or "archive"
        files, skip = _expand_one_zip(att.get("data"), stem)
        if skip:
            print(f"  skipped zip {name} ({skip})")
            out.append(att)                       # leave as-is -> record still flags
        elif not files:
            print(f"  skipped zip {name} (no readable files inside)")
            out.append(att)
        else:
            print(f"  unzipped {name} -> {len(files)} file(s)")
            out.extend(files)
    return out


def _sanitize_filename(name):
    name = re.sub(r'[\\/:*?"<>|]', "", name)
    name = re.sub(r"\s+", " ", name).strip().strip(".")
    return name[:180] or "invoice"


def _short_store(store):
    return re.sub(r"\s*pharmacy\s*$", "", (store or "").strip(), flags=re.I).strip() or "Unknown"


def _fmt_amount(amount):
    try:
        return "$" + format(float(amount), ",.2f")
    except (TypeError, ValueError):
        return ""


def _fmt_date_dots(d):
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", str(d or ""))
    if m:
        y, mo, dd = m.groups()
        return f"{int(mo)}.{int(dd)}.{y}"
    return str(d or "")


def build_invoice_filename(store, vendor, number, amount, date_iso):
    """'Store - Vendor - Invoice # - Amount - Date', e.g.
    '3M - Prodermix - 9035 - $1,603.62 - 7.22.2026'."""
    parts = [
        _short_store(store), (vendor or "").strip(), str(number or "").strip(),
        _fmt_amount(amount), _fmt_date_dots(date_iso),
    ]
    return _sanitize_filename(" - ".join(p for p in parts if p))


def _match_attachment(source_document, attachments):
    """Find the attachment a record was read from, given the model's
    source_document value. None means 'email body' (or no match)."""
    sd = (source_document or "").strip().lower()
    if not sd or sd == "email_body":
        return None
    for att in attachments:
        if (att.get("filename") or "").strip().lower() == sd:
            return att
    for att in attachments:
        fn = (att.get("filename") or "").strip().lower()
        if fn and (fn in sd or sd in fn):
            return att
    return None


def _extract_xlsx_text(data, max_rows_per_sheet=300, max_sheets=10):
    """Plain-text rendering of an .xlsx/.xlsm workbook's cell values — sheet by
    sheet, row by row, tab-separated — so a spreadsheet attachment can be inlined
    into the prompt the same way a CSV already is. Formulas are read as their LAST
    CALCULATED VALUE (data_only=True), not the formula text. Truncates a very
    large sheet/workbook rather than blowing up the prompt, and says so, so the
    model knows the data may be incomplete rather than assuming it saw everything."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as e:
        return f"[could not open as a spreadsheet: {e}]"
    parts = []
    try:
        for si, ws in enumerate(wb.worksheets):
            if si >= max_sheets:
                parts.append(f"... ({len(wb.worksheets) - max_sheets} more sheet(s) not shown)")
                break
            parts.append(f"--- Sheet: {ws.title} ---")
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                if all(v is None for v in row):
                    continue
                parts.append("\t".join("" if v is None else str(v) for v in row))
                row_count += 1
                if row_count >= max_rows_per_sheet:
                    parts.append(f"... (more rows not shown, sheet has more than {max_rows_per_sheet})")
                    break
    finally:
        wb.close()
    return "\n".join(parts) if parts else "[empty spreadsheet]"


def build_attachment_blocks(attachments):
    """Turn extracted attachments into Claude message content blocks, each
    preceded by a '[Attachment N: filename]' label so the model can cite which
    file a record came from (source_document). PDFs/images go natively; plain
    text/CSV and .xlsx/.xlsm spreadsheets are inlined as text; anything else is
    flagged as unprocessed."""
    blocks = []
    skipped = []
    for i, att in enumerate(attachments, 1):
        media_type = _effective_media_type(att)
        label = f"[Attachment {i}: {att['filename']}]"
        if media_type in SUPPORTED_DOCUMENT_TYPES:
            blocks.append({"type": "text", "text": label})
            blocks.append({
                "type": "document",
                "source": {"type": "base64", "media_type": media_type,
                           "data": base64.b64encode(att["data"]).decode("ascii")},
            })
        elif media_type in SUPPORTED_IMAGE_TYPES:
            # Downscale over-large images so the API doesn't reject the request for
            # dimensions exceeding its 8000-px limit (a permanent 400, not transient).
            img_data = _downscale_image_for_api(att["data"], media_type)
            blocks.append({"type": "text", "text": label})
            blocks.append({
                "type": "image",
                "source": {"type": "base64", "media_type": media_type,
                           "data": base64.b64encode(img_data).decode("ascii")},
            })
        elif media_type in INLINABLE_TEXT_TYPES:
            text = att["data"].decode("utf-8", errors="replace")
            blocks.append({"type": "text", "text": f"{label}\n{text}"})
        elif media_type in SUPPORTED_SPREADSHEET_TYPES:
            text = _extract_xlsx_text(att["data"])
            blocks.append({"type": "text", "text": f"{label}\n{text}"})
        else:
            skipped.append(att["filename"])
    return blocks, skipped


# Review-reason substrings that mark a queued item as "couldn't read the attachment/
# zip" — the target of the one-off --retry-unreadable pass (case-insensitive).
_UNREADABLE_REASON_PATTERNS = (
    "zip", "could not be read", "could not be opened", "unable to read", "unable to open",
    "couldn't read", "couldn't open", "unreadable", "reading failed", "automatic reading failed",
)


def _is_unreadable_review_item(row):
    """True if a review_queue row (load_review_queue dict) was flagged for an unreadable
    attachment/zip — either the whole-email 'unreadable' type or a reason matching the
    patterns above (or an attachment+read phrasing)."""
    if (row.get("Email Type") or "").strip().lower() == "unreadable":
        return True
    q = (row.get("Question (AI review reason)") or "").lower()
    if any(p in q for p in _UNREADABLE_REASON_PATTERNS):
        return True
    return "attachment" in q and "read" in q


def record_flagged(record):
    """True if this individual record (invoice/credit memo/payment) needs human
    review. Review is per-record: a clean invoice in a batch is processed even
    if another invoice scanned with it is flagged."""
    return bool(record.get("needs_human_review"))


def routes_to_review(rtype, rec):
    """Whether a flagged record actually goes to the HUMAN review queue. A record
    the model flagged normally does — EXCEPT an excluded-vendor invoice/credit memo:
    it's never entered in QuickBooks, so review would be wasted. Those are still
    recorded (they show in Browse/History) but are not queued. Vendor exclusion is
    NEVER applied to payments — a payment is always entered in QB regardless of
    vendor, so is_excluded_vendor() is not consulted for rtype 'payment'."""
    if rtype == "invoice" and rec.get("is_sales_order"):
        return False                          # sales orders are HELD, never reviewed
    if not record_flagged(rec):
        return False
    if rtype in ("invoice", "credit_memo") and is_excluded_vendor(rec.get("vendor_matched")):
        return False
    return True


def review_suppressed_by_exclusion(rtype, rec):
    """True only for a record the model flagged but that we hold back from review
    because its vendor is QB-excluded (invoice/credit memo only). Used to annotate
    the recorded row so Browse shows why it's here yet not in the queue."""
    return record_flagged(rec) and not routes_to_review(rtype, rec)


_EXCLUDED_SKIP_NOTE = ("[review skipped: excluded vendor — not entered in QuickBooks, "
                       "so no human review needed]")


def _record_needs_file(record, record_type):
    """Records that get a saved source document: non-skip invoices and credit
    memos (they land in qb_ready.xlsx), plus every payment (kept as proof of
    payment even though payments aren't in qb_ready). No file for skip-QB
    invoices/credit memos, informational emails, or records flagged for review
    (those wait in the review queue until resolved)."""
    if record_flagged(record):
        return False
    if record_type == "payment":
        return True
    return not record.get("skip_qb_import")


def _parse_pages(spec):
    """Parse a 1-based page spec like '1', '2-3', '1,4' into a sorted unique
    list of ints. Returns [] for null/blank/unparseable."""
    if not spec:
        return []
    pages = set()
    for part in str(spec).replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            try:
                a, b = part.split("-", 1)
                for p in range(int(a), int(b) + 1):
                    pages.add(p)
            except ValueError:
                return []
        else:
            try:
                pages.add(int(part))
            except ValueError:
                return []
    return sorted(pages)


def _offset_page_spec(spec, offset):
    """Re-express a 1-based page spec ('1', '2-3', '1,4') from CHUNK-LOCAL numbering
    into the ORIGINAL full PDF's numbering by adding `offset` to every page. Returns
    a comma-joined spec of the shifted pages, or None if the spec is null/blank/
    unparseable (the merge caller then bounds the record to the chunk's own range)."""
    pages = _parse_pages(spec)
    if not pages:
        return None
    return ",".join(str(p + offset) for p in pages)


def _offset_rotations(page_rotations, offset):
    """Shift page_rotations keyed by 1-based CHUNK-LOCAL page into the ORIGINAL PDF's
    numbering (dict form). A positional list is returned unchanged — it aligns to the
    record's own pages in order, which merging preserves — as is null/anything else."""
    if isinstance(page_rotations, dict):
        shifted = {}
        for k, v in page_rotations.items():
            try:
                shifted[str(int(str(k).strip()) + offset)] = v
            except (TypeError, ValueError):
                continue
        return shifted
    return page_rotations


def _parse_page_list(value):
    """A list of 1-based page numbers from the model (blank_pages / page_order). Accepts a JSON
    list of ints (or int-like strings), preserving the GIVEN order (page_order relies on order,
    so this never sorts). Non-integers are dropped. A comma/range string is accepted as a
    courtesy (via _parse_pages, which sorts — fine for the order-agnostic blank_pages). Returns
    [] for null/blank/unparseable."""
    if value is None:
        return []
    if isinstance(value, (list, tuple)):
        out = []
        for v in value:
            try:
                out.append(int(str(v).strip()))
            except (TypeError, ValueError):
                continue
        return out
    return _parse_pages(value)


def _offset_page_list(value, offset):
    """Shift a list of 1-based CHUNK-LOCAL page numbers (blank_pages / page_order) into the
    ORIGINAL PDF's numbering, preserving order. Returns a list of ints, or None when empty (so
    the merge caller leaves the field unset). Mirrors _offset_page_spec/_offset_rotations for
    the two new page fields, so chunked (oversized) scans get correct global page numbers."""
    pages = _parse_page_list(value)
    if not pages:
        return None
    return [p + offset for p in pages]


def _apply_page_order(selected, page_order):
    """Reading order for `selected` (this record's pages, ascending): the model's page_order iff
    it's an EXACT permutation of `selected` (same set, same length); otherwise ascending
    `selected` unchanged. Same safe-fallback contract as source_pages/page_rotations — a wrong or
    absent value never reorders on a guess, worst case is today's physical order."""
    po = _parse_page_list(page_order)
    if po and len(po) == len(selected) and set(po) == set(selected):
        return po
    return list(selected)


def _parse_rotations(page_rotations, selected_pages):
    """Normalize the model's page_rotations to {1-based source page: degrees}, with
    degrees limited to {90, 180, 270} (0 / unknown / invalid are dropped = no turn).
    Accepts a dict keyed by 1-based source page number (preferred), or a list aligned
    positionally to `selected_pages` (this record's pages, in order)."""
    def _deg(v):
        try:
            d = int(v) % 360
        except (TypeError, ValueError):
            return 0
        return d if d in (90, 180, 270) else 0

    out = {}
    if not page_rotations:
        return out
    if isinstance(page_rotations, dict):
        for k, v in page_rotations.items():
            try:
                pg = int(str(k).strip())
            except (TypeError, ValueError):
                continue
            if _deg(v):
                out[pg] = _deg(v)
    elif isinstance(page_rotations, (list, tuple)):
        for i, v in enumerate(page_rotations):
            if i < len(selected_pages) and _deg(v):
                out[selected_pages[i]] = _deg(v)
    return out


def _extract_pdf_pages(pdf_bytes, page_spec, page_rotations=None,
                       blank_pages=None, page_order=None):
    """Return a new PDF (bytes) for this record: only its pages (when a multi-record PDF is split
    by page_spec), in correct reading order (page_order, when the model gave a valid permutation),
    with entirely-blank pages dropped (blank_pages) and each remaining page rotated UPRIGHT per
    page_rotations. Returns None to mean 'keep the original file as-is' — nothing to split,
    reorder, drop, or rotate — or on any error, so the caller safely falls back to the unmodified
    source. Every one of the four page fields degrades safely: a wrong/absent value never loses
    data — worst case the page stays exactly as it is today."""
    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        total = len(reader.pages)
        if total < 1:
            return None
        pages = _parse_pages(page_spec)
        # A genuine subset (valid, multi-page, not the whole doc) -> split; else the
        # whole document (still reordered/rotated/blank-dropped below if needed).
        if pages and total > 1 and all(1 <= p <= total for p in pages) and len(pages) != total:
            selected = pages
        else:
            selected = list(range(1, total + 1))

        # Reading order: model's page_order iff an exact permutation of selected, else ascending.
        ordered = _apply_page_order(selected, page_order)
        # Drop entirely-blank pages (only within this record's pages). Never drop EVERY page —
        # an all-blank claim is clearly wrong, so keep them (safe fallback).
        blanks = set(_parse_page_list(blank_pages)) & set(selected)
        if blanks and len(blanks) < len(ordered):
            ordered = [p for p in ordered if p not in blanks]

        rotmap = _parse_rotations(page_rotations, ordered)

        # Rebuild whenever the final page sequence differs from the untouched original (fewer
        # pages, or reordered), or any page needs rotating.
        need_change = ordered != list(range(1, total + 1))
        need_rotate = any(rotmap.get(p) for p in ordered)
        if not need_change and not need_rotate:
            return None  # nothing to do — keep the original bytes

        writer = PdfWriter()
        for p in ordered:
            writer.add_page(reader.pages[p - 1])
            deg = rotmap.get(p, 0)
            if deg:
                writer.pages[-1].rotate(deg)  # pypdf: clockwise, multiple of 90
        out = io.BytesIO()
        writer.write(out)
        return out.getvalue()
    except Exception:
        return None


_PIL_FORMAT = {"image/png": "PNG", "image/jpeg": "JPEG", "image/gif": "GIF", "image/webp": "WEBP"}


def _rotate_image_bytes(img_bytes, media_type, page_rotations):
    """Return image bytes rotated UPRIGHT for review, or None to keep the original
    unchanged (nothing to do, Pillow missing, or any error — a safe, display-only
    fallback that never fails the save). Normalizes phone-photo EXIF orientation
    first (exif_transpose), then applies the model's clockwise rotation on top. A
    single-image attachment is treated as page 1 (source_pages is null for images)."""
    try:
        from PIL import Image, ImageOps
    except Exception:
        return None  # Pillow not installed — save the image as-is

    rotmap = _parse_rotations(page_rotations, [1])
    deg = rotmap.get(1, 0)
    if not deg and len(rotmap) == 1:      # model keyed the single image differently
        deg = next(iter(rotmap.values()))

    try:
        im = Image.open(io.BytesIO(img_bytes))
        orient = im.getexif().get(0x0112)          # 0x0112 = EXIF Orientation
        has_exif_orient = orient not in (None, 1)
        if not deg and not has_exif_orient:
            return None                             # nothing to normalize or rotate

        im = ImageOps.exif_transpose(im)            # bake in phone EXIF orientation
        if deg:
            im = im.rotate(-deg, expand=True)       # PIL rotates CCW; negate for clockwise

        fmt = _PIL_FORMAT.get(media_type, "PNG")
        save_kwargs = {}
        if fmt == "JPEG":
            im = im.convert("RGB")                  # JPEG has no alpha
            save_kwargs["quality"] = 95
        out = io.BytesIO()
        im.save(out, format=fmt, **save_kwargs)
        return out.getvalue()
    except Exception:
        return None


# The Anthropic API rejects an image whose longest edge exceeds 8000 px (400
# invalid_request_error) and internally downscales anything past ~1568 px anyway,
# so we cap the long edge at 1568: no legibility loss vs. what the API would use,
# no wasted tokens, and safely clear of the hard limit.
MAX_IMAGE_EDGE = 1568


def _downscale_image_for_api(img_bytes, media_type):
    """If an image's longest edge exceeds MAX_IMAGE_EDGE, return downscaled bytes
    (aspect ratio preserved) so the request isn't rejected for oversized dimensions.
    Returns the ORIGINAL bytes unchanged when no resize is needed, Pillow is missing,
    or on any error (safe fallback — same non-failing contract as the rotate helper)."""
    try:
        from PIL import Image, ImageOps
    except Exception:
        return img_bytes  # Pillow not installed — send as-is
    try:
        im = Image.open(io.BytesIO(img_bytes))
        im = ImageOps.exif_transpose(im)       # measure the true (display) dimensions
        w, h = im.size
        if max(w, h) <= MAX_IMAGE_EDGE:
            return img_bytes                    # already within the usable size
        scale = MAX_IMAGE_EDGE / float(max(w, h))
        im = im.resize((max(1, round(w * scale)), max(1, round(h * scale))), Image.LANCZOS)
        fmt = _PIL_FORMAT.get(media_type, "PNG")
        if fmt == "JPEG" and im.mode not in ("RGB", "L"):
            im = im.convert("RGB")             # JPEG has no alpha
        out = io.BytesIO()
        im.save(out, format=fmt)
        return out.getvalue()
    except Exception:
        return img_bytes


def _log_no_source_file(message):
    """Append a persistent, timestamped line to output/no_source_file.log so a later
    'no source document' report is diagnosable without re-running (the console print scrolls
    away). Best-effort — never fails the save."""
    try:
        import time as _time
        OUTPUT_DIR.mkdir(exist_ok=True)
        with open(OUTPUT_DIR / "no_source_file.log", "a", encoding="utf-8") as fh:
            fh.write(f"{_time.strftime('%Y-%m-%d %H:%M:%S')}  {message}\n")
    except Exception:
        pass


def _write_record_document(record, store, vendor, number, amount, date_iso, target_dir, attachments, text, html_body=None, cid_images=None):
    """Save ONE record's source document to target_dir, named
    'Store - Vendor - Invoice # - Amount - Date': the matched attachment (split
    to just this record's pages when a PDF holds several records), or a PDF
    render of the email body. Returns the saved path (str), or '' if nothing
    could be produced."""
    target_dir.mkdir(exist_ok=True)
    stem = build_invoice_filename(store, vendor, number, amount, date_iso)
    att = _match_attachment(record.get("source_document"), attachments)
    if att is not None:
        media_type = _effective_media_type(att)
        ext = MEDIA_TYPE_TO_EXT.get(media_type) or (Path(att["filename"]).suffix.lower() or ".bin")
        data = att["data"]
        if media_type == "application/pdf":
            extracted = _extract_pdf_pages(data, record.get("source_pages"),
                                           record.get("page_rotations"),
                                           record.get("blank_pages"),
                                           record.get("page_order"))
            if extracted is not None:
                data = extracted
        elif media_type in SUPPORTED_IMAGE_TYPES:
            rotated = _rotate_image_bytes(data, media_type, record.get("page_rotations"))
            if rotated is not None:
                data = rotated
        dest = _unique_path(target_dir, stem, ext)
        dest.write_bytes(data)
        return str(dest.resolve())
    # No matched attachment — render the email body to PDF. Prefer the real
    # HTML body (looks like the actual email); fall back to the plaintext
    # render whenever HTML is missing or the browser render fails.
    fd, tmp = tempfile.mkstemp(suffix=".pdf")
    os.close(fd)
    tmp = Path(tmp)
    ok = render_html_to_pdf(html_body, cid_images or {}, tmp) if html_body else False
    if not ok:
        ok = render_text_to_pdf(text, tmp)
    if not ok:
        tmp.unlink(missing_ok=True)
        # No file could be produced -> this record shows "No source file recorded" in review.
        # Log WHICH side failed so it's diagnosable: the model didn't attribute this record to a
        # readable attachment (source_document below), AND the email-body render fell through too.
        reason = (f"NO SOURCE FILE for {vendor} #{number}: source_document="
                  f"{record.get('source_document')!r} matched no attachment "
                  f"({', '.join(a.get('filename', '?') for a in attachments) or 'none'}); "
                  f"html_render={'skipped(no html)' if not html_body else 'failed'}, "
                  f"text_render=failed.")
        print(f"  {reason}")
        _log_no_source_file(reason)          # persist it so the failure is diagnosable later
        return ""
    dest = _unique_path(target_dir, stem, ".pdf")
    shutil.copy2(tmp, dest)
    tmp.unlink(missing_ok=True)
    return str(dest.resolve())


def _record_fields(rtype, rec):
    """(store, vendor, number, amount, date) for naming this record's file."""
    if rtype == "invoice":
        return (rec.get("store_matched"), rec.get("vendor_matched"),
                rec.get("invoice_number"), rec.get("invoice_amount"), rec.get("invoice_date"))
    if rtype == "credit_memo":
        return (rec.get("store_matched"), rec.get("vendor_matched"),
                rec.get("credit_memo_number"), rec.get("credit_amount"), rec.get("credit_memo_date"))
    return (rec.get("store_matched"), rec.get("vendor_matched"),
            rec.get("reference_number"), rec.get("payment_amount"), rec.get("payment_date"))


def assign_invoice_files(source_path, text, attachments, classification, html_body=None, cid_images=None):
    """Save one source document for each NON-flagged qb-bound record: invoices &
    credit memos → invoice_files/, payments → payment_files/. Sets
    record['_file_path']; returns the list of saved paths (for logging)."""
    saved = []
    for rtype, rec in _all_records(classification):
        if not _record_needs_file(rec, rtype):
            rec["_file_path"] = ""
            continue
        store, vendor, number, amount, date_iso = _record_fields(rtype, rec)
        target_dir = PAYMENT_FILES_DIR if rtype == "payment" else INVOICE_FILES_DIR
        path = _write_record_document(rec, store, vendor, number, amount, date_iso, target_dir, attachments, text, html_body, cid_images)
        rec["_file_path"] = path
        if path:
            saved.append(path)
    return saved


def _find_document_sibling(rtype, rec, classification):
    """Another record in the SAME classification (i.e. same source email) with the SAME
    vendor + number + amount (ALL three) as `rec` — the recorded/flagged twin of a
    duplicate-detection extraction (order-confirmation emails classified as BOTH the real
    invoice AND a possible-duplicate check produce exactly this pair). Returns the sibling
    record dict, or None. Strict match so we never attach the WRONG document."""
    _s, vendor, number, amount, _d = _record_fields(rtype, rec)
    if not (vendor and number and amount is not None):
        return None
    for srt, srec in _all_records(classification):
        if srec is rec:
            continue
        _s2, s_vendor, s_number, s_amount, _d2 = _record_fields(srt, srec)
        if (_norm(vendor) == _norm(s_vendor)
                and number and s_number and _norm(str(number)) == _norm(str(s_number))
                and _amounts_equal(amount, s_amount)):
            return srec
    return None


def _recover_review_file_from_sibling(rtype, rec, store, vendor, number, amount, date_iso,
                                      classification, attachments, text, html_body, cid_images):
    """Last resort when a flagged record's own document generation failed: reuse the document of
    a same-email sibling with the SAME vendor+number+amount (a duplicate/twin of THIS record, so
    its document is exactly what a reviewer needs to compare). Prefers copying the sibling's saved
    local file; if that's gone (e.g. already swapped for a Drive URL), re-derives from the
    sibling's own source_document attribution (the attachment is still in `attachments`). Returns
    the new review_pending path, or '' if no usable sibling exists. Never fabricates — only reuses
    an already-legitimately-produced document."""
    sib = _find_document_sibling(rtype, rec, classification)
    if sib is None:
        return ""
    stem = build_invoice_filename(store, vendor, number, amount, date_iso)
    # 1) Copy the sibling's saved local file, if it still exists on disk.
    for key in ("_review_file", "_file_path"):
        sp = sib.get(key)
        if sp and not str(sp).lower().startswith("http") and os.path.exists(sp):
            ext = Path(sp).suffix or ".pdf"
            dest = _unique_path(REVIEW_PENDING_DIR, stem, ext)
            shutil.copy2(sp, dest)
            print(f"  recovered source file for {vendor} #{number} by copying sibling's "
                  f"document ({os.path.basename(sp)})")
            _log_no_source_file(f"RECOVERED {vendor} #{number} from sibling file {os.path.basename(sp)}")
            return str(dest.resolve())
    # 2) Re-derive from the sibling's document attribution (its attachment is in this same email).
    if _match_attachment(sib.get("source_document"), attachments) is not None:
        merged = dict(rec)
        for f in ("source_document", "source_pages", "page_rotations", "blank_pages", "page_order"):
            merged[f] = sib.get(f)
        path = _write_record_document(merged, store, vendor, number, amount, date_iso,
                                      REVIEW_PENDING_DIR, attachments, text, html_body, cid_images)
        if path:
            print(f"  recovered source file for {vendor} #{number} by re-extracting sibling's "
                  f"document ({sib.get('source_document')!r})")
            _log_no_source_file(f"RECOVERED {vendor} #{number} by re-extracting sibling "
                                f"source_document {sib.get('source_document')!r}")
            return path
    return ""


def save_review_files(source_path, text, attachments, classification, html_body=None, cid_images=None):
    """For each FLAGGED record, save its own standalone split document to
    review_pending/ and set record['_review_file']. This is the single-record
    file that gets reprocessed (alone) once you answer it — the clean records
    in the same batch are never touched. Returns the list of saved paths."""
    saved = []
    for rtype, rec in _all_records(classification):
        if not routes_to_review(rtype, rec):
            rec["_review_file"] = ""
            continue
        # Marks that this record's review document was already attempted here (a failure is logged
        # by _write_record_document), so the late-flag backstop doesn't retry it.
        rec["_review_attempted"] = True
        store, vendor, number, amount, date_iso = _record_fields(rtype, rec)
        path = _write_record_document(rec, store, vendor, number, amount, date_iso, REVIEW_PENDING_DIR, attachments, text, html_body, cid_images)
        if not path:
            # Own document generation failed — before leaving the reviewer with nothing, reuse a
            # same-email twin's document (recorded or flagged) with the same vendor+number+amount.
            path = _recover_review_file_from_sibling(
                rtype, rec, store, vendor, number, amount, date_iso,
                classification, attachments, text, html_body, cid_images)
        rec["_review_file"] = path
        if path:
            saved.append(path)
    return saved


def _create_message(client, *, model, max_tokens, system, messages, timeout):
    """Send a Messages API request via raw httpx, returning an object shaped like an SDK response
    (.content[].text/.type, .usage.*, .stop_reason).

    WHY NOT client.messages.create: on this machine the anthropic SDK (0.116.0 + httpx 0.28.1)
    HANGS INDEFINITELY on every request, ignoring its own timeout, while a direct httpx POST with
    the SAME key + CA bundle returns normally in ~2s (verified with test_api.py vs test_api_raw.py).
    This routes around the SDK hang. The request/response schema is identical, so the rest of the
    pipeline is unchanged. The CA bundle (corporate-proxy TLS) is still honored via verify=."""
    from types import SimpleNamespace
    key = getattr(client, "api_key", None) or (os.environ.get("ANTHROPIC_API_KEY") or "").strip()
    verify = str(CA_BUNDLE) if CA_BUNDLE.exists() else True
    headers = {"x-api-key": key, "anthropic-version": "2023-06-01", "content-type": "application/json"}
    body = {"model": model, "max_tokens": max_tokens, "system": system, "messages": messages}
    resp = httpx.post("https://api.anthropic.com/v1/messages", headers=headers, json=body,
                      timeout=httpx.Timeout(timeout, connect=15.0), verify=verify)
    if resp.status_code != 200:
        # Mirror the SDK's error wording so _is_auth_error / _is_invalid_request_error classify it
        # (they match "authentication_error"/"invalid_request_error" or "error code: 401"/"400").
        raise RuntimeError(f"API error code: {resp.status_code} - {resp.text[:500]}")
    data = resp.json()
    content = [SimpleNamespace(type=b.get("type"), text=b.get("text", ""))
               for b in (data.get("content") or [])]
    u = data.get("usage") or {}
    usage = SimpleNamespace(
        cache_creation_input_tokens=u.get("cache_creation_input_tokens", 0) or 0,
        cache_read_input_tokens=u.get("cache_read_input_tokens", 0) or 0,
        input_tokens=u.get("input_tokens", 0) or 0,
        output_tokens=u.get("output_tokens", 0) or 0,
    )
    return SimpleNamespace(content=content, usage=usage, stop_reason=data.get("stop_reason"))


def classify_email(client, filename, text, attachments, learnings_text="", reviewer_answer=""):
    attachment_blocks, skipped = build_attachment_blocks(attachments)

    # If this email came from a known self-scan mailbox, tell the model its store
    # decisively (from the editable sender_store_map).
    text += _sender_store_hint(text)
    # And when a To/Cc recipient identifies the store (third-party biller emailing us).
    text += _contact_store_hint(text)

    if skipped:
        text += (
            "\n\n[Note: this email had attachment(s) that this pipeline could not "
            f"read and are NOT reflected above: {', '.join(skipped)}. Treat any "
            "invoice data as incomplete and set needs_human_review to true.]"
        )

    if reviewer_answer:
        text += (
            "\n\n[HUMAN REVIEWER ANSWER — this email was previously flagged for review, "
            "and a human reviewer has provided the following authoritative answer. Apply "
            "it, and unless it raises a NEW unresolved question, set needs_human_review to "
            f"false:\n{reviewer_answer}]"
        )

    system = get_system_prompt()
    if learnings_text:
        system += (
            "\n\nLEARNED RULES (from prior human review — apply these; they override your "
            f"default assumptions when relevant):\n{learnings_text}"
        )

    # The API rejects an empty text block ("text content blocks must be non-empty").
    # A text-less message (image-only/scanned PDF, no email body) with no reviewer
    # answer would leave `text` empty, so send a minimal instruction and let the model
    # classify from the attached document(s) instead of 400-ing.
    if not text.strip():
        text = "(No extractable text in this message; classify from the attached document(s) below.)"
    content = [{"type": "text", "text": text}] + attachment_blocks

    # Prompt caching: the system content is byte-identical across every email in a run
    # (get_system_prompt() + learnings_text are both built once per run), so mark it as a
    # cache breakpoint. The first call of a run writes the cache (cache_creation_input_tokens);
    # every later call reads it at ~10% cost (cache_read_input_tokens). Default 5-minute
    # ephemeral TTL is enough because each cache READ refreshes the TTL, so a back-to-back run
    # keeps it warm as long as no single gap between calls exceeds 5 min (true here — one email
    # classify, even with a PDF render, is well under that). Same content, same output — only
    # billing/usage metadata changes.
    # Raw-httpx call (see _create_message): the anthropic SDK hangs forever on this machine; a
    # direct POST with the same key + CA bundle works in ~2s. 120s timeout → a stall fails here
    # and hits classify_with_retry's visible retry instead of a silent freeze.
    response = _create_message(
        client,
        model=MODEL,
        max_tokens=MAX_TOKENS,
        system=[{"type": "text", "text": system, "cache_control": {"type": "ephemeral"}}],
        messages=[{"role": "user", "content": content}],
        timeout=120.0,
    )
    # Cache telemetry: on the first call of a run the system prompt is WRITTEN
    # (cache_creation_input_tokens > 0, cache_read = 0); every later call in the same
    # run READS it (cache_read_input_tokens > 0, cache_creation = 0) at ~10% cost. This
    # one line makes the caching visible/verifiable each run.
    _u = getattr(response, "usage", None)
    if _u is not None:
        _cc = getattr(_u, "cache_creation_input_tokens", 0) or 0
        _cr = getattr(_u, "cache_read_input_tokens", 0) or 0
        print(f"  cache: {_cc} written, {_cr} read "
              f"(input={getattr(_u, 'input_tokens', 0)}, output={getattr(_u, 'output_tokens', 0)})")
    stop_reason = getattr(response, "stop_reason", None)
    raw_text = "".join(
        block.text for block in response.content if block.type == "text"
    ).strip()

    if raw_text.startswith("```"):
        raw_text = raw_text.strip("`")
        if raw_text.startswith("json"):
            raw_text = raw_text[4:]
        raw_text = raw_text.strip()

    if not raw_text:
        # Model returned no text. stop_reason tells us why — most often
        # "max_tokens" (output budget exhausted, common on big multi-invoice
        # scans) vs "end_turn" (a genuine empty reply).
        parsed = None
        error = f"Empty response from model (stop_reason={stop_reason})"
    else:
        try:
            parsed = json.loads(raw_text)
            error = None
        except json.JSONDecodeError as e:
            parsed = None
            hint = " — response was cut off at max_tokens" if stop_reason == "max_tokens" else ""
            error = f"Failed to parse JSON response: {e} (stop_reason={stop_reason}){hint}"

    return {
        "filename": filename,
        "classification": parsed,
        "raw_response": raw_text if parsed is None else None,
        "skipped_attachments": skipped,
        "error": error,
    }


def _is_auth_error(msg):
    m = (msg or "").lower()
    return "authentication_error" in m or "invalid x-api-key" in m or "error code: 401" in m


def _is_invalid_request_error(msg):
    """A 400 invalid_request_error is a DETERMINISTIC client-side problem with this
    request (e.g. an attachment the API can't accept, oversized image dimensions) —
    unlike a connection/network blip, retrying it can never succeed."""
    m = (msg or "").lower()
    return "invalid_request_error" in m or "error code: 400" in m


def classify_with_retry(client, filename, text, attachments, learnings_text="", reviewer_answer="", attempts=3):
    """Call classify_email up to `attempts` times. Empty/unparseable responses
    are one-off (retry usually fixes them). A 401 auth error means the API key
    is bad — fail the whole run fast rather than mangling every email. Other API
    errors (connection/network) are flagged infra_error so the caller leaves the
    email in test_emails/ to retry, instead of dumping it in the review queue."""
    last = None
    for i in range(attempts):
        try:
            result = classify_email(client, filename, text, attachments, learnings_text, reviewer_answer)
        except Exception as e:
            if _is_auth_error(str(e)):
                raise SystemExit(
                    "\nANTHROPIC_API_KEY appears INVALID (401 authentication_error). "
                    "No emails were changed — fix the key (check for a placeholder or a "
                    "trailing space) and rerun.\n")
            if _is_invalid_request_error(str(e)):
                # Deterministic 400 — retrying can't help. Return a PLAIN error (no
                # infra_error), so the caller routes it to the review queue as
                # "unreadable" for a human instead of leaving it to retry forever.
                return {
                    "filename": filename, "classification": None, "raw_response": None,
                    "skipped_attachments": [], "error": f"API rejected the request: {e}",
                }
            last = {
                "filename": filename, "classification": None, "raw_response": None,
                "skipped_attachments": [], "error": f"API call failed: {e}", "infra_error": True,
            }
            # Make a network/timeout stall VISIBLE (this exception path used to retry silently,
            # so a stalled request looked like a total freeze with no output).
            if i + 1 < attempts:
                print(f"  attempt {i + 1} failed (API call failed: {str(e)[:120]}); retrying...")
            continue
        if not result.get("error"):
            return result
        last = result
        if i + 1 < attempts:
            print(f"  attempt {i + 1} failed ({result.get('error')}); retrying...")
    return last


def _count_nonblank_pdf_pages(pdf_bytes, dpi=40, white_fraction=0.90):
    """Count pages that aren't essentially blank (a scanner's blank duplex back), by
    rendering each page at low resolution and checking how much of it is near-white.
    A page is 'blank' when at least `white_fraction` of its pixels are near-white.
    Returns None (caller falls back to the raw page count) if PyMuPDF isn't installed or
    rendering fails for any reason — this is purely an accuracy improvement on top of the
    existing page-count check, never a hard requirement.

    Thresholds calibrated on the real review #2299 duplex scan: its blank backs measured
    ~0.98-0.99 near-white in the 235-255 band while its 5 content pages measured <=0.73, a
    wide, clean gap. A scanner's blank back is NOT pure white (speckle/noise), so a stricter
    band/fraction (e.g. 250-255 @ 0.995) fails to detect it. Erring toward calling a page
    blank is the SAFE direction here: a lower count only makes chunking LESS likely, and it's
    over-eager chunking of a duplex doc that caused the bug."""
    try:
        import pymupdf as fitz  # PyMuPDF (>=1.24 exposes `pymupdf`; `fitz` is the legacy alias)
    except ImportError:
        try:
            import fitz  # older PyMuPDF
        except ImportError:
            return None
    try:
        from PIL import Image
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        count = 0
        for page in doc:
            pix = page.get_pixmap(dpi=dpi)
            img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples).convert("L")
            hist = img.histogram()
            near_white = sum(hist[235:256])
            total = sum(hist) or 1
            if (near_white / total) < white_fraction:
                count += 1
        return count
    except Exception:
        return None


def _lone_oversized_pdf(attachments):
    """Return the single attachment when it is the ONLY attachment, is a PDF, and
    has MORE than CHUNK_THRESHOLD_PAGES REAL (non-blank) pages — the 'one big scanned
    stack of separate invoices' case that should be page-chunked before classification.
    A duplex scan padded with blank backs (content page, blank, content page, blank...)
    can have double its real page count without needing to be chunked at all — chunking
    it anyway risks splitting mid-document and losing continuity across the split
    (confirmed cause of review #2299's false 'pages missing' flag). Otherwise None, so
    every other email falls through to the normal single call unchanged."""
    if len(attachments) != 1:
        return None
    att = attachments[0]
    if _effective_media_type(att) != "application/pdf":
        return None
    try:
        n = len(PdfReader(io.BytesIO(att["data"])).pages)
    except Exception:
        return None
    if n <= CHUNK_THRESHOLD_PAGES:
        return None
    real_n = _count_nonblank_pdf_pages(att["data"])
    if real_n is not None and real_n <= CHUNK_THRESHOLD_PAGES:
        return None
    return att


def _split_pdf_into_chunks(pdf_bytes, chunk_size):
    """Split a PDF into sequential page-range chunks of at most `chunk_size` pages.
    Returns a list of (start_page_1based, end_page_1based, chunk_pdf_bytes), in order."""
    reader = PdfReader(io.BytesIO(pdf_bytes))
    total = len(reader.pages)
    chunks = []
    start = 1
    while start <= total:
        end = min(start + chunk_size - 1, total)
        writer = PdfWriter()
        for p in range(start, end + 1):
            writer.add_page(reader.pages[p - 1])
        buf = io.BytesIO()
        writer.write(buf)
        chunks.append((start, end, buf.getvalue()))
        start = end + 1
    return chunks


def _classify_oversized_scan(client, filename, text, pdf_att, learnings_text="", reviewer_answer=""):
    """Classify one oversized scanned PDF by splitting it into sequential page-range
    chunks (CHUNK_SIZE_PAGES each), classifying each chunk on its own, and merging
    the per-chunk classifications into ONE classification whose source_pages /
    page_rotations are re-expressed in the ORIGINAL PDF's page numbering (downstream
    re-opens the ORIGINAL full bytes and splits by source_pages, so these MUST be
    global page numbers). Returns a result dict shaped exactly like classify_email's.

    If ANY chunk fails to classify after retries, the WHOLE result fails with a
    descriptive error (carrying infra_error through when the chunk's failure was a
    network blip), so it routes exactly like a normal failure instead of silently
    losing the records on the other pages."""
    chunks = _split_pdf_into_chunks(pdf_att["data"], CHUNK_SIZE_PAGES)
    n = len(chunks)
    total_pages = chunks[-1][1] if chunks else 0
    print(f"  oversized scan ({total_pages} pages) - classifying in {n} page-chunk(s) "
          f"of up to {CHUNK_SIZE_PAGES}")
    merged = None
    review_flag = False
    for k, (start, end, chunk_bytes) in enumerate(chunks, 1):
        chunk_att = [{"filename": pdf_att["filename"], "mimetype": "application/pdf", "data": chunk_bytes}]
        note = (f"\n\n[This is part {k} of {n} of a multi-page scan (pages {start}-{end} of the "
                f"original document); classify ONLY the invoices/records visible on these pages. "
                f"Do not assume anything is missing — the other pages are handled separately.]")
        res = classify_with_retry(client, filename, text + note, chunk_att,
                                  learnings_text=learnings_text, reviewer_answer=reviewer_answer)
        c = res.get("classification")
        if not c:
            # This chunk couldn't be classified even after retries — fail the whole
            # merged result rather than dropping this page range's records.
            out = {
                "filename": filename, "classification": None,
                "raw_response": res.get("raw_response"),
                "skipped_attachments": res.get("skipped_attachments", []),
                "error": f"page {start}-{end} of scan failed to classify: {res.get('error')}",
            }
            if res.get("infra_error"):
                out["infra_error"] = True     # network blip → leave for retry, like today
            return out
        offset = start - 1
        for key in ("invoices", "credit_memos", "payments"):
            for rec in c.get(key) or []:
                shifted = _offset_page_spec(rec.get("source_pages"), offset)
                if shifted is None:
                    # No/blank page spec for a chunk record → it spans the whole chunk;
                    # bound the saved file to THIS chunk's global pages, never the whole scan.
                    shifted = f"{start}-{end}"
                rec["source_pages"] = shifted
                rec["page_rotations"] = _offset_rotations(rec.get("page_rotations"), offset)
                # Same chunk->original page-number shift for the two new page fields, or the
                # blank-drop / reorder would target the wrong pages on any chunked (oversized) scan.
                rec["blank_pages"] = _offset_page_list(rec.get("blank_pages"), offset)
                rec["page_order"] = _offset_page_list(rec.get("page_order"), offset)
        if merged is None:
            merged = c                         # base: keep the first chunk's top-level fields
        else:
            for key in ("invoices", "credit_memos", "payments", "statement_invoices"):
                if c.get(key):
                    merged.setdefault(key, [])
                    merged[key].extend(c[key])
        if c.get("needs_human_review"):
            review_flag = True
    merged["needs_human_review"] = bool(merged.get("needs_human_review")) or review_flag
    return {
        "filename": filename, "classification": merged, "raw_response": None,
        "skipped_attachments": [], "error": None,
    }


def _write_xlsx(path, fieldnames, rows, sheet_title):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title
    ws.append(fieldnames)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for row in rows:
        ws.append([(row.get(col) if row.get(col) != "" else None) for col in fieldnames])
    for col_cells in ws.columns:
        width = max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(width + 2, 10), 60)
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def load_history():
    if USE_DB:
        return _database().load_history()
    if not HISTORY_PATH.exists():
        return []
    wb = load_workbook(HISTORY_PATH, read_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    history = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))
        if row_dict.get("vendor_matched") is None and row_dict.get("source_file") is None:
            continue
        amount = row_dict.get("invoice_amount")
        invoice_number = row_dict.get("invoice_number")
        invoice_date = row_dict.get("invoice_date")
        history.append({
            "vendor_matched": row_dict.get("vendor_matched"),
            "store_matched": row_dict.get("store_matched"),
            "invoice_number": str(invoice_number) if invoice_number is not None else None,
            "invoice_date": str(invoice_date) if invoice_date is not None else None,
            "invoice_amount": float(amount) if amount is not None else None,
            "source_file": row_dict.get("source_file"),
            "source_email_id": row_dict.get("source_email_id") or row_dict.get("source_file"),
            "date_added": row_dict.get("date_added"),
            "from_address": None,   # Excel history has no sender; store-preference is DB-mode only
        })
    wb.close()
    return history


def save_history(history):
    if USE_DB:
        _database().save_history(history)
    if WRITE_EXCEL:
        _write_xlsx(HISTORY_PATH, HISTORY_FIELDNAMES, history, "History")


def load_payment_history():
    if USE_DB:
        return _database().load_payment_history()
    if not PAYMENT_HISTORY_PATH.exists():
        return []
    wb = load_workbook(PAYMENT_HISTORY_PATH, read_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        if d.get("vendor_matched") is None and d.get("source_file") is None:
            continue
        amount = d.get("amount")
        ref = d.get("reference_number")
        out.append({
            "vendor_matched": d.get("vendor_matched"), "store_matched": d.get("store_matched"),
            "reference_number": str(ref) if ref is not None else None,
            "invoice_number": str(d["invoice_number"]) if d.get("invoice_number") is not None else None,
            "payment_date": str(d.get("payment_date")) if d.get("payment_date") is not None else None,
            "amount": float(amount) if amount is not None else None,
            "source_file": d.get("source_file"),
            "source_email_id": d.get("source_email_id") or d.get("source_file"),
            "date_added": d.get("date_added"),
        })
    wb.close()
    return out


def save_payment_history(history):
    if USE_DB:
        _database().save_payment_history(history)
    if WRITE_EXCEL:
        _write_xlsx(PAYMENT_HISTORY_PATH, PAYMENT_HISTORY_FIELDNAMES, history, "Payment History")


def _read_xlsx_rows(path):
    if not path.exists():
        return []
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))
        if all(v is None for v in row):
            continue
        rows.append(row_dict)
    wb.close()
    return rows


def _rows_to_learnings_text(rows):
    text = "\n".join(f"- {r.get('Learned Rule')}" for r in rows if r.get("Learned Rule"))
    return rows, text


def load_learnings():
    """Returns (learnings_rows, learnings_text). learnings_text is the bullet
    list injected into the system prompt on every run."""
    rows = _database().load_learnings() if USE_DB else _read_xlsx_rows(LEARNINGS_PATH)
    return _rows_to_learnings_text(rows)


def save_learnings(rows):
    if USE_DB:
        _database().save_learnings(rows)
    if WRITE_EXCEL:
        _write_xlsx(LEARNINGS_PATH, LEARNINGS_FIELDNAMES, rows, "Learnings")


def add_learning(learnings_rows, rule, from_email):
    """Append a new learned rule, skipping exact duplicates."""
    rule = (rule or "").strip()
    if not rule:
        return
    if any((r.get("Learned Rule") or "").strip() == rule for r in learnings_rows):
        return
    learnings_rows.append({
        "Learned Rule": rule,
        "Date Added": TODAY,
        "From Email": from_email or "",
    })


def load_review_queue():
    if USE_DB:
        return _database().load_review_queue()
    return _read_xlsx_rows(REVIEW_QUEUE_PATH)


def save_review_queue(rows):
    if USE_DB:
        _database().save_review_queue(rows)
    if WRITE_EXCEL:
        _write_xlsx(REVIEW_QUEUE_PATH, REVIEW_QUEUE_FIELDNAMES, rows, "Review Queue")


def _norm(value):
    return (value or "").strip().lower()


def _applied_invoice_from_notes(notes):
    """Fallback when the model stated the applied invoice in free-text notes but left the
    structured invoice_number field empty -- regex-extracts a '...against invoice <NUM>'
    mention (the same phrasing the extraction prompt already asks the model to use, e.g.
    'Payment against invoice #TRINV-685027' or 'Payment applied against invoice
    TRINV-688744.'). None if no match."""
    if not notes:
        return None
    m = re.search(r"against\s+invoice\s*#?\s*([A-Za-z0-9][A-Za-z0-9\-]*)", notes, re.IGNORECASE)
    return m.group(1).rstrip(".,;:") if m else None


# Common legal-entity suffixes, stripped ONLY for duplicate-comparison vendor matching
# (so "Schwartz & Company CPAs" and "Schwartz & Company CPAs PLLC" compare equal). Periods
# are removed before this membership test, so "l.l.c." / "co." also match.
_DEDUPE_ENTITY_SUFFIXES = {
    "pllc", "llc", "llp", "lllp", "inc", "incorporated", "corp", "corporation",
    "co", "company", "pc", "ltd", "cpa", "cpas",
}


def _norm_vendor_dedupe(value):
    """Vendor normalizer for DUPLICATE comparison ONLY: like _norm(), but also strips common
    trailing legal-entity suffixes (PLLC / LLC / L.L.C. / Inc / Inc. / Incorporated / Corp /
    Corp. / Corporation / Co / Co. / PC / P.C. / Ltd / Ltd. / LLP / L.L.P.), from the END only,
    case-insensitive and punctuation-tolerant, possibly several in a row ("Co., Inc."). Used
    solely inside check_duplicate()/check_duplicate_payment(); the general _norm() and
    is_excluded_vendor() are deliberately left unchanged."""
    name = _norm(value)
    while name:
        m = re.search(r"[\s,]+([a-z.]+)$", name)
        if not m or m.group(1).replace(".", "") not in _DEDUPE_ENTITY_SUFFIXES:
            break
        name = name[:m.start()].rstrip(" ,.")
    return name


_DEDUPE_VENDOR_CACHE = {}


def _dedupe_vendor_key(value):
    """Vendor key for DUPLICATE identity: the vendor's STANDARD name (resolve_standard_vendor --
    the same exact-then-fuzzy vendor_name_map lookup standardization uses) when it resolves, else
    the raw spelling; then _norm_vendor_dedupe. Before 2026-10-08 duplicate checks compared the raw
    vendor_matched spelling only, so 'Costco (Same-Day/Instacart)' vs 'Costco', 'MaxiHealth' vs
    'Maxi Health', 'Prestige Sales Inc.' vs 'A&S Sales' never matched and 37 same-invoice pairs
    stayed live on /qb (27 with one copy already entered in QB). A fuzzy resolve can only ever
    cause an auto-ignore together with an identical invoice # AND amount, so the risk of merging
    two genuinely different vendors is negligible. Cached per loaded VENDOR_STD."""
    if not value:
        return _norm_vendor_dedupe(value)
    ck = (id(VENDOR_STD), value)
    if ck not in _DEDUPE_VENDOR_CACHE:
        std = None
        try:
            std = resolve_standard_vendor(value)
        except Exception:
            std = None
        _DEDUPE_VENDOR_CACHE[ck] = _norm_vendor_dedupe(std or value)
    return _DEDUPE_VENDOR_CACHE[ck]


def _stable_review_key(source_email, vendor_matched, number):
    """A review item's identity that survives reprocessing, unlike its generated display
    filename (which can change: store corrected, or a naming collision bumps the suffix). Same
    formula every time the SAME underlying record is flagged, however many times — so notes and
    manual attachments keyed to it are never orphaned by a rename. Keyed by source email +
    normalized vendor + invoice/reference number."""
    v = _norm_vendor_dedupe(vendor_matched) or ""
    n = _norm(number) or ""
    return f"{source_email}::{v}::{n}" if (v or n) else source_email


def is_excluded_vendor(vendor_matched):
    """True if this vendor is on the QuickBooks exclusion list (skip_qb_import).
    The list is every spelling of each Setup › Vendor names row ticked Exclude from QB
    (database.load_excluded_vendors; the seed defaults in Excel mode) — the SAME list
    spliced into the model-facing prompt,
    so there's one source of truth. Read fresh each call so a live UI edit in the
    long-running web app takes effect immediately (this is only called by the UI's
    skip_qb re-evaluation, not in the CLI's hot loop). Matches by normalized
    substring either direction so 'Anda New'/'Anda' and 'BluPax Pharma'/'BluPax'
    both resolve."""
    name = _norm(vendor_matched)
    if not name:
        return False
    for excluded in load_excluded_vendors():
        ex = _norm(excluded)
        if ex and (ex in name or name in ex):
            return True
    return False


def _amounts_equal(a, b):
    """Tolerant numeric equality for invoice amounts (handles float vs string vs
    None). None on either side is NOT equal."""
    if a is None or b is None:
        return False
    try:
        return abs(float(a) - float(b)) < 0.005
    except (TypeError, ValueError):
        return _norm(a) == _norm(b)


def _earlier_store_copy(invoice, history, vendor_prior, source_id):
    """For check_duplicate's store-preference: an earlier STORE-SCANNED copy of this invoice already
    in the ledger (same vendor + number + amount + store, from another email, sent by a mailbox in
    sender_store_map), other than `vendor_prior` (the vendor's own send). None if there's none."""
    vendor = _dedupe_vendor_key(invoice.get("vendor_matched"))
    num = _norm(invoice.get("invoice_number"))
    for p in history:
        if p is vendor_prior:
            continue
        if (_dedupe_vendor_key(p.get("vendor_matched")) == vendor and _norm(p.get("invoice_number")) == num
                and _amounts_equal(p.get("invoice_amount"), invoice.get("invoice_amount"))
                and _norm(p.get("store_matched")) == _norm(invoice.get("store_matched"))
                and (p.get("source_email_id") or p.get("source_file")) != source_id
                and SENDER_STORE_MAP.get(_norm(p.get("from_address")))):
            return p
    return None


def check_duplicate(invoice, history, source_id=None, source_file=None):
    """Compare against previously recorded invoices (vendor + invoice number).
    The self-match gate is AMOUNT-based, not filename-based: a prior from the SAME
    source email with the SAME amount is "self" (the literal same document, or the same
    page scanned twice — nothing to flag). A prior from the same email with a DIFFERENT
    amount is NOT self — it's a revision candidate and falls through to 'possible' (this
    is what lets a duplicate-then-revision pair bundled in one scan be detected).
    Filenames can't be used to tell documents apart because they're GENERATED from
    (store, vendor, invoice#, amount, date) and collide by design. Returns (status, reason, dup_of):
      'none'      — no prior match.
      'self'      — only matched a same-email + same-amount prior → NOT a duplicate.
      'duplicate' — REAL duplicate (same vendor + # + amount, a DIFFERENT email) OR a
                    REDUNDANT re-record (SAME email, DIFFERENT file, IDENTICAL amount —
                    same page scanned twice) → auto-ignore (out of QB + review).
      'possible'  — a prior matches vendor + invoice # but the amount differs (or is
                    missing) → flag for review as a likely REVISION. The reason may
                    include not-received/short-shipped markup evidence (and whether it
                    explains the amount delta) when either copy carries it — guidance
                    only; the keep/ignore decision still needs a human.
    dup_of is the matched prior record's source-email identity (for a Browse link),
    or None. A 4th value is the matched PRIOR dict itself (or None) — used by the caller to
    read/propagate its duplicate_group_id. Real duplicate wins over possible, which wins over
    redundant, which wins over self."""
    vendor = _dedupe_vendor_key(invoice.get("vendor_matched"))
    inv_num = _norm(invoice.get("invoice_number"))
    if not vendor or not inv_num:
        return "none", None, None, None

    self_match = False
    redundant_of = None  # SAME email, DIFFERENT recorded file, IDENTICAL amount -> redundant copy
    redundant_prior = None
    possible = None      # (reason, dup_of, prior) remembered if no exact real dup is found
    for prior in history:
        if _dedupe_vendor_key(prior.get("vendor_matched")) != vendor or _norm(prior.get("invoice_number")) != inv_num:
            continue
        prior_id = prior.get("source_email_id") or prior.get("source_file")
        prior_file = prior.get("source_file")
        same_email = bool(source_id and prior_id and prior_id == source_id)
        amounts_match = _amounts_equal(prior.get("invoice_amount"), invoice.get("invoice_amount"))
        dup_of = prior.get("source_file") or prior_id

        if same_email and amounts_match:
            # Same message, same amount: either the literal same document seen again,
            # or the same page attached/scanned twice into one job (redundant). Neither
            # is a revision and neither needs a human -- just note which for the audit
            # trail. NOT gated on filename: generated filenames collide by design when
            # vendor/number/amount/date all match, so filename can't distinguish "same
            # document" from "different document with identical amount."
            self_match = True
            if prior_file != source_file:
                redundant_of = prior_file
                redundant_prior = prior
            continue

        if amounts_match:
            # Same vendor + invoice # + amount, but recorded to a DIFFERENT store: this is a
            # store reclassification, not an ordinary duplicate. Handled specially by the caller
            # (auto-ignore the EARLIER one). Same (reason, dup_of) shape as 'duplicate'.
            if _norm(prior.get("store_matched")) != _norm(invoice.get("store_matched")):
                return "store_conflict", (
                    f"Invoice #{invoice.get('invoice_number')} from {invoice.get('vendor_matched')} "
                    f"was already recorded to {prior.get('store_matched')} from {prior.get('source_file')} "
                    f"— this copy has it under {invoice.get('store_matched')}."
                ), dup_of, prior
            # STORE-PREFERENCE (forward-only): on an exact duplicate to the SAME store, prefer the
            # STORE-SCANNED copy over the vendor's own send, regardless of arrival order. Applies
            # ONLY when BOTH senders are known (never guess on a missing sender) and EXACTLY ONE
            # resolves to a store via sender_store_map. Otherwise fall through to first-arrived.
            new_fa = _norm(invoice.get("from_address"))
            prior_fa = _norm(prior.get("from_address"))
            if new_fa and prior_fa:
                new_store = SENDER_STORE_MAP.get(new_fa)
                prior_store = SENDER_STORE_MAP.get(prior_fa)
                if bool(new_store) != bool(prior_store):
                    if new_store:
                        # ...unless a store-scanned copy is ALREADY on file — then THIS is a second
                        # scan, a plain duplicate of that copy. The loop reaches the vendor's send
                        # first, so without this the second scan was "preferred" too and both scans
                        # stayed live (A&S #17496, Boiron #FI26144209 — found 2026-10-08).
                        store_copy = _earlier_store_copy(invoice, history, prior, source_id)
                        if store_copy is not None:
                            return "duplicate", (
                                f"Invoice #{invoice.get('invoice_number')} from {invoice.get('vendor_matched')} "
                                f"for ${store_copy.get('invoice_amount')} was already recorded from "
                                f"{store_copy.get('source_file')} (an earlier store-scanned copy) — this "
                                f"second scan is auto-ignored [first-arrived]."
                            ), store_copy.get("source_file") or store_copy.get("source_email_id"), store_copy
                        # THIS copy is the store scan, the earlier one is the vendor's send: keep
                        # THIS and auto-ignore the earlier vendor copy (caller supersedes it).
                        return "store_preferred", (
                            f"Invoice #{invoice.get('invoice_number')} from {invoice.get('vendor_matched')} "
                            f"for ${invoice.get('invoice_amount')} matches an earlier copy from "
                            f"{prior.get('source_file')} (the vendor's own send). Keeping THIS "
                            f"store-scanned copy ({new_store}); auto-ignoring the earlier vendor copy "
                            f"[store-preferred]."
                        ), dup_of, prior
                    # The EARLIER copy is the store scan (first-arrived already = store-preferred):
                    # keep it, auto-ignore this vendor copy.
                    return "duplicate", (
                        f"Invoice #{invoice.get('invoice_number')} from {invoice.get('vendor_matched')} "
                        f"for ${prior.get('invoice_amount')} was already recorded from "
                        f"{prior.get('source_file')} (the store-scanned copy, {prior_store}) — this "
                        f"vendor copy is auto-ignored [store-preferred]."
                    ), dup_of, prior
            return "duplicate", (
                f"Invoice #{invoice.get('invoice_number')} from {invoice.get('vendor_matched')} "
                f"for ${prior.get('invoice_amount')} was already recorded from "
                f"{prior.get('source_file')} — this is a duplicate copy, auto-ignored [first-arrived]."
            ), dup_of, prior

        # Amount differs from this prior -- a likely revision, whether it arrived in
        # the same email/scan (a duplicate-then-correction bundled together) or
        # separately. This is the case that was previously getting swallowed as "self."
        if possible is None:
            same_batch = " (from the same email/scan)" if same_email else ""
            reason = (
                f"Invoice #{invoice.get('invoice_number')} from {invoice.get('vendor_matched')} "
                f"was already recorded from {prior.get('source_file')}{same_batch} for "
                f"${prior.get('invoice_amount')}, but this copy shows "
                f"${invoice.get('invoice_amount')} — likely a REVISED/CORRECTED invoice, "
                f"verify the final amount before entering."
            )
            # Markup evidence: if either copy shows items marked not-received/short-shipped,
            # surface it and (when the dollar impact is known) check whether it explains the
            # delta — as guidance only; the ignore/keep decision still needs a human click.
            note = invoice.get("not_received_note") or prior.get("not_received_note")
            note_amount = (invoice.get("not_received_amount") if invoice.get("not_received_note")
                           else prior.get("not_received_amount"))
            if note:
                reason += f" Document notes items marked not-received/short: {note}."
                try:
                    delta = abs(float(prior.get("invoice_amount")) - float(invoice.get("invoice_amount")))
                    if note_amount is not None:
                        if abs(delta - float(note_amount)) <= 0.02:
                            reason += (f" This matches the ${delta:.2f} difference — consistent with "
                                       f"the lower amount being the corrected invoice after those items "
                                       f"were backed out. Still verify before entering.")
                        else:
                            reason += (f" This does NOT cleanly match the ${delta:.2f} difference "
                                       f"(marked items total ${float(note_amount):.2f}) — double-check "
                                       f"before entering.")
                except (TypeError, ValueError):
                    pass
            possible = (reason, dup_of, prior)

    if possible is not None:
        return "possible", possible[0], possible[1], possible[2]
    if redundant_of:
        return "duplicate", (
            f"Invoice #{invoice.get('invoice_number')} from {invoice.get('vendor_matched')} is a "
            f"redundant re-record of the same document (already recorded as {redundant_of}) — "
            f"auto-ignored."), redundant_of, redundant_prior
    if self_match:
        return "self", None, None, None
    return "none", None, None, None


def check_duplicate_payment(payment, history, source_id=None, source_file=None):
    """Mirror of check_duplicate for PAYMENTS (they can be double-recorded too).
    Match key is vendor + reference_number (the check/txn/card ref) + amount.
    Returns (status, reason, dup_of):
      'self'      — a prior from the SAME source email + SAME file matches → not a duplicate.
      'duplicate' — DIFFERENT source with vendor + reference_number + amount ALL
                    matching → auto-ignore. ONLY when a reference_number is present;
                    a payment is NEVER auto-ignored on amount alone. ALSO a REDUNDANT
                    re-record (SAME source email, DIFFERENT recorded file, identical ref+amount).
      'possible'  — reference matches but amount differs, OR (no reference_number)
                    vendor + amount + payment_date match → flag, do NOT auto-ignore.
      'none'      — unique.
    History rows carry `amount`, `reference_number`, `payment_date`, `source_file`,
    `source_email_id`. A 4th return value is the matched PRIOR dict itself (or None), for the
    caller's duplicate_group_id propagation."""
    vendor = _dedupe_vendor_key(payment.get("vendor_matched"))
    ref = _norm(payment.get("reference_number"))
    amount = payment.get("payment_amount")
    pdate = _norm(payment.get("payment_date"))
    if not vendor:
        return "none", None, None, None
    redundant_of = None   # SAME email, DIFFERENT recorded file, identical ref + amount
    redundant_prior = None

    self_match = False
    possible = None
    for prior in history:
        if _dedupe_vendor_key(prior.get("vendor_matched")) != vendor:
            continue
        prior_id = prior.get("source_email_id") or prior.get("source_file")
        same_source = bool(source_id and prior_id and prior_id == source_id)
        dup_of = prior.get("source_file") or prior_id
        pref = _norm(prior.get("reference_number"))
        amt_eq = _amounts_equal(prior.get("amount"), amount)

        inv_num = _norm(payment.get("invoice_number"))
        prior_inv_num = _norm(prior.get("invoice_number"))
        # When both sides state which invoice this payment was applied to, that has
        # to match too — a shared payment/batch reference legitimately repeats across
        # several different invoices it settled in one transaction, so reference_number
        # alone isn't enough to call two payments the same document. Fall back to
        # reference-only matching when either side doesn't state an invoice_number
        # (not every payment document ties itself to one).
        inv_num_conflicts = bool(inv_num and prior_inv_num and inv_num != prior_inv_num)

        if ref and pref and ref == pref and not inv_num_conflicts:   # reference is the key when present
            if same_source:
                self_match = True
                # Redundant re-record: same email, DIFFERENT recorded file, identical amount.
                if (prior.get("source_file") and source_file
                        and prior.get("source_file") != source_file and amt_eq):
                    redundant_of = prior.get("source_file")
                    redundant_prior = prior
                continue
            if amt_eq:
                return "duplicate", (
                    f"Payment ref {payment.get('reference_number')} from "
                    f"{payment.get('vendor_matched')} for ${prior.get('amount')} was already "
                    f"recorded from {prior.get('source_file')} — duplicate payment, auto-ignored."
                ), dup_of, prior
            if possible is None:
                possible = ((
                    f"Payment ref {payment.get('reference_number')} from "
                    f"{payment.get('vendor_matched')} was already recorded from "
                    f"{prior.get('source_file')} for ${prior.get('amount')}, but this shows "
                    f"${payment.get('payment_amount')} — check before entering."
                ), dup_of, prior)
        elif not ref:                                   # no reference: soft match only
            if amt_eq and pdate and _norm(prior.get("payment_date")) == pdate:
                if same_source:
                    self_match = True
                    continue
                if possible is None:
                    possible = ((
                        f"A payment from {payment.get('vendor_matched')} for "
                        f"${payment.get('payment_amount')} dated {payment.get('payment_date')} was "
                        f"already recorded from {prior.get('source_file')} (no reference # to "
                        f"confirm) — check before entering."
                    ), dup_of, prior)

    if possible is not None:
        return "possible", possible[0], possible[1], possible[2]
    if redundant_of:
        return "duplicate", (
            f"Payment ref {payment.get('reference_number')} from {payment.get('vendor_matched')} "
            f"is a redundant re-record of the same document (already recorded as {redundant_of}) — "
            f"auto-ignored."), redundant_of, redundant_prior
    if self_match:
        return "self", None, None, None
    return "none", None, None, None


def _all_records(classification):
    """All line-item records in an email, tagged with their type."""
    out = []
    for inv in classification.get("invoices") or []:
        out.append(("invoice", inv))
    for cm in classification.get("credit_memos") or []:
        out.append(("credit_memo", cm))
    for pay in classification.get("payments") or []:
        out.append(("payment", pay))
    return out


def result_has_review(result):
    """True if this email has anything needing human review — a flagged record,
    or (for record-less emails) a top-level flag. Errors/unreadable are handled
    separately, not here."""
    c = result.get("classification")
    if not c:
        return False
    records = _all_records(c)
    if records:
        return any(routes_to_review(rtype, rec) for rtype, rec in records)
    return bool(c.get("needs_human_review"))


def _record_source_id(filename):
    """The STABLE source-email identity for a result filename — the original email
    (carried for a reprocessed split file via ORIGINAL_EMAIL) else the filename."""
    return ORIGINAL_EMAIL.get(filename) or filename


def _base_filename(fn):
    """Strip a trailing '__N' reprocess-churn suffix (before the extension), so
    'X__2.pdf' and 'X.pdf' compare equal. '' for a blank name."""
    if not fn:
        return ""
    p = str(fn)
    stem, dot, ext = p.rpartition(".")
    if not dot:
        stem, ext = p, ""
    stem = re.sub(r"__\d+$", "", stem)
    return stem + (("." + ext) if ext else "")


# (table, number column, whether the table has a duplicate_ignored column). credit_memos
# has none — it's never dedup-ignored, so treat every row as non-ignored.
_DEDUPE_TABLES = [("invoices", "invoice_number", True),
                  ("credit_memos", "credit_memo_number", False),
                  ("payments", "reference_number", True)]


def dedupe_records():
    """One-off maintenance (--dedupe-records, opt-in): remove redundant re-record rows left
    by the churned-filename bug. Within each recorded table, groups NON-ignored rows by
    content-key (vendor + number + amount); when ≥2 in a group share the same base filename
    (email_filename or file_path, ignoring a '__N' suffix — i.e. the same document), keeps
    the LOWEST id and DELETES the rest, printing each. Different-source true duplicates are
    already `duplicate_ignored` (out of totals) so they're skipped, not deleted."""
    if not USE_DB:
        print("--dedupe-records needs DATABASE_URL (DB mode). Nothing to do.")
        return
    import psycopg2.extras
    conn = _database().connect()
    conn.autocommit = False
    removed = 0
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            for table, numcol, has_ignored in _DEDUPE_TABLES:
                ign = "duplicate_ignored" if has_ignored else "FALSE AS duplicate_ignored"
                cur.execute(f"SELECT id, email_filename, file_path, vendor_matched, {numcol} AS num, "
                            f"amount, {ign} FROM {table} ORDER BY id")
                groups = {}
                for r in cur.fetchall():
                    if r.get("duplicate_ignored"):
                        continue                       # already out of totals
                    amt = round(float(r["amount"]), 2) if r.get("amount") is not None else None
                    key = (_norm(r.get("vendor_matched")), _norm(r.get("num")), amt)
                    groups.setdefault(key, []).append(r)
                for grp in groups.values():
                    if len(grp) < 2:
                        continue
                    grp.sort(key=lambda x: x["id"])
                    keeper = grp[0]
                    kbe, kbf = _base_filename(keeper.get("email_filename")), _base_filename(keeper.get("file_path"))
                    for r in grp[1:]:
                        same_doc = ((kbe and _base_filename(r.get("email_filename")) == kbe)
                                    or (kbf and _base_filename(r.get("file_path")) == kbf))
                        if not same_doc:
                            continue                   # same content but a different document — leave it
                        cur.execute(f"DELETE FROM {table} WHERE id = %s", (r["id"],))
                        removed += 1
                        print(f"  removed {table} id {r['id']} "
                              f"(vendor={r.get('vendor_matched')}, #={r.get('num')}, "
                              f"amount={r.get('amount')}) — redundant copy of id {keeper['id']}")
        conn.commit()
        print(f"--dedupe-records: removed {removed} redundant re-record row(s).")
    except Exception:
        conn.rollback()
        print("--dedupe-records: ERROR — rolled back, no rows removed.")
        raise
    finally:
        conn.close()


def _supersede_reprocessed_records(results, reprocessed_files, history, payment_history):
    """ROOT FIX for the re-record duplicate: a reprocessed record may have been recorded
    before under a DIFFERENT (churned, e.g. '__2') email_filename, which save_results'
    delete-by-filename can't replace. Find those PRIOR recordings (SAME source-email
    identity + IDENTICAL content, DIFFERENT source_file), drop their ledger entries (so the
    reprocess records cleanly rather than self-matching them), and RETURN one save_results
    replace_records spec per stale record (its old email_filename + vendor + number + amount)
    → exactly ONE recorded row. Scoped to the record, not the filename: the old filename can be
    a whole-email recording that the email's other records share, and deleting the whole
    filename took them too (2026-10-08). Only reprocessed records are considered; a normal
    same-filename reprocess finds no different-file prior and is untouched."""
    superseded = []
    if not reprocessed_files:
        return superseded
    for r in results:
        c = r.get("classification")
        if not c or r.get("filename") not in reprocessed_files:
            continue
        source_id = _record_source_id(r["filename"])
        cur_file = r["filename"]
        for inv in c.get("invoices") or []:
            if inv.get("is_sales_order") or record_flagged(inv):
                continue
            v, n = _norm(inv.get("vendor_matched")), _norm(inv.get("invoice_number"))
            if not v or not n:
                continue
            keep = []
            for h in history:
                pid = h.get("source_email_id") or h.get("source_file")
                if (pid == source_id and h.get("source_file") != cur_file
                        and _norm(h.get("vendor_matched")) == v and _norm(h.get("invoice_number")) == n
                        and _amounts_equal(h.get("invoice_amount"), inv.get("invoice_amount"))):
                    superseded.append({"filename": h.get("source_file"), "record_type": "invoice",
                                       "vendor": h.get("vendor_matched"), "number": h.get("invoice_number"),
                                       "amount": h.get("invoice_amount")})   # the stale row
                else:
                    keep.append(h)
            history[:] = keep
        for pay in c.get("payments") or []:
            if record_flagged(pay):
                continue
            v = _norm(pay.get("vendor_matched"))
            if not v:
                continue
            ref = _norm(pay.get("reference_number"))
            keep = []
            for h in payment_history:
                pid = h.get("source_email_id") or h.get("source_file")
                if (pid == source_id and h.get("source_file") != cur_file
                        and _norm(h.get("vendor_matched")) == v and _norm(h.get("reference_number")) == ref
                        and _amounts_equal(h.get("amount"), pay.get("payment_amount"))):
                    superseded.append({"filename": h.get("source_file"), "record_type": "payment",
                                       "vendor": h.get("vendor_matched"), "number": h.get("reference_number"),
                                       "amount": h.get("amount")})
                else:
                    keep.append(h)
            payment_history[:] = keep
    return [s for s in superseded if s["filename"]]


def _replaced_by(row, replace_records):
    """True when a results row is one of the records a save_results replace_records spec names — the
    Excel-mode mirror of database._scoped_where (same filename, type, vendor, number, amount)."""
    for s in replace_records or ():
        if (row.get("Filename") == s.get("filename") and row.get("Record Type") == s.get("record_type")
                and _norm_vendor_dedupe(row.get("Vendor Matched")) == _norm_vendor_dedupe(s.get("vendor"))
                and _norm(str(row.get("Invoice #") or "")) == _norm(str(s.get("number") or ""))
                and (s.get("amount") in (None, "") or _amounts_equal(row.get("Amount"), s.get("amount")))):
            return True
    return False


def _result_from_sales_order_snapshot(snap):
    """Rebuild a single-invoice `result` (for the normal save path) from a held order's
    snapshot, so releasing it records the invoice exactly as it would have been."""
    inv = dict(snap.get("invoice") or {})
    inv["is_sales_order"] = False            # record normally now
    fn = snap.get("filename") or ""
    if snap.get("gmail_link"):
        GMAIL_LINKS[fn] = snap["gmail_link"]
    if snap.get("source_email"):
        ORIGINAL_EMAIL[fn] = snap["source_email"]
    return {"filename": fn, "classification": {
        "email_type": snap.get("email_type"), "confidence": snap.get("confidence"),
        "needs_human_review": bool(inv.get("needs_human_review")), "review_reason": None,
        "route_to": snap.get("route_to"), "routing_reason": snap.get("routing_reason"),
        "invoices": [inv], "credit_memos": [], "payments": []}}


def _flag_orders_already_recorded(results, history):
    """Before a held sales order is released: when its OWN email already has a recorded invoice from
    the same vendor for the same amount under a DIFFERENT number, it's almost certainly this order
    read twice (once by its order number, once by an invoice number). Recording it would double it,
    so it goes to review with a note instead. (Releasing used to delete the email's other records,
    which happened to remove that other copy — see save_results' replace_records.) Returns how many
    were flagged."""
    flagged = 0
    for r in results:
        source_id = _record_source_id(r["filename"])
        for inv in (r.get("classification") or {}).get("invoices") or []:
            vendor = _dedupe_vendor_key(inv.get("vendor_matched"))
            if not vendor or inv.get("invoice_amount") is None:
                continue
            for h in history:
                if ((h.get("source_email_id") or h.get("source_file")) == source_id
                        and _dedupe_vendor_key(h.get("vendor_matched")) == vendor
                        and _norm(h.get("invoice_number")) != _norm(inv.get("invoice_number"))
                        and _amounts_equal(h.get("invoice_amount"), inv.get("invoice_amount"))):
                    inv["needs_human_review"] = True
                    _append_text(inv, "review_reason", (
                        f"This sales order's email already has invoice #{h.get('invoice_number')} from "
                        f"{inv.get('vendor_matched')} for the same amount — probably this same order read "
                        f"twice. If so, Ignore this item; if it's a separate order, answer to record it."))
                    flagged += 1
                    break
    return flagged


def apply_released_order_check(classification):
    """Record an order whose hold already ran out instead of holding it AGAIN. A held order that's
    released goes to review when it needs a person; answering it reprocesses a document that still
    reads as a sales order, so it used to be held once more and come back after another hold window
    (the answered item cycled). When a RELEASED order with the same vendor + number exists, clear
    is_sales_order so the record goes through the normal gates and is recorded (or reviewed) as an
    invoice, with a note saying why. Superseded orders are left alone — their invoice is recorded, so
    a re-hold is superseded at once. Runs before apply_review_gates. DB-only."""
    if not USE_DB or not classification:
        return
    for inv in classification.get("invoices") or []:
        if not inv.get("is_sales_order"):
            continue
        try:
            rel = _database().find_released_sales_order(inv.get("vendor_matched"), inv.get("invoice_number"))
        except Exception as e:
            print(f"  (could not check for a released sales order: {e})")
            continue
        if not rel:
            continue
        inv["is_sales_order"] = False
        held = rel.get("held_at")
        _append_text(inv, "notes", (
            f"[Sales order #{inv.get('invoice_number')} was already released from its hold"
            + (f" (held {held:%m/%d/%Y})" if held else "")
            + " — recorded as an invoice instead of being held again.]"))
        print(f"  {inv.get('vendor_matched')} #{inv.get('invoice_number')}: already released from a "
              f"sales-order hold — treated as an invoice, not held again.")


def reconcile_sales_orders(run_id):
    """After this run's records are saved: for each HELD sales order, SUPERSEDE it if a
    real invoice (same vendor + normalized number) is now recorded; else RELEASE it (record
    it via the normal save path + queue it if it needs review) once the hold window elapsed;
    else keep it held. DB-only; a no-op with no held orders."""
    if not USE_DB:
        return
    db = _database()
    held = db.load_held_sales_orders()
    if not held:
        return
    now = datetime.now(timezone.utc)
    superseded_n = 0
    release = []
    for o in held:
        inv_ref = db.find_recorded_invoice_ref(o.get("vendor_matched"), o.get("number"))
        if inv_ref:
            db.set_sales_order_status(o["id"], "superseded", matched_invoice=inv_ref)
            superseded_n += 1
            continue
        ra = o.get("release_after")
        if ra is not None and now >= ra:
            release.append(o)

    released_review_rows = []
    # Someone acted on a record in an order's email in the last few minutes (save_results'
    # backstop would skip recording it, yet the order would still be marked released and lost):
    # leave that order held for the next run instead.
    busy = db.recently_human_touched([(o.get("snapshot") or {}).get("filename") for o in release])
    if busy:
        release = [o for o in release if (o.get("snapshot") or {}).get("filename") not in busy]
        print(f"Sales orders: kept {len(busy)} email(s)' order(s) held one more run — someone is "
              f"working on that email right now.")
    if release:
        results = [_result_from_sales_order_snapshot(o.get("snapshot") or {}) for o in release]
        # An order held during a run without Drive upload carries a local file path in its snapshot;
        # upload it now (a Drive link in the snapshot isn't a local file, so it's left as is).
        for res in results:
            upload_record_files(res["classification"])
        history = load_history()
        payment_history = load_payment_history()
        _flag_orders_already_recorded(results, history)
        apply_duplicate_checks(results, history, payment_history)
        new_rows, review_dropped = build_result_rows(results)
        # A late gate here would otherwise leave a history-only orphan; strip those ledger rows
        # (the dropped records are still surfaced below via result_has_review -> review rows).
        route_dropped_to_review(review_dropped, results, history, payment_history)
        # Replace only the released order itself under its filename: that filename is usually the
        # order's original email, whose other records must stay (deleting the whole filename here
        # used to take them too — 2026-10-08).
        replace = [{"filename": r["filename"], "record_type": "invoice",
                    "vendor": inv.get("vendor_matched"), "number": inv.get("invoice_number"),
                    "amount": inv.get("invoice_amount")}
                   for r in results for inv in (r["classification"].get("invoices") or [])]
        _protected, recorded = db.save_results(new_rows, set(), run_id, replace_records=replace)
        # Released order records are freshly recorded too — let them resolve any waiting QB
        # export rows (same local-import rationale as the main save_results path).
        if recorded:
            import qb_match
            n = qb_match.resolve_unmatched_for_new_records(recorded, "auto-resolved on arrival")
            if n:
                print(f"  QB-UNMATCHED: {n} newly-recorded record(s) matched a previously-uploaded "
                      f"QB export row that had nothing to match at the time.")
        save_history(history)
        save_payment_history(payment_history)
        for res in results:
            if result_has_review(res):
                released_review_rows.extend(build_review_rows(res))
        for row in released_review_rows:
            row["Run ID"] = run_id
        if released_review_rows:
            db.append_review_rows(released_review_rows, run_id)
        for o in release:
            db.set_sales_order_status(o["id"], "released")

    if superseded_n or release:
        print(f"Sales orders: superseded {superseded_n} (an invoice arrived), "
              f"released {len(release)} (hold elapsed, recorded now).")


def _group_id_for(prior, record_type):
    """Return the shared group id for `prior`, minting one if it has none. Reusing an
    existing id propagates to a 3rd+ match automatically (the same `prior` dict is
    mutated in `history`/`payment_history`, later persisted whole by save_history).

    The ORIGINAL record only ever gets its group tag via the retroactive live-row stamp
    below (it had no prior when first recorded, so nothing set the tag on it directly).
    That stamp runs on BOTH mint AND reuse, and self-heals: it re-stamps whenever the
    prior's live DB row is missing/differs from the gid. (Reuse used to skip it, so a
    single mint-time miss — e.g. find_live_invoice couldn't resolve the original's row
    then because a churned-filename reprocess had it under a different email_filename —
    stayed permanently unstamped; every later match short-circuited on the existing gid.)
    None when there's no prior."""
    if prior is None:
        return None
    gid = prior.get("duplicate_group_id")
    if not gid:
        import secrets
        gid = secrets.token_hex(4).upper()   # e.g. "A1B2C3D4"
        prior["duplicate_group_id"] = gid
    if USE_DB:
        key = prior.get("invoice_number") if record_type == "invoice" else prior.get("reference_number")
        finder = (_database().find_live_invoice if record_type == "invoice"
                  else _database().find_live_payment)
        live = finder(prior.get("source_file"), prior.get("vendor_matched"), key)
        # Stamp only when the live row lacks (or disagrees with) the gid — avoids a
        # redundant UPDATE on every reuse. '' (the insert-time default) counts as unset.
        if live and (live.get("duplicate_group_id") or "") != gid:
            _database().set_group_id(record_type, live["id"], gid)
    return gid


def _load_waiting_reminders():
    """{(vendor, number): [reminder record]} for the live payment-reminder records still waiting for
    their invoice (database.list_open_reminder_records) — one query per duplicate-check pass. {} in
    Excel mode or if the read fails."""
    if not USE_DB:
        return {}
    try:
        rows = _database().list_open_reminder_records()
    except Exception as e:
        print(f"  (could not read waiting payment reminders: {e})")
        return {}
    out = {}
    for r in rows:
        key = (_dedupe_vendor_key(r.get("vendor_matched")), _norm(r.get("invoice_number")))
        if all(key):
            out.setdefault(key, []).append(r)
    return out


def _resolve_waiting_reminders(inv, filename, waiting):
    """The invoice a payment reminder was about has arrived. A reminder that came FIRST is recorded
    live with a 'possible missed invoice' note and kept out of the ledger (apply_duplicate_checks), so
    nothing matched the real invoice to it when it arrived and both stayed live — a double entry
    waiting to happen (MasterKo #35045, found 2026-10-08). Now the waiting reminder is auto-ignored as
    a duplicate of this invoice; if it was already entered in QuickBooks, this invoice is flagged as a
    possible duplicate instead (never override a completed QB entry). `waiting` is
    _load_waiting_reminders(); a resolved reminder is removed from it. Never raises."""
    key = (_dedupe_vendor_key(inv.get("vendor_matched")), _norm(inv.get("invoice_number")))
    for rem in list(waiting.get(key) or []):
        if rem.get("email_filename") == filename:
            continue                                  # this same email, re-read — not a second copy
        label = f"invoice #{inv.get('invoice_number')} from {inv.get('vendor_matched')}"
        if rem.get("entered_in_qb"):
            inv["is_duplicate_suspect"] = True
            inv["duplicate_reason"] = (f"A payment reminder for {label} was recorded earlier from "
                                       f"{rem.get('email_filename')} and is already entered in QuickBooks — "
                                       f"check QuickBooks before entering this invoice.")
            print(f"  {label}: a payment reminder for it is already entered in QuickBooks -> flagged.")
            continue
        try:
            _database().supersede_duplicate_invoice(
                rem["id"], filename,
                f"Payment reminder for {label}; the invoice itself has now arrived (from {filename}) — "
                f"auto-ignored so it isn't entered twice.", rem.get("duplicate_group_id"))
            waiting[key].remove(rem)
            print(f"  {label}: the invoice for an earlier payment reminder arrived -> ignored the reminder "
                  f"(id={rem['id']}).")
        except Exception as e:
            print(f"  (could not resolve the earlier payment reminder id={rem.get('id')}: {e})")


def _flag_ignored_prior(inv, dup_of, gid):
    """store_conflict / store_preferred when the matched earlier copy is ALREADY ignored: a person (or
    an earlier run) kept a different copy of this invoice. Don't "supersede" the ignored one again and
    keep this copy clean beside the kept one (Incento #12598, 2026-10-11) — flag it for a person."""
    inv["is_duplicate_suspect"] = True
    inv["duplicate_reason"] = (
        f"Invoice #{inv.get('invoice_number')} from {inv.get('vendor_matched')} for "
        f"${inv.get('invoice_amount')} matches an earlier copy from {dup_of}, which was already ignored "
        f"as a duplicate of another copy — so this invoice is probably on file already. Check which copy "
        f"to keep before entering this one.")
    inv["duplicate_ignored"] = False
    inv["duplicate_of"] = dup_of
    inv["duplicate_group_id"] = gid
    print(f"  {inv.get('vendor_matched')} #{inv.get('invoice_number')}: the earlier copy it matched was "
          f"already ignored -> flagged this copy for a person.")


def apply_duplicate_checks(results, history, payment_history=None):
    """Runs on the records that will land in results. Flagged records are skipped
    (held in the review queue). Sets per-record duplicate flags from the three-way
    check_duplicate (invoices) / check_duplicate_payment (payments), keyed by
    source-email identity so a record never duplicates itself. REAL duplicates are
    auto-ignored and are NOT added to history (so they can't later make the ORIGINAL
    look like a duplicate)."""
    if payment_history is None:
        payment_history = []
    # Sender per source email (filename -> from_address), for check_duplicate's store-vs-vendor
    # preference. DB-mode only; empty otherwise -> the preference simply never fires.
    sender_map = {}
    if USE_DB:
        try:
            sender_map = _database().email_from_addresses()
        except Exception as e:
            print(f"(could not load email senders for store-preference: {e})")
    # Payment reminders recorded before their invoice arrived (resolved when it does).
    waiting_reminders = _load_waiting_reminders()
    for r in results:
        c = r.get("classification")
        if not c:
            continue
        source_id = _record_source_id(r["filename"])
        email_from = sender_map.get(r["filename"])
        for inv in c.get("invoices") or []:
            if inv.get("is_sales_order"):
                continue                      # held, not recorded/dup-checked while pending
            if record_flagged(inv):
                continue
            inv["from_address"] = email_from   # so check_duplicate can resolve this copy's sender
            status, reason, dup_of, prior = check_duplicate(inv, history, source_id, r["filename"])
            if inv.get("is_payment_reminder"):
                # This "invoice" is really a payment-reminder / past-due notice about an EARLIER
                # invoice, not a fresh bill. Two outcomes, both deterministic here (the model only
                # identified it as a reminder; it never sees our records):
                #   - the referenced invoice IS on file (any prior match) -> auto-resolve: mark
                #     ignored as a reminder, no review, and don't add it to history as a new bill.
                #   - NO matching invoice on file (status 'none') -> do NOT auto-resolve; leave it
                #     as a live record with a clear note so a human sees it (possible missed invoice).
                if status != "none":
                    inv["is_duplicate_suspect"] = False
                    inv["duplicate_ignored"] = True
                    inv["duplicate_of"] = dup_of
                    inv["duplicate_group_id"] = _group_id_for(prior, "invoice")
                    inv["duplicate_reason"] = (
                        f"Payment reminder for invoice #{inv.get('invoice_number')} from "
                        f"{inv.get('vendor_matched')}, which is already on file"
                        + (f" ({dup_of})" if dup_of else "")
                        + " — auto-resolved as a reminder; no action needed.")
                    inv["needs_human_review"] = False
                    inv["review_reason"] = None
                    print(f"  payment reminder for {inv.get('vendor_matched')} "
                          f"#{inv.get('invoice_number')}: invoice already on file -> auto-resolved.")
                else:
                    inv["is_duplicate_suspect"] = False
                    inv["duplicate_ignored"] = False
                    inv["duplicate_of"] = None
                    inv["duplicate_group_id"] = None
                    inv["duplicate_reason"] = None
                    note = (
                        f"Payment reminder for invoice #{inv.get('invoice_number')} from "
                        f"{inv.get('vendor_matched')}, but NO matching invoice is on file — "
                        f"possible missed invoice; verify whether it was ever received/entered.")
                    inv["notes"] = ((inv.get("notes") + " ") if inv.get("notes") else "") + note
                    print(f"  payment reminder for {inv.get('vendor_matched')} "
                          f"#{inv.get('invoice_number')}: NO invoice on file -> surfaced "
                          f"(possible missed invoice).")
                continue   # a reminder is never added to invoice history as a bill of its own
            if status == "store_conflict":
                # Same invoice + amount previously recorded to a DIFFERENT store. Until 2026-10-12 the
                # LATER copy always won and the earlier one was auto-ignored, meant for a copy Joel
                # forwards with a reclassify instruction. But none of the 3 real cases were that: two
                # were store scans whose read store disagreed with the vendor's own email (Garden of
                # Life #52080970 kept 3M from REFUAH's scanner), and one copy with no store at all
                # superseded Maple (Lantev #9GH5739). Now the later copy wins on its own only when a
                # reviewer's answer produced it (a person decided the store). Otherwise both copies are
                # marked possible duplicates -- off QB entry, side by side on /duplicates -- and a
                # person keeps the right one. Never overrides a copy already entered in QuickBooks.
                gid = _group_id_for(prior, "invoice")   # tie both copies into one group
                # Look the prior up by ITS OWN spelling + number: the two copies can be spelled
                # differently (matched by standard name), and the DB lookup can't always resolve that.
                prior_live = USE_DB and _database().find_live_invoice(
                    dup_of, (prior or {}).get("vendor_matched") or inv.get("vendor_matched"),
                    (prior or {}).get("invoice_number") or inv.get("invoice_number"))
                if prior_live and prior_live.get("duplicate_ignored"):
                    _flag_ignored_prior(inv, dup_of, gid)   # someone already kept another copy
                    append_to_history = False
                elif prior_live and not prior_live.get("entered_in_qb") and not c.get("_reviewer_answered"):
                    # The copies disagree about the store and no person has said which is right.
                    pick = (f" Two copies of invoice #{inv.get('invoice_number')} from "
                            f"{inv.get('vendor_matched')} name different stores ({(prior or {}).get('store_matched')} "
                            f"and {inv.get('store_matched')}) — open Possible duplicates and Keep this one on "
                            f"the copy with the right store; the other is then ignored.")
                    _database().flag_store_conflict_prior(
                        prior_live["id"],
                        f"A later copy ({r['filename']}) has this invoice under {inv.get('store_matched')}." + pick,
                        gid)
                    inv["is_duplicate_suspect"] = True
                    inv["duplicate_reason"] = reason + pick
                    inv["duplicate_ignored"] = False
                    inv["duplicate_of"] = dup_of
                    inv["duplicate_group_id"] = gid
                    append_to_history = False   # earlier copy stays in the ledger; a person picks
                    print(f"  store conflict on {inv.get('vendor_matched')} #{inv.get('invoice_number')}: "
                          f"{(prior or {}).get('store_matched')} (id={prior_live['id']}) vs "
                          f"{inv.get('store_matched')} -> both marked possible duplicates for a person.")
                elif prior_live and not prior_live.get("entered_in_qb"):
                    _database().supersede_duplicate_invoice(
                        prior_live["id"], r["filename"],
                        f"Superseded by a later reclassification to {inv.get('store_matched')} "
                        f"(from {r['filename']}, a reviewer's answer).", gid)
                    print(f"  store conflict on {inv.get('vendor_matched')} "
                          f"#{inv.get('invoice_number')}: auto-ignored the EARLIER copy "
                          f"(recorded to a different store; now superseded, id={prior_live['id']}) "
                          f"in favor of this one -> {inv.get('store_matched')}. Any saved file for "
                          f"the earlier copy in invoice_files/ is now stale (not auto-deleted).")
                    # This copy survives clean — clear all dup flags, treat like a fresh record.
                    inv["is_duplicate_suspect"] = False
                    inv["duplicate_reason"] = None
                    inv["duplicate_ignored"] = False
                    inv["duplicate_of"] = None
                    inv["duplicate_group_id"] = gid
                    append_to_history = True   # so a THIRD copy later still dedupes vs this survivor
                else:
                    # Already in QB (or DB unavailable / prior not found): do NOT auto-override.
                    # Flag THIS copy as a possible duplicate for a human to resolve by hand.
                    inv["is_duplicate_suspect"] = True
                    inv["duplicate_reason"] = (
                        reason + " The earlier copy is already entered in QuickBooks — verify "
                        "manually before entering this one." if prior_live else reason)
                    inv["duplicate_ignored"] = False
                    inv["duplicate_of"] = dup_of
                    inv["duplicate_group_id"] = gid
                    append_to_history = False   # earlier copy stays live; don't add a competing row
            elif status == "store_preferred":
                # Exact duplicate to the SAME store where THIS copy is the store scan and the
                # earlier one is the vendor's send: keep THIS, auto-ignore the earlier vendor copy
                # (store-preferred override of first-arrived). Same guardrails as store_conflict —
                # never override a prior already entered in QuickBooks.
                gid = _group_id_for(prior, "invoice")
                prior_live = USE_DB and _database().find_live_invoice(   # by the prior's own spelling, as above
                    dup_of, (prior or {}).get("vendor_matched") or inv.get("vendor_matched"),
                    (prior or {}).get("invoice_number") or inv.get("invoice_number"))
                if prior_live and prior_live.get("duplicate_ignored"):
                    _flag_ignored_prior(inv, dup_of, gid)   # someone already kept another copy
                    append_to_history = False
                elif prior_live and not prior_live.get("entered_in_qb"):
                    _database().supersede_duplicate_invoice(prior_live["id"], r["filename"], reason, gid)
                    print(f"  store-preferred on {inv.get('vendor_matched')} "
                          f"#{inv.get('invoice_number')}: kept THIS store-scanned copy, auto-ignored "
                          f"the earlier vendor copy (id={prior_live['id']}).")
                    inv["is_duplicate_suspect"] = False
                    inv["duplicate_reason"] = None
                    inv["duplicate_ignored"] = False
                    inv["duplicate_of"] = None
                    inv["duplicate_group_id"] = gid
                    append_to_history = True
                else:
                    inv["is_duplicate_suspect"] = True
                    inv["duplicate_reason"] = (
                        reason + " The earlier copy is already entered in QuickBooks — verify "
                        "manually before entering this one." if prior_live else reason)
                    inv["duplicate_ignored"] = False
                    inv["duplicate_of"] = dup_of
                    inv["duplicate_group_id"] = gid
                    append_to_history = False
            else:
                inv["is_duplicate_suspect"] = status in ("duplicate", "possible")
                inv["duplicate_reason"] = reason
                inv["duplicate_ignored"] = (status == "duplicate")
                inv["duplicate_of"] = dup_of if status == "duplicate" else None
                # A matched cluster (real dup or possible) shares a group id, stamped on the
                # original too — independent of the suspect/ignored flags. None otherwise.
                inv["duplicate_group_id"] = (_group_id_for(prior, "invoice")
                                             if status in ("duplicate", "possible") else None)
                # Don't add a real duplicate (or a redundant self-match) to history.
                append_to_history = status in ("none", "possible")
                if status == "none":
                    _resolve_waiting_reminders(inv, r["filename"], waiting_reminders)
            if append_to_history:
                history.append({
                    "vendor_matched": inv.get("vendor_matched"),
                    "store_matched": inv.get("store_matched"),
                    "invoice_number": inv.get("invoice_number"),
                    "invoice_date": inv.get("invoice_date"),
                    "invoice_amount": inv.get("invoice_amount"),
                    "source_file": r["filename"],
                    "source_email_id": source_id,
                    "date_added": TODAY,
                    "duplicate_group_id": inv.get("duplicate_group_id"),
                    "not_received_note": inv.get("not_received_note"),
                    "not_received_amount": inv.get("not_received_amount"),
                    "from_address": inv.get("from_address"),   # carry sender for store-preference
                })
        for pay in c.get("payments") or []:
            if record_flagged(pay):
                continue
            # The applied-invoice guard in check_duplicate_payment keys on the structured
            # invoice_number; if the model put it only in free-text notes (as happened with 4 MMS
            # rows sharing a reference #), backfill it from notes so two payments on the same ref
            # settling DIFFERENT invoices aren't mistaken for revisions of each other. Same dict is
            # later read by build_result_rows, so this also persists into the invoice_number column.
            if not pay.get("invoice_number"):
                pay["invoice_number"] = _applied_invoice_from_notes(pay.get("notes"))
            status, reason, dup_of, prior = check_duplicate_payment(pay, payment_history, source_id, r["filename"])
            pay["is_duplicate_suspect"] = status in ("duplicate", "possible")
            pay["duplicate_reason"] = reason
            pay["duplicate_ignored"] = (status == "duplicate")
            pay["duplicate_of"] = dup_of if status == "duplicate" else None
            pay["duplicate_group_id"] = (_group_id_for(prior, "payment")
                                         if status in ("duplicate", "possible") else None)
            if status in ("none", "possible"):
                payment_history.append({
                    "vendor_matched": pay.get("vendor_matched"),
                    "store_matched": pay.get("store_matched"),
                    "reference_number": pay.get("reference_number"),
                    "invoice_number": pay.get("invoice_number"),   # disambiguates a shared ref
                    "payment_date": pay.get("payment_date"),
                    "amount": pay.get("payment_amount"),
                    "source_file": r["filename"],
                    "source_email_id": source_id,
                    "date_added": TODAY,
                    "duplicate_group_id": pay.get("duplicate_group_id"),
                })


def resolve_store_preferred_duplicates(actor, dry_run=False):
    """Post-save backstop for the store-preference rule (keep the store-SCANNED copy, auto-ignore the
    vendor copy on an EXACT duplicate). check_duplicate's in-pipeline store_preferred supersede relies
    on find_live_invoice (a DB lookup), so it silently fails when BOTH copies arrive in the SAME run —
    the prior isn't persisted yet. This runs AFTER save, when every copy IS in the DB, and resolves
    those. Conservative — acts ONLY on a cluster that:
      - shares normalized vendor + invoice number + amount (rounded) + store (an exact duplicate);
      - has every member still LIVE: not duplicate_ignored, not entered_in_qb, not skip_qb_import;
      - includes >=1 duplicate_suspect member (an actual 'possible duplicate', per Joel's rule);
      - has EXACTLY ONE member from a self-scan sender (sender_store_map) and >=1 non-self-scan copy.
    Otherwise it SKIPS (two scans, no internal copy, or nothing flagged) — never guesses. Keeps the
    store scan clean and supersedes the vendor copy(ies) (duplicate_ignored, shared group), logged to
    edit_log. Idempotent: once a vendor copy is ignored it drops out of the live set. Returns
    {"resolved": [...], "skipped": [...]}. dry_run reports without writing."""
    if not USE_DB:
        return {"resolved": [], "skipped": []}
    import secrets as _secrets
    import psycopg2.extras as _extras
    smap = load_sender_store_map()            # {normalized sender email: store}
    db = _database()
    conn = db.connect()
    conn.autocommit = False
    resolved, skipped = [], []
    try:
        with conn.cursor(cursor_factory=_extras.RealDictCursor) as cur:
            cur.execute(
                "SELECT i.id, i.vendor_matched, i.store_matched, i.invoice_number, i.amount, "
                "       i.email_filename, i.duplicate_suspect, i.duplicate_group_id, e.from_address "
                "  FROM invoices i LEFT JOIN emails e ON e.filename = i.email_filename "
                " WHERE i.duplicate_ignored IS NOT TRUE AND i.entered_in_qb IS NOT TRUE "
                "   AND i.skip_qb_import IS NOT TRUE "
                "   AND i.invoice_number IS NOT NULL AND i.invoice_number <> '' "
                "   AND i.amount IS NOT NULL")
            clusters = {}
            for r in cur.fetchall():
                # Standard-name vendor key, the same one check_duplicate matches on (2026-10-12; the
                # printed spelling split 'MaxiHealth' / 'Maxi Health' copies into separate clusters).
                key = (_dedupe_vendor_key(r["vendor_matched"] or ""),
                       (r["invoice_number"] or "").strip().lower(),
                       round(float(r["amount"]), 2),
                       _norm(r["store_matched"]))
                clusters.setdefault(key, []).append(r)

            for key, members in clusters.items():
                if len(members) < 2 or not any(m["duplicate_suspect"] for m in members):
                    continue
                scans = [m for m in members if _norm(m.get("from_address")) in smap]
                others = [m for m in members if _norm(m.get("from_address")) not in smap]
                if len(scans) != 1 or not others:
                    skipped.append({"key": key, "n_scans": len(scans), "n_other": len(others)})
                    continue
                keep = scans[0]
                gid = (keep.get("duplicate_group_id")
                       or next((m["duplicate_group_id"] for m in members if m.get("duplicate_group_id")), None)
                       or _secrets.token_hex(4).upper())
                reason = (f"Superseded by the store-scanned copy ({keep['email_filename']}, "
                          f"store-preferred) — auto-ignored. Resolved by the post-save store-preference "
                          f"backstop (the copies arrived together, so the in-run supersede couldn't).")
                resolved.append({"keep_id": keep["id"], "ignored_ids": [v["id"] for v in others],
                                 "vendor": keep.get("vendor_matched"), "number": keep.get("invoice_number"),
                                 "store": keep.get("store_matched")})
                if dry_run:
                    continue
                cur.execute("UPDATE invoices SET duplicate_suspect=FALSE, duplicate_ignored=FALSE, "
                            "duplicate_reason=NULL, duplicate_of=NULL, duplicate_group_id=%s, "
                            "edited_at=now() WHERE id=%s", (gid, keep["id"]))
                for v in others:
                    cur.execute("UPDATE invoices SET duplicate_suspect=TRUE, duplicate_ignored=TRUE, "
                                "duplicate_reason=%s, duplicate_of=%s, duplicate_group_id=%s, "
                                "edited_at=now() WHERE id=%s",
                                (reason, keep["email_filename"], gid, v["id"]))
                    cur.execute(
                        "INSERT INTO edit_log (table_name, record_id, field, old_value, new_value, "
                        "changed_by) VALUES ('invoices',%s,'duplicate_ignored',%s,'true',%s)",
                        (v["id"], str(bool(v.get("duplicate_ignored"))).lower(), actor))
        conn.commit() if not dry_run else conn.rollback()
    finally:
        conn.close()
    return {"resolved": resolved, "skipped": skipped}


def enrich_flagged_with_duplicate_info(results, history):
    """Read-only: for records already flagged for review, check them against
    invoice history too (this is normally skipped for flagged records in
    apply_duplicate_checks) and, if a match is found, prepend a clear 'multiple
    versions exist' notice + the shared duplicate_group_id to the review reason —
    so a human opening the review item sees this immediately instead of having to go
    find the other copy by hand. NEVER auto-resolves the flagged record or touches
    any other record's QB/record state: the match here only informs the person
    answering this flagged item, since it's already flagged for its own reason and
    might not be a real revision at all. (_group_id_for may stamp the shared group id
    on the prior's row so the Group link can show both versions — benign metadata.)"""
    for r in results:
        c = r.get("classification")
        if not c:
            continue
        source_id = _record_source_id(r["filename"])
        for inv in c.get("invoices") or []:
            if inv.get("is_sales_order") or not record_flagged(inv):
                continue
            status, reason, dup_of, prior = check_duplicate(inv, history, source_id, r["filename"])
            if status in ("duplicate", "possible", "store_conflict"):
                gid = _group_id_for(prior, "invoice")
                inv["duplicate_group_id"] = gid
                notice = (f"Multiple versions of invoice #{inv.get('invoice_number')} from "
                          f"{inv.get('vendor_matched')} exist — see the Group link on this row "
                          f"for all of them before answering. ")
                inv["review_reason"] = notice + (inv.get("review_reason") or "")


def _review_number_amount(rtype, rec):
    if rtype == "invoice":
        return rec.get("invoice_number"), rec.get("invoice_amount")
    if rtype == "credit_memo":
        return rec.get("credit_memo_number"), rec.get("credit_amount")
    return rec.get("reference_number"), rec.get("payment_amount")


def build_review_rows(result):
    """One review-queue row PER flagged record. Each row's Filename is that
    record's OWN standalone split file (in review_pending/), so answering it
    reprocesses only that record. For a record-less email flagged at the top
    level, a single row keyed on the email itself.

    Source Email and the Review Key use the ORIGINAL email's identity (ORIGINAL_EMAIL, set when a
    review item's split file is reprocessed), not the split file this result was read from. A
    re-flagged item used to get the split PDF as its "Source Email", so its review key changed
    (internal notes and uploads no longer attached), its next retry couldn't recover the original
    email's From/To/Cc (store hints lost), and its records link / Gmail receive time went blank."""
    c = result.get("classification") or {}
    unit_file = result["filename"]                  # the file this result was read from
    source_email = _record_source_id(unit_file)     # the original email it came from
    records = _all_records(c)
    rows = []

    flagged = [(t, rec) for t, rec in records if routes_to_review(t, rec)]
    if not flagged and not records and c.get("needs_human_review"):
        flagged = [(None, None)]  # record-less email flagged as a whole

    for rtype, rec in flagged:
        row = {k: "" for k in REVIEW_QUEUE_FIELDNAMES}
        if rec is None:
            # Whole record-less email — the .msg itself is the review unit.
            review_path = str((REVIEW_PENDING_DIR / unit_file).resolve())
            row.update({
                "Filename": unit_file,
                "Review Key": _stable_review_key(source_email, c.get("vendor_matched"),
                                                 c.get("reference_number")),
                "Vendor Matched": c.get("vendor_matched") or "",
                "Store Matched": c.get("store_matched") or "",
                "Invoice #": c.get("reference_number") or "",
                "Question (AI review reason)": c.get("review_reason") or "",
                "Duplicate Group": c.get("duplicate_group_id") or "",
            })
        else:
            number, amount = _review_number_amount(rtype, rec)
            review_path = rec.get("_review_file") or ""
            review_name = Path(review_path).name if review_path else unit_file
            row.update({
                "Filename": review_name,
                "Review Key": _stable_review_key(
                    source_email, rec.get("vendor_matched") or c.get("vendor_matched"), number),
                "Vendor Matched": rec.get("vendor_matched") or c.get("vendor_matched") or "",
                "Store Matched": rec.get("store_matched") or c.get("store_matched") or "",
                "Invoice #": number or "",
                "Amount": amount if amount is not None else "",
                "Question (AI review reason)": rec.get("review_reason") or c.get("review_reason") or "",
                "Duplicate Group": rec.get("duplicate_group_id") or "",
            })
        row.update({
            "Source Email": source_email,
            "Date Flagged": TODAY,
            "Email Type": c.get("email_type") or (rtype or ""),
            "File Path": review_path,
            "Gmail Link": GMAIL_LINKS.get(source_email) or GMAIL_LINKS.get(unit_file, ""),
        })
        rows.append(row)
    return rows


def carry_review_assignments(new_rows, old_rows, unit_file):
    """Keep who an item is assigned to when a reprocess re-flags it. A re-flag writes a FRESH queue
    row (build_review_rows), which used to start unassigned. Each new row takes the assignee of the
    old row with the same review key; failing that (the answer changed the vendor or number), the one
    assignee shared by the old row(s) of the file that was reprocessed (`unit_file`). Only fills rows
    with no assignee; never guesses between two different assignees. Returns how many got one."""
    by_key = {}
    for r in old_rows or []:
        who = (r.get("Assigned To") or "").strip()
        if who and r.get("Review Key"):
            by_key.setdefault(r["Review Key"], who)
    parent = {(r.get("Assigned To") or "").strip() for r in old_rows or [] if r.get("Filename") == unit_file}
    fallback = next(iter(parent)) if len(parent) == 1 else ""
    carried = 0
    for row in new_rows:
        if (row.get("Assigned To") or "").strip():
            continue
        who = by_key.get(row.get("Review Key")) or fallback
        if who:
            row["Assigned To"] = who
            carried += 1
    return carried


def carry_review_ids(new_rows, old_rows, unit_file, taken):
    """Give a re-flagged item's fresh row the queue id its old row had, so /review/<id> links and
    --retry-id values keep pointing at the same item (save_review_queue keeps a row's id). The old
    row is the queue row for the file that was reprocessed (`unit_file`); its id goes to the new row
    with the same review key, else to the only new row — never guessed between two. `taken` holds the
    ids other rows already keep (updated in place). Returns how many rows got their old id."""
    carried = 0
    for old in old_rows or []:
        oid = old.get("Id")
        if old.get("Filename") != unit_file or oid in (None, "") or oid in taken:
            continue
        open_rows = [r for r in new_rows if r.get("Id") in (None, "")]
        target = next((r for r in open_rows
                       if old.get("Review Key") and r.get("Review Key") == old.get("Review Key")), None)
        if target is None and len(new_rows) == 1 and open_rows:
            target = open_rows[0]
        if target is not None:
            target["Id"] = oid
            taken.add(oid)
            carried += 1
    return carried


def build_summary(results):
    total = len(results)
    type_counts = {}
    needs_review = 0
    invoice_count = 0
    skip_qb_count = 0
    low_confidence = []

    for r in results:
        c = r.get("classification")
        if not c:
            reason = r.get("error") or "no classification (parse error)"
            low_confidence.append((r["filename"], reason))
            continue

        email_type = c.get("email_type", "unknown")
        type_counts[email_type] = type_counts.get(email_type, 0) + 1

        if result_has_review(r):
            needs_review += 1

        if c.get("confidence") == "low":
            low_confidence.append((r["filename"], "overall confidence: low"))

        for inv in c.get("invoices") or []:
            invoice_count += 1
            if inv.get("skip_qb_import"):
                skip_qb_count += 1
            if inv.get("vendor_confidence") == "low" or inv.get("store_confidence") == "low":
                low_confidence.append((r["filename"], f"invoice confidence low: {inv.get('vendor_name_raw')}"))

    return {
        "total": total,
        "type_counts": type_counts,
        "needs_review": needs_review,
        "invoice_count": invoice_count,
        "skip_qb_count": skip_qb_count,
        "low_confidence": low_confidence,
    }


def print_summary(summary):
    print()
    print("=" * 50)
    print("EMAIL CLASSIFICATION SUMMARY")
    print("=" * 50)
    print(f"Total emails processed: {summary['total']}")
    print()
    print("Breakdown by email_type:")
    for email_type, count in sorted(summary["type_counts"].items()):
        print(f"  {email_type}: {count}")
    print()
    print(f"Emails needing human review: {summary['needs_review']}")
    print()
    print(f"Invoices found: {summary['invoice_count']}")
    print(f"  Flagged skip_qb_import: {summary['skip_qb_count']}")
    print(f"  QB-ready: {summary['invoice_count'] - summary['skip_qb_count']}")
    print()
    if summary["low_confidence"]:
        print("Low-confidence results needing attention:")
        for filename, reason in summary["low_confidence"]:
            print(f"  {filename}: {reason}")
    else:
        print("No low-confidence results.")
    print("=" * 50)


def _yn(value):
    if value is None:
        return ""
    return "Yes" if value else "No"


def load_existing_results():
    """Returns (existing_rows, successfully_processed_filenames). Filenames that
    previously errored are NOT included, so they get retried rather than
    silently skipped forever. In DB mode there is no sheet to preserve, so
    existing_rows is empty and processed filenames come from the emails table."""
    if USE_DB:
        return [], _database().load_processed_filenames()
    if not RESULTS_PATH.exists():
        return [], set()
    wb = load_workbook(RESULTS_PATH, read_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    wb.close()

    failed_filenames = {r["Filename"] for r in rows if r.get("Error")}
    all_filenames = {r["Filename"] for r in rows if r.get("Filename")}
    processed_filenames = all_filenames - failed_filenames
    return rows, processed_filenames


def _total_source(rec):
    """The record's total_source as saved: 'stated' (read from a printed total), 'derived' (summed
    from line items), or '' when the model didn't say. Stored so the amount check's input is visible
    on the record afterwards."""
    v = (rec.get("total_source") or "").strip().lower()
    return v if v in ("stated", "derived") else ""


def build_result_rows(results):
    # review_dropped: records the gates here flip to needs_review AFTER apply_review_gates
    # already ran upstream (process_source_email). These are dropped from results.xlsx below;
    # the caller must still give each one a review row and keep it out of invoice_history, or
    # it becomes a history-only orphan with a frozen email status (the /reconcile backlog bug).
    rows = []
    review_dropped = []
    for r in results:
        c = r.get("classification")
        base = {k: "" for k in RESULTS_FIELDNAMES}
        base["Filename"] = r["filename"]
        base["Date Added"] = TODAY
        base["Gmail Link"] = GMAIL_LINKS.get(r["filename"], "")
        base["Error"] = r.get("error") or ""
        base["Raw Response"] = (r.get("raw_response") or "")[:500]
        base["Skipped Attachments"] = ", ".join(r.get("skipped_attachments") or [])

        if not c:
            rows.append(base)
            continue

        base["Email Type"] = c.get("email_type") or ""
        base["Confidence"] = c.get("confidence") or ""
        base["Needs Review"] = _yn(c.get("needs_human_review"))
        base["Review Reason"] = c.get("review_reason") or ""
        base["Route To"] = c.get("route_to") or ""
        base["Routing Reason"] = c.get("routing_reason") or ""

        invoices = c.get("invoices") or []
        credit_memos = c.get("credit_memos") or []
        payments = c.get("payments") or []

        if not invoices and not credit_memos and not payments:
            # Informational email (statement, shipment notice, correspondence,
            # etc.). If flagged at the email level it goes to the review queue,
            # not results; otherwise surface its top-level vendor/store/ref.
            if c.get("needs_human_review"):
                continue
            row = dict(base)
            row.update({
                "Record Type": c.get("email_type") or "",
                "Vendor Matched": c.get("vendor_matched") or "",
                "Store Matched": c.get("store_matched") or "",
                "Invoice #": c.get("reference_number") or "",
            })
            rows.append(row)
            continue

        # Only NON-flagged records land in results.xlsx; flagged ones are routed
        # to the review queue (per record) so they don't hold back the batch.
        for inv in invoices:
            # Card last-4 -> store (safety net) before the review-routing check, so
            # a corrected store also reaches any review row / split-file name.
            if inv.get("is_sales_order"):
                continue                      # HELD in pending_sales_orders, not recorded here
            apply_card_store_map(inv)
            apply_store_review_gate(inv)   # unresolved store -> force review, never silent to QB
            apply_scoped_learning_gate(inv)   # Carepax→BMW leak onto a non-Sageteq vendor -> force review
            apply_date_review_gate(inv, "invoice_date")   # no date -> force review, never silent to QB
            apply_amount_review_gate(inv, "invoice_amount")   # no/derived total -> force review
            apply_class_rules(inv)   # fill default QB class (store now resolved) if none stated
            apply_vendor_exclusion(inv)   # deterministic: excluded vendor -> skip_qb_import True
            apply_unlisted_skip_gate(inv, "invoice", c)   # ...and a skip for an unlisted vendor -> review
            apply_vendor_standardization(inv)   # canonical vendor/account + store-specific QB names
            if routes_to_review("invoice", inv):
                review_dropped.append((r, "invoice", inv))
                continue
            suppressed = review_suppressed_by_exclusion("invoice", inv)
            notes = inv.get("notes") or ""
            if suppressed:
                # Recorded but never queued (excluded vendor → never entered in QB).
                notes = (notes + " " if notes else "") + _EXCLUDED_SKIP_NOTE
            row = dict(base)
            row.update({
                "Record Type": "invoice",
                "File Path": inv.get("_file_path") or "",
                # It's not actually pending review — it was suppressed by exclusion.
                "Needs Review": "No" if suppressed else _yn(inv.get("needs_human_review")),
                "Review Reason": inv.get("review_reason") or "",
                "Vendor Raw": inv.get("vendor_name_raw") or "",
                "Vendor Matched": inv.get("vendor_matched") or "",
                "Vendor Confidence": inv.get("vendor_confidence") or "",
                "Skip QB Import": _yn(inv.get("skip_qb_import")),
                "Store Raw": inv.get("store_raw") or "",
                "Store Matched": inv.get("store_matched") or "",
                "Store Confidence": inv.get("store_confidence") or "",
                "Invoice #": inv.get("invoice_number") or "",
                "Invoice Date": inv.get("invoice_date") or "",
                "Amount": inv.get("invoice_amount") if inv.get("invoice_amount") is not None else "",
                "Ship Charge": inv.get("shipping_charge") if inv.get("shipping_charge") is not None else "",
                "Sales Tax": inv.get("sales_tax") if inv.get("sales_tax") is not None else "",
                "Duplicate Suspect": _yn(inv.get("is_duplicate_suspect")),
                "Duplicate Reason": inv.get("duplicate_reason") or "",
                "Duplicate Ignored": _yn(inv.get("duplicate_ignored")),
                "Duplicate Of": inv.get("duplicate_of") or "",
                "Duplicate Group": inv.get("duplicate_group_id") or "",
                "Notes": notes,
                "Card Last4": inv.get("card_last4") or "",
                "Class": inv.get("class") or "",
                "Total Source": _total_source(inv),
                "Standard Vendor": inv.get("standard_vendor") or "",
                "Standard Account": inv.get("standard_account") or "",
                "QB Vendor": inv.get("qb_vendor") or "",
                "QB Account": inv.get("qb_account") or "",
            })
            rows.append(row)

        for cm in credit_memos:
            apply_card_store_map(cm)
            apply_store_review_gate(cm)   # unresolved store -> force review, never silent to QB
            apply_scoped_learning_gate(cm)   # Carepax→BMW leak onto a non-Sageteq vendor -> force review
            apply_date_review_gate(cm, "credit_memo_date")   # no date -> force review, never silent to QB
            apply_amount_review_gate(cm, "credit_amount")   # no/derived total -> force review
            apply_class_rules(cm)
            apply_vendor_exclusion(cm)   # deterministic: excluded vendor -> skip_qb_import True
            apply_unlisted_skip_gate(cm, "credit_memo", c)
            apply_vendor_standardization(cm)   # canonical vendor/account + store-specific QB names
            if routes_to_review("credit_memo", cm):
                review_dropped.append((r, "credit_memo", cm))
                continue
            suppressed = review_suppressed_by_exclusion("credit_memo", cm)
            notes = cm.get("notes") or ""
            if suppressed:
                notes = (notes + " " if notes else "") + _EXCLUDED_SKIP_NOTE
            row = dict(base)
            row.update({
                "Record Type": "credit_memo",
                "File Path": cm.get("_file_path") or "",
                "Needs Review": "No" if suppressed else _yn(cm.get("needs_human_review")),
                "Review Reason": cm.get("review_reason") or "",
                "Vendor Raw": cm.get("vendor_name_raw") or "",
                "Vendor Matched": cm.get("vendor_matched") or "",
                "Vendor Confidence": cm.get("vendor_confidence") or "",
                "Skip QB Import": _yn(cm.get("skip_qb_import")),
                "Store Raw": cm.get("store_raw") or "",
                "Store Matched": cm.get("store_matched") or "",
                "Store Confidence": cm.get("store_confidence") or "",
                "Invoice #": cm.get("credit_memo_number") or "",
                "Invoice Date": cm.get("credit_memo_date") or "",
                "Amount": cm.get("credit_amount") if cm.get("credit_amount") is not None else "",
                "Notes": notes,
                "Card Last4": cm.get("card_last4") or "",
                "Class": cm.get("class") or "",
                "Total Source": _total_source(cm),
                "Standard Vendor": cm.get("standard_vendor") or "",
                "Standard Account": cm.get("standard_account") or "",
                "QB Vendor": cm.get("qb_vendor") or "",
                "QB Account": cm.get("qb_account") or "",
            })
            rows.append(row)

        for pay in payments:
            # Card last-4 -> store (safety net), same helper as invoices/credit memos.
            apply_card_store_map(pay)
            apply_store_review_gate(pay, is_payment=True)   # payments always gated (all vendors enter QB)
            apply_date_review_gate(pay, "payment_date")   # no date -> force review, never silent to QB
            apply_amount_review_gate(pay, "payment_amount")   # no/derived total -> force review
            apply_class_rules(pay)
            # Payments are NEVER suppressed by vendor exclusion — a flagged payment
            # still goes to review; a normal one is recorded and entered in QB.
            if routes_to_review("payment", pay):
                review_dropped.append((r, "payment", pay))
                continue
            row = dict(base)
            row.update({
                "Record Type": "payment",
                "File Path": pay.get("_file_path") or "",
                "Needs Review": _yn(pay.get("needs_human_review")),
                "Review Reason": pay.get("review_reason") or "",
                "Vendor Raw": pay.get("vendor_name_raw") or "",
                "Vendor Matched": pay.get("vendor_matched") or "",
                "Skip QB Import": _yn(pay.get("skip_qb_import")),
                "Store Matched": pay.get("store_matched") or "",
                "Invoice #": pay.get("reference_number") or "",
                # The specific invoice this payment settled, when the remittance states it — distinct
                # from reference_number (above), which legitimately repeats across invoices a single
                # payment covers. Disambiguates the payment dedup. (Own key: "Invoice #" already
                # carries the payment's reference_number for payments.)
                "Applied Invoice #": pay.get("invoice_number") or "",
                "Invoice Date": pay.get("payment_date") or "",
                "Amount": pay.get("payment_amount") if pay.get("payment_amount") is not None else "",
                "Duplicate Suspect": _yn(pay.get("is_duplicate_suspect")),
                "Duplicate Reason": pay.get("duplicate_reason") or "",
                "Duplicate Ignored": _yn(pay.get("duplicate_ignored")),
                "Duplicate Of": pay.get("duplicate_of") or "",
                "Duplicate Group": pay.get("duplicate_group_id") or "",
                "Notes": pay.get("notes") or "",
                "Card Last4": pay.get("card_last4") or "",
                "Class": pay.get("class") or "",
                "Total Source": _total_source(pay),
            })
            rows.append(row)

    return rows, review_dropped


def _find_source_file(filename):
    """Where a result's source file sits during a run: still in test_emails/ (a new email — moves are
    deferred to the end of the run), review_pending/ (a split file being reprocessed), or
    processed_emails/ (an archived email, e.g. a released sales order's original). None if not found."""
    if not filename:
        return None
    for d in (EMAILS_DIR, REVIEW_PENDING_DIR, PROCESSED_DIR):
        p = d / filename
        if p.exists():
            return p
    return None


def _late_review_document(result, rtype, rec):
    """Give a LATE-flagged record (flagged at record time, after save_review_files ran — e.g. a held
    sales order released into review, or any future record-time check) its own review_pending/ split
    file, so its review item has a preview and reprocesses as that one record instead of the whole
    source email. Prefers COPYING the document the record already got (invoice_files/ /
    payment_files/, or the local copy behind a Drive link); otherwise re-derives it from the source
    email with the record's own source_document/source_pages — but only when that document is in the
    email (or the record came from the email body), never a stand-in. Sets rec['_review_file'] and
    returns the path, or '' — never raises, so a failure just leaves the whole-email review row (the
    old behavior). Records save_review_files already handled (_review_attempted) are left alone."""
    if rec.get("_review_file") or rec.get("_review_attempted"):
        return rec.get("_review_file") or ""
    store, vendor, number, amount, date_iso = _record_fields(rtype, rec)
    stem = build_invoice_filename(store, vendor, number, amount, date_iso)
    try:
        REVIEW_PENDING_DIR.mkdir(exist_ok=True)
        for local in (rec.get("_local_file_path"), rec.get("_file_path")):
            if local and not str(local).lower().startswith("http") and os.path.exists(local):
                dest = _unique_path(REVIEW_PENDING_DIR, stem, Path(local).suffix or ".pdf")
                shutil.copy2(local, dest)
                rec["_review_file"] = str(dest.resolve())
                print(f"  late-flagged {vendor} #{number}: copied its saved document into review_pending/")
                return rec["_review_file"]
        src = _find_source_file(result.get("filename"))
        if src is None:
            return ""
        text, attachments, html_body = read_email_content(src)
        cid_images = {a["content_id"]: (a.get("data"), a.get("mimetype"))
                      for a in attachments if a.get("content_id")}
        attachments = expand_zip_attachments(attachments)
        sd = (rec.get("source_document") or "").strip().lower()
        if sd and sd != "email_body" and _match_attachment(sd, attachments) is None:
            return ""      # its document isn't in this email (portal page / manual upload) — don't stand in
        path = _write_record_document(rec, store, vendor, number, amount, date_iso, REVIEW_PENDING_DIR,
                                      attachments, text, html_body, cid_images)
        if path:
            rec["_review_file"] = path
            print(f"  late-flagged {vendor} #{number}: rebuilt its document from {src.name}")
        return path
    except Exception as e:
        print(f"  late-flagged {vendor} #{number}: could not prepare a review document ({e})")
        return ""


def route_dropped_to_review(review_dropped, results, history, payment_history, newly_flagged=None):
    """Make records the late gates dropped in build_result_rows end up in review, not in limbo.

    apply_duplicate_checks runs BEFORE build_result_rows and adds non-flagged records to the
    in-memory history ledger. When a gate inside build_result_rows then flips one to needs_review,
    build_result_rows drops it from results.xlsx (no live invoice) but its history row stays — a
    history-only orphan whose email status never advances (the /reconcile backlog bug). Here we
    (1) strip those just-added ledger rows so the reprocess/record can't be shadowed, (2) make
    sure the owning result is in `newly_flagged` so it gets a review row + a stamped email status, and
    (3) give each genuinely LATE record (one save_review_files never saw as flagged) its own
    review_pending/ document, so its review item has a preview and reprocesses on its own.

    `review_dropped` holds EVERY record build_result_rows routed to review, including ones flagged at
    classification time (those already have their split file). Returns {"late": n, "documents": m}.

    Matching mirrors apply_duplicate_checks' own keys: source_file + the record's own number.
    """
    if not review_dropped:
        return {"late": 0, "documents": 0}
    # Only invoices and payments are ever added to the ledgers (credit memos never are), and
    # apply_duplicate_checks keys an invoice row by invoice_number, a payment row by
    # reference_number. Skip empty numbers so we never match an unrelated blank-number row.
    drop_inv = {(r["filename"], (rec.get("invoice_number") or "").strip().lower())
                for (r, rt, rec) in review_dropped
                if rt == "invoice" and (rec.get("invoice_number") or "").strip()}
    drop_pay = {(r["filename"], (rec.get("reference_number") or "").strip().lower())
                for (r, rt, rec) in review_dropped
                if rt == "payment" and (rec.get("reference_number") or "").strip()}
    if drop_inv and history is not None:
        history[:] = [h for h in history
                      if (h.get("source_file"), (h.get("invoice_number") or "").strip().lower())
                      not in drop_inv]
    if drop_pay and payment_history is not None:
        payment_history[:] = [h for h in payment_history
                              if (h.get("source_file"), (h.get("reference_number") or "").strip().lower())
                              not in drop_pay]
    if newly_flagged is not None:
        for (r, _rt, _rec) in review_dropped:
            if r not in newly_flagged:
                newly_flagged.append(r)
    late = [(r, rt, rec) for (r, rt, rec) in review_dropped if not rec.get("_review_attempted")]
    documents = sum(1 for (r, rt, rec) in late if _late_review_document(r, rt, rec))
    return {"late": len(late), "documents": documents}


def write_results_xlsx(rows):
    _write_xlsx(RESULTS_PATH, RESULTS_FIELDNAMES, rows, "Results")
    return len(rows)


def load_qb_ready_preserved():
    """Returns {(filename, invoice_number): {"Entered in QB": ..., "Date Added": ...}}.
    qb_ready.xlsx is fully regenerated from results on every run, but the
    "Entered in QB" checkmark is filled in manually by hand and "Date Added"
    should reflect when a row first appeared — neither may be recomputed."""
    if not QB_READY_PATH.exists():
        return {}
    wb = load_workbook(QB_READY_PATH, read_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    preserved = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = dict(zip(headers, row))
        key = (row_dict.get("Filename"), row_dict.get("Invoice #"))
        preserved[key] = {
            "Entered in QB": row_dict.get("Entered in QB"),
            "Date Added": row_dict.get("Date Added"),
        }
    wb.close()
    return preserved


def write_qb_ready_xlsx(all_result_rows):
    preserved = load_qb_ready_preserved()
    rows = []
    for row in all_result_rows:
        rtype = row.get("Record Type")
        # Only invoices and credit memos go to QB here (both are "bills");
        # credit memos are entered as NEGATIVE. Skip-QB records are excluded.
        if rtype not in ("invoice", "credit_memo"):
            continue
        if row.get("Skip QB Import") == "Yes":
            continue

        amount = row.get("Amount")
        if rtype == "credit_memo" and amount not in (None, ""):
            try:
                amount = -abs(float(amount))
            except (TypeError, ValueError):
                pass

        possible_dup = "Yes" if row.get("Duplicate Suspect") == "Yes" else "No"
        filename = row.get("Filename") or ""
        key = (filename, row.get("Invoice #"))
        prior = preserved.get(key, {})
        rows.append({
            "Filename": filename,
            "Store": row.get("Store Matched") or "",
            "Vendor": row.get("Vendor Matched") or "",
            "Invoice #": row.get("Invoice #") or "",
            "Invoice Date": row.get("Invoice Date") or "",
            "Amount": amount if amount not in (None, "") else "",
            "Possible Duplicate": possible_dup,
            "Entered in QB": prior.get("Entered in QB") or "",
            "Date Added": prior.get("Date Added") or TODAY,
            "File Path": row.get("File Path") or "",
        })

    _write_xlsx(QB_READY_PATH, QB_READY_FIELDNAMES, rows, "QB Ready")
    return len(rows)


def build_error_review_row(result):
    """A review-queue row for an email the pipeline could not auto-read even
    after retries — so a human opens the saved file and enters it by hand."""
    row = {k: "" for k in REVIEW_QUEUE_FIELDNAMES}
    row.update({
        "Filename": result["filename"],
        "Source Email": result["filename"],
        "Date Flagged": TODAY,
        "Email Type": "unreadable",
        # A max_tokens failure is NOT a missing/corrupt attachment — the AI read the
        # document fine but ran out of output budget mid-response. "Enter it manually"
        # is the wrong instruction there (and prompts a counterproductive re-upload of
        # files that are already being read), so give accurate, specific guidance.
        "Question (AI review reason)": (
            "The AI's response got too long to finish reading this document in one pass "
            "(ran out of output budget) — this usually means it's a large or highly detailed "
            "document (many line items, multiple photographed pages), not that the attachment "
            "is missing or corrupted. Do NOT re-attach the same file — the attachment(s) are "
            "already being read correctly, adding more will only make this worse. Retry this "
            "item as-is (--retry-id), or if it keeps failing, put the invoice details directly "
            "in Answer and rerun."
        ) if "max_tokens" in (result.get("error") or "") else (
            "Automatic reading failed (" + (result.get("error") or "unknown error")
            + "). Open the email in review_pending/ and enter this one manually (or put the "
            "details in Answer and rerun)."
        ),
        "File Path": str((REVIEW_PENDING_DIR / result["filename"]).resolve()),
    })
    return row


def _delete_unused_review_files(paths, queue_rows):
    """Delete superseded review_pending/ files at the end of a run — except any a queue row still
    points at (by Filename or File Path), e.g. a re-flag that came back with no new split file of its
    own (record-less, or its document couldn't be produced) and so keeps using the file it was read
    from. Returns the paths actually deleted."""
    in_use = ({row.get("Filename") for row in queue_rows}
              | {Path(row.get("File Path") or "").name for row in queue_rows})
    deleted = []
    for src in paths:
        if src.name in in_use:
            continue
        try:
            src.unlink(missing_ok=True)
            deleted.append(src)
        except OSError:
            pass
    return deleted


def _safe_move(src, dst):
    """Move a file, but never let one problematic file abort the whole run —
    a missing/locked file logs a warning and is skipped."""
    try:
        shutil.move(str(src), str(dst))
        return True
    except FileNotFoundError:
        print(f"  WARNING: source file not found, skipping move: {src}")
        return False
    except OSError as e:
        print(f"  WARNING: could not move {os.path.basename(str(src))}: {e}")
        return False


def process_source_email(client, path, learnings_text, reviewer_answer="", extra_attachments=None,
                         header_source_path=None):
    """Read → classify (with retry) → save one source document per QB-bound
    record. Shared by the new-email pass and the review-reprocess pass. On a
    classification failure, saves the raw docs as a fallback and marks the
    result 'unreadable' so main routes it to the review queue.

    `extra_attachments` (opt-in; empty for normal processing) is a list of file paths a
    reviewer uploaded to the item — e.g. PDFs hand-extracted from a zip the pipeline
    couldn't open. They're read and APPENDED to the email's own attachments so the model
    sees the real document. The unreadable original stays where it is (still skipped).

    `header_source_path` (opt-in) is the ORIGINAL source email whose From/To/Cc headers to
    recover when `path` is a split PDF/PNG with no headers of its own (a review reprocess) —
    so sender/contact store hints still fire on retry."""
    have_source = False
    html_body = None
    try:
        text, attachments, html_body = read_email_content(path)
        have_source = True
        # A split review file (PDF/PNG) has no headers of its own → text is empty. Recover
        # the From/To/Cc header block from the original source email so _sender_store_hint /
        # _contact_store_hint still work on a retry. Only fills when text is otherwise empty.
        if not text and header_source_path is not None and header_source_path.exists():
            text = _read_headers_only(header_source_path)
    except Exception as e:
        # No readable source. If the reviewer uploaded documents, THOSE are the source of
        # truth — classify from them (below) instead of throwing them away. Only when there
        # are no uploads do we return the read_error (today's behavior).
        if not (extra_attachments or []):
            return {
                "filename": path.name, "classification": None, "raw_response": None,
                "error": f"Failed to read file: {e}", "invoice_file_paths": [], "read_error": True,
            }
        print(f"  original source unreadable ({e}); classifying from uploaded document(s).")
        text = ("[Original source attachment was unreadable/missing; classifying from the "
                "reviewer-uploaded document(s) below.]")
        attachments = []

    # Inline-image map (cid → (bytes, mimetype)) for high-fidelity HTML fallback
    # rendering, built once per email from its own attachments before zip
    # expansion touches them.
    cid_images = {}
    for a in attachments:
        cid = a.get("content_id")
        if cid:
            cid_images[cid] = (a.get("data"), a.get("mimetype"))

    # Expand any ZIP attachments into their contained PDFs/images BEFORE anything else,
    # so the model + the source-document saving both see the real documents. A zip that
    # can't be safely expanded is left as-is (today's behavior).
    attachments = expand_zip_attachments(attachments)

    # Merge reviewer-uploaded manual attachments, deduped by CONTENT HASH against the
    # email's OWN attachments — so re-uploading files the email already has doesn't send
    # the same image bytes to the model twice (which blew past max_tokens on an otherwise
    # normal 4-image request). Hash (not filename) so a renamed re-upload is still caught.
    # This does NOT dedupe two genuinely different manual uploads against each other.
    def _content_hash(data):
        return hashlib.sha256(data).hexdigest()

    existing_hashes = {_content_hash(a["data"]) for a in attachments if a.get("data")}
    for ep in (extra_attachments or []):
        try:
            data = Path(ep).read_bytes()
            h = _content_hash(data)
            if h in existing_hashes:
                print(f"  skipping manual attachment {Path(ep).name} — identical to an "
                      f"existing attachment (already in the email)")
                continue
            attachments.append({"filename": Path(ep).name, "mimetype": None, "data": data})
            existing_hashes.add(h)
            print(f"  + manual attachment: {Path(ep).name}")
        except OSError as e:
            print(f"  could not read manual attachment {ep}: {e}")

    # Link-only invoice email (no attachment) from a VERIFIED portal (LINK_INVOICE_DOMAINS,
    # e.g. BillerGenie): fetch the invoice PAGE itself and render it to a PDF, fed into the
    # normal classification path so the model reads real invoice data — not just the email
    # body. Any failure (no browser, timeout, site changed) leaves attachments empty and the
    # email falls through to today's email-body render; never a hard external dependency.
    if not _has_usable_document(attachments) and html_body:
        link = _find_invoice_link(html_body)
        if link:
            fd, tmp = tempfile.mkstemp(suffix=".pdf")
            os.close(fd)
            tmp = Path(tmp)
            if render_url_to_pdf(link, tmp):
                attachments.append({"filename": "invoice_portal_page.pdf",
                                    "mimetype": "application/pdf", "data": tmp.read_bytes()})
                print(f"  fetched invoice page from {link} -> rendered to PDF for classification")
            tmp.unlink(missing_ok=True)

    # Give the model the AUTHORITATIVE receive time so its invoice-date sanity check
    # doesn't fire on the scanner-clock 'Date:' header. Only when there's a readable source
    # (a missing source has no header to read).
    received = _received_date_str(path) if have_source else ""
    if received:
        text += (
            "\n\n[EMAIL RECEIVED — authoritative server receive time (use THIS as the "
            f"email's date; the 'Date:' header may be wrong due to scanner clock skew): {received}]"
        )

    if attachments:
        print(f"  found {len(attachments)} attachment(s): {', '.join(a['filename'] for a in attachments)}")

    # A lone oversized scanned PDF (a stack of separate invoices) is classified in
    # page-range chunks so it can't blow the output-token budget; everything else
    # (the vast majority of emails) takes the single classify call unchanged. Either
    # way `attachments` stays the ORIGINAL full-PDF list so the downstream save path
    # splits each record's file from the real original bytes.
    _oversized = _lone_oversized_pdf(attachments)
    if _oversized is not None:
        result = _classify_oversized_scan(
            client, path.name, text, _oversized,
            learnings_text=learnings_text, reviewer_answer=reviewer_answer,
        )
    else:
        result = classify_with_retry(
            client, path.name, text, attachments,
            learnings_text=learnings_text, reviewer_answer=reviewer_answer,
        )

    if not result.get("classification"):
        result["invoice_file_paths"] = []
        if result.get("infra_error"):
            # Couldn't REACH the API (connection/network) — transient. Leave the
            # email in test_emails/ (main won't move it) so a later run retries;
            # do NOT dump it in the review queue as "unreadable".
            print(f"  API unreachable ({result.get('error')}) — left for retry")
        else:
            # The model responded but the output was unusable (empty/unparseable)
            # even after retries — a genuine "can't auto-read": route to review.
            result["unreadable"] = True
            print(f"  could not auto-read after retries ({result.get('error')}) — routing to review queue")
    else:
        c = result["classification"]
        # Whether a reviewer's answer is being applied: apply_unlisted_skip_gate then trusts the
        # model's skip_qb_import (it followed the person) instead of sending the record back to review.
        c["_reviewer_answered"] = bool((reviewer_answer or "").strip())
        # An order already released from its hold is an invoice now (don't hold it again).
        apply_released_order_check(c)
        # Make the review-routing decision FINAL here (before file-saving / newly_flagged /
        # duplicate-checks assume a record's fate) — otherwise a gate-flagged record vanishes.
        apply_review_gates(c)
        # A credit memo already recorded from another email goes to review (no ledger covers them).
        apply_credit_memo_duplicate_gate(c, path.name)
        # Clean qb-bound records → invoice_files/ or payment_files/.
        saved = assign_invoice_files(path, text, attachments, c, html_body, cid_images)
        result["invoice_file_paths"] = saved
        # Optionally push those to Google Drive and swap file paths for links.
        upload_record_files(c)
        # Flagged records → their own standalone split file in review_pending/,
        # so each is reviewed/reprocessed on its own.
        review_saved = save_review_files(path, text, attachments, c, html_body, cid_images)
        if saved:
            print(f"  saved {len(saved)} invoice file(s)")
        if review_saved:
            print(f"  saved {len(review_saved)} flagged record(s) to review_pending/")
        if not saved and not review_saved:
            # Nothing saved AND nothing queued. This is legitimate for an
            # informational email or an all-excluded/held batch — but it USED to
            # look identical in the log to a real record silently vanishing. Say
            # explicitly why nothing was produced, and shout if a record fell
            # through the routing without any explanation (the "other" bucket).
            records = _all_records(c)
            if not records:
                print("  no invoice/credit-memo/payment records found (informational/non-invoice email)")
            else:
                def _excluded(rt, rec):
                    return rt in ("invoice", "credit_memo") and is_excluded_vendor(rec.get("vendor_matched"))

                def _held(rt, rec):
                    return rt == "invoice" and rec.get("is_sales_order")

                excluded = sum(1 for rt, rec in records if _excluded(rt, rec))
                held = sum(1 for rt, rec in records if _held(rt, rec))
                # "other" = records explained by NEITHER exclusion NOR a sales-order hold.
                # Counted per-record (not len - excluded - held) because those two buckets
                # OVERLAP — a held sales order from an excluded vendor is in both — so the
                # subtraction could go negative and false-fire the UNEXPECTED anomaly.
                other = sum(1 for rt, rec in records if not _excluded(rt, rec) and not _held(rt, rec))
                parts = []
                if excluded:
                    parts.append(f"{excluded} excluded-vendor")
                if held:
                    parts.append(f"{held} sales-order-held")
                if other:
                    parts.append(f"{other} UNEXPECTED — investigate")
                print(f"  {len(records)} record(s) found, none needed a file/review "
                      f"({', '.join(parts)}) — recorded only")

    return result


def main():
    # One-off maintenance: dedupe redundant re-record rows, then exit (no classify run,
    # no API key needed). Opt-in via --dedupe-records; default runs never reach this.
    if DEDUPE_RECORDS:
        print("=== --dedupe-records (removing redundant re-record rows) ===")
        dedupe_records()
        return

    # .strip() is critical: a leading/trailing space in the key makes an ILLEGAL
    # HTTP header, which the Anthropic SDK reports as a misleading "Connection
    # error" on every call. Also reject an empty/placeholder key up front.
    api_key = (os.environ.get("ANTHROPIC_API_KEY") or "").strip()
    if not api_key:
        print("ERROR: ANTHROPIC_API_KEY is not set.")
        sys.exit(1)
    if "your-key" in api_key or not api_key.startswith("sk-ant-"):
        print(f"ERROR: ANTHROPIC_API_KEY looks invalid (starts with "
              f"'{api_key[:7]}...'). Set a real key (should start with 'sk-ant-').")
        sys.exit(1)

    # Explicit per-request timeout: without one the SDK waits out its 10-min default (×retries)
    # on a stalled request — a silent ~20-30 min freeze. Bound it so a stall fails fast and hits
    # the retry below VISIBLY instead of hanging. 180s read is far above a normal classification
    # (seconds) yet well under the non-streaming server limit. Fewer SDK retries (1) so the app's
    # own 3-attempt retry is the main authority and total time stays bounded.
    _REQ_TIMEOUT = httpx.Timeout(180.0, connect=15.0)
    if CA_BUNDLE.exists():
        client = Anthropic(api_key=api_key, timeout=_REQ_TIMEOUT, max_retries=1,
                           http_client=httpx.Client(verify=str(CA_BUNDLE), timeout=_REQ_TIMEOUT))
    else:
        client = Anthropic(api_key=api_key, timeout=_REQ_TIMEOUT, max_retries=1)

    OUTPUT_DIR.mkdir(exist_ok=True)
    PROCESSED_DIR.mkdir(exist_ok=True)
    REVIEW_PENDING_DIR.mkdir(exist_ok=True)

    print(f"Storage: {'PostgreSQL' if USE_DB else 'Excel'}"
          f"{' (+Excel export)' if USE_DB and WRITE_EXCEL else ''}"
          f"{' (+Drive uploads)' if USE_DRIVE else ''}")
    if USE_DB and not USE_DRIVE:
        print("NOTE: Drive upload is OFF for this run (--no-drive / UPLOAD_TO_DRIVE=0) — saved documents "
              "stay on this PC only. Catch them up later with: python upload_local_files_to_drive.py --commit")
    # One id per classify run, for traceability (which run touched which email /
    # flagged which review item). DB-only; Excel mode is unaffected.
    run_id = datetime.now().strftime("%Y%m%dT%H%M%S")
    if USE_DB:
        _database().create_run(run_id)
        print(f"Run: {run_id}")

    global GMAIL_LINKS, SENDER_STORE_MAP, CONTACT_STORE_MAP, CARD_STORE_MAP, CLASS_RULES, VENDOR_STD, RECEIVED_AT
    global STORE_SCOPE_RULES
    GMAIL_LINKS = load_gmail_links()
    # Self-scan mailbox -> store, for the decisive store signal (sender_store_map).
    SENDER_STORE_MAP = load_sender_store_map()
    # Recipient (To/Cc) -> store, for third-party billers emailing us (contact_store_map).
    CONTACT_STORE_MAP = load_contact_store_map()
    # Paying card last-4 -> store, for resolving a payment's store (card_store_map).
    CARD_STORE_MAP = load_card_store_map()
    # Default QuickBooks class by store/vendor, filled in when the doc didn't state one.
    CLASS_RULES = load_class_rules()
    # Vendor-scoped store rules (Setup › Store rules) for apply_scoped_learning_gate.
    STORE_SCOPE_RULES = load_store_scope_rules()
    # Vendor/account standardization -> canonical + store-specific QB names (DB-only; {} else).
    VENDOR_STD = load_vendor_standardization()
    # filename -> authoritative received time, for the invoice-date sanity check.
    RECEIVED_AT = load_received_dates()

    learnings_rows, _ = load_learnings()
    existing_rows, processed_filenames = load_existing_results()

    to_results = []          # results destined for results.xlsx (+ history)
    newly_flagged = []       # this run's flagged items -> review queue
    queue_carryover = []     # review rows kept in the queue (unanswered / errored retry)
    error_review_rows = []   # unreadable emails routed to the review queue
    unreadable_results = []  # kept only for the run summary counts
    # Source-email file moves are DEFERRED and applied only AFTER all DB writes
    # succeed, so a crash/interruption mid-run leaves every email where it was
    # for a clean retry — never moved-but-unsaved (no orphans).
    pending_moves = []       # (src Path, dst Path) to move after save
    pending_deletes = []     # split review files superseded on reprocess, deleted after save

    # ---- Phase 1: resolve answered items from the review queue ----
    review_rows = load_review_queue()
    answered = [r for r in review_rows if (r.get("Answer") or "").strip()]
    queue_carryover = [r for r in review_rows if not (r.get("Answer") or "").strip()]
    reprocessed_files = set()

    # Retroactively clear excluded-vendor invoices/credit memos queued by an OLDER
    # run (before the vendor was excluded). routes_to_review() blocks NEW ones, but
    # carryover never re-checked. Reprocessing records them to Browse (review-skipped)
    # and drops them from the queue. Payments are never excluded, so an excluded-vendor
    # payment reprocesses back as still-flagged and correctly stays queued.
    exclusion_clear = [r for r in queue_carryover if is_excluded_vendor(r.get("Vendor Matched"))]
    if exclusion_clear:
        queue_carryover = [r for r in queue_carryover if not is_excluded_vendor(r.get("Vendor Matched"))]

    # Items a reviewer uploaded PDFs/images to (unconsumed) reprocess this run even with
    # NO answer — the uploaded documents are the trigger. Pull them out of carryover so
    # they aren't also whole-queue re-checked, and reprocess them with the files appended
    # (loaded per-item in the resolve loop). DB-only (manual_attachments is a DB table).
    manual_items = []
    if USE_DB:
        manual_items = [r for r in queue_carryover
                        if _database().load_manual_attachments(r.get("Review Key") or r.get("Filename"))]
        if manual_items:
            _mids = {id(r) for r in manual_items}
            queue_carryover = [r for r in queue_carryover if id(r) not in _mids]

    # One-off retry of already-queued items (opt-in flags only; default runs skip this
    # entirely). --retry-unreadable targets unreadable-attachment/zip flags so the new
    # auto-unzip clears the backlog; --retry-open-all retries every remaining open item.
    # Reprocessed through the SAME resolve loop → adopt result: resolves → recorded and
    # dropped; still unreadable (error) → kept in the queue unchanged (manual-upload path).
    retry_items = []
    if RETRY_OPEN_ALL:
        retry_items = list(queue_carryover)
    elif RETRY_UNREADABLE:
        retry_items = [r for r in queue_carryover if _is_unreadable_review_item(r)]
    elif RETRY_ID:
        retry_items = [r for r in queue_carryover if str(r.get("Id")) == str(RETRY_ID)]
        if not retry_items:
            print(f"--retry-id={RETRY_ID}: no OPEN review item with that id was found "
                  f"— it may already be resolved or ignored (check the Review page).")
    if retry_items:
        _tids = {id(r) for r in retry_items}
        queue_carryover = [r for r in queue_carryover if id(r) not in _tids]
    retry_filenames = {r.get("Filename") for r in retry_items}

    # The unanswered, non-excluded carryover — the target set for the GATED whole-queue
    # re-check below. Captured NOW (before the resolve loop appends any failed
    # answered/exclusion rows back onto queue_carryover) so those aren't re-checked.
    recheck_candidates = list(queue_carryover)

    if answered:
        # Fold every answer's "Notes to Learn" into the durable learnings first,
        # so all of them are in play when we reprocess.
        for r in answered:
            add_learning(learnings_rows, r.get("Notes to Learn"), r.get("Filename"))
    _, learnings_text = _rows_to_learnings_text(learnings_rows)

    # Reprocess answered items AND excluded-vendor carryover through the SAME loop.
    # exclusion rows have no answer, so r.get("Answer") is "" and they reprocess with
    # no reviewer answer — an excluded invoice/credit memo then routes_to_review()=False
    # and is recorded to Browse (removed from queue); an excluded-vendor payment stays
    # flagged and is regenerated as a fresh review row (nothing silently dropped).
    resolve_list = answered + exclusion_clear + manual_items + retry_items
    if resolve_list:
        # Each row IS its own standalone split file (one flagged record), so we
        # reprocess just that file — the clean records in the same original batch
        # are never touched and never re-marked as duplicates.
        if exclusion_clear:
            print(f"Clearing {len(exclusion_clear)} excluded-vendor item(s) still queued from older runs...")
        if answered:
            print(f"Resolving {len(answered)} answered review item(s)...")
        if manual_items:
            print(f"Reprocessing {len(manual_items)} item(s) with reviewer-uploaded document(s)...")
        if retry_items:
            print(f"Retrying {len(retry_items)} "
                  f"{'open' if RETRY_OPEN_ALL else 'unreadable-attachment' if RETRY_UNREADABLE else f'#{RETRY_ID}'} item(s)...")
        for r in resolve_list:
            filename = r.get("Filename")
            review_key = r.get("Review Key") or filename   # stable key for notes/attachments
            # Load the reviewer's uploaded documents FIRST — when present they're the
            # source of truth, so we reprocess from them even if the original split file
            # is gone/unreadable (empty for the normal case).
            extra = _database().load_manual_attachments(review_key) if USE_DB else []
            src = REVIEW_PENDING_DIR / (filename or "")
            if not src.exists():
                src = PROCESSED_DIR / (filename or "")  # fallback for older items
            if not filename or (not src.exists() and not extra):
                # Nothing to work with (no source file AND no uploads) — keep in queue.
                print(f"  {filename}: review file not found in {REVIEW_PENDING_DIR.name}/ — keeping in queue.")
                queue_carryover.append(r)
                continue
            if not src.exists():
                # Source file is gone but the reviewer uploaded documents — use the queue
                # filename purely as the record NAME; process_source_email classifies from
                # the uploads (it tolerates the missing source when extras are present).
                src = REVIEW_PENDING_DIR / filename
            # Carry the ORIGINAL email's Gmail link onto the reprocessed record.
            # GMAIL_LINKS is keyed by the source-email filename, but here we read the
            # per-record SPLIT file (its name isn't in that map), so without this the
            # reprocessed record would lose its Gmail link. The review_queue row saved
            # the source email's link — reuse it under the split filename. (file_path/
            # Drive link is regenerated by process_source_email's assign_invoice_files.)
            if r.get("Gmail Link"):
                GMAIL_LINKS[filename] = r.get("Gmail Link")
            # Keep the ORIGINAL email's identity on the reprocessed record so the
            # duplicate check treats it as the same email (self, not a duplicate).
            if r.get("Source Email"):
                ORIGINAL_EMAIL[filename] = r.get("Source Email")
            print(f"Reprocessing {filename} with your answer...")
            # Recover From/To/Cc headers from the ORIGINAL email (still in processed_emails/)
            # when the reprocess source is a headerless split file, so store hints still fire.
            header_src = (PROCESSED_DIR / r["Source Email"]) if r.get("Source Email") else None
            result = process_source_email(client, src, learnings_text,
                                          reviewer_answer=r.get("Answer"), extra_attachments=extra,
                                          header_source_path=header_src)
            if result.get("error"):
                # Reprocess failed (API error / still unreadable) — keep the row and the
                # review file for next run. Manual attachments are LEFT unconsumed so an
                # upload can be retried; only a read failure keeps the files.
                print(f"  error reprocessing {filename}: {result['error']} — keeping in queue.")
                queue_carryover.append(r)
                continue
            # The reprocess actually read a result (resolved or still-flagged) — the
            # uploaded files served their purpose, so consume them (prevents re-trigger).
            if extra:
                _database().mark_manual_attachments_consumed(review_key)
            reprocessed_files.add(filename)
            to_results.append(result)
            # A split record file (Filename != Source Email) is superseded on
            # reprocess — its resolved doc now lives in invoice_files/ (or a fresh
            # review file if still flagged); the whole-email case is archived/kept.
            is_split = bool(r.get("Source Email")) and r.get("Filename") != r.get("Source Email")
            if result_has_review(result):
                newly_flagged.append(result)
                if is_split and src.parent == REVIEW_PENDING_DIR:
                    pending_deletes.append(src)
            else:
                # RESOLVED: its invoice/payment record is already committed. Delete its
                # review_queue row NOW (out-of-band) rather than waiting for the end-of-run
                # save_review_queue() full-sync — so an interrupt in the meantime can't leave
                # a stale "ghost" row for an item that's already recorded.
                if USE_DB:
                    _database().delete_review_row(filename)
                if src.parent == REVIEW_PENDING_DIR:
                    if is_split:
                        pending_deletes.append(src)
                    else:
                        pending_moves.append((src, PROCESSED_DIR / src.name))

    # Drop any leftover queue rows for files we just reprocessed — they're
    # regenerated fresh (still-flagged) or resolved (gone).
    queue_carryover = [r for r in queue_carryover if r.get("Filename") not in reprocessed_files]

    # One-off retry summary. Resolved = read + recorded cleanly (gone from the queue);
    # still stuck = the reprocess still couldn't read it (errored → kept for manual upload).
    if retry_items:
        reflagged = {res.get("filename") for res in newly_flagged}
        resolved_n = len([f for f in retry_filenames if f in reprocessed_files and f not in reflagged])
        stuck_n = len([f for f in retry_filenames if f not in reprocessed_files])
        print(f"Retried {len(retry_items)} "
              f"{'open' if RETRY_OPEN_ALL else 'unreadable-attachment' if RETRY_UNREADABLE else f'#{RETRY_ID}'} item(s): "
              f"resolved {resolved_n}, still stuck {stuck_n} (need manual upload).")

    # ---- Phase 1b: GATED whole-queue re-check against updated learnings ----
    # A queued item can only resolve DIFFERENTLY once a rule changed, so re-checking
    # every run would just waste API calls. Gate it: re-check the remaining unanswered
    # carryover ONLY when the learnings differ from the previous run — covering both a
    # rule edited in the Learnings tab between runs AND one added via a review answer
    # THIS run (learnings_text, hashed below, already reflects both). ADOPT-ONLY-IF-
    # RESOLVED: a reprocess that now resolves is recorded and the row leaves the queue;
    # one that stays flagged (or errors) is DISCARDED and the ORIGINAL row is kept
    # untouched (flagged_at / run lineage preserved). DB-only (needs the runs table to
    # compare run-to-run); Excel mode carries over exactly as before.
    learnings_hash_now = None
    if USE_DB:
        learnings_hash_now = hashlib.sha256((learnings_text or "").encode("utf-8")).hexdigest()
        prev_hash = _database().get_last_learnings_hash(run_id)
        if prev_hash and prev_hash != learnings_hash_now and recheck_candidates:
            resolved_ids = set()
            for r in recheck_candidates:
                filename = r.get("Filename")
                src = REVIEW_PENDING_DIR / (filename or "")
                if not src.exists():
                    src = PROCESSED_DIR / (filename or "")  # fallback for older items
                if not filename or not src.exists():
                    continue  # can't reprocess — leave the row exactly as-is
                # Same identity carries as the answered-item loop so the reprocessed
                # record keeps its Gmail link and self-duplicate identity.
                if r.get("Gmail Link"):
                    GMAIL_LINKS[filename] = r.get("Gmail Link")
                if r.get("Source Email"):
                    ORIGINAL_EMAIL[filename] = r.get("Source Email")
                header_src = (PROCESSED_DIR / r["Source Email"]) if r.get("Source Email") else None
                result = process_source_email(client, src, learnings_text, reviewer_answer="",
                                              header_source_path=header_src)
                if result.get("error") or result_has_review(result):
                    # Reprocess errored OR still flagged → discard; keep the original
                    # row unchanged (adopt only if the updated rule actually resolves it).
                    continue
                # Resolved by the updated learnings → record it, drop it from the queue.
                to_results.append(result)
                resolved_ids.add(id(r))
                # Out-of-band delete of the now-resolved row (same ghost-row safety as the
                # answered-item loop) — the record is committed; don't rely only on the
                # end-of-run full-sync, which an interrupt could skip.
                _database().delete_review_row(filename)
                is_split = bool(r.get("Source Email")) and r.get("Filename") != r.get("Source Email")
                if src.parent == REVIEW_PENDING_DIR:
                    if is_split:
                        pending_deletes.append(src)
                    else:
                        pending_moves.append((src, PROCESSED_DIR / src.name))
            if resolved_ids:
                queue_carryover = [r for r in queue_carryover if id(r) not in resolved_ids]
            n = len(recheck_candidates)
            print(f"Re-checking {n} queued item(s) against updated learnings — "
                  f"resolved {len(resolved_ids)}, kept {n - len(resolved_ids)}.")

    # ---- Phase 2: process new emails from test_emails/ ----
    all_files = load_emails()
    # Skip anything already handled — by filename OR by Gmail message id. The message-id check
    # catches a message already recorded under a DIFFERENT filename (a manual import, or any other
    # pathway) that the Gmail scanner later re-pulls as 'Maple Scanner - <id>.eml': filename alone
    # misses that and reprocesses it as a brand-new record (the manual-import duplication bug).
    email_mids = _database().email_message_ids() if USE_DB else {}
    processed_mids = {email_mids[fn] for fn in processed_filenames if email_mids.get(fn)}
    skipped_files, files, mid_dupes = [], [], []
    for p in all_files:
        if p.name in processed_filenames:
            skipped_files.append(p)                    # already handled under this same name
        elif email_mids.get(p.name) and email_mids[p.name] in processed_mids:
            skipped_files.append(p); mid_dupes.append(p)   # same Gmail message, different filename
        else:
            files.append(p)

    if skipped_files:
        print(f"Skipping {len(skipped_files)} already-handled email(s)"
              + (f" ({len(mid_dupes)} by Gmail message id — re-pull of an already-recorded "
                 f"message, archived)" if mid_dupes else "") + ".")
    # Archive the message-id duplicates so they don't linger in test_emails/ (re-skipping every run)
    # and stamp their emails row handled. Filename-skips keep the existing leave-in-place behavior.
    if USE_DB and mid_dupes:
        _database().stamp_handled_emails(
            [(p.name, "processed", "duplicate_repull") for p in mid_dupes], run_id)
        for p in mid_dupes:
            pending_moves.append((p, PROCESSED_DIR / p.name))

    if files:
        print(f"\nFound {len(files)} new email(s) to process.\n")
    total = len(files)
    for i, path in enumerate(files, 1):
        print(f"Processing {path.name}... ({i}/{total})")
        result = process_source_email(client, path, learnings_text)

        if result.get("read_error"):
            # Couldn't even read the file — record an error row and LEAVE it in
            # test_emails/ so it's retried automatically next run.
            to_results.append(result)
            continue

        if result.get("infra_error"):
            # Transient API/connection failure — leave the file in test_emails/
            # (don't move, don't record, don't queue) so a later run retries it.
            continue

        if result.get("unreadable"):
            # Whole email couldn't be classified — will move to review_pending/ (after save).
            pending_moves.append((path, REVIEW_PENDING_DIR / path.name))
            error_review_rows.append(build_error_review_row(result))
            unreadable_results.append(result)
            continue

        # Classified. Its NON-flagged records land in results (build_result_rows
        # emits only those); each flagged record was already saved as its own
        # standalone split file in review_pending/ (by save_review_files).
        to_results.append(result)
        if result_has_review(result):
            newly_flagged.append(result)
            if _all_records(result["classification"]):
                # Flagged records are standalone now — archive the whole email (after save).
                pending_moves.append((path, PROCESSED_DIR / path.name))
            else:
                # Record-less email flagged as a whole — it IS the review unit.
                pending_moves.append((path, REVIEW_PENDING_DIR / path.name))
        else:
            pending_moves.append((path, PROCESSED_DIR / path.name))

    # ---- Duplicate checks: mutates the invoice dicts for results; the history
    # ledger itself is persisted AFTER results (below), so if a later save fails
    # and the email is retried it isn't double-counted into history. ----
    history = load_history()
    payment_history = load_payment_history()
    _dup_results = [r for r in to_results if not r.get("error")]
    # ROOT FIX: replace (not duplicate) a reprocessed record whose email_filename churned
    # — delete its stale prior row and drop its ledger entry before the dup check runs.
    superseded_records = _supersede_reprocessed_records(
        _dup_results, reprocessed_files, history, payment_history)
    if superseded_records:
        print(f"Replacing {len(superseded_records)} superseded re-recorded record(s) "
              f"from earlier reprocess(es).")
    apply_duplicate_checks(_dup_results, history, payment_history)

    # ---- Persist results ----
    new_rows, review_dropped = build_result_rows(to_results)
    # AIRTIGHT ROUTING: if a deterministic gate inside build_result_rows flips a record to
    # needs_review that apply_review_gates (upstream) didn't already catch, that record is dropped
    # from new_rows (no live invoice). Without this it would linger as a history-only orphan whose
    # email status never advances — the /reconcile backlog bug. route_dropped_to_review strips the
    # just-added ledger row AND adds the owning email to newly_flagged, so it gets a review row and
    # a stamped status below regardless of WHERE the flag fired. No record can vanish into limbo.
    late_stats = route_dropped_to_review(review_dropped, to_results, history, payment_history, newly_flagged)
    if late_stats["late"]:
        print(f"Routed {late_stats['late']} record(s) to review that a record-time check flagged after "
              f"dedup (kept out of the ledger; {late_stats['documents']} given their own review document).")
    # The superseded (churned) records go in as replace_records so save_results DELETES just
    # those stale rows — the reprocessed record then replaces them (exactly one recorded row)
    # and any other record filed under the same old filename stays.
    touched_filenames = {r["filename"] for r in to_results}
    if USE_DB:
        _protected, recorded = _database().save_results(
            new_rows, touched_filenames, run_id, replace_records=superseded_records)
        print(f"\nSaved {len(new_rows)} record row(s) to PostgreSQL "
              f"(replaced {len(touched_filenames)} re-run filename(s)).")
        # Resolve any previously-uploaded QB export rows that had nothing to match at the time but
        # now match a freshly-recorded record (same Tier-1 vendor+number+amount+store logic). Local
        # import: qb_match imports classify_emails, so a top-level import here would be circular.
        if recorded:
            import qb_match
            n = qb_match.resolve_unmatched_for_new_records(recorded, "auto-resolved on arrival")
            if n:
                print(f"  QB-UNMATCHED: {n} newly-recorded record(s) matched a previously-uploaded "
                      f"QB export row that had nothing to match at the time.")
    # Annotate flagged "multiple versions exist" items + attach a Group link BEFORE save_history,
    # so the duplicate_group_id _group_id_for() stamps onto the matched PRIOR (in `history`)
    # actually persists in this run's save_history() write. (It used to run AFTER save_history, so
    # the history-side stamp was silently dropped and every Group link on these items was dead.)
    # Read-only cross-check: flagged records are skipped by apply_duplicate_checks, so check_duplicate
    # runs for them here purely to annotate the reason + group link. Never auto-records/resolves.
    enrich_flagged_with_duplicate_info(newly_flagged, history)
    save_history(history)
    save_payment_history(payment_history)

    # Store-preference backstop: resolve EXACT invoice duplicates where a store scan AND the vendor's
    # own copy both landed this run — the in-pipeline store_preferred supersede couldn't fire because
    # the prior copy wasn't in the DB yet when the later one was processed. Runs now that both ARE
    # saved: keeps the store-scanned copy, auto-ignores the vendor copy. Conservative (see the fn).
    if USE_DB:
        sp = resolve_store_preferred_duplicates("auto-resolved (store-preferred backstop)")
        if sp["resolved"]:
            print(f"  STORE-PREFERRED: auto-resolved {len(sp['resolved'])} exact-duplicate set(s) "
                  f"(kept the store-scanned copy, ignored the vendor copy):")
            for r in sp["resolved"]:
                print(f"    {r['vendor']} #{r['number']} @ {r['store']}: kept id {r['keep_id']}, "
                      f"ignored {r['ignored_ids']}")

    # Status-update safety net: any email SUCCESSFULLY handled this run (classified without error —
    # recorded, flagged-to-review, or informational) must not stay stuck on a stale status='error'
    # from an earlier failed run, even if its records took a path that didn't reach save_results'
    # _upsert_email. Only clears rows currently 'error' (never disturbs other statuses). This closes
    # the gap that left EuroPharma INV0800123 'error' for 2 weeks after it was actually handled.
    if USE_DB:
        handled = {r["filename"] for r in to_results
                   if r.get("filename") and not (r.get("error") or r.get("read_error")
                                                  or r.get("infra_error"))}
        cleared = _database().clear_stale_error_status(handled, run_id)
        if cleared:
            print(f"Cleared stale 'error' status on {cleared} previously-errored email(s) now handled.")

    # ---- HOLD sales orders: diverted from recording/review (build_result_rows +
    # apply_duplicate_checks skip is_sales_order), snapshot them so a later invoice can
    # supersede them, else they record themselves after the hold window (reconcile). ----
    if USE_DB:
        held_n = 0
        release_after = datetime.now(timezone.utc) + timedelta(days=SALES_ORDER_HOLD_DAYS)
        for r in to_results:
            c = r.get("classification")
            if not c:
                continue
            fn = r["filename"]
            for inv in c.get("invoices") or []:
                if not inv.get("is_sales_order"):
                    continue
                snapshot = {
                    "filename": fn, "gmail_link": GMAIL_LINKS.get(fn, ""),
                    "source_email": _record_source_id(fn),
                    "email_type": c.get("email_type"), "confidence": c.get("confidence"),
                    "route_to": c.get("route_to"), "routing_reason": c.get("routing_reason"),
                    "invoice": dict(inv),
                }
                if _database().insert_sales_order({
                        "vendor_matched": inv.get("vendor_matched"), "store_matched": inv.get("store_matched"),
                        "number": inv.get("invoice_number"), "amount": inv.get("invoice_amount"),
                        "order_date": inv.get("invoice_date"), "source_email": snapshot["source_email"],
                        "source_file": fn, "file_path": inv.get("_file_path") or "",
                        "gmail_link": snapshot["gmail_link"], "snapshot": snapshot,
                        "release_after": release_after}):
                    held_n += 1
        if held_n:
            print(f"Held {held_n} sales order(s) for ~{SALES_ORDER_HOLD_DAYS:g}-day window (not yet recorded).")

    # kept_existing_rows preserves un-re-run rows when rewriting the whole sheet
    # (Excel mode). In DB mode we re-read the FULL dataset from the DB so the
    # Excel export is a complete snapshot, not just this run.
    all_rows = [r for r in existing_rows if r.get("Filename") not in touched_filenames
                and not _replaced_by(r, superseded_records)] + new_rows
    if WRITE_EXCEL:
        excel_rows = _database().load_all_result_rows() if USE_DB else all_rows
        result_row_count = write_results_xlsx(excel_rows)
        qb_row_count = write_qb_ready_xlsx(excel_rows)
        print(f"Wrote {result_row_count} row(s) to {RESULTS_PATH} and "
              f"{qb_row_count} to {QB_READY_PATH}.")

    # ---- Write review queue (carryover + this run's new per-record flags + unreadable) ----
    # (enrich_flagged_with_duplicate_info already ran above, before save_history, so its
    # group-id stamp on the matched prior persists; the flagged rows below carry the same id.)
    new_review_rows = []
    kept_assignee = 0
    kept_ids = {r.get("Id") for r in queue_carryover if r.get("Id") not in (None, "")}
    for r in newly_flagged:
        rows = build_review_rows(r)
        # A re-flagged item's fresh row keeps the assignee and the queue id its old row had
        # (review_rows = the queue as loaded at the start of this run).
        kept_assignee += carry_review_assignments(rows, review_rows, r["filename"])
        carry_review_ids(rows, review_rows, r["filename"], kept_ids)
        new_review_rows.extend(rows)
    if kept_assignee:
        print(f"Kept the assignee on {kept_assignee} re-flagged review item(s).")
    # Stamp this run's id on the NEW rows (flagged_at is left blank so the DB sets
    # it to now()); carryover rows already carry their original run_id/flagged_at.
    for row in new_review_rows + error_review_rows:
        row["Run ID"] = run_id
    queue_out = queue_carryover + new_review_rows + error_review_rows
    save_review_queue(queue_out)

    # Guarantee an emails-row status write for every email HANDLED this run that produced NO recorded
    # row — flagged entirely to review (all records flagged, or a record-less model-flagged email) or
    # unreadable. _upsert_email only ran for recorded rows, so these would otherwise stay frozen at
    # their fetch-time status (the 237-email backlog). Emails with at least one recorded row are left
    # as save_results set them ('processed').
    if USE_DB:
        recorded_fns = {row.get("Filename") for row in new_rows}
        stamp = []
        seen = set()
        for r in newly_flagged:
            fn = r.get("filename")
            if fn and fn not in recorded_fns and fn not in seen:
                seen.add(fn)
                stamp.append((fn, "review", (r.get("classification") or {}).get("email_type")))
        for r in unreadable_results:
            fn = r.get("filename")
            if fn and fn not in recorded_fns and fn not in seen:
                seen.add(fn)
                stamp.append((fn, "review", "unreadable"))
        # ROOT-CAUSE FIX for the never-classified backlog: an INFORMATIONAL email (classified
        # cleanly but producing no record and not flagged — order acks, shipment notices,
        # statements, correspondence) is moved to processed_emails/ but _upsert_email never ran
        # (no recorded row) and it's in neither newly_flagged nor unreadable. Without this its
        # emails row stayed frozen at 'fetched'/run_id-NULL forever even though it was fully
        # handled — exactly how 64 emails accumulated unprocessed-looking in processed_emails/.
        # Stamp them 'processed' with their classified email_type. (The pending_moves backstop
        # below is the by-construction net; this fixes it at the source so the net stays at 0.)
        for r in to_results:
            fn = r.get("filename")
            if not fn or fn in recorded_fns or fn in seen:
                continue
            if r.get("error") or r.get("read_error") or r.get("infra_error") or not r.get("classification"):
                continue                      # errors/unreadable: retried or handled elsewhere
            seen.add(fn)
            stamp.append((fn, "processed", (r.get("classification") or {}).get("email_type")))
        n = _database().stamp_handled_emails(stamp, run_id)
        if n:
            print(f"Stamped emails status for {n} handled email(s) with no recorded row "
                  f"(review-routed, unreadable, or informational).")

    save_learnings(learnings_rows)

    # ---- Reconcile HELD sales orders: supersede (an invoice arrived) or release (hold
    # elapsed → record now, queue if it needs review). Runs AFTER this run's records are
    # saved so it sees invoices recorded this run. Append-only to the queue just written. ----
    reconcile_sales_orders(run_id)

    # ---- Statement reconciliation: ingest each processed statement's listed invoices into
    # statement_lines, then cross-check every line against recorded invoices (matched /
    # missing / amount_mismatch / in_flight / pre_cutoff). Runs AFTER records + the review
    # queue are saved so it sees this run's invoices and in-flight items. Exceptions are
    # surfaced in the web Review queue at read time (NOT stored in review_queue). ----
    if USE_DB:
        stmt_lines_n = 0
        for r in to_results:
            c = r.get("classification")
            if not c or (c.get("email_type") or "").lower() != "statement":
                continue
            listed = c.get("statement_invoices") or []
            if not listed:
                continue
            fn = r["filename"]
            default_vendor = c.get("vendor_matched")
            norm = [{"vendor": (ln.get("vendor") or default_vendor),
                     "invoice_number": ln.get("invoice_number"),
                     "invoice_date": ln.get("invoice_date"),
                     "amount": ln.get("amount")} for ln in listed]
            stmt_lines_n += _database().ingest_statement_lines(
                statement_email=_record_source_id(fn), statement_file=fn,
                lines=norm, run_id=run_id)
        recon = _database().reconcile_statement_lines(run_id)
        if stmt_lines_n or recon.get("flagged"):
            print(f"Statement reconciliation: ingested {stmt_lines_n} listed invoice(s); "
                  f"{recon.get('missing', 0)} missing, {recon.get('amount_mismatch', 0)} mismatch, "
                  f"{recon.get('matched', 0)} matched, {recon.get('in_flight', 0)} in-flight, "
                  f"{recon.get('pre_cutoff', 0)} pre-cutoff.")

    print(f"Review queue: {len(queue_out)} open item(s) "
          f"({len(new_review_rows) + len(error_review_rows)} new this run"
          f"{f', {len(exclusion_clear)} excluded-vendor item(s) cleared' if exclusion_clear else ''}). "
          f"Learnings: {len(learnings_rows)} rule(s).")

    # ---- All DB writes succeeded — NOW move the source-email files ----
    # If anything above raised, we never get here, so nothing was moved and every
    # email is still in test_emails/ (or review_pending/) for a clean retry.
    for src, dst in pending_moves:
        _safe_move(src, dst)
    _delete_unused_review_files(pending_deletes, queue_out)

    # BY-CONSTRUCTION status backstop: every source email that just LEFT test_emails/ has been
    # handled this run, so its emails row must be off the fetch-time status. The known-good writes
    # above already do this (so this normally corrects 0); if it EVER corrects >0, some code path
    # moved a file out of the inbox without stamping it — the same gap that stranded records in
    # invoice_history / froze emails at 'fetched'. Fixed here regardless, and logged so it's visible.
    if USE_DB and pending_moves:
        def _review_dest(d):
            return d.parent == REVIEW_PENDING_DIR or REVIEW_PENDING_DIR in d.parents
        queue_exits = [(src.name, "review" if _review_dest(dst) else "processed")
                       for src, dst in pending_moves if src.parent == EMAILS_DIR]
        corrected = _database().stamp_moved_from_queue(queue_exits, run_id)
        if corrected:
            print(f"WARNING: {corrected} email(s) left test_emails/ without a status write this run "
                  f"— backstopped to their destination status. A code path recorded/routed them "
                  f"without stamping emails.status; investigate so it's fixed at the source.")

    # ---- Stamp the run row with completion time + counts ----
    if USE_DB:
        n_emails = len(touched_filenames)
        n_review = len(new_review_rows) + len(error_review_rows)
        # Record this run's learnings hash so the NEXT run can detect a change and
        # gate its whole-queue re-check (Phase 1b).
        _database().finish_run(run_id, n_emails, len(new_rows), n_review,
                               learnings_hash=learnings_hash_now)

    # newly_flagged results are already in to_results (their clean records landed
    # in results); only unreadable ones are separate.
    summary = build_summary(to_results + unreadable_results)
    print_summary(summary)
    if DRIVE_FALLBACKS:
        print(f"\nDRIVE: {len(DRIVE_FALLBACKS)} document(s) could not be uploaded this run and stay on this PC"
              f"{' (' + _DRIVE_UNAVAILABLE + ')' if _DRIVE_UNAVAILABLE else ''}. Upload them later with:\n"
              f"  python upload_local_files_to_drive.py --commit")


if __name__ == "__main__":
    main()
