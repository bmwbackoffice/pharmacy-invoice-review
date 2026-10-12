"""
Phase 2B — web review UI over the existing PostgreSQL database.

Additive: reads/writes the SAME tables the CLI pipeline uses, plus the web-only
users/edit_log/audit columns (created at startup by web_db.init_web_schema).
It records reviewer answers, QB-entry flags, and field edits — it does NOT run
classification (the next `classify_emails.py` run resolves answered items).

Run it with:  python run_web.py   (prints the LAN URL + firewall rule)
Config (env): DATABASE_URL, SECRET_KEY, ADMIN_USER, ADMIN_PASSWORD, PORT.

Stack: FastAPI + Jinja2 + Starlette signed-cookie sessions + uvicorn. No state
lives in the process — everything is in PostgreSQL — so moving to a cloud VM is
just: run this same app pointing at the same DB, add HTTPS + real logins.
"""

import csv
import io
import os
import re
from datetime import date, datetime
from pathlib import Path
from urllib.parse import quote

from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.concurrency import run_in_threadpool
from starlette.middleware.base import BaseHTTPMiddleware
from starlette.middleware.sessions import SessionMiddleware

import chat_qa
import qb_match
import web_db

BASE_DIR = Path(__file__).parent
SECRET_KEY = os.environ.get("SECRET_KEY", "dev-insecure-change-me")

app = FastAPI(title="Pharmacy Invoice Review")

# Read-only users (role 'viewer', 2026-10-12) can open every page and export, but every request that
# could change something (anything but GET/HEAD) is refused here, in one place, instead of in each of
# the ~56 POST routes. Exceptions: signing in/out, Chat (it only reads data), and the user's own QB-entry
# column layout (a personal display preference; its forms carry class "pref-form" so base.html shows
# them). Added BEFORE the session middleware so that one wraps this and request.session is available.
_READ_ONLY_ALLOWED = {"/login", "/logout", "/chat/ask", "/qb/columns"}
READ_ONLY_MESSAGE = "You have read-only access, so you can look but not change anything. Ask an admin if you need to."


async def _read_only_guard(request: Request, call_next):
    if request.method not in ("GET", "HEAD", "OPTIONS") and request.url.path not in _READ_ONLY_ALLOWED:
        uid = request.session.get("user_id")
        user = await run_in_threadpool(web_db.get_user_by_id, uid) if uid else None
        if user and user.get("role") == "viewer":
            wants_json = ("application/json" in request.headers.get("accept", "")
                          or request.headers.get("content-type", "").startswith("application/json"))
            if wants_json:
                return JSONResponse({"error": READ_ONLY_MESSAGE}, status_code=403)
            resp = render(request, "message.html", user, message=READ_ONLY_MESSAGE)
            resp.status_code = 403
            return resp
    return await call_next(request)


app.add_middleware(BaseHTTPMiddleware, dispatch=_read_only_guard)
app.add_middleware(SessionMiddleware, secret_key=SECRET_KEY, https_only=False)

templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


# --------------------------------------------------------------------------- #
# Display-only date filters: everything shown in the UI renders as M/D/YY
# (e.g. 8/27/26), timestamps as M/D/YY H:MM (24-hour). Centralized here so no
# template hardcodes a format. NOTE: built by hand — strftime('%-m')/('%-d') is
# not portable (fails on Windows, the host). Stored values are never touched.
# --------------------------------------------------------------------------- #

def _coerce_dt(value):
    """A date/datetime for `value`, or None if it isn't (and can't be parsed) one."""
    if value is None or value == "":
        return None
    if isinstance(value, (datetime, date)):
        return value
    s = str(value).strip()
    if not s:
        return None
    try:
        return datetime.fromisoformat(s)
    except ValueError:
        pass
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    return None


def us_date(value):
    """None/'' -> ''; a date/datetime or parseable string -> 'M/D/YY' (8/27/26);
    an unparseable string is returned unchanged."""
    if value is None or value == "":
        return ""
    dt = _coerce_dt(value)
    if dt is None:
        return str(value)
    return f"{dt.month}/{dt.day}/{dt.year % 100:02d}"


def us_datetime(value):
    """Like us_date, plus ' H:MM' (24-hour) for datetimes; a date-only value (or a
    string with no time) renders as us_date. Unparseable strings pass through."""
    if value is None or value == "":
        return ""
    dt = _coerce_dt(value)
    if dt is None:
        return str(value)
    d = f"{dt.month}/{dt.day}/{dt.year % 100:02d}"
    if isinstance(dt, datetime) and (dt.hour or dt.minute or dt.second):
        return f"{d} {dt.hour}:{dt.minute:02d}"
    return d


templates.env.filters["us_date"] = us_date
templates.env.filters["us_datetime"] = us_datetime
templates.env.filters["store_label"] = web_db.store_label


def _diff_against(value, other):
    """`value` rendered as HTML with the characters that DIFFER from `other` wrapped in
    <span class="diffchar">. Character-level (difflib, autojunk off) so e.g. 'U13944' vs
    'U-13944' reddens only the inserted '-' rather than the whole value. Both sides coerced to
    str; output is escaped then marked safe. Used on the QB match-review cards to pinpoint the
    one differing field (see qb_reconcile_review.html)."""
    import difflib
    from markupsafe import Markup, escape
    a = "" if value is None else str(value)
    b = "" if other is None else str(other)
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    parts = []
    for tag, i1, i2, _j1, _j2 in sm.get_opcodes():
        seg = str(escape(a[i1:i2]))
        if not seg:
            continue
        parts.append(seg if tag == "equal" else f'<span class="diffchar">{seg}</span>')
    return Markup("".join(parts))


templates.env.filters["diff_against"] = _diff_against

# Read-only static mounts for inline file previews.
for name in ("invoice_files", "payment_files", "review_pending"):
    d = BASE_DIR / name
    d.mkdir(exist_ok=True)
    app.mount(f"/{name}", StaticFiles(directory=str(d)), name=name)


@app.on_event("startup")
def _startup():
    web_db.init_web_schema()
    web_db.seed_admin(os.environ.get("ADMIN_USER"), os.environ.get("ADMIN_PASSWORD"))


# --------------------------------------------------------------------------- #
# Auth helpers
# --------------------------------------------------------------------------- #

def current_user(request: Request):
    uid = request.session.get("user_id")
    if not uid:
        return None
    user = web_db.get_user_by_id(uid)
    if not user or not user["active"]:
        return None
    return user


def _login_redirect():
    return RedirectResponse("/login", status_code=302)


def render(request, template, user, **ctx):
    # Inject the open Tracked count so base.html can label the nav tab "Tracked (N)".
    # Cheap COUNT; wrapped so a DB hiccup never breaks page rendering.
    if user and "tracked_open" not in ctx:
        try:
            ctx["tracked_open"] = web_db.tracked_open_count()
        except Exception:
            ctx["tracked_open"] = None
    # Sitewide maintenance banner on first render (polling in base.html keeps it live after).
    # Wrapped so a DB hiccup never breaks page rendering.
    if "notice" not in ctx:
        try:
            ctx["notice"] = web_db.get_active_notice()
        except Exception:
            ctx["notice"] = None
    return templates.TemplateResponse(request, template, {"user": user, **ctx})


# The invoice mailbox the Gmail deep-links must open. Stored links written before
# the account-fix (or by a run whose GMAIL_LINKS was stale) use the bare /u/N/#…
# form that opens whichever account is first in the viewer's browser. fix_gmail_link
# repairs those at DISPLAY time (idempotent — links that already carry authuser are
# left alone), a belt-and-suspenders over backfill_gmail_links.py.
AP_MAILBOX = os.environ.get("GMAIL_MAILBOX", "maplegroupap@lakewoodrx.com")
_GMAIL_UNPINNED = re.compile(r"(https://mail\.google\.com/mail/u/\d+/)#")


def fix_gmail_link(link):
    if not link or "authuser=" in link:
        return link
    return _GMAIL_UNPINNED.sub(
        lambda m: f"{m.group(1)}?authuser={quote(AP_MAILBOX, safe='@')}#", link, count=1)


def file_preview_url(file_path):
    """Map a stored file_path to something previewable: a Drive/HTTP URL as-is,
    or a local file under one of the mounted folders by basename."""
    if not file_path:
        return None
    if str(file_path).startswith(("http://", "https://")):
        return file_path
    p = str(file_path).replace("\\", "/")
    base = p.rsplit("/", 1)[-1]
    for folder in ("review_pending", "payment_files", "invoice_files"):
        if f"/{folder}/" in p or p.endswith(f"/{base}") and folder in p:
            return f"/{folder}/{base}"
    # fall back by folder name appearing anywhere in the path
    for folder in ("review_pending", "payment_files", "invoice_files"):
        if folder in p:
            return f"/{folder}/{base}"
    return None


# Manual-upload handling: PDFs/images a reviewer attaches to a review item.
MANUAL_UPLOAD_EXT = {".pdf", ".png", ".jpg", ".jpeg", ".gif", ".webp"}
MANUAL_UPLOAD_CT = {"application/pdf", "image/png", "image/jpeg", "image/gif", "image/webp",
                    "application/octet-stream"}   # some browsers send octet-stream; extension gates it
REVIEW_PENDING_DIR = BASE_DIR / "review_pending"


def _safe_upload_name(name):
    """Sanitize an uploaded filename to a bare, safe basename (no paths/traversal)."""
    base = os.path.basename((name or "").replace("\\", "/"))
    base = re.sub(r'[^A-Za-z0-9._ -]', "_", base).strip().strip(".")
    return base[:180] or "upload"


def _manual_preview_url(file_path):
    """URL for a manual attachment saved under review_pending/ (keeps the sub-path,
    unlike file_preview_url which is basename-only)."""
    p = str(file_path or "").replace("\\", "/")
    i = p.lower().rfind("review_pending/")
    return ("/" + p[i:]) if i >= 0 else file_preview_url(file_path)


app.state.file_preview_url = file_preview_url


# --------------------------------------------------------------------------- #
# Sitewide maintenance notice (banner): first render via render()'s `notice`,
# kept live by base.html polling /api/notice. Admin can set/clear here or via notice.py.
# --------------------------------------------------------------------------- #

@app.get("/api/notice")
def api_notice(request: Request):
    """Current sitewide notice for the base.html poller — so someone already sitting on a page
    sees a new banner appear without navigating. Returns {notice: {...}|null}."""
    user = current_user(request)
    if not user:
        return JSONResponse({"notice": None})
    try:
        return JSONResponse({"notice": web_db.get_active_notice()})
    except Exception:
        return JSONResponse({"notice": None})


# Shown when the admin clicks "Show" without typing a message — the common case is a quick restart,
# so one click with no typing posts this.
DEFAULT_NOTICE = "Maintenance Alert — please SAVE any entry in progress; back in a moment."


@app.post("/admin/notice")
def admin_set_notice(request: Request, message: str = Form(""), minutes: int = Form(1)):
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    # Blank message -> the default (one-click restart warning). minutes <= 0 -> until cleared.
    msg = (message or "").strip() or DEFAULT_NOTICE
    web_db.set_notice(msg, (minutes if minutes and minutes > 0 else None), user["username"])
    # Stay on the page the form was on (the dashboard) rather than bouncing to /review.
    return RedirectResponse(request.headers.get("referer") or "/dashboard", status_code=302)


@app.post("/admin/notice/clear")
def admin_clear_notice(request: Request):
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    web_db.clear_notice(user["username"])
    return RedirectResponse(request.headers.get("referer") or "/dashboard", status_code=302)


# --------------------------------------------------------------------------- #
# Login / logout
# --------------------------------------------------------------------------- #

@app.get("/login", response_class=HTMLResponse)
def login_form(request: Request, error: str = ""):
    return render(request, "login.html", None, error=error)


@app.post("/login")
def login_submit(request: Request, username: str = Form(...), password: str = Form(...)):
    user = web_db.get_user(username)
    if not user or not user["active"] or not web_db.verify_password(password, user["password_hash"]):
        return render(request, "login.html", None, error="Invalid username or password.")
    request.session["user_id"] = user["id"]
    return RedirectResponse("/review", status_code=302)


@app.get("/logout")
def logout(request: Request):
    request.session.clear()
    return _login_redirect()


@app.get("/")
def root(request: Request):
    return RedirectResponse("/review" if current_user(request) else "/login", status_code=302)


# --------------------------------------------------------------------------- #
# 0) To-do dashboard — every "needs action" queue in one place (counts only;
#    per-user breakdown exists solely for the review queue, the one queue with
#    an assignee field — everything else is a shared pool).
# --------------------------------------------------------------------------- #

@app.get("/dashboard", response_class=HTMLResponse)
def dashboard(request: Request):
    user = current_user(request)
    if not user:
        return _login_redirect()
    d = web_db.dashboard_counts()
    setup = d["qb_invoices_setup"]
    # Tiles merged so each page is represented ONCE (counts/notes combine the sub-queues).
    tiles = [
        {"label": "Review queue", "count": d["review"], "href": "/review",
         "note": f"{d['review']} awaiting an answer", "breakdown": d["review_by"]},
        {"label": "Statement exceptions", "count": d["statement"], "href": "/statements",
         "note": "missing/mismatched invoices from a statement"},
        {"label": "Possible duplicates", "count": d["duplicates"], "href": "/duplicates",
         "note": "sets to resolve"},
        {"label": "QB entry", "count": d["qb_invoices"] + d["qb_credits"], "href": "/qb",
         "note": (f"{d['qb_invoices']} invoices"
                  + (f" ({setup} need vendor/account setup)" if setup else "")
                  + f", {d['qb_credits']} credit memos")},
        {"label": "Payments to enter", "count": d["payments"], "href": "/payments",
         "note": "not yet entered"},
        {"label": "Flags", "count": d["flags_pending"] + d["flags_approved"] + d["investigations"],
         "href": "/flags",
         "note": (f"{d['flags_pending']} pending, {d['flags_approved']} approved for fix"
                  + (f", {d['investigations']} investigation(s) running/failed" if d["investigations"] else ""))},
        {"label": "Change requests", "count": d["requests_pending"] + d["requests_approved"],
         "href": "/requests",
         "note": f"{d['requests_pending']} pending, {d['requests_approved']} approved for build"},
        {"label": "Sales orders (held)", "count": d["sales_orders"], "href": "/sales-orders",
         "note": "held, awaiting an invoice"},
        {"label": "Cards → store", "count": d["cards"], "href": "/cards",
         "note": "card last-4 with no store mapping"},
        {"label": "Failed emails", "count": d["failed_emails"], "href": "/emails?disposition=error",
         "note": "processing error"},
        {"label": "Orphaned ledger rows", "count": d.get("orphaned_history", 0), "href": "/reconcile",
         "note": "extracted but never recorded"},
        {"label": "Bills with no live copy", "count": d.get("no_live_copy", 0), "href": "/no-live-copy",
         "note": "every copy ignored as a duplicate, none entered"},
        {"label": "Never classified", "count": d.get("never_classified", 0),
         "href": "/emails?disposition=never_classified",
         "note": "fetched >1 day ago, no run ever processed it"},
        {"label": "QB match review", "count": d.get("qb_pending", 0), "href": "/qb-reconcile/review",
         "note": "doubtful QuickBooks-reconcile matches to confirm"},
    ]
    return render(request, "dashboard.html", user, tiles=tiles)


@app.get("/reconcile", response_class=HTMLResponse)
def reconcile(request: Request):
    """Read-only drill-down for the reconcile check: invoice_history / payment_history rows that
    never became a live invoices / payments row ('extracted but never recorded'). Each is likely a
    recovery that silently failed — recover via recover_errored_email.py."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    return render(request, "reconcile.html", user, rows=web_db.orphaned_history_records())


@app.get("/no-live-copy", response_class=HTMLResponse)
def no_live_copy(request: Request):
    """Invoice/payment groups where every copy is ignored as a duplicate and none is entered -- a bill
    that will never show for entry. Restore a copy (existing duplicate route), or confirm the group."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    groups = web_db.no_live_copy_groups()
    return render(request, "no_live_copy.html", user,
                  needs=[g for g in groups if not g["by_person"] and not g["confirmed"]],
                  person=[g for g in groups if g["by_person"] and not g["confirmed"]],
                  confirmed=[g for g in groups if g["confirmed"]])


@app.post("/no-live-copy/confirm")
def no_live_copy_confirm(request: Request, record_type: str = Form(...), vendor_key: str = Form(...),
                         number_key: str = Form(...), undo: str = Form("")):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.confirm_no_live_copy(record_type, vendor_key, number_key, user["username"], confirm=not undo)
    return RedirectResponse("/no-live-copy", status_code=302)


# --------------------------------------------------------------------------- #
# 1) Review queue
# --------------------------------------------------------------------------- #

@app.get("/review", response_class=HTMLResponse)
def review_list(request: Request, run: str = "", sort: str = "", direction: str = "desc",
                store: str = "", vendor: str = "", search: str = "", show_answered: str = "",
                assigned: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    include_answered = show_answered in ("1", "true", "on", "yes")
    search = (search or "").strip()
    items = web_db.list_open_reviews(run_id=run or None, sort=sort, direction=direction,
                                     store=store or None, vendor=vendor or None,
                                     search=search or None, include_answered=include_answered,
                                     assigned=assigned or None)
    # Statement-reconciliation exceptions now live on their own /statements page, not here.
    counts = web_db.get_review_note_counts()   # separate query; never JOINed into the queue
    pending_upload = web_db.pending_manual_filenames()   # items with unconsumed uploads
    for it in items:
        it["gmail_link"] = fix_gmail_link(it.get("gmail_link"))  # repair stale /u/0/ links
        # Key notes/uploads by the STABLE review_key (survives a reprocess rename); fall back to
        # filename for any legacy row that predates review_key.
        rkey = it.get("review_key") or it.get("filename")
        it["note_count"] = counts.get(rkey, 0)
        it["has_pending_upload"] = rkey in pending_upload
    # Sink items that carry internal notes to the bottom. Python's sort is STABLE, so
    # the DB order (default or column-sorted) is preserved WITHIN the noted/not-noted groups.
    items.sort(key=lambda it: 1 if it.get("note_count") else 0)
    # Group by assignment as the top-level order: mine first, then unassigned, then
    # everyone else's. Stable, so the note-sink + DB order above still governs ordering
    # within each group. (Statement-recon exceptions have no assigned_to -> unassigned.)
    my_username = user["username"]
    def _assign_rank(it):
        a = (it.get("assigned_to") or "").strip()
        if not a:
            return 1
        return 0 if a == my_username else 2
    items.sort(key=_assign_rank)
    opts = web_db.review_filter_options(run_id=run or None)
    active_users = [u for u in web_db.list_users() if u.get("active")]
    return render(request, "review_list.html", user, items=items,
                  runs=web_db.list_runs(), run=run,
                  stores=opts["stores"], vendors=opts["vendors"],
                  store=store, vendor=vendor, search=search, show_answered=include_answered,
                  assigned=assigned, users=active_users,
                  answered_n=web_db.answered_count(run_id=run or None, store=store or None,
                                                   vendor=vendor or None, search=search or None,
                                                   assigned=assigned or None),
                  ignored_n=web_db.ignored_count(),
                  params=dict(run=run, sort=sort, direction=direction, store=store,
                              vendor=vendor, search=search, assigned=assigned,
                              show_answered=("1" if include_answered else "")))


@app.get("/statements", response_class=HTMLResponse)
def statements_list(request: Request, run: str = "", vendor: str = "",
                    search: str = "", assigned: str = ""):
    """Statement-reconciliation exceptions on their own page (one grouped row per statement),
    split out of the Review queue so they don't clutter the regular review list."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    search = (search or "").strip()
    items = web_db.list_open_statement_exceptions(
        run_id=run or None, vendor=vendor or None, search=search or None,
        assigned=assigned or None)
    for it in items:
        it["gmail_link"] = fix_gmail_link(it.get("gmail_link"))   # pin the AP mailbox
    my_username = user["username"]
    def _assign_rank(it):
        a = (it.get("assigned_to") or "").strip()
        if not a:
            return 1
        return 0 if a == my_username else 2
    items.sort(key=_assign_rank)
    opts = web_db.statement_filter_options(run_id=run or None)
    active_users = [u for u in web_db.list_users() if u.get("active")]
    return render(request, "statements.html", user, items=items,
                  runs=web_db.list_runs(), run=run, vendor=vendor,
                  search=search, assigned=assigned, users=active_users,
                  vendors=opts["vendors"],
                  params=dict(run=run, vendor=vendor, search=search, assigned=assigned))


@app.get("/review/{review_id}", response_class=HTMLResponse)
def review_detail(request: Request, review_id: int, flagged: str = "", fix_error: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    item = web_db.get_review(review_id)
    if not item:
        return RedirectResponse("/review", status_code=302)
    item["gmail_link"] = fix_gmail_link(item.get("gmail_link"))  # repair stale /u/0/ link
    rkey = item.get("review_key") or item.get("filename")   # stable key (fallback for legacy rows)
    manual = web_db.list_manual_attachments(rkey)
    for m in manual:
        m["preview_url"] = _manual_preview_url(m.get("file_path"))
    has_pending_upload = any(not m.get("consumed_at") for m in manual)
    return render(request, "review_detail.html", user, item=item,
                  preview=file_preview_url(item.get("file_path")),
                  siblings=web_db.list_email_siblings(item.get("source_email")),
                  notes=web_db.list_review_notes(rkey),
                  users=[u for u in web_db.list_users() if u.get("active")],
                  manual_attachments=manual, has_pending_upload=has_pending_upload,
                  flagged=flagged, stores=web_db.store_names(),
                  fix_kinds=quick_fix_kinds(item.get("question")), fix_error=fix_error)


@app.post("/review/{review_id}/attach")
async def review_attach(request: Request, review_id: int,
                        files: list[UploadFile] = File(default=[])):
    """Upload one or more PDFs/images to a review item (e.g. hand-extracted from a zip
    the pipeline couldn't open). Saved under review_pending/manual/<item-stem>/, recorded
    in manual_attachments; the NEXT classify run reprocesses the item WITH these files.
    Validates extension + content-type and sanitizes names. Saves nothing else."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    item = web_db.get_review(review_id)
    if not item or not item.get("filename"):
        return RedirectResponse("/review", status_code=302)
    review_filename = item["filename"]                          # on-disk upload folder name
    review_key = item.get("review_key") or review_filename      # stable storage key
    stem = _safe_upload_name(Path(review_filename).stem)
    dest_dir = REVIEW_PENDING_DIR / "manual" / stem
    saved_names = []
    for f in files or []:
        if not f or not (f.filename or "").strip():
            continue
        safe = _safe_upload_name(f.filename)
        ext = Path(safe).suffix.lower()
        ct = (f.content_type or "").lower()
        if ext not in MANUAL_UPLOAD_EXT or (ct and ct not in MANUAL_UPLOAD_CT):
            continue                                  # reject non-PDF/image quietly
        data = await f.read()
        if not data:
            continue
        dest_dir.mkdir(parents=True, exist_ok=True)
        dest = dest_dir / safe
        n = 1
        while dest.exists():                          # never clobber an existing upload
            dest = dest_dir / f"{Path(safe).stem}_{n}{ext}"; n += 1
        dest.write_bytes(data)
        web_db.add_manual_attachment(review_key, str(dest), safe, user["username"])
        saved_names.append(safe)
    if saved_names:
        # Document the action on the item's Notes thread (audit trail + note badge), so it
        # visibly reflects "something was done" rather than sitting untouched in the queue.
        web_db.add_review_note(
            review_key,
            f"Uploaded {len(saved_names)} file(s) for reprocess: {', '.join(saved_names)}",
            user["username"])
    return RedirectResponse(f"/review/{review_id}", status_code=302)


@app.post("/review/{review_id}/note")
def review_add_note(request: Request, review_id: int, note: str = Form(...)):
    user = current_user(request)
    if not user:
        return _login_redirect()
    item = web_db.get_review(review_id)
    if item:  # key notes by the STABLE review_key (survives a reprocess rename), not the filename
        web_db.add_review_note(item.get("review_key") or item.get("filename"), note, user["username"])
    return RedirectResponse(f"/review/{review_id}", status_code=302)


@app.post("/review/{review_id}/note/{note_id}/delete")
def review_note_delete(request: Request, review_id: int, note_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    if user["role"] != "admin":                 # deletes are admin-only
        return render(request, "message.html", user, message="Deleting internal notes is admin-only.")
    web_db.delete_review_note(note_id, user["username"])
    return RedirectResponse(f"/review/{review_id}", status_code=302)


@app.post("/statement-line/{line_id}/dismiss")
def statement_line_dismiss(request: Request, line_id: int):
    """Resolve/dismiss a SINGLE statement-reconciliation exception line. It leaves the Review
    queue and won't reappear (reconcile skips dismissed lines). Kept as a building block; the
    grouped Review row uses /statement-line/dismiss-group below."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.dismiss_statement_line(line_id, user["username"])
    return RedirectResponse("/review", status_code=302)


@app.post("/statement-line/dismiss-group")
def statement_line_dismiss_group(request: Request, line_ids: str = Form(...)):
    """Resolve/dismiss ALL open lines of one statement at once (one grouped Review row).
    `line_ids` is the comma-joined list of statement_lines ids from that row. They leave the
    Review queue and won't reappear (reconcile skips dismissed lines)."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    ids = [int(x) for x in line_ids.split(",") if x.strip().isdigit()]
    if ids:
        web_db.dismiss_statement_lines(ids, user["username"])
    return RedirectResponse("/review", status_code=302)


@app.post("/statement-line/assign")
def statement_lines_assign(request: Request, line_ids: str = Form(""),
                           assigned_to: str = Form("")):
    """Assign (or clear, with blank) who should work ALL open lines of one statement (one grouped
    Review row). `line_ids` is the comma-joined list of statement_lines ids from that row."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    ids = [int(x) for x in line_ids.split(",") if x.strip().isdigit()]
    web_db.assign_statement_lines(ids, assigned_to, user["username"])
    return RedirectResponse("/review", status_code=302)


# Quick fix: which one-click fixes a review item's reason calls for. The fixes themselves are written
# as plain lines into the Answer (compose_quick_fix_answer) and applied in code on the next run
# (classify_emails.apply_reviewer_fixes), so the gate that flagged the item can't flag it again.
_QUICK_FIX_SIGNS = {
    "store": ("store could not be identified", "store rules", "confirm the real store"),
    "date": ("date found",),
    "amount": ("no usable amount", "calculated from line items"),
    "enter": ("marked this bill not for quickbooks",),
}
_QUICK_FIX_LINE = re.compile(r"^\s*(Store|Date|Amount|Enter in QuickBooks):.*$\n?", re.I | re.M)


def quick_fix_kinds(question):
    """The Quick fix controls to show first for this reason text (others sit under 'More fixes')."""
    q = (question or "").lower()
    return [k for k, signs in _QUICK_FIX_SIGNS.items() if any(s in q for s in signs)]


def compose_quick_fix_answer(existing, store="", date="", amount="", enter=False, note="", stores=()):
    """The Answer text for a Quick fix: one line per fix, then the reviewer's note (if any), then any
    earlier free-text answer with its old fix lines removed. Returns (answer, errors)."""
    lines, errors = [], []
    if store:
        if store in stores:
            lines.append(f"Store: {store}")
        else:
            errors.append("Pick a store from the list.")
    if date:
        try:
            datetime.strptime(date, "%Y-%m-%d")
            lines.append(f"Date: {date}")
        except ValueError:
            errors.append("The date isn't valid.")
    if amount:
        try:
            amt = float(str(amount).replace(",", "").replace("$", "").strip())
            if amt <= 0:
                raise ValueError
            lines.append(f"Amount: {amt:.2f} (confirmed)")
        except ValueError:
            errors.append("The amount must be a positive number.")
    if enter:
        lines.append("Enter in QuickBooks: yes")
    if not lines and not errors:
        errors.append("Choose at least one fix.")
    rest = _QUICK_FIX_LINE.sub("", existing or "").strip()
    parts = ["\n".join(lines)] + [p for p in ((note or "").strip(), rest) if p]
    return "\n".join(parts), errors


@app.post("/review/{review_id}/quick-fix")
def review_quick_fix(request: Request, review_id: int, store: str = Form(""), fix_date: str = Form(""),
                     amount: str = Form(""), enter: str = Form(""), note: str = Form("")):
    """Save a Quick fix as this item's Answer (keeping its Notes to Learn). Nothing is reprocessed here;
    the next classify run applies the values in code."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    item = web_db.get_review(review_id)
    if not item:
        return RedirectResponse("/review", status_code=302)
    answer, errors = compose_quick_fix_answer(item.get("answer"), store.strip(), fix_date.strip(),
                                              amount.strip(), bool(enter), note, web_db.store_names())
    if errors:
        return RedirectResponse(f"/review/{review_id}?fix_error={quote(' '.join(errors))}", status_code=302)
    web_db.submit_review_answer(review_id, answer, item.get("notes_to_learn"), user["username"])
    return RedirectResponse("/review", status_code=302)


@app.post("/review/{review_id}/answer")
def review_answer(request: Request, review_id: int,
                  answer: str = Form(...), notes_to_learn: str = Form(""),
                  track_outcome: str = Form("")):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.submit_review_answer(review_id, answer, notes_to_learn, user["username"])
    if track_outcome:                            # opt-in: keep it in the Tracked tab
        web_db.track_review(review_id, answer, user["username"])
    return RedirectResponse("/review", status_code=302)


@app.post("/review/{review_id}/draft-learning")
async def review_draft_learning(request: Request, review_id: int):
    """AI-assist: turn a reviewer's rough note into ONE generalizable learning rule,
    using this record's context. READ-ONLY — saves nothing (the reviewer approves it
    via the Answer form or the save-learning endpoint). Supports multi-turn refinement."""
    user = current_user(request)
    if not user:
        return JSONResponse({"error": "not signed in"}, status_code=401)
    item = web_db.get_review(review_id)
    if not item:
        return JSONResponse({"error": "review item not found"}, status_code=404)
    try:
        body = await request.json()
    except Exception:
        body = {}
    rough = (body.get("rough") or "").strip()
    history = body.get("history") or []
    if not rough and not history:
        return JSONResponse({"error": "Describe the rule in a few words first."}, status_code=400)
    context = {k: item.get(k) for k in ("vendor_matched", "store_matched", "invoice_number",
                                        "amount", "question", "source_email", "email_type")}
    try:
        rules = [r["learned_rule"] for r in web_db.list_learnings()]
    except Exception:
        rules = []
    try:
        # Blocking Anthropic call → threadpool so it never ties up the event loop.
        draft = await run_in_threadpool(
            chat_qa.draft_learning, rough, context,
            history if isinstance(history, list) else [], rules)
    except Exception as e:                       # config/API errors → friendly message, never 500
        return JSONResponse({"error": f"Draft unavailable: {e}"}, status_code=200)
    answer, learning = (draft or {}).get("answer", ""), (draft or {}).get("learning", "")
    if not answer and not learning:
        return JSONResponse({"error": "Couldn't draft — try rephrasing your note."}, status_code=200)
    return JSONResponse({"answer": answer, "learning": learning})


@app.post("/review/{review_id}/save-learning")
async def review_save_learning(request: Request, review_id: int):
    """Approve a drafted rule by saving it to the learnings table NOW — decoupled from
    answering this record (a general rule can be added without resolving the item).
    Same login checks as the other learnings routes."""
    user = current_user(request)
    if not user:
        return JSONResponse({"error": "not signed in"}, status_code=401)
    try:
        body = await request.json()
    except Exception:
        body = {}
    rule = (body.get("rule") or "").strip()
    if not rule:
        return JSONResponse({"error": "Nothing to save."}, status_code=400)
    web_db.add_learning_web(rule, user["username"])
    return JSONResponse({"ok": True})


@app.post("/review/{review_id}/ignore")
def review_ignore(request: Request, review_id: int, reason: str = Form("")):
    """One-click dismiss: snapshot the item + remove it from the queue instantly (no
    reprocess, no API call). Restorable from /ignored. Web-only; CLI untouched."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.ignore_review(review_id, user["username"], reason)
    return RedirectResponse("/review", status_code=302)


@app.post("/review/{review_id}/assign")
def review_assign(request: Request, review_id: int, assigned_to: str = Form("")):
    """Set/clear who should work this review item (assigned_to='' clears it).
    Redirects back to wherever the change was made (list or detail page)."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.assign_review(review_id, assigned_to, user["username"])
    ref = request.headers.get("referer") or "/review"
    return RedirectResponse(ref, status_code=302)


@app.get("/ignored", response_class=HTMLResponse)
def ignored_list(request: Request):
    user = current_user(request)
    if not user:
        return _login_redirect()
    return render(request, "ignored.html", user, items=web_db.list_ignored_reviews())


@app.post("/ignored/{ignored_id}/restore")
def ignored_restore(request: Request, ignored_id: int):
    """Undo an ignore: re-insert the snapshot into the review queue (fresh id)."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.restore_ignored_review(ignored_id, user["username"])
    return RedirectResponse("/ignored", status_code=302)


@app.get("/tracked", response_class=HTMLResponse)
def tracked_list(request: Request):
    """Answered items being outcome-tracked. Each row's outcome is derived live from
    the results tables (no API call), so it updates on its own as classify runs happen."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    return render(request, "tracked.html", user, items=web_db.list_tracked("open"))


@app.post("/tracked/{tracked_id}/done")
def tracked_done(request: Request, tracked_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.mark_tracked_done(tracked_id, user["username"])
    return RedirectResponse("/tracked", status_code=302)


@app.get("/sales-orders", response_class=HTMLResponse)
def sales_orders(request: Request):
    """Sales orders being held for a maturity window before recording (or superseded by
    an invoice / already released). Read-only, plus per-row Release-now / Ignore-now."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    held = web_db.list_sales_orders("held")
    for o in held:
        o["file_preview"] = file_preview_url(o.get("file_path"))
        o["gmail_link"] = fix_gmail_link(o.get("gmail_link"))
    superseded = web_db.list_sales_orders("superseded", limit=100)
    released = web_db.list_sales_orders("released", limit=100)
    return render(request, "sales_orders.html", user,
                  held=held, superseded=superseded, released=released)


@app.post("/sales-orders/{order_id}/release")
def sales_order_release(request: Request, order_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.release_sales_order_now(order_id, user["username"])
    return RedirectResponse("/sales-orders", status_code=302)


@app.post("/sales-orders/{order_id}/ignore")
def sales_order_ignore(request: Request, order_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.ignore_sales_order(order_id, user["username"])
    return RedirectResponse("/sales-orders", status_code=302)


@app.get("/l/{code}")
def short_link(request: Request, code: str):
    """Resolve an internal short link and 302 to its target. 'file' targets are already
    public (Drive / the app's read-only static mounts) so no login; 'email' and 'records'
    are login-required, so remote access stays gated by the app's auth. Unknown → /review."""
    link = web_db.get_link(code)
    if not link or not link.get("target_url"):
        return RedirectResponse("/review", status_code=302)
    if link.get("kind") in ("email", "records") and not current_user(request):
        return _login_redirect()
    target = link["target_url"]
    # Normalize APP-OWNED targets (local file previews + /browse pages) to the host serving
    # THIS request (or LINK_BASE_URL when set) at CLICK time — whether the target was stored
    # app-relative (new mints) OR absolutized under an old host (localhost / a LAN IP, before
    # the mint-relative fix). This makes every code durable across a host/domain move without
    # a data migration. Genuinely external targets (Drive files, Gmail links) are left as-is:
    # their paths don't match our own routes so the regex won't touch them.
    m = re.match(r"^(?:https?://[^/]+)?(/(?:review_pending|payment_files|invoice_files|browse)(?:[/?].*)?)$", target)
    if m:
        base = web_db.LINK_BASE_URL or str(request.base_url).rstrip("/")
        target = base + m.group(1)
    return RedirectResponse(target, status_code=302)


# --------------------------------------------------------------------------- #
# 2) QB entry
# --------------------------------------------------------------------------- #

# The canonical column set + DEFAULT order the accountant enters into QuickBooks. This
# stays the master list of AVAILABLE columns; a per-user pref may pick a SUBSET + order.
# The default export keeps EXACTLY this order (import-ready), and "Standard QB order"
# always emits it regardless of prefs.
# "Data" columns: the configurable/reorderable set — drives the main on-screen columns,
# CSV export, and "Copy visible rows" by default.
QB_DATA_COLUMNS = ["Vendor", "Account", "Store", "Invoice #", "Date", "Amount", "Sales Tax",
                   "Shipping", "Class", "Memo"]
# Columns shown ON SCREEN but withheld from "Copy Selected" + CSV export: the store isn't part of
# the QuickBooks import (store maps to class/location separately), so it's kept visible for context
# but stripped from the copied/exported column set.
QB_COPY_EXPORT_HIDE = {"Store"}
# "Reference" columns: always rendered on-screen at the end of each row for quick access,
# excluded from CSV export / copy unless the user opts a column in (qb_export_include_refs).
QB_REFERENCE_COLUMNS = ["Possible Duplicate", "File Link", "Gmail Link"]
# Column -> sortable ORDER-BY key (only these have clickable sort headers).
QB_SORT_KEY = {"Vendor": "vendor", "Store": "store", "Invoice #": "number", "Date": "date",
               "Amount": "amount", "Sales Tax": "sales_tax", "Shipping": "ship"}
QB_NUMERIC_COLS = {"Amount", "Sales Tax", "Shipping"}
# Column -> the row's data-* attribute name, for the client-side "Copy visible rows".
QB_DATA_ATTR = {"Vendor": "vendor", "Account": "account", "Store": "store", "Invoice #": "number",
                "Date": "date", "Amount": "amount", "Sales Tax": "tax", "Shipping": "ship",
                "Class": "class", "Possible Duplicate": "dup", "File Link": "file",
                "Gmail Link": "gmail", "Memo": "memo"}
def _qb_memo(short_file, short_gmail):
    """One combined, paste-ready memo string. The doc link comes first with no label
    (it's the primary link); the email link is prefixed 'E:' to distinguish it. A missing
    half is dropped; neither -> ''. e.g. 'http://host/l/AAA | E: http://host/l/BBB'."""
    parts = []
    if short_file:
        parts.append(short_file)
    if short_gmail:
        parts.append(f"E: {short_gmail}")
    return " | ".join(parts)


def _sanitize_qb_columns(cols):
    """Keep only known DATA columns, de-duplicated, order preserved. Empty/invalid ->
    the full canonical order (so a user with no/blank prefs sees today's layout)."""
    if not isinstance(cols, list):
        return list(QB_DATA_COLUMNS)
    seen, out = set(), []
    for c in cols:
        if c in QB_DATA_COLUMNS and c not in seen:
            seen.add(c); out.append(c)
    return out or list(QB_DATA_COLUMNS)


def _sanitize_qb_refs(cols):
    """Subset of QB_REFERENCE_COLUMNS the user opted into export/copy, canonical order,
    de-duplicated. Not-a-list or unknown entries silently dropped (default = none)."""
    if not isinstance(cols, list):
        return []
    chosen = set(c for c in cols if c in QB_REFERENCE_COLUMNS)
    return [c for c in QB_REFERENCE_COLUMNS if c in chosen]


def _qb_cell_value(col, r, absfn, shorten=None):
    """The CSV/plain value for one canonical column of a qb-ready row. `shorten(url, kind)`
    (optional) compacts the File/Gmail links to internal /l/<code> short links so they fit
    in a QuickBooks memo — applied only to the two link columns; everything else raw."""
    if col == "Vendor":            return r.get("qb_vendor") or r.get("vendor_matched") or ""
    if col == "Account":           return r.get("qb_account") or ""
    if col == "Store":             return web_db.store_label(r.get("store_matched")) or ""
    if col == "Invoice #":         return r.get("number") or ""
    if col == "Date":              return us_date(r.get("record_date")) or ""
    if col == "Amount":            return "" if r.get("amount") is None else r["amount"]
    if col == "Sales Tax":         return "" if r.get("sales_tax") is None else r["sales_tax"]
    if col == "Shipping":          return "" if r.get("ship_charge") is None else r["ship_charge"]
    if col == "Class":             return r.get("class") or ""
    if col == "Possible Duplicate": return "Yes" if r.get("duplicate_suspect") else ""
    if col == "File Link":
        # Key the short-link lookup on the RAW (app-relative or Drive-absolute) target,
        # matching how it was registered — NOT host-absolutized. If unshortened, still
        # emit an absolute url (absfn) so the CSV cell stays clickable.
        rel = file_preview_url(r.get("file_path"))
        return shorten(rel, "file") if (shorten and rel) else absfn(rel)
    if col == "Gmail Link":
        g = r.get("gmail_link") or ""
        return shorten(g, "email") if (shorten and g) else g
    if col == "Memo":
        rel = file_preview_url(r.get("file_path"))
        sf = shorten(rel, "file") if (shorten and rel) else absfn(rel)
        g = r.get("gmail_link") or ""
        sg = shorten(g, "email") if (shorten and g) else g
        return _qb_memo(sf, sg)
    return ""


PAGE_SIZE_DEFAULT = 150
PAGE_SIZE_OPTIONS = [150, 500]   # selectable sizes; "all" is the escape hatch


def _paginate(rows_all, page, page_size, prefs, user_id, pref_key):
    """Slice `rows_all` (the FULL filtered set) to one page + return pagination meta.
    `page_size` is a string: an explicit choice ('150'/'500'/'all') persists to prefs;
    blank falls back to the user's last choice (default 150). 'all' renders everything
    (the escape hatch). Returns (rows_slice, {total, page, pages, page_size, size})."""
    total = len(rows_all)
    if page_size:                                    # explicit choice -> remember it
        web_db.set_user_pref(user_id, pref_key, page_size)
    else:
        page_size = str(prefs.get(pref_key) or PAGE_SIZE_DEFAULT)
    if page_size == "all":
        return rows_all, {"total": total, "page": 1, "pages": 1, "page_size": "all", "size": total}
    size = int(page_size) if (str(page_size).isdigit() and int(page_size) > 0) else PAGE_SIZE_DEFAULT
    pages = max(1, (total + size - 1) // size)
    try:
        page = int(page)
    except (TypeError, ValueError):
        page = 1
    page = max(1, min(page, pages))
    return (rows_all[(page - 1) * size:page * size],
            {"total": total, "page": page, "pages": pages, "page_size": str(size), "size": size})


@app.get("/qb", response_class=HTMLResponse)
def qb_list(request: Request, entered: str = "no", run: str = "", email: str = "", store: str = "",
            sort: str = "", direction: str = "desc", dup: str = "", ignored: str = "", group: str = "",
            vendor: str = "", search: str = "", date_from: str = "", date_to: str = "",
            amount_min: str = "", amount_max: str = "", amount: str = "", record_type: str = "",
            needs_action: str = "", action: str = "", page: str = "1", page_size: str = "",
            missing: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    prefs = web_db.get_user_prefs(user["id"])
    # Store-focus: remember the last chosen store. An explicit store in the URL (including
    # "" = All stores) is a deliberate choice → remember it. Fresh nav with no store param
    # at all → fall back to the last used store.
    if "store" in request.query_params:
        web_db.set_user_pref(user["id"], "qb_last_store", store)
    else:
        store = prefs.get("qb_last_store") or ""
    rows_all = web_db.list_qb_ready(entered=entered, run_id=run or None,
                                    email_filename=email or None, store=store or None,
                                    sort=sort, direction=direction, dup=dup or None,
                                    ignored=ignored or None, group=group or None,
                                    vendor=vendor or None, search=search or None,
                                    date_from=date_from or None, date_to=date_to or None,
                                    amount_min=amount_min, amount_max=amount_max,
                                    amount=amount or None,
                                    record_type=record_type or None,
                                    needs_action=needs_action or None,
                                    action=action or None, missing=missing or None)
    # Records saved before the store/date/amount review checks existed can still lack one; the
    # banner counts them (all stores, not entered) and links to the missing=1 filter.
    try:
        missing_count = web_db.missing_info_counts()["qb"]
    except Exception:
        missing_count = 0
    # Paginate: render only the current page's rows (lighter DOM); Export CSV still uses the
    # full filtered set. Short links are then built only for the rows actually shown.
    rows, pg = _paginate(rows_all, page, page_size, prefs, user["id"], "qb_page_size")
    # Compact internal short links for the File / Gmail / Records links (so the copied
    # rows + CSV paste small into QB memos). Resolved in ONE db round-trip for all rows
    # (deduped by URL, so codes are stable across loads).
    base = str(request.base_url).rstrip("/")

    def _abs(link):
        return (base + link) if link and link.startswith("/") else (link or "")
    row_urls, pairs = [], []
    for r in rows:
        # Register short-link targets as APP-RELATIVE paths (host resolved fresh at click
        # time in /l/{code}); a Drive/http file url stays absolute (file_preview_url returns
        # those as-is). So the stored target survives a future host/domain change.
        fu = file_preview_url(r.get("file_path"))
        gu = r.get("gmail_link") or ""
        ru = f"/browse?email={quote(r.get('email_filename') or '')}" if r.get("email_filename") else ""
        row_urls.append((fu, gu, ru))
        if fu:
            pairs.append((fu, "file"))
        if gu:
            pairs.append((gu, "email"))
        if ru:
            pairs.append((ru, "records"))
    smap = web_db.short_links_for(pairs, base)
    for r, (fu, gu, ru) in zip(rows, row_urls):
        r["short_file"] = smap.get(fu, "")
        r["short_gmail"] = smap.get(gu, "")
        r["short_records"] = smap.get(ru, "")
        # Combined one-cell memo from the SAME batched short links (no extra queries).
        r["memo"] = _qb_memo(r["short_file"], r["short_gmail"])
    view_columns = _sanitize_qb_columns(prefs.get("qb_view_columns"))
    export_follow = prefs.get("qb_export_follow_view", True)
    export_columns = _sanitize_qb_columns(prefs.get("qb_export_columns"))
    export_include_refs = _sanitize_qb_refs(prefs.get("qb_export_include_refs"))
    # Copy Selected must match what Export CSV actually produces: follow-view uses the on-screen
    # columns, otherwise the user's separate custom export-column layout (same resolution as qb_export).
    copy_columns = [c for c in ((view_columns if export_follow else export_columns) + export_include_refs)
                    if c not in QB_COPY_EXPORT_HIDE]
    return render(request, "qb.html", user, rows=rows, entered=entered,
                  file_url=file_preview_url, runs=web_db.list_runs(), run=run, email=email,
                  store=store, stores=web_db.store_names(), dup=dup,
                  ignored=ignored, group=group,
                  vendor=vendor, search=search, date_from=date_from, date_to=date_to,
                  amount_min=amount_min, amount_max=amount_max, amount=amount, record_type=record_type,
                  needs_action=needs_action, action=action,
                  missing=missing, missing_count=missing_count,
                  record_types=["invoice", "credit_memo"],
                  view_columns=view_columns, all_columns=QB_DATA_COLUMNS,
                  reference_columns=QB_REFERENCE_COLUMNS, export_include_refs=export_include_refs,
                  copy_columns=copy_columns,
                  qb_sort_key=QB_SORT_KEY, qb_numeric=list(QB_NUMERIC_COLS),
                  qb_data_attr=QB_DATA_ATTR,
                  export_follow=export_follow, export_columns=export_columns,
                  pg=pg, page_size_options=PAGE_SIZE_OPTIONS,
                  params=dict(entered=entered, run=run, store=store,
                              email=email, dup=dup, ignored=ignored, group=group,
                              sort=sort, direction=direction,
                              vendor=vendor, search=search, date_from=date_from, date_to=date_to,
                              amount_min=amount_min, amount_max=amount_max, amount=amount,
                              record_type=record_type, needs_action=needs_action,
                              action=action, missing=missing))


@app.post("/qb/columns")
async def qb_columns(request: Request):
    """Persist this user's QB-entry column view (visibility + order) and CSV-export
    preference. 'reset' restores the full canonical order + follow-view. Web-only."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    form = await request.form()
    if form.get("reset"):
        web_db.set_user_pref(user["id"], "qb_view_columns", list(QB_DATA_COLUMNS))
        web_db.set_user_pref(user["id"], "qb_export_follow_view", True)
        web_db.set_user_pref(user["id"], "qb_export_include_refs", [])
        return RedirectResponse(request.headers.get("referer", "/qb"), status_code=302)
    web_db.set_user_pref(user["id"], "qb_view_columns",
                         _sanitize_qb_columns(list(form.getlist("view_col"))))
    web_db.set_user_pref(user["id"], "qb_export_include_refs",
                         _sanitize_qb_refs(list(form.getlist("include_ref"))))
    follow = form.get("export_follow") == "1"
    web_db.set_user_pref(user["id"], "qb_export_follow_view", follow)
    if not follow:
        web_db.set_user_pref(user["id"], "qb_export_columns",
                             _sanitize_qb_columns(list(form.getlist("export_col"))))
    return RedirectResponse(request.headers.get("referer", "/qb"), status_code=302)


def _qb_xlsx_response(cols, table):
    """Same rows/columns as the CSV, as a real .xlsx. Why: Excel opening the CSV turns long
    numeric invoice #s into numbers -- 20751521485459184 (Costco) became 2.07515E+16 and lost its
    last digits (Excel keeps 15 significant digits), and leading zeros vanish. Here every non-number
    column is stored as TEXT ('@'), so Invoice # stays exactly as recorded, and saving this file as
    CSV from Excel keeps it intact. Amounts stay real numbers; Date becomes a real date (m/d/yyyy)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "QB Export"
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for ri, vals in enumerate(table, start=2):
        for ci, (col, v) in enumerate(zip(cols, vals), start=1):
            cell = ws.cell(row=ri, column=ci)
            if col in QB_NUMERIC_COLS and v not in ("", None):
                cell.value = float(v)
                cell.number_format = "#,##0.00"
            elif col == "Date" and v:
                parsed = None
                for fmt in ("%m/%d/%y", "%m/%d/%Y"):   # us_date gives M/D/YY
                    try:
                        parsed = datetime.strptime(str(v), fmt); break
                    except ValueError:
                        pass
                if parsed:
                    cell.value = parsed; cell.number_format = "m/d/yyyy"
                else:
                    cell.value = str(v); cell.number_format = "@"
            else:
                cell.value = "" if v is None else str(v)
                cell.number_format = "@"
    for ci, col in enumerate(cols, start=1):
        width = max([len(str(col))] + [len(str(r[ci - 1] or "")) for r in table[:500]]) + 2
        ws.column_dimensions[get_column_letter(ci)].width = min(max(width, 8), 70)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    fname = f"qb_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(content=out.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@app.get("/qb/export.csv")
@app.get("/qb/export.xlsx")
def qb_export(request: Request, entered: str = "no", run: str = "", email: str = "", store: str = "",
              dup: str = "", ignored: str = "", group: str = "", layout: str = "",
              vendor: str = "", search: str = "", date_from: str = "", date_to: str = "",
              amount_min: str = "", amount_max: str = "", amount: str = "",
              record_type: str = "", needs_action: str = "", action: str = "", missing: str = "",
              ids: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    # SAME filters as the page, so the download is exactly the visible rows.
    rows = web_db.list_qb_ready(entered=entered, run_id=run or None,
                                email_filename=email or None, store=store or None, dup=dup or None,
                                ignored=ignored or None, group=group or None,
                                vendor=vendor or None, search=search or None,
                                date_from=date_from or None, date_to=date_to or None,
                                amount_min=amount_min, amount_max=amount_max,
                                amount=amount or None,
                                record_type=record_type or None,
                                needs_action=needs_action or None,
                                action=action or None, missing=missing or None)
    # If the page passed ticked rows, export ONLY those; no ids = full filtered set (unchanged default).
    if ids:
        wanted = {(rt, int(rid)) for rt, _, rid in
                  (tok.partition(":") for tok in ids.split(","))
                  if rt and rid.isdigit()}
        if wanted:
            rows = [r for r in rows if (r.get("record_type"), r.get("record_id")) in wanted]
    base = str(request.base_url).rstrip("/")

    def _abs(link):  # local preview paths -> absolute so the CSV link is clickable
        return (base + link) if link and link.startswith("/") else (link or "")

    # Column set: 'standard' ignores prefs (always the canonical QB-import order — the
    # safe fallback). Otherwise follow the user's export pref: match the on-screen view,
    # or a separately-configured export layout. A user with no prefs gets the canonical
    # order either way, so today's CSV is unchanged.
    if layout == "standard":
        cols = list(QB_DATA_COLUMNS)
    else:
        prefs = web_db.get_user_prefs(user["id"])
        if prefs.get("qb_export_follow_view", True):
            cols = _sanitize_qb_columns(prefs.get("qb_view_columns"))
        else:
            cols = _sanitize_qb_columns(prefs.get("qb_export_columns"))
        cols = cols + _sanitize_qb_refs(prefs.get("qb_export_include_refs"))   # opted-in refs

    # The store isn't part of the QuickBooks import — drop it from the CSV (same as Copy Selected).
    cols = [c for c in cols if c not in QB_COPY_EXPORT_HIDE]

    # Compact the File/Gmail links to internal short links so they fit in QB memos —
    # resolve them ALL in one db round-trip, then look them up per cell.
    link_pairs = []
    for r in rows:
        fu = file_preview_url(r.get("file_path"))   # app-relative (or Drive-absolute) — not host-baked
        if fu:
            link_pairs.append((fu, "file"))
        if r.get("gmail_link"):
            link_pairs.append((r["gmail_link"], "email"))
    smap = web_db.short_links_for(link_pairs, base)

    def _short(url, kind):
        return smap.get(url, url)   # fall back to the raw url if somehow unresolved

    table = [[_qb_cell_value(c, r, _abs, _short) for c in cols] for r in rows]
    if request.url.path.endswith(".xlsx"):
        return _qb_xlsx_response(cols, table)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cols)
    for vals in table:
        w.writerow(vals)
    fname = f"qb_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    # Lead with a UTF-8 BOM so Excel opens it as UTF-8 (without it Excel assumes ANSI/cp1252 and
    # mangles non-ASCII like "·" into "Â·").
    return Response(content="﻿" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@app.post("/qb/{record_type}/{record_id}/entered")
def qb_toggle(request: Request, record_type: str, record_id: int, entered: str = Form("")):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.set_entered_in_qb(record_type, record_id, entered == "1", user["username"])
    return RedirectResponse(request.headers.get("referer", "/qb"), status_code=302)


@app.post("/records/{record_type}/{record_id}/ignore")
def record_ignore_toggle(request: Request, record_type: str, record_id: int,
                         ignored: str = Form("")):
    """Manual per-record Ignore / Restore -- excludes (or re-includes) this one record from
    QuickBooks import, independent of vendor rules or duplicate detection. Available on any record
    type, wherever the record is listed."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    if record_type not in ("invoice", "credit_memo", "payment"):
        return RedirectResponse(request.headers.get("referer", "/browse"), status_code=302)
    web_db.set_record_ignored(record_type, record_id, ignored == "1", user["username"])
    return RedirectResponse(request.headers.get("referer", "/browse"), status_code=302)


@app.post("/records/{record_type}/{record_id}/duplicate")
def record_duplicate_toggle(request: Request, record_type: str, record_id: int,
                            action: str = Form(...), duplicate_of: str = Form("")):
    """Manual duplicate management (invoices/payments only):
      action='ignore'  -> mark a confirmed duplicate, excluded from QB (duplicate_of optional:
                          the email_filename of the record it duplicates, for the browse ↗ link).
      action='restore' -> clear the ignore (back to a normal record).
      action='clear'   -> 'not actually a duplicate': drop the possible-duplicate suspicion.
    Redirects back to wherever the action came from (browse / detail / duplicates page)."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    back = request.headers.get("referer", "/browse")
    if record_type not in ("invoice", "payment"):
        return RedirectResponse(back, status_code=302)
    if action == "clear":
        web_db.clear_duplicate_flag(record_type, record_id, user["username"])
    else:
        web_db.set_duplicate_ignored(record_type, record_id, action == "ignore",
                                     user["username"], duplicate_of or None)
    return RedirectResponse(back, status_code=302)


@app.post("/records/{record_type}/{record_id}/duplicate/keep")
async def record_duplicate_keep(request: Request, record_type: str, record_id: int):
    """Resolve a whole duplicate set in one click: keep this record, ignore every
    other one listed in the 'others' form field (record_type:record_id tokens)."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    back = request.headers.get("referer", "/duplicates")
    if record_type not in ("invoice", "payment"):
        return RedirectResponse(back, status_code=302)
    form = await request.form()
    others = []
    for token in form.getlist("others"):
        rtype, _, rid = str(token).partition(":")
        if rtype in ("invoice", "payment") and rid.isdigit():
            others.append((rtype, int(rid)))
    web_db.keep_and_ignore_others(record_type, record_id, others, user["username"])
    return RedirectResponse(back, status_code=302)


@app.get("/duplicates", response_class=HTMLResponse)
def duplicates_page(request: Request, flagged: str = ""):
    """The working queue for duplicate decisions: every still-flagged (duplicate_suspect,
    not-ignored) invoice/payment, paired with its best-guess counterpart."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    return render(request, "duplicates.html", user,
                  dup_groups=web_db.list_duplicate_candidates(), file_url=file_preview_url,
                  flagged=bool(flagged))


@app.post("/duplicates/flag_set")
async def duplicates_flag_set(request: Request):
    """Flag a whole possible-duplicate set for review in one entry (not one flag per
    record) — see web_db.add_duplicate_set_flag for why: the investigating bot has no
    DB access, so every record's data has to be embedded in the ONE snapshot it gets."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    form = await request.form()
    note = (form.get("note") or "").strip()
    back = request.headers.get("referer", "/duplicates")
    if not note:
        return RedirectResponse(back, status_code=302)
    records = []
    for token in form.getlist("records"):
        rtype, _, rid = str(token).partition(":")
        if rtype in ("invoice", "payment") and rid.isdigit():
            records.append((rtype, int(rid)))
    try:
        web_db.add_duplicate_set_flag(records, note, user["username"])
        flagged = True
    except ValueError:
        flagged = False
    return RedirectResponse(f"{back}{'?flagged=1' if flagged else ''}", status_code=302)


@app.post("/qb/bulk_entered")
async def qb_bulk_entered(request: Request):
    """Mark every selected row as Entered in QB in one click. Each checkbox value is
    'record_type:record_id'; we reuse set_entered_in_qb (sets entered_in_qb +
    entered_qb_by/at) per row, so the audit stamp is identical to the single toggle."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    form = await request.form()
    for token in form.getlist("selected"):
        rtype, _, rid = str(token).partition(":")
        if rtype in ("invoice", "credit_memo", "payment") and rid.isdigit():
            web_db.set_entered_in_qb(rtype, int(rid), True, user["username"])
    return RedirectResponse(request.headers.get("referer", "/qb"), status_code=302)


# --------------------------------------------------------------------------- #
# 2c) QuickBooks reconciliation import (qb_match.py). Upload a QB register export
#     before a QB-entry session: auto-mark records already in QB (so they aren't
#     entered twice), flag doubtful matches for a human, and resolve 'missing'
#     statement lines QB actually has. Reusable — re-run every entry session.
# --------------------------------------------------------------------------- #

@app.get("/qb-reconcile", response_class=HTMLResponse)
def qb_reconcile(request: Request):
    user = current_user(request)
    if not user:
        return _login_redirect()
    return render(request, "qb_reconcile.html", user, batches=web_db.qb_list_batches())


@app.post("/qb-reconcile/upload")
async def qb_reconcile_upload(request: Request, file: UploadFile = File(...)):
    user = current_user(request)
    if not user:
        return _login_redirect()
    name = (file.filename or "").strip()
    if not name.lower().endswith(".xlsx"):
        return RedirectResponse("/qb-reconcile?flash=Please%20upload%20an%20.xlsx%20QB%20export",
                                status_code=302)
    data = await file.read()
    if not data:
        return RedirectResponse("/qb-reconcile?flash=Empty%20file", status_code=302)
    # Blocking (parse + DB matching) → run off the event loop.
    batch_id = await run_in_threadpool(qb_match.process_qb_export, data, name, user["username"])
    return RedirectResponse(f"/qb-reconcile/review?batch_id={batch_id}", status_code=302)


@app.get("/qb-reconcile/review", response_class=HTMLResponse)
def qb_reconcile_review(request: Request, batch_id: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    bid = int(batch_id) if batch_id.isdigit() else None
    return render(request, "qb_reconcile_review.html", user,
                  candidates=web_db.qb_list_candidates(batch_id=bid),
                  batch=web_db.qb_batch(bid) if bid else None)


@app.post("/qb-reconcile/review/recheck")
async def qb_reconcile_recheck(request: Request):
    """Re-evaluate every still-pending doubtful candidate against the CURRENT vendor list (e.g. after
    adding a spelling alias) — resolves Tier-1 matches and drops stale ones without a re-upload."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    msg = await run_in_threadpool(qb_match.recheck_pending_candidates, user["username"])
    return RedirectResponse(f"/qb-reconcile/review?msg={quote(msg)}", status_code=302)


@app.post("/qb-reconcile/review/{candidate_id}/confirm")
def qb_reconcile_confirm(request: Request, candidate_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.qb_confirm_candidate(candidate_id, user["username"])
    return RedirectResponse(request.headers.get("referer") or "/qb-reconcile/review", status_code=302)


@app.post("/qb-reconcile/review/{candidate_id}/reject")
def qb_reconcile_reject(request: Request, candidate_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.qb_reject_candidate(candidate_id, user["username"])
    return RedirectResponse(request.headers.get("referer") or "/qb-reconcile/review", status_code=302)


@app.get("/qb-reconcile/unmatched", response_class=HTMLResponse)
def qb_reconcile_unmatched(request: Request, status: str = "pending"):
    """QB export rows that matched no app record at upload — they wait here and resolve
    automatically when the matching email arrives. Manual dismiss for ones that never will."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    status = status if status in ("pending", "resolved", "dismissed", "all") else "pending"
    return render(request, "qb_reconcile_unmatched.html", user,
                  rows=web_db.qb_list_unmatched(status=status), status=status,
                  pending_count=web_db.qb_unmatched_pending_count())


@app.post("/qb-reconcile/unmatched/{unmatched_id}/dismiss")
def qb_reconcile_unmatched_dismiss(request: Request, unmatched_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.qb_dismiss_unmatched(unmatched_id, user["username"])
    return RedirectResponse(request.headers.get("referer") or "/qb-reconcile/unmatched", status_code=302)


@app.post("/qb-reconcile/unmatched/{unmatched_id}/restore")
def qb_reconcile_unmatched_restore(request: Request, unmatched_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.qb_restore_unmatched(unmatched_id, user["username"])
    return RedirectResponse(request.headers.get("referer") or "/qb-reconcile/unmatched", status_code=302)


# --------------------------------------------------------------------------- #
# 2b) Payments — dedicated tab. Normally every payment is QB-entered, but a
#     reference-only payment (skip_qb_import=True — e.g. a per-invoice remittance
#     breakdown that repeats cash already captured by the real bank receipt) is
#     HIDDEN from the to-enter queue by default, reachable via the 'skipped' toggle.
#     Reuses the QB toggle + bulk-entered routes (record_type='payment').
# --------------------------------------------------------------------------- #

# Payments search fields (2026-10-12): query-string name -> list_payments keyword.
PAYMENT_SEARCH_FIELDS = {"q": "q", "vendor": "vendor", "reference": "reference", "applied": "applied",
                         "amount_min": "amount_min", "amount_max": "amount_max",
                         "added_from": "added_from", "added_to": "added_to"}


def _payment_search(request):
    """The Payments search fields from the query string: (list_payments kwargs, params for the template)."""
    vals = {k: (request.query_params.get(k) or "").strip() for k in PAYMENT_SEARCH_FIELDS}
    return {PAYMENT_SEARCH_FIELDS[k]: v or None for k, v in vals.items()}, vals


@app.get("/payments", response_class=HTMLResponse)
def payments_page(request: Request, store: str = "", entered: str = "", run: str = "",
                  date_from: str = "", date_to: str = "", sort: str = "", direction: str = "desc",
                  dup: str = "", ignored: str = "", group: str = "", skipped: str = "",
                  page: str = "1", page_size: str = "", missing: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    prefs = web_db.get_user_prefs(user["id"])
    srt = sort or "date_added"
    search, search_params = _payment_search(request)
    rows_all = web_db.list_payments(store=store or None, entered=entered or None, run_id=run or None,
                                    date_from=date_from or None, date_to=date_to or None,
                                    sort=srt, direction=direction, dup=dup or None,
                                    show_ignored=(ignored == "1"), group=group or None,
                                    show_skipped=(skipped == "1"), missing=missing or None, **search)
    rows, pg = _paginate(rows_all, page, page_size, prefs, user["id"], "payments_page_size")
    try:
        missing_count = web_db.missing_info_counts()["payments"]
    except Exception:
        missing_count = 0
    return render(request, "payments.html", user, rows=rows, file_url=file_preview_url,
                  runs=web_db.list_runs(), stores=web_db.store_names(), dup=dup,
                  group=group, pg=pg, page_size_options=PAGE_SIZE_OPTIONS, skipped=skipped,
                  missing_count=missing_count, vendors=web_db.payment_vendor_names(),
                  search_keys=list(PAYMENT_SEARCH_FIELDS),
                  params=dict(store=store, entered=entered, run=run, date_from=date_from,
                              date_to=date_to, dup=dup, ignored=ignored, group=group,
                              skipped=skipped, sort=srt, direction=direction, missing=missing,
                              **search_params))


# Column order for the Payments CSV export / "Copy visible rows" (kept in sync with payments.html).
PAYMENT_EXPORT_COLUMNS = ["Vendor", "Store", "Amount", "Payment Date", "Reference / Method",
                          "Possible Duplicate", "File Link", "Gmail Link"]


@app.get("/payments/export.csv")
def payments_export(request: Request, store: str = "", entered: str = "", run: str = "",
                    date_from: str = "", date_to: str = "", sort: str = "", direction: str = "desc",
                    dup: str = "", ignored: str = "", group: str = "", skipped: str = "",
                    missing: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    rows = web_db.list_payments(store=store or None, entered=entered or None, run_id=run or None,
                                date_from=date_from or None, date_to=date_to or None,
                                sort=sort or "date_added", direction=direction, dup=dup or None,
                                show_ignored=(ignored == "1"), group=group or None,
                                show_skipped=(skipped == "1"), missing=missing or None,
                                **_payment_search(request)[0])
    base = str(request.base_url).rstrip("/")

    def _abs(link):
        return (base + link) if link and link.startswith("/") else (link or "")

    def _dup_label(r):
        return "Duplicate — ignored" if r.get("duplicate_ignored") else ("Possible" if r.get("duplicate_suspect") else "")

    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(PAYMENT_EXPORT_COLUMNS)
    for r in rows:
        w.writerow([
            r.get("vendor_matched") or "", r.get("store_matched") or "",
            "" if r.get("amount") is None else r["amount"],
            us_date(r.get("payment_date")) or "", r.get("reference_number") or "", _dup_label(r),
            _abs(file_preview_url(r.get("file_path"))), r.get("gmail_link") or "",
        ])
    fname = f"payments_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    # UTF-8 BOM so Excel opens it as UTF-8 (see qb_export note).
    return Response(content="﻿" + buf.getvalue(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="{fname}"'})


# --------------------------------------------------------------------------- #
# 2c) Chat Q&A — READ-ONLY, grounded in the DB via Claude tool-use (chat_qa.py).
# --------------------------------------------------------------------------- #

@app.get("/chat", response_class=HTMLResponse)
def chat_page(request: Request):
    user = current_user(request)
    if not user:
        return _login_redirect()
    return render(request, "chat.html", user)


@app.post("/chat/ask")
async def chat_ask(request: Request):
    user = current_user(request)
    if not user:
        return JSONResponse({"error": "not signed in"}, status_code=401)
    try:
        body = await request.json()
    except Exception:
        body = {}
    question = (body.get("question") or "").strip()
    history = body.get("history") or []
    if not question:
        return JSONResponse({"answer": "Ask a question about the invoices, payments, or review items."})
    try:
        # chat_qa.ask() is fully synchronous (blocking Anthropic + DB calls); run it in a
        # threadpool so it never ties up the async event loop and hangs the request.
        result = await run_in_threadpool(
            chat_qa.ask, question, history if isinstance(history, list) else [])
    except Exception as e:                       # config/API errors → friendly message, never 500
        return JSONResponse({"answer": f"Chat is unavailable: {e}", "tools_used": []})
    return JSONResponse(result)


# --------------------------------------------------------------------------- #
# 3) Edit fields
# --------------------------------------------------------------------------- #

@app.get("/edit/{record_type}/{record_id}", response_class=HTMLResponse)
def edit_form(request: Request, record_type: str, record_id: int, saved: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    if record_type not in web_db.EDITABLE:
        return RedirectResponse("/browse", status_code=302)
    record = web_db.get_record(record_type, record_id)
    if not record:
        return RedirectResponse("/browse", status_code=302)
    return render(request, "edit.html", user, record=record, record_type=record_type,
                  fields=web_db.EDITABLE[record_type], log=web_db.get_edit_log(record_type, record_id),
                  preview=file_preview_url(record.get("file_path")), saved=saved,
                  stores=web_db.store_names())


@app.post("/edit/{record_type}/{record_id}")
async def edit_submit(request: Request, record_type: str, record_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    if record_type not in web_db.EDITABLE:
        return RedirectResponse("/browse", status_code=302)
    form = await request.form()
    record = web_db.get_record(record_type, record_id)
    # Update only fields that actually changed, so the edit_log stays meaningful.
    for ui_field, col in web_db.EDITABLE[record_type].items():
        if ui_field not in form:
            continue
        new_val = (form.get(ui_field) or "").strip()
        old_val = record.get(col)
        old_str = "" if old_val is None else str(old_val)
        if new_val != old_str:
            web_db.update_record_field(record_type, record_id, ui_field, new_val, user["username"])
    _sync_account_class_rules("record edit")
    return RedirectResponse(f"/edit/{record_type}/{record_id}?saved=1", status_code=302)


# --------------------------------------------------------------------------- #
# 3b) Per-record "how it was processed" detail view (read-only). Surfaces
#     sender / received / processed / run + a derived outcome from data that
#     already exists (via web_db.get_record's enriched row + the `outcome` macro).
# --------------------------------------------------------------------------- #

_DETAIL_NUM = {"invoice": "invoice_number", "credit_memo": "credit_memo_number",
               "payment": "reference_number", "review": "invoice_number"}
_DETAIL_DATE = {"invoice": "invoice_date", "credit_memo": "credit_memo_date",
                "payment": "payment_date", "review": "date_flagged"}


def _record_detail_view(record_type, r):
    """Normalize a raw enriched record row into uniform keys the template + the
    `outcome` macro use (per-type number/date column names differ)."""
    return {
        "record_type": record_type, "id": r.get("id"),
        "vendor": r.get("vendor_matched"), "store": r.get("store_matched"),
        # QuickBooks-bound vendor (standardized) + the standardized name it resolved to. Only
        # invoices/credit_memos carry these columns; payments/review rows return None. Surfaced
        # so the detail page can show the QB vendor and flag when it diverges from the matched one.
        "standard_vendor": r.get("standard_vendor"), "qb_vendor": r.get("qb_vendor"),
        # Warn only when the matched name isn't a spelling Setup lists for that standard vendor (a fuzzy
        # guess, or a stale mapping) — not merely because the two names aren't identical (2026-10-12).
        "vendor_mapping_doubt": (record_type in ("invoice", "credit_memo")
                                 and web_db.vendor_mapping_doubt(r.get("vendor_matched"), r.get("standard_vendor"))),
        "number": r.get(_DETAIL_NUM[record_type]), "amount": r.get("amount"),
        "record_date": r.get(_DETAIL_DATE[record_type]), "klass": r.get("email_type"),
        "sender": r.get("sender"), "received_at": r.get("received_at"),
        "date_added": r.get("date_added") or r.get("date_flagged"),
        "run_id": r.get("run_id"), "run_started_at": r.get("run_started_at"),
        "run_finished_at": r.get("run_finished_at"), "email_status": r.get("email_status"),
        "email_filename": r.get("email_filename") or r.get("source_email") or r.get("filename"),
        "file_path": r.get("file_path"), "gmail_link": r.get("gmail_link"),
        "skip_qb_import": r.get("skip_qb_import"), "needs_review": r.get("needs_review"),
        # Manual-ignore audit stamps (invoices/credit_memos/payments); None on review rows. Lets the
        # detail page + outcome() distinguish a manual ignore from a vendor-rule skip and offer Restore.
        "ignored_by": r.get("ignored_by"), "ignored_at": r.get("ignored_at"),
        "duplicate_suspect": r.get("duplicate_suspect"), "duplicate_ignored": r.get("duplicate_ignored"),
        "duplicate_of": r.get("duplicate_of"), "duplicate_reason": r.get("duplicate_reason"),
        "review_reason": r.get("review_reason") or r.get("question"),
        "entered_in_qb": r.get("entered_in_qb"), "notes_field": r.get("notes"),
        # 'stated' / 'derived' / None (records saved before 2026-10-08, or the model didn't say).
        "total_source": r.get("total_source"),
        # The invoice a payment settles, when the remittance states it (payments.invoice_number).
        "applied_invoice": r.get("invoice_number") if record_type == "payment" else None,
    }


@app.get("/records/{record_type}/{record_id}", response_class=HTMLResponse)
def record_detail(request: Request, record_type: str, record_id: int, flagged: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    if record_type not in ("invoice", "credit_memo", "payment", "review"):
        return RedirectResponse("/browse", status_code=302)
    raw = web_db.get_record(record_type, record_id)
    if not raw:
        return RedirectResponse("/browse", status_code=302)
    rec = _record_detail_view(record_type, raw)
    rec["gmail_link"] = fix_gmail_link(rec.get("gmail_link"))     # repair stale /u/0/ links
    note_key = ((raw.get("review_key") or raw.get("filename")) if record_type == "review"
                else raw.get("email_filename"))
    return render(request, "record_detail.html", user, rec=rec,
                  preview=file_preview_url(rec.get("file_path")),
                  notes=web_db.list_review_notes(note_key), flagged=flagged)


# --------------------------------------------------------------------------- #
# 3b) Flag for review (any user) + admin /flags queue + admin-triggered
#     read-only automated investigation. See investigate.py for the runner.
# --------------------------------------------------------------------------- #

@app.post("/records/{record_type}/{record_id}/flag")
async def flag_record(request: Request, record_type: str, record_id: int):
    """Any authenticated user flags a record for review with a free-text note."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    if record_type not in ("invoice", "credit_memo", "payment", "review"):
        return RedirectResponse("/browse", status_code=302)
    form = await request.form()
    note = (form.get("note") or "").strip()
    try:
        web_db.add_flag(record_type, record_id, note, user["username"])
    except ValueError:
        return RedirectResponse("/browse", status_code=302)
    # Review-queue items live at /review/<id>; there is no /records/review/<id> page.
    dest = f"/review/{record_id}" if record_type == "review" else f"/records/{record_type}/{record_id}"
    return RedirectResponse(f"{dest}?flagged=1", status_code=302)


@app.get("/flags", response_class=HTMLResponse)
def flags_page(request: Request):
    """Admin: every flag + its latest investigation. Non-admin: only flags they submitted."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    is_admin = user["role"] == "admin"
    flags = web_db.list_flags(username=user["username"], is_admin=is_admin)
    # Pre-build the copy-paste fix prompt for every approved flag so the /flags page can offer a
    # re-copyable "Copy fix prompt" button (pure text — triggers nothing).
    for f in flags:
        if f.get("status") == "approved_for_fix":
            f["fix_prompt"] = web_db.build_fix_prompt(f)
    import investigate
    return render(request, "flags.html", user, flags=flags, is_admin=is_admin,
                  busy=web_db.any_investigation_running(),
                  max_turns=investigate.MAX_TURNS)


@app.post("/flags/{flag_id}/dismiss")
def flag_dismiss(request: Request, flag_id: int):
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    web_db.set_flag_status(flag_id, "dismissed", user["username"])
    return RedirectResponse("/flags", status_code=302)


@app.post("/flags/{flag_id}/approve")
def flag_approve(request: Request, flag_id: int):
    """Label the flag 'approved_for_fix' — a status only. This phase does NOT trigger any
    automated fix; Joel still handles fixes manually."""
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    web_db.set_flag_status(flag_id, "approved_for_fix", user["username"])
    return RedirectResponse("/flags", status_code=302)


@app.post("/flags/{flag_id}/investigate")
def flag_investigate(request: Request, flag_id: int):
    """Admin-only: kick off ONE read-only Claude Code investigation of this flag (queued
    one-at-a-time). Returns immediately; the /flags page shows running -> completed/failed."""
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    import investigate
    ok, msg = investigate.enqueue("flag", flag_id, user["username"])
    return RedirectResponse(f"/flags?msg={quote(msg)}", status_code=302)


# --------------------------------------------------------------------------- #
# 3b2) Change requests — any user asks for an app change/feature; same lifecycle
#      as flags (pending -> dismiss/approve -> optional investigation -> build ->
#      done), but about the app itself, not an accounting record.
# --------------------------------------------------------------------------- #

@app.get("/requests", response_class=HTMLResponse)
def requests_page(request: Request):
    """Submission form + list. Admin: every request + its latest investigation. Non-admin: only
    requests they submitted."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    is_admin = user["role"] == "admin"
    reqs = web_db.list_change_requests(username=user["username"], is_admin=is_admin)
    for r in reqs:
        if r.get("status") == "approved_for_build":
            r["build_prompt"] = web_db.build_request_prompt(r)   # pure text; triggers nothing
    import investigate
    return render(request, "requests.html", user, requests=reqs, is_admin=is_admin,
                  busy=web_db.any_investigation_running(), max_turns=investigate.MAX_TURNS)


@app.post("/requests/submit")
async def request_submit(request: Request, title: str = Form(...), description: str = Form(...),
                         category: str = Form("")):
    user = current_user(request)
    if not user:
        return _login_redirect()
    if (title or "").strip() and (description or "").strip():
        web_db.add_change_request(title, description, category, user["username"])
    return RedirectResponse("/requests", status_code=302)


@app.post("/requests/{request_id}/dismiss")
def request_dismiss(request: Request, request_id: int):
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    web_db.set_request_status(request_id, "dismissed", user["username"])
    return RedirectResponse("/requests", status_code=302)


@app.post("/requests/{request_id}/approve")
def request_approve(request: Request, request_id: int):
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    web_db.set_request_status(request_id, "approved_for_build", user["username"])
    return RedirectResponse("/requests", status_code=302)


@app.post("/requests/{request_id}/done")
def request_done(request: Request, request_id: int):
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    web_db.set_request_status(request_id, "done", user["username"])
    return RedirectResponse("/requests", status_code=302)


@app.post("/requests/{request_id}/investigate")
def request_investigate(request: Request, request_id: int):
    """Admin-only: kick off ONE read-only Claude Code investigation of this request."""
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    import investigate
    ok, msg = investigate.enqueue("request", request_id, user["username"])
    return RedirectResponse(f"/requests?msg={quote(msg)}", status_code=302)


# --------------------------------------------------------------------------- #
# 3c) Emails ledger — every ingested email + what the system did with it.
# --------------------------------------------------------------------------- #

@app.get("/emails", response_class=HTMLResponse)
def emails_page(request: Request, disposition: str = "", email_type: str = "", run: str = "",
                date_from: str = "", date_to: str = "", search: str = "", page: int = 1):
    user = current_user(request)
    if not user:
        return _login_redirect()
    page_size = 200
    rows, total = web_db.list_emails(
        disposition=disposition or None, email_type=email_type or None, run_id=run or None,
        date_from=date_from or None, date_to=date_to or None, search=search or None,
        page=page, page_size=page_size)
    for r in rows:                          # repair any stale /u/0/ Gmail links at display time
        r["gmail_link"] = fix_gmail_link(r.get("gmail_link"))
    pages = max(1, (total + page_size - 1) // page_size)
    return render(request, "emails.html", user, rows=rows, total=total, page=page, pages=pages,
                  dispositions=web_db.EMAIL_DISPOSITIONS, email_types=web_db._email_types(),
                  runs=web_db.list_runs(),
                  params=dict(disposition=disposition, email_type=email_type, run=run,
                              date_from=date_from, date_to=date_to, search=search, page=page))


# --------------------------------------------------------------------------- #
# 4) Records — merged Browse + History ledger under one tab (source toggle).
#    Records = the recorded invoices/credit-memos/payments (+ in-review items) from
#    all_records; History = the read-only invoice_history dedup ledger. Both keep
#    their full column + filter sets, plus a shared Excluded-vendors filter.
# --------------------------------------------------------------------------- #

@app.get("/browse", response_class=HTMLResponse)
def browse(request: Request, source: str = "records", excluded: str = "exclude",
           store: str = "", vendor: str = "", record_type: str = "", status: str = "",
           entered: str = "", date_from: str = "", date_to: str = "", search: str = "",
           run: str = "", email: str = "", dup: str = "", ignored: str = "", group: str = "",
           invoice_number: str = "", amount_min: str = "", amount_max: str = "", amount: str = "",
           page: int = 1, sort: str = "", direction: str = "desc"):
    user = current_user(request)
    if not user:
        return _login_redirect()
    if excluded not in ("include", "exclude", "only"):
        excluded = "exclude"
    stores = web_db.store_names()

    if source == "history":
        page_size = 50
        srt = sort or "date_added"
        rows, total = web_db.list_invoice_history(
            vendor=vendor or None, invoice_number=invoice_number or None, store=store or None,
            date_from=date_from or None, date_to=date_to or None,
            amount_min=amount_min or None, amount_max=amount_max or None, amount=amount or None,
            sort=srt, direction=direction, page=page, page_size=page_size, excluded=excluded)
        pages = max(1, (total + page_size - 1) // page_size)
        params = dict(source="history", excluded=excluded, vendor=vendor,
                      invoice_number=invoice_number, store=store, date_from=date_from,
                      date_to=date_to, amount_min=amount_min, amount_max=amount_max, amount=amount,
                      sort=srt, direction=direction, page=page)
        return render(request, "browse.html", user, source="history", rows=rows, total=total,
                      page=page, pages=pages, excluded=excluded, file_url=file_preview_url,
                      stores=stores, params=params)

    srt = sort or "date_added"
    rows = web_db.list_all_records(
        store=store or None, vendor=vendor or None, record_type=record_type or None,
        status=status or None, entered=entered or None, date_from=date_from or None,
        date_to=date_to or None, search=search or None, sort=srt, direction=direction,
        run_id=run or None, email_filename=email or None, excluded=excluded,
        dup=dup or None, ignored=ignored or None, group=group or None,
        amount_min=amount_min or None, amount_max=amount_max or None, amount=amount or None)
    params = dict(source="records", excluded=excluded, store=store, vendor=vendor,
                  record_type=record_type, status=status, entered=entered, date_from=date_from,
                  date_to=date_to, search=search, run=run, email=email, dup=dup, ignored=ignored,
                  group=group, amount_min=amount_min, amount_max=amount_max, amount=amount,
                  sort=srt, direction=direction)
    return render(request, "browse.html", user, source="records", rows=rows, excluded=excluded,
                  file_url=file_preview_url, stores=stores,
                  statuses=web_db.distinct_values("email_status"), runs=web_db.list_runs(),
                  params=params)


@app.get("/history")
def history_redirect(request: Request):
    """Old History tab merged into Records — keep the URL working."""
    return RedirectResponse("/browse?source=history", status_code=302)


# --------------------------------------------------------------------------- #
# 5) Learnings editor  (same `learnings` table the pipeline injects each run)
# --------------------------------------------------------------------------- #

@app.get("/learnings", response_class=HTMLResponse)
def learnings_page(request: Request):
    user = current_user(request)
    if not user:
        return _login_redirect()
    return render(request, "learnings.html", user, items=web_db.list_learnings())


@app.post("/learnings/add")
def learnings_add(request: Request, rule: str = Form(...)):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.add_learning_web(rule, user["username"])
    return RedirectResponse("/learnings", status_code=302)


@app.post("/learnings/edit")
def learnings_edit(request: Request, old_rule: str = Form(...), rule: str = Form(...)):
    user = current_user(request)
    if not user:
        return _login_redirect()
    web_db.edit_learning_web(old_rule, rule, user["username"])
    return RedirectResponse("/learnings", status_code=302)


@app.post("/learnings/delete")
def learnings_delete(request: Request, rule: str = Form(...)):
    user = current_user(request)
    if not user:
        return _login_redirect()
    if user["role"] != "admin":            # deletes are admin-only
        return render(request, "message.html", user, message="Deleting learnings is admin-only.")
    web_db.delete_learning_web(rule, user["username"])
    return RedirectResponse("/learnings", status_code=302)


# --------------------------------------------------------------------------- #
# 5b) Reference data editors (stores / vendor names / sender_store_map / …).
# These are the config lists the pipeline used to hardcode; edited here, the
# next classify_emails.py run reads them. One generic page/route family, driven
# by REFERENCE_PAGES. Deletes are admin-only + confirmed (see the template).
# --------------------------------------------------------------------------- #

REFERENCE_PAGES = {
    "stores": {
        "title": "Stores",
        "help": ("The pharmacy entities. The classifier rebuilds its store list "
                 "(names, aliases/DBAs, addresses, disambiguation notes) from this table "
                 "on every run."),
        "columns": [
            {"name": "store_name", "label": "Store", "type": "text", "required": True},
            {"name": "display_name", "label": "Display name (shown instead of Store, if set)", "type": "text"},
            {"name": "aliases", "label": "Aliases / DBAs / legal name", "type": "textarea"},
            {"name": "address", "label": "Address", "type": "text"},
            {"name": "notes", "label": "Notes (extra guidance)", "type": "textarea"},
        ],
    },
    "vendor_name_map": {
        "title": "Vendor names",
        "help": ("Standardized vendor name (column A), any OTHER spellings/variants "
                 "of that same vendor (one per line — not tied to any specific "
                 "store, e.g. an OCR misread or an older letterhead name), plus "
                 "each store's own spelling of that vendor for QuickBooks entry. "
                 "Blank a cell if that vendor doesn't bill that entity. BMW / "
                 "Lakewood Rx columns are blank for now — fill in as needed. "
                 "The two Exclude checkboxes replace the old Excluded Vendors / "
                 "Statement-recon lists: “Exclude from QB” keeps this vendor's "
                 "invoices/credit memos out of QuickBooks (skip_qb_import; payments "
                 "still enter); “Exclude from statement recon” only silences "
                 "statement-reconciliation missing/mismatch noise. Exclusion matches "
                 "this vendor's standard name AND every alias/store spelling. "
                 "“Trust line-item totals” is for vendors whose bills only list line "
                 "items (no printed total): an amount added up from the line items is "
                 "accepted instead of sent to review (a missing or zero amount still is). "
                 "It applies to documents read after you tick it, matched on this "
                 "vendor's exact spellings."),
        "columns": [
            {"name": "standard_vendor_name", "label": "Standard vendor", "type": "text", "required": True},
            {"name": "exclude_from_qb", "label": "Exclude from QB", "type": "checkbox"},
            {"name": "exclude_from_statement_recon", "label": "Exclude from stmt recon", "type": "checkbox"},
            {"name": "trust_line_item_totals", "label": "Trust line-item totals", "type": "checkbox"},
            {"name": "aliases", "label": "Other spellings / variants", "type": "textarea"},
            {"name": "name_3m", "label": "3M", "type": "text"},
            {"name": "name_madison", "label": "Madison", "type": "text"},
            {"name": "name_maple", "label": "Maple", "type": "text"},
            {"name": "name_refuah", "label": "Refuah", "type": "text"},
            {"name": "name_westgate", "label": "Westgate", "type": "text"},
            {"name": "name_bmw", "label": "BMW", "type": "text"},
            {"name": "name_lakewood_rx", "label": "Lakewood Rx", "type": "text"},
        ],
    },
    "account_name_map": {
        "title": "Account names",
        "help": "Standardized QuickBooks account name plus each entity's own spelling of it.",
        "columns": [
            {"name": "standard_account_name", "label": "Standard account", "type": "text", "required": True},
            {"name": "name_3m", "label": "3M", "type": "text"},
            {"name": "name_madison", "label": "Madison", "type": "text"},
            {"name": "name_maple", "label": "Maple", "type": "text"},
            {"name": "name_refuah", "label": "Refuah", "type": "text"},
            {"name": "name_westgate", "label": "Westgate", "type": "text"},
            {"name": "name_bmw", "label": "BMW", "type": "text"},
            {"name": "name_lakewood_rx", "label": "Lakewood Rx", "type": "text"},
        ],
    },
    "vendor_account_assignment": {
        "title": "Vendor → account",
        "help": ("Which standardized account each standardized vendor's bills get coded "
                 "to. A vendor missing here is flagged for review even if it has a row "
                 "in Vendor names."),
        # Both columns are selects over large lists (660 vendors + 252 accounts); a per-row
        # <select> duplicates the whole list on every row (629 rows -> ~576k <option>, 34 MB).
        # Render as text inputs backed by ONE shared <datalist> each instead.
        "use_datalist": True,
        "columns": [
            {"name": "standard_vendor_name", "label": "Vendor", "type": "select",
             "options_from": "vendor_names", "required": True},
            {"name": "standard_account_name", "label": "Account", "type": "select",
             "options_from": "account_names", "required": True},
        ],
    },
    "sender_store_map": {
        "title": "Sender → store",
        "help": ("Self-scan mailboxes. An email arriving from one of these is treated as "
                 "that store with high confidence (decisive for self-scans); unmapped "
                 "third-party senders still resolve by the address on the invoice."),
        "columns": [
            {"name": "sender_email", "label": "Sender email", "type": "text", "required": True},
            {"name": "store_matched", "label": "Store", "type": "select",
             "options_from": "stores", "required": True},
            {"name": "note", "label": "Note", "type": "text"},
        ],
    },
    "contact_store_map": {
        "title": "Contact → store",
        "help": ("Known external contacts (a buyer's own email, or a mailbox like "
                 "vitamins@maplerx.com) that identify a store when they appear in the "
                 "To/Cc of an email FROM a third-party vendor — e.g. a biller who "
                 "emails a specific person at a specific store. Different from "
                 "Sender→store, which is for OUR OWN self-scan mailboxes."),
        "columns": [
            {"name": "contact_email", "label": "Contact email", "type": "text", "required": True},
            {"name": "store_matched", "label": "Store", "type": "select",
             "options_from": "stores", "required": True},
            {"name": "note", "label": "Note", "type": "text"},
        ],
    },
    "card_store_map": {
        "title": "Card → store",
        "help": ("The card a payment was made on identifies the store. A payment whose "
                 "card ends in a listed last-4 is treated as that store with high "
                 "confidence and not flagged. Discover unmapped cards on the Cards tab."),
        "columns": [
            {"name": "card_last4", "label": "Card last 4", "type": "text", "required": True},
            {"name": "store_matched", "label": "Store", "type": "select",
             "options_from": "stores", "required": True},
            {"name": "note", "label": "Note", "type": "text"},
        ],
    },
    "store_scope_rules": {
        "title": "Store rules",
        "help": ("Vendor-specific store rules, enforced in code. Standing learnings reach the AI on "
                 "every email, so a rule meant for one vendor (e.g. 'Carepax means BMW on Sageteq "
                 "bills') can leak onto another vendor's invoice — each rule here sends such a record "
                 "to review instead. “only these vendors”: a record resolved to the store from a "
                 "mention of the reference, from any other vendor, is flagged. “every vendor except "
                 "these”: a record mentioning the reference but resolved to a different store is "
                 "flagged, except for the listed vendors. The reference is matched in the billed-to / "
                 "ship-to text (any case); vendors one per line, part of the name is enough. Takes "
                 "effect on the next classify run."),
        "columns": [
            {"name": "reference", "label": "Reference on the document", "type": "text", "required": True},
            {"name": "store_matched", "label": "Means store", "type": "select",
             "options_from": "stores", "required": True},
            {"name": "applies_to", "label": "Applies to", "type": "select",
             "options_from": "store_rule_kinds", "required": True},
            {"name": "vendors", "label": "Vendors (one per line)", "type": "textarea"},
            {"name": "note", "label": "Note", "type": "text"},
        ],
    },
    "class_rules": {
        "title": "Class rules",
        "help": ("Default QuickBooks class by store and/or vendor. Filled in when the document "
                 "didn't state a class — an explicit class on the email always wins. Leave Store "
                 "or Vendor blank for 'any'; when several rules match, the more specific wins "
                 "(both fields > vendor-only > store-only). "
                 "ACCOUNT rules: fill Account and the class is FORCED on every not-yet-entered "
                 "record with that account (at that store/vendor, if set) — it overrides any other "
                 "class, and follows the record if its account changes later. Records already "
                 "entered in QB are never changed."),
        "columns": [
            {"name": "store_matched", "label": "Store (blank = any)", "type": "select",
             "options_from": "stores"},
            {"name": "vendor_matched", "label": "Vendor (blank = any)", "type": "text"},
            {"name": "account", "label": "Account (blank = not an account rule)", "type": "select",
             "options_from": "account_names"},
            {"name": "class", "label": "Class", "type": "text", "required": True},
            {"name": "note", "label": "Note", "type": "text"},
        ],
    },
}


def _reference_options(page):
    """Build {options_from: [values]} for any select columns on this page, sourced from
    web_db.REFERENCE_OPTION_SOURCES. Generic — future select columns just name a source."""
    out = {}
    for c in page["columns"]:
        src = c.get("options_from")
        if c.get("type") == "select" and src and src not in out:
            fn = web_db.REFERENCE_OPTION_SOURCES.get(src)
            out[src] = fn() if fn else []
    return out


@app.get("/reference/{table}", response_class=HTMLResponse)
def reference_page(request: Request, table: str, error: str = "", prefill: str = "",
                   resolved: str = "", stmt_resolved: str = "", vlookup: str = "", store: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    page = REFERENCE_PAGES.get(table)
    if not page:
        return RedirectResponse("/review", status_code=302)
    show_vendor_lookup = table in _VENDOR_LOOKUP_TABLES
    lookup = web_db.vendor_lookup(vlookup) if (show_vendor_lookup and vlookup) else None
    # Arriving from a /qb "add vendor" (prefill = the raw unmatched spelling): suggest existing
    # vendor_name_map entries this spelling is likely the same as, so a human can attach it in one
    # click instead of retyping. Does NOT change classify's auto-apply thresholds.
    suggestions = (web_db.suggest_vendor_matches(prefill)
                   if (table == "vendor_name_map" and prefill) else None)
    return render(request, "reference.html", user, table=table, page=page,
                  items=web_db.list_reference(table), options=_reference_options(page),
                  error=error, prefill=prefill, resolved=resolved, stmt_resolved=stmt_resolved,
                  show_vendor_lookup=show_vendor_lookup, vlookup=vlookup, lookup=lookup,
                  suggestions=suggestions, store=store)


# Setup tables that feed vendor standardization — a save here can newly resolve existing
# unentered records, so we re-run the backfill and surface how many cleared.
_VENDOR_STD_TABLES = {"vendor_name_map", "vendor_account_assignment"}

# The three vendor-related Setup pages that carry the cross-table vendor/account lookup box.
_VENDOR_LOOKUP_TABLES = {"vendor_name_map", "vendor_account_assignment", "account_name_map"}


def _sync_account_class_rules(why):
    """Re-apply account-based class rules (database.apply_account_class_rules) after anything that
    can change a record's account/store/vendor or the rules themselves. Never raises."""
    try:
        import database
        database.apply_account_class_rules(f"web ({why})")
    except Exception as e:
        print(f"(account class rules after {why} failed: {e})")


def _reference_saved_redirect(table):
    """Redirect back to a reference page after a successful save. For the vendor-standardization
    tables, first re-resolve existing unentered records against the now-updated mappings and pass
    the resolved count through for a one-line flash."""
    _sync_account_class_rules(table)
    if table in _VENDOR_STD_TABLES:
        try:
            import classify_emails as ce
            import database
            _checked, resolved = ce.reapply_vendor_standardization_backfill()
            # A new alias can also clear open statement exceptions (recon matches vendor by
            # standardized name in its second pass) — re-run it right away, same as the QB backfill.
            # Report how many exceptions THIS save newly cleared: flagged-before minus flagged-after
            # (reconcile recomputes from scratch, so its own alias_resolved counts every rescue each
            # run; the delta is what actually changed for the user).
            # Toggling a vendor_name_map "Exclude from QB" checkbox is now the exclusion action, so
            # also backfill skip_qb_import on already-recorded records (only ever flips toward
            # exclusion — unchecking never un-skips, per Joel's 2026-09-24 decision).
            excluded_n = _backfill_excluded_skips() if table == "vendor_name_map" else 0
            stmt_resolved = 0
            try:
                flagged_before = database.count_flagged_statement_lines()
                recon = database.reconcile_statement_lines()
                stmt_resolved = max(0, flagged_before - recon.get("flagged", 0))
            except Exception as e:
                print(f"(statement reconcile after {table} save failed: {e})")
            return RedirectResponse(
                f"/reference/{table}?resolved={resolved}&stmt_resolved={stmt_resolved}"
                + (f"&excluded={excluded_n}" if excluded_n else ""),
                status_code=302)
        except Exception as e:
            print(f"(vendor/account backfill after {table} save failed: {e})")
    return RedirectResponse(f"/reference/{table}", status_code=302)


@app.post("/reference/{table}/add")
async def reference_add(request: Request, table: str):
    user = current_user(request)
    if not user:
        return _login_redirect()
    page = REFERENCE_PAGES.get(table)
    if not page:
        return RedirectResponse("/review", status_code=302)
    form = await request.form()
    values = {c["name"]: form.get(c["name"], "") for c in page["columns"]}
    try:
        web_db.add_reference(table, values, user["username"])
    except ValueError as e:
        return RedirectResponse(f"/reference/{table}?error={quote(str(e))}", status_code=302)
    return _reference_saved_redirect(table)


@app.post("/reference/vendor-name-map/attach-spelling")
async def vendor_attach_spelling(request: Request):
    """One-click from an add-vendor suggestion: attach the raw invoice spelling to an EXISTING
    vendor's row (per-store column when the store is known, else aliases) instead of creating a
    new standard vendor — then re-run the standardization backfill so the invoice resolves."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    form = await request.form()
    web_db.attach_vendor_spelling(form.get("standard_vendor_name", ""), form.get("raw", ""),
                                  form.get("store", ""), user["username"])
    return _reference_saved_redirect("vendor_name_map")


def _safe_return_to(return_to, referer, default="/qb"):
    """A same-site relative redirect target (guards against open-redirect): must start with a single
    '/'. Prefers an explicit return_to, then the Referer, then the default."""
    for cand in (return_to, referer, default):
        c = (cand or "").strip()
        if c.startswith("/") and not c.startswith("//"):
            return c
    return default


def _redirect_with_flash(url, msg):
    """302 back to `url`, appending a one-line `flash` message (base.html renders it on any page)."""
    sep = "&" if ("?" in url) else "?"
    return RedirectResponse(f"{url}{sep}flash={quote(msg)}", status_code=302)


def _backfill_excluded_skips():
    """Run the excluded-vendor skip backfill; return how many already-recorded records were newly
    skipped. Never lets a backfill error break the calling save/redirect."""
    try:
        import classify_emails as ce
        return ce.reapply_excluded_vendor_backfill()
    except Exception as e:
        print(f"(excluded-vendor backfill failed: {e})")
        return 0


@app.post("/vendors/exclude")
async def vendor_exclude(request: Request):
    """One-click 'Ignore this vendor': tick Exclude from QB for a vendor (known from context, no
    retyping) on vendor_name_map so its invoices never enter QuickBooks. Idempotent. Reusable from anywhere a
    vendor name shows (add-vendor flow, /qb Account cell, record detail). Redirects back with a flash."""
    user = current_user(request)
    if not user:
        return _login_redirect()
    form = await request.form()
    vendor = (form.get("vendor_name") or "").strip()
    return_to = _safe_return_to(form.get("return_to"), request.headers.get("referer"))
    if not vendor:
        return _redirect_with_flash(return_to, "No vendor name to exclude.")
    result = web_db.exclude_vendor(vendor, user["username"])
    if result == "added":
        # Backfill: flip skip_qb_import on already-recorded, not-yet-entered records now excluded.
        n = _backfill_excluded_skips()
        extra = (f" {n} already-recorded record{'' if n == 1 else 's'} updated." if n else "")
        msg = f"Excluded “{vendor}” — its invoices will no longer enter QuickBooks.{extra}"
    elif result == "exists":
        msg = f"“{vendor}” is already excluded from QuickBooks."
    else:
        msg = "No vendor name to exclude."
    return _redirect_with_flash(return_to, msg)


@app.post("/reference/{table}/{record_id}/edit")
async def reference_edit(request: Request, table: str, record_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    page = REFERENCE_PAGES.get(table)
    if not page:
        return RedirectResponse("/review", status_code=302)
    form = await request.form()
    values = {c["name"]: form.get(c["name"], "") for c in page["columns"] if c["name"] in form}
    try:
        web_db.edit_reference(table, record_id, values, user["username"])
    except ValueError as e:
        return RedirectResponse(f"/reference/{table}?error={quote(str(e))}", status_code=302)
    return _reference_saved_redirect(table)


@app.post("/reference/{table}/{record_id}/delete")
def reference_delete(request: Request, table: str, record_id: int):
    user = current_user(request)
    if not user:
        return _login_redirect()
    if table not in REFERENCE_PAGES:
        return RedirectResponse("/review", status_code=302)
    if user["role"] != "admin":            # deletes are admin-only
        return render(request, "message.html", user,
                      message="Deleting reference data is admin-only.")
    web_db.delete_reference(table, record_id, user["username"])
    return RedirectResponse(f"/reference/{table}", status_code=302)


# --------------------------------------------------------------------------- #
# Cards discovery tab — every card last-4 seen on payments, with an inline
# "assign store" that upserts into card_store_map (read by the pipeline).
# --------------------------------------------------------------------------- #

@app.get("/cards", response_class=HTMLResponse)
def cards_page(request: Request, error: str = ""):
    user = current_user(request)
    if not user:
        return _login_redirect()
    return render(request, "cards.html", user, cards=web_db.list_seen_cards(),
                  stores=web_db.store_names(), error=error)


@app.post("/cards/assign")
async def cards_assign(request: Request):
    user = current_user(request)
    if not user:
        return _login_redirect()
    form = await request.form()
    try:
        web_db.assign_card_store(form.get("card_last4", ""),
                                 form.get("store_matched", ""), user["username"])
    except ValueError as e:
        return RedirectResponse(f"/cards?error={quote(str(e))}", status_code=302)
    return RedirectResponse("/cards", status_code=302)


# --------------------------------------------------------------------------- #
# (Invoice history is now the "History ledger" source of the Records tab above.)
# --------------------------------------------------------------------------- #


# --------------------------------------------------------------------------- #
# Admin — users
# --------------------------------------------------------------------------- #

@app.get("/admin/users", response_class=HTMLResponse)
def users_page(request: Request, error: str = "", sort: str = "", direction: str = "asc",
               edit: int = 0):
    user = current_user(request)
    if not user:
        return _login_redirect()
    if user["role"] != "admin":
        return render(request, "message.html", user, message="Admins only.")
    return render(request, "users.html", user, users=web_db.list_users(sort=sort, direction=direction),
                  error=error, edit_id=edit,
                  params=dict(sort=sort, direction=direction))


@app.post("/admin/users/create")
def users_create(request: Request, username: str = Form(...), password: str = Form(...),
                 display_name: str = Form(""), role: str = Form("reviewer")):
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    if web_db.get_user(username):
        return RedirectResponse("/admin/users?error=Username+already+exists", status_code=302)
    web_db.create_user(username, password, display_name or username,
                       role if role in web_db.ROLES else "reviewer")
    return RedirectResponse("/admin/users", status_code=302)


@app.post("/admin/users/{user_id}/active")
def users_toggle(request: Request, user_id: int, active: str = Form("")):
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    web_db.set_user_active(user_id, active == "1")
    return RedirectResponse("/admin/users", status_code=302)


@app.post("/admin/users/{user_id}/edit")
def users_edit(request: Request, user_id: int, display_name: str = Form(""),
               role: str = Form(""), active: str = Form(""), new_password: str = Form("")):
    user = current_user(request)
    if not user or user["role"] != "admin":
        return _login_redirect()
    try:
        web_db.update_user(
            user_id, user["username"],
            display_name=display_name,
            role=role if role in web_db.ROLES else None,
            active=(active == "1"),
            new_password=(new_password or None))
    except ValueError as e:
        return RedirectResponse(f"/admin/users?error={quote(str(e))}", status_code=302)
    return RedirectResponse("/admin/users", status_code=302)
