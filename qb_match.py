"""
QuickBooks reconciliation import.

A user exports a QB register (one worksheet per store) before a QB data-entry session and uploads it
here. For every Bill/Credit row we check the app's not-yet-entered invoices/credit_memos so nothing
gets entered into QB twice, and we auto-resolve stale 'missing' statement flags that QB shows aren't
actually missing. Reuses the app's OWN vendor/number/amount matching (classify_emails) — no second,
parallel algorithm.

  process_qb_export(file_bytes, filename, uploaded_by) -> batch_id

Matching reuses: classify_emails.resolve_standard_vendor (same standard_vendor already on each
record), classify_emails._norm (the number normalization check_duplicate uses), and
classify_emails._amounts_equal (0.005 tolerance). openpyxl (already a dependency) — not pandas.

Rows that match nothing are kept as pending unmatched rows only when they could still match a later
email (UnmatchedKeeper), and every uploaded file is kept in the qb_imports folder (keep_upload_copy).
"""
import io
import os
import re
from collections import Counter
from datetime import date, datetime, timedelta

import psycopg2.extras
from openpyxl import load_workbook

import classify_emails as ce
import database

# Worksheet tab name -> store_matched (confirmed against live data). Matched case-insensitively.
STORE_MAP = {
    "3m": "3M Pharmacy", "madison": "Madison Pharmacy", "maple": "Maple Pharmacy",
    "refuah": "Refuah Pharmacy", "westgate": "Westgate Pharmacy",
    "bmw": "BMW", "lakewood rx": "Lakewood Rx", "carepax": "Carepax",
}
_WANTED_COLS = ("type", "date", "num", "name", "amount")


def _store_for_tab(tab_name):
    return STORE_MAP.get((tab_name or "").strip().lower())


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _num_str(v):
    """Stringify a Num cell for normalized comparison — openpyxl reads an integer invoice number as
    a float (3266014701.0); drop the trailing .0 so it matches the stored text '3266014701'."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _parse_date(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%m-%d-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _std_eq(a, b):
    return bool(a and b and ce._norm(a) == ce._norm(b))


def _digit_runs(v):
    """Every run of digits in v once punctuation no longer splits it, each with leading zeros
    stripped, after folding the classic OCR letter/digit confusions (I/i <-> 1, O/o <-> 0) -- so
    '86I0721' reads as '8610721', matching the vendor's own numeric invoice number. The fold applies
    ONLY to a letter flanked by digits on BOTH sides (genuine mid-number OCR noise); a letter affix
    such as the 'IN' in '1618282-IN' is left alone, so it isn't glued on as a digit. Punctuation is
    removed first so an internal '/' or '-' doesn't split one number into two runs (e.g.
    'INV/2026/47830' -> ['202647830'], matching 'INV/202647830'). Empty if there are no digits."""
    v = (v or "").strip()
    alnum = re.sub(r"[^0-9A-Za-z]", "", v)
    fold = {"I": "1", "i": "1", "O": "0", "o": "0"}
    chars = list(alnum)
    for i, ch in enumerate(chars):
        if ch in fold and 0 < i < len(chars) - 1 and chars[i - 1].isdigit() and chars[i + 1].isdigit():
            chars[i] = fold[ch]
    return [r.lstrip("0") or "0" for r in re.findall(r"\d+", "".join(chars))]


def _digit_core(v):
    """The longest of _digit_runs(v) -- the index bucket _num_bucket uses. Empty if no digits."""
    runs = _digit_runs(v)
    return max(runs, key=len) if runs else ""


def _num_eq(a, b):
    # Exact (normalized) match first; then fall back to the digits so cosmetic-only differences
    # (vendor prefix, leading zeros, punctuation/"IN"/"INV" affixes) still match, while genuinely
    # different numbers do not. Covers Tier-1 auto-match, Tier-2 candidate detection, the statement
    # pass, and the recheck button (all route through here).
    # EVERY digit run must agree, not just the longest (2026-10-12): letters inside a number that
    # split its digits are part of the number, so 'SUB2026M08' / 'SUB2026M09' (a monthly bill at a
    # fixed amount) or 'S1234A01' / 'S1234A02' no longer match on '2026' / '1234' alone -- a false
    # Tier-1 match marks the wrong bill entered and the real one silently never shows for entry. A
    # number this turns away still reaches the review list on vendor + amount + date.
    na, nb = ce._norm(_num_str(a)), ce._norm(_num_str(b))
    if na == "" or nb == "":
        return False
    if na == nb:
        return True
    runs_a = _digit_runs(na)
    return bool(runs_a) and runs_a == _digit_runs(nb)


# --------------------------------------------------------------------------- #
# Which "no match" rows are kept as pending unmatched rows (qb_unmatched_rows).
#
# A pending row waits for a later email: when that invoice is recorded, it's marked entered in QB on
# arrival instead of being entered a second time. But the export is year-to-date, and most of its
# no-match rows aren't waiting for anything: they're QuickBooks history from before the app recorded
# invoices, or documents the app already has (an entered, excluded or ignored record never reaches the
# matching pool, so its QB row comes out "no match"). The Oct 8, 2026 upload had 10,888 no-match rows
# and only ~580 were neither. So a no-match row is kept only when:
#   - it's dated on/after UNMATCHED_SINCE (the day the app started recording invoices) and at most
#     UNMATCHED_WINDOW_DAYS before the upload (99% of the app's records arrived within ~4 months of
#     their invoice date — measured Oct 2026);
#   - no app record in the same store, in ANY status, is the same document (record_on_file), and no
#     earlier review-list match linked the row to a record (a confirmed qb_match_candidates row); and
#   - an earlier upload didn't already save the same row, in any status (so a row someone dismissed
#     doesn't come back as pending when the next export repeats it).
# --------------------------------------------------------------------------- #
UNMATCHED_SINCE = date(2026, 8, 7)
UNMATCHED_WINDOW_DAYS = 120


def _num_bucket(v):
    """Index bucket for a number: its digit core (what _num_eq falls back to), or the normalized text
    when it has no digits — two numbers _num_eq calls equal always land in the same bucket."""
    n = ce._norm(_num_str(v))
    core = _digit_core(n)
    return ("core", core) if core else ("text", n)


def unmatched_row_key(store, qb_type, qb_date, qb_num, qb_name, qb_amount):
    """The same QB row across uploads (every year-to-date export repeats the earlier rows)."""
    return (store or "", str(qb_type or "").strip().lower(), qb_date, ce._norm(_num_str(qb_num)),
            " ".join(str(qb_name or "").lower().split()),
            round(float(qb_amount), 2) if qb_amount is not None else None)


class UnmatchedKeeper:
    """Decides which no-match QB rows to keep (rule above). Build one per upload (or backfill) inside
    its transaction: it loads every invoice/credit memo and every saved unmatched row once."""

    def __init__(self, cur, upload_day):
        self.since = max(UNMATCHED_SINCE, upload_day - timedelta(days=UNMATCHED_WINDOW_DAYS))
        self.records = {}   # (table, store, number bucket) -> records, any status
        for table, num_col in (("invoices", "invoice_number"), ("credit_memos", "credit_memo_number")):
            dup = "duplicate_ignored" if table == "invoices" else "FALSE AS duplicate_ignored"
            cur.execute(f"SELECT id, store_matched, standard_vendor, {num_col} AS number, amount, "
                        f"entered_in_qb, skip_qb_import, {dup} FROM {table}")
            for r in cur.fetchall():
                self.records.setdefault((table, r["store_matched"], _num_bucket(r["number"])), []).append(r)
        cur.execute("SELECT qb_store, qb_type, qb_date, qb_num, qb_name, qb_amount FROM qb_unmatched_rows")
        self.saved = Counter(unmatched_row_key(u["qb_store"], u["qb_type"], u["qb_date"], u["qb_num"],
                                               u["qb_name"], u["qb_amount"]) for u in cur.fetchall())
        # QB rows already linked to a record on the review list (confirmed by a person or the recheck).
        # The link can rest on the vendor + amount + date rule, which record_on_file doesn't use.
        cur.execute("SELECT qb_store, qb_type, qb_date, qb_num, qb_name, qb_amount FROM qb_match_candidates "
                    "WHERE status = 'confirmed'")
        self.confirmed = {unmatched_row_key(c["qb_store"], c["qb_type"], c["qb_date"], c["qb_num"],
                                            c["qb_name"], c["qb_amount"]) for c in cur.fetchall()}

    def record_on_file(self, table, store, qb_std, qb_num, qb_amount):
        """The app record (any status, same store) that is this QB row's document, or None: the same
        number plus the same amount or the same standard vendor — the Tier-1/Tier-2 evidence minus the
        weak vendor + amount + date rule. An exact vendor + number + amount match is returned first."""
        if not ce._norm(_num_str(qb_num)):
            return None
        near = None
        for r in self.records.get((table, store, _num_bucket(qb_num)), ()):
            if not _num_eq(qb_num, r["number"]):
                continue
            v_eq, a_eq = _std_eq(qb_std, r["standard_vendor"]), ce._amounts_equal(qb_amount, r["amount"])
            if v_eq and a_eq:
                return r
            if (v_eq or a_eq) and near is None:
                near = r
        return near

    def skip_reason(self, table, store, qb_type, qb_date, qb_num, qb_name, qb_std, qb_amount):
        """Why a no-match row is NOT kept — 'no date' | 'older' | 'in the app' | 'already saved' — or
        None to keep it. A row found already saved is counted off, so an export listing the same row
        twice needs two saved copies to skip both."""
        if qb_date is None:
            return "no date"
        if qb_date < self.since:
            return "older"
        key = unmatched_row_key(store, qb_type, qb_date, qb_num, qb_name, qb_amount)
        if key in self.confirmed or self.record_on_file(table, store, qb_std, qb_num, qb_amount):
            return "in the app"
        if self.saved[key] > 0:
            self.saved[key] -= 1
            return "already saved"
        return None


QB_IMPORTS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "qb_imports")


def keep_upload_copy(batch_id, filename, file_bytes):
    """Save an uploaded export as qb_imports\\batch_<id> - <name>, so the upload can be re-read later
    (the Oct 8, 2026 upload wasn't kept and had to be backfilled from a Downloads copy). Best effort:
    a failure is printed, never raised — the import has already committed. Returns the path or None."""
    name = re.sub(r'[\\/:*?"<>|]+', "_", os.path.basename(filename or "")).strip() or "export.xlsx"
    path = os.path.join(QB_IMPORTS_DIR, f"batch_{batch_id} - {name}")
    try:
        os.makedirs(QB_IMPORTS_DIR, exist_ok=True)
        with open(path, "wb") as f:
            f.write(file_bytes)
        return path
    except OSError as e:
        print(f"(could not keep a copy of QB export batch {batch_id}: {e})")
        return None


# --------------------------------------------------------------------------- #
# Auto-correct per-store vendor spellings in vendor_name_map from QB recon history.
#
# Every QB export row's matching implicitly tells us whether our mapping agrees with what
# QuickBooks itself calls a vendor at a store. Two confidence_reason buckets carry the signal:
#   A) "number + amount match; vendor did not resolve to the same standard vendor" — STRONG: the
#      bill is unquestionably the same (number+amount matched), but resolve_standard_vendor() could
#      not map the QB row's raw name to that standard vendor — direct evidence the mapping is wrong/
#      missing for that (standard_vendor, store). May FILL an empty slot OR OVERRIDE an existing one.
#   B) "vendor + amount match, date within 10 days; number missing/mismatched" — WEAK: the vendor
#      already resolved correctly (the mismatch is the NUMBER), so this only ever FILLS an EMPTY slot.
# Rules (conservative — skip rather than guess):
#   - Only status <> 'rejected' candidates count. Group by (standard_vendor, qb_store), per bucket.
#   - If a bucket has >1 DISTINCT qb_name for a group, skip that bucket for that group (never choose).
#   - Fill an EMPTY name_<store>: either bucket, any single occurrence.
#   - OVERRIDE an EXISTING name_<store>: bucket A only, and only when >= 2 DISTINCT batches agree on
#     the exact spelling (a one-off QB data-entry typo in a single batch is thereby excluded).
#   - Every applied change is logged to edit_log (table_name='vendor_name_map').
# --------------------------------------------------------------------------- #
STORE_COL = {
    "3M Pharmacy": "name_3m", "Madison Pharmacy": "name_madison",
    "Maple Pharmacy": "name_maple", "Refuah Pharmacy": "name_refuah",
    "Westgate Pharmacy": "name_westgate", "BMW": "name_bmw",
    "Lakewood Rx": "name_lakewood_rx", "Carepax": "name_carepax",
}
_OVERRIDE_ELIGIBLE_REASON = "number + amount match; vendor did not resolve to the same standard vendor"
_FILL_ONLY_REASON = "vendor + amount match, date within 10 days; number missing/mismatched"
_MIN_BATCHES_TO_OVERRIDE = 2


def _vendor_name_evidence(cur):
    """{(confidence_reason, standard_vendor, qb_store): {qb_name: set(batch_id)}} across every
    non-rejected qb_match_candidates row in either relevant bucket."""
    cur.execute("""
        SELECT c.confidence_reason, c.qb_name, c.qb_store, c.batch_id,
               COALESCE(i.standard_vendor, cm.standard_vendor) AS standard_vendor
        FROM qb_match_candidates c
        LEFT JOIN invoices i ON c.record_type = 'invoice' AND i.id = c.record_id
        LEFT JOIN credit_memos cm ON c.record_type = 'credit_memo' AND cm.id = c.record_id
        WHERE c.confidence_reason = ANY(%s) AND c.status <> 'rejected'
    """, ([_OVERRIDE_ELIGIBLE_REASON, _FILL_ONLY_REASON],))
    evidence = {}
    for row in cur.fetchall():
        sv, store, name = row["standard_vendor"], row["qb_store"], (row["qb_name"] or "").strip()
        if not sv or not store or not name:
            continue
        key = (row["confidence_reason"], sv, store)
        evidence.setdefault(key, {}).setdefault(name, set()).add(row["batch_id"])
    return evidence


def apply_vendor_name_corrections_from_qb_history(actor, dry_run=False):
    """Uses the FULL history of qb_match_candidates (every batch ever uploaded, not just the
    latest) to fill or correct vendor_name_map's per-store spellings. See the module-level rule
    comment above — deliberately conservative: it will skip a vendor/store rather than guess.
    Returns {"applied": [...], "skipped_ambiguous": [...], "skipped_thin": [...]} for logging/review.
    dry_run=True reports what it WOULD do without writing (used by the one-time backfill script;
    process_qb_export always calls with dry_run=False)."""
    conn = database.connect()
    conn.autocommit = False
    applied, skipped_ambiguous, skipped_thin = [], [], []
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            evidence = _vendor_name_evidence(cur)
            cur.execute("SELECT * FROM vendor_name_map")
            vmap = {r["standard_vendor_name"]: r for r in cur.fetchall()}

            groups = {}
            for (reason, sv, store), names in evidence.items():
                g = groups.setdefault((sv, store), {"override_names": {}, "fill_names": {}})
                bucket = "override_names" if reason == _OVERRIDE_ELIGIBLE_REASON else "fill_names"
                g[bucket] = names

            for (sv, store), g in sorted(groups.items()):
                col = STORE_COL.get(store)
                row = vmap.get(sv)
                if not col or not row:
                    continue
                current = (row.get(col) or "").strip() or None

                override_names, fill_names = g["override_names"], g["fill_names"]
                chosen, n_batches, source = None, 0, None
                if len(override_names) == 1:
                    (name, batches), = override_names.items()
                    chosen, n_batches, source = name, len(batches), "override-eligible"
                elif len(override_names) > 1:
                    skipped_ambiguous.append((sv, store, "override", sorted(override_names)))

                if chosen is None and len(fill_names) == 1:
                    (name, batches), = fill_names.items()
                    chosen, n_batches, source = name, len(batches), "fill-only"
                elif chosen is None and len(fill_names) > 1:
                    skipped_ambiguous.append((sv, store, "fill", sorted(fill_names)))

                if chosen is None or chosen == current:
                    continue
                if current is not None and (source != "override-eligible" or n_batches < _MIN_BATCHES_TO_OVERRIDE):
                    skipped_thin.append((sv, store, current, chosen, source, n_batches))
                    continue

                applied.append((sv, store, col, current, chosen, source, n_batches))
                if not dry_run:
                    cur.execute(f"UPDATE vendor_name_map SET {col} = %s WHERE standard_vendor_name = %s",
                                (chosen, sv))
                    cur.execute(
                        "INSERT INTO edit_log (table_name, record_id, field, old_value, new_value, changed_by) "
                        "VALUES ('vendor_name_map', %s, %s, %s, %s, %s)",
                        (row["id"], col, current, chosen, actor))
            # Push the corrected store spellings onto records not yet entered in QB (otherwise
            # they keep the name snapshotted when they were first processed).
            if applied and not dry_run:
                import web_db
                web_db.refresh_qb_names(cur, "vendor_name_map", actor,
                                        standard_names=sorted({a[0] for a in applied}))
        conn.commit() if not dry_run else conn.rollback()
    finally:
        conn.close()
    return {"applied": applied, "skipped_ambiguous": skipped_ambiguous, "skipped_thin": skipped_thin}


def _parse_rows(file_bytes):
    """Every Bill/Credit row across the store tabs: {store, type, date, num, name, amount}. Columns
    are mapped by HEADER TEXT (row 1, case-insensitive) so column order can shift between exports."""
    wb = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    rows = []
    try:
        for ws in wb.worksheets:
            store = _store_for_tab(ws.title)
            if not store:
                continue
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
            if not header:
                continue
            colmap = {}
            for idx, h in enumerate(header):
                key = str(h).strip().lower() if h is not None else ""
                if key in _WANTED_COLS and key not in colmap:
                    colmap[key] = idx
            if not all(k in colmap for k in ("type", "name", "amount")):
                continue
            for r in ws.iter_rows(min_row=2, values_only=True):
                if r is None or all(c is None for c in r):
                    continue

                def cell(k, _r=r):
                    i = colmap.get(k)
                    return _r[i] if i is not None and i < len(_r) else None

                rows.append({"store": store, "type": cell("type"), "date": cell("date"),
                             "num": cell("num"), "name": cell("name"), "amount": cell("amount")})
    finally:
        wb.close()
    return rows


def process_qb_export(file_bytes, filename, uploaded_by):
    """Parse the QB export, match each row against not-yet-entered records (Tier 1 auto-mark /
    Tier 2 doubtful candidate), resolve 'missing' statement lines that QB actually has, and write
    the batch summary. Every row with no candidate counts in n_no_match; only the ones
    UnmatchedKeeper keeps become pending unmatched rows. The file is kept in qb_imports\\ after the
    commit. Returns the new qb_import_batches id."""
    # Load vendor standardization FIRST, on its own autocommit connection, BEFORE the transaction
    # below — load_vendor_standardization -> init_db takes ACCESS-EXCLUSIVE locks and would deadlock
    # inside an open txn (same gotcha as recover_errored_email.py).
    std = ce.load_vendor_standardization()
    rows = _parse_rows(file_bytes)
    actor = f"{uploaded_by} via QB reconcile import"

    conn = database.connect()
    conn.autocommit = False
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute("INSERT INTO qb_import_batches (filename, uploaded_by) VALUES (%s,%s) "
                        "RETURNING id", (filename, uploaded_by))
            batch_id = cur.fetchone()["id"]
            keeper = UnmatchedKeeper(cur, date.today())

            n_rows = n_auto = n_doubt = n_nomatch = n_kept = 0
            bill_rows = []   # for the statement pass: {std, num, amount}

            for row in rows:
                qtype = (str(row["type"]).strip().lower() if row["type"] is not None else "")
                if qtype == "bill":
                    table, rtype, num_col, date_col = "invoices", "invoice", "invoice_number", "invoice_date"
                    qb_amount = abs(_num(row["amount"])) if _num(row["amount"]) is not None else None
                elif qtype == "credit":
                    table, rtype, num_col, date_col = "credit_memos", "credit_memo", "credit_memo_number", "credit_memo_date"
                    qb_amount = _num(row["amount"])
                else:
                    continue  # skip any other Type

                n_rows += 1
                qb_num, qb_name, store = row["num"], row["name"], row["store"]
                qb_date = _parse_date(row["date"])
                qb_std = ce.resolve_standard_vendor(qb_name, std_tables=std)
                if qtype == "bill":
                    bill_rows.append({"std": qb_std, "num": qb_num, "amount": qb_amount})

                # credit_memos has no duplicate_ignored column (only invoices/payments have the
                # duplicate-suspect workflow) — so that clause is invoices-only, else every row throws
                # UndefinedColumn and the whole upload 500s on the first Credit row.
                # Skip records the user marked never-to-enter-QB (reference-only): matching them
                # against the export is noise in both tiers — the outcome doesn't matter either way.
                dup_clause = " AND duplicate_ignored IS NOT TRUE" if table == "invoices" else ""
                cur.execute(
                    f"SELECT id, standard_vendor, {num_col} AS number, amount, {date_col} AS rdate "
                    f"FROM {table} WHERE store_matched = %s AND entered_in_qb = FALSE "
                    f"AND skip_qb_import IS NOT TRUE"
                    f"{dup_clause}", (store,))
                cands = cur.fetchall()

                # Tier 1 — auto-resolve: standard_vendor + number + amount all match.
                t1 = next((c for c in cands
                           if _std_eq(qb_std, c["standard_vendor"]) and _num_eq(qb_num, c["number"])
                           and ce._amounts_equal(qb_amount, c["amount"])), None)
                if t1:
                    cur.execute(f"UPDATE {table} SET entered_in_qb = TRUE, entered_qb_by = %s, "
                                f"entered_qb_at = now() WHERE id = %s", (actor, t1["id"]))
                    cur.execute(
                        "INSERT INTO edit_log (table_name, record_id, field, old_value, new_value, "
                        "changed_by) VALUES (%s,%s,'entered_in_qb','false',%s,%s)",
                        (table, t1["id"],
                         f"true (found in QuickBooks — {row['type']} Num={_num_str(qb_num)} "
                         f"Amount={qb_amount} Date={qb_date})", actor))
                    n_auto += 1
                    continue

                # Tier 2 — doubtful, needs a human. First matching rule wins.
                t2 = t2_reason = None
                for c in cands:
                    v_eq = _std_eq(qb_std, c["standard_vendor"])
                    n_eq = _num_eq(qb_num, c["number"])
                    a_eq = ce._amounts_equal(qb_amount, c["amount"])
                    # invoice_date/credit_memo_date are stored as TEXT — parse before the date diff.
                    rdate = _parse_date(c["rdate"])
                    date_close = bool(qb_date and rdate and abs((qb_date - rdate).days) <= 10)
                    if v_eq and a_eq and not n_eq and date_close:
                        t2, t2_reason = c, "vendor + amount match, date within 10 days; number missing/mismatched"
                        break
                    if n_eq and a_eq and not v_eq:
                        t2, t2_reason = c, "number + amount match; vendor did not resolve to the same standard vendor"
                        break
                    if v_eq and n_eq and not a_eq:
                        t2, t2_reason = c, "vendor + number match; amount differs beyond tolerance"
                        break
                if t2:
                    cur.execute(
                        "INSERT INTO qb_match_candidates (batch_id, record_type, record_id, qb_type, "
                        "qb_date, qb_num, qb_name, qb_amount, qb_store, confidence_reason) "
                        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                        (batch_id, rtype, t2["id"], row["type"], qb_date, _num_str(qb_num) or None,
                         qb_name, qb_amount, store, t2_reason))
                    n_doubt += 1
                else:
                    # No candidate at all — freeze the QB row so a later-arriving email (or the
                    # recheck button) can match it, unless it's older history, a document the app
                    # already has, or a row an earlier upload saved (UnmatchedKeeper). qb_std was
                    # resolved above; no new lookup.
                    n_nomatch += 1
                    if keeper.skip_reason(table, store, row["type"], qb_date, qb_num, qb_name, qb_std,
                                          qb_amount) is None:
                        cur.execute(
                            "INSERT INTO qb_unmatched_rows (batch_id, qb_type, qb_date, qb_num, qb_name, "
                            "qb_std_vendor, qb_amount, qb_store) VALUES (%s,%s,%s,%s,%s,%s,%s,%s)",
                            (batch_id, row["type"], qb_date, _num_str(qb_num) or None, qb_name, qb_std,
                             qb_amount, store))
                        n_kept += 1

            # Statement-reconciliation pass: resolve 'missing' lines QB actually has (Tier-1-quality
            # match against Bill rows anywhere in the workbook — statement_lines carry no store, so
            # not store-scoped). standard_vendor + number + amount.
            n_stmt = 0
            cur.execute("SELECT id, vendor, invoice_number, amount FROM statement_lines "
                        "WHERE status = 'missing'")
            for sl in cur.fetchall():
                sl_std = ce.resolve_standard_vendor(sl["vendor"], std_tables=std)
                hit = next((b for b in bill_rows
                            if _std_eq(sl_std, b["std"]) and _num_eq(sl["invoice_number"], b["num"])
                            and ce._amounts_equal(sl["amount"], b["amount"])), None)
                if hit:
                    cur.execute("UPDATE statement_lines SET status = 'dismissed', dismissed_by = %s, "
                                "dismissed_at = now(), updated_at = now() WHERE id = %s",
                                (uploaded_by, sl["id"]))
                    cur.execute(
                        "INSERT INTO edit_log (table_name, record_id, field, old_value, new_value, "
                        "changed_by) VALUES ('statement_lines',%s,'status','missing',%s,%s)",
                        (sl["id"], "dismissed (found directly in QuickBooks, entered outside the app)", actor))
                    n_stmt += 1

            cur.execute("UPDATE qb_import_batches SET n_rows=%s, n_auto_matched=%s, n_doubtful=%s, "
                        "n_statement_resolved=%s, n_no_match=%s WHERE id=%s",
                        (n_rows, n_auto, n_doubt, n_stmt, n_nomatch, batch_id))
        conn.commit()
        print(f"QB import batch {batch_id}: {n_nomatch} row(s) matched nothing; kept {n_kept} as pending "
              f"unmatched (dated {keeper.since:%m/%d/%Y} or later, not already in the app or saved).")
        keep_upload_copy(batch_id, filename, file_bytes)
        # Learn per-store vendor spellings from this batch's (now-committed) candidates + all prior
        # history. Fresh connection inside the helper, so it sees the just-committed rows.
        corr = apply_vendor_name_corrections_from_qb_history(actor)
        if corr["applied"]:
            print(f"Vendor names: corrected {len(corr['applied'])} per-store spelling(s) "
                  f"from QB reconciliation history.")
            for sv, store, col, old, new, source, n in corr["applied"]:
                print(f"  {sv!r} @ {store}: {col} {old!r} -> {new!r} ({source}, {n} batch(es))")
        return batch_id
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def recheck_pending_candidates(triggered_by):
    """Re-resolve every pending qb_match_candidates row's vendor against the CURRENT
    vendor_name_map (a user may have just added/fixed an alias) and re-test Tier 1 (exact
    vendor+number+amount). Anything that now matches exactly is auto-resolved the same way the
    upload path resolves Tier 1; anything whose linked record is no longer a valid candidate
    (already entered, now skip_qb_import, now duplicate_ignored) is dismissed as stale. Everything
    else is left untouched. Returns a one-line summary string for the flash message."""
    std = ce.load_vendor_standardization()
    actor = f"{triggered_by} via QB recheck"
    conn = database.connect()
    conn.autocommit = False
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            cur.execute("SELECT * FROM qb_match_candidates WHERE status = 'pending'")
            pending = cur.fetchall()
            n_resolved = n_stale = n_unchanged = 0
            for c in pending:
                table = "invoices" if c["record_type"] == "invoice" else "credit_memos"
                num_col = "invoice_number" if table == "invoices" else "credit_memo_number"
                dup_clause = " AND duplicate_ignored IS NOT TRUE" if table == "invoices" else ""
                cur.execute(
                    f"SELECT id, standard_vendor, {num_col} AS number, amount, "
                    f"entered_in_qb, skip_qb_import FROM {table} "
                    f"WHERE id = %s AND entered_in_qb = FALSE AND skip_qb_import IS NOT TRUE"
                    f"{dup_clause}", (c["record_id"],))
                rec = cur.fetchone()
                if not rec:
                    # record left the eligible pool since this candidate was written (already
                    # entered, now skip_qb_import, now duplicate_ignored, etc.) -> dismiss as stale.
                    cur.execute("UPDATE qb_match_candidates SET status = 'rejected', "
                                "decided_by = %s, decided_at = now() WHERE id = %s",
                                (f"{triggered_by} (recheck — no longer applicable)", c["id"]))
                    n_stale += 1
                    continue
                qb_std = ce.resolve_standard_vendor(c["qb_name"], std_tables=std)
                if (_std_eq(qb_std, rec["standard_vendor"])
                        and _num_eq(c["qb_num"], rec["number"])
                        and ce._amounts_equal(c["qb_amount"], rec["amount"])):
                    cur.execute(f"UPDATE {table} SET entered_in_qb = TRUE, entered_qb_by = %s, "
                                f"entered_qb_at = now() WHERE id = %s", (actor, rec["id"]))
                    cur.execute(
                        "INSERT INTO edit_log (table_name, record_id, field, old_value, "
                        "new_value, changed_by) VALUES (%s,%s,'entered_in_qb','false',"
                        "'true (resolved on recheck — vendor list updated)',%s)",
                        (table, rec["id"], actor))
                    cur.execute("UPDATE qb_match_candidates SET status = 'confirmed', "
                                "decided_by = %s, decided_at = now() WHERE id = %s",
                                (actor, c["id"]))
                    n_resolved += 1
                else:
                    n_unchanged += 1

            # Also sweep QB rows that matched NOTHING at upload (qb_unmatched_rows) against the
            # CURRENT eligible record pool: a vendor the user just standardized may now resolve the
            # same way, turning an old "no match" into a real Tier-1 hit. Same test + risk profile
            # as the candidate sweep above.
            n_unmatched_resolved = 0
            cur.execute("SELECT * FROM qb_unmatched_rows WHERE status = 'pending'")
            for u in cur.fetchall():
                matched = None
                for table, rtype, num_col in (("invoices", "invoice", "invoice_number"),
                                              ("credit_memos", "credit_memo", "credit_memo_number")):
                    dup_clause = " AND duplicate_ignored IS NOT TRUE" if table == "invoices" else ""
                    cur.execute(
                        f"SELECT id, standard_vendor, {num_col} AS number, amount FROM {table} "
                        f"WHERE store_matched = %s AND entered_in_qb = FALSE AND skip_qb_import IS NOT TRUE"
                        f"{dup_clause}", (u["qb_store"],))
                    rec = next((r for r in cur.fetchall()
                                if _std_eq(u["qb_std_vendor"], r["standard_vendor"])
                                and _num_eq(u["qb_num"], r["number"])
                                and ce._amounts_equal(u["qb_amount"], r["amount"])), None)
                    if rec:
                        matched = (table, rtype, rec)
                        break
                if not matched:
                    continue
                table, rtype, rec = matched
                cur.execute(f"UPDATE {table} SET entered_in_qb = TRUE, entered_qb_by = %s, "
                            f"entered_qb_at = now() WHERE id = %s", (actor, rec["id"]))
                cur.execute(
                    "INSERT INTO edit_log (table_name, record_id, field, old_value, new_value, "
                    "changed_by) VALUES (%s,%s,'entered_in_qb','false',"
                    "'true (resolved on recheck — matched a previously-unmatched QB export row)',%s)",
                    (table, rec["id"], actor))
                cur.execute("UPDATE qb_unmatched_rows SET status = 'resolved', "
                            "resolved_record_type = %s, resolved_record_id = %s, resolved_at = now() "
                            "WHERE id = %s", (rtype, rec["id"], u["id"]))
                n_unmatched_resolved += 1
        conn.commit()
        return (f"Re-checked {len(pending)} pending match(es): {n_resolved} resolved, "
                f"{n_stale} no longer applicable, {n_unchanged} still need review."
                + (f" Also resolved {n_unmatched_resolved} previously-unmatched QB export row(s)."
                   if n_unmatched_resolved else ""))
    except Exception:
        conn.rollback(); raise
    finally:
        conn.close()


def resolve_unmatched_for_new_records(filenames, actor):
    """Call right after save_results() persists new invoice/credit_memo rows. Checks each
    newly-recorded, not-yet-entered, non-skip_qb_import record against qb_unmatched_rows
    (QB export rows that had no candidate at all when uploaded) for an exact
    vendor+number+amount+store match. On a hit, marks the record entered_in_qb = TRUE
    (same as a Tier-1 upload-time match) and the unmatched row resolved. Returns the
    number resolved (0 if filenames is empty or nothing matches)."""
    if not filenames:
        return 0
    conn = database.connect()
    conn.autocommit = False
    try:
        with conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor) as cur:
            n_resolved = 0
            for table, rtype, num_col in (
                    ("invoices", "invoice", "invoice_number"),
                    ("credit_memos", "credit_memo", "credit_memo_number")):
                dup_clause = " AND duplicate_ignored IS NOT TRUE" if table == "invoices" else ""
                cur.execute(
                    f"SELECT id, standard_vendor, {num_col} AS number, amount, store_matched "
                    f"FROM {table} WHERE email_filename = ANY(%s) AND entered_in_qb = FALSE "
                    f"AND skip_qb_import IS NOT TRUE{dup_clause}", (list(filenames),))
                for rec in cur.fetchall():
                    cur.execute(
                        "SELECT * FROM qb_unmatched_rows WHERE status = 'pending' "
                        "AND qb_store = %s", (rec["store_matched"],))
                    hit = next((u for u in cur.fetchall()
                                if _std_eq(u["qb_std_vendor"], rec["standard_vendor"])
                                and _num_eq(u["qb_num"], rec["number"])
                                and ce._amounts_equal(u["qb_amount"], rec["amount"])), None)
                    if not hit:
                        continue
                    cur.execute(
                        f"UPDATE {table} SET entered_in_qb = TRUE, entered_qb_by = %s, "
                        f"entered_qb_at = now() WHERE id = %s", (actor, rec["id"]))
                    cur.execute(
                        "INSERT INTO edit_log (table_name, record_id, field, old_value, "
                        "new_value, changed_by) VALUES (%s,%s,'entered_in_qb','false',"
                        "'true (found in a previously-uploaded QB export, matched on arrival "
                        "— no action needed)',%s)", (table, rec["id"], actor))
                    cur.execute(
                        "UPDATE qb_unmatched_rows SET status = 'resolved', "
                        "resolved_record_type = %s, resolved_record_id = %s, "
                        "resolved_at = now() WHERE id = %s",
                        (rtype, rec["id"], hit["id"]))
                    n_resolved += 1
        conn.commit()
        return n_resolved
    except Exception:
        conn.rollback(); raise
    finally:
        conn.close()
