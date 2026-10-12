"""
One-time migration of the existing Excel outputs into PostgreSQL, preserving
history. Safe to re-run (each table is replaced from the sheet).

    python migrate_from_excel.py

Reads output/{results,invoice_history,review_queue,learnings,qb_ready}.xlsx and
loads them via database.py. The qb_ready sheet's "Entered in QB" checkmarks are
carried onto the matching invoice/credit-memo rows.
"""

from pathlib import Path

from openpyxl import load_workbook

import database as db

OUTPUT_DIR = Path(__file__).parent / "output"


def _read(path):
    if not path.exists():
        print(f"  (skip) {path.name} not found")
        return []
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(headers, row)))
    wb.close()
    return rows


def _apply_entered_in_qb(qb_rows):
    """Carry the manual 'Entered in QB' flag from qb_ready onto invoices/credit_memos."""
    marked = [(r.get("Filename"), r.get("Invoice #"))
              for r in qb_rows if str(r.get("Entered in QB") or "").strip().lower() in ("yes", "true", "1")]
    if not marked:
        return 0
    conn = db.connect()
    try:
        with conn.cursor() as cur:
            for filename, inv_no in marked:
                cur.execute(
                    "UPDATE invoices SET entered_in_qb = TRUE "
                    "WHERE email_filename = %s AND invoice_number = %s",
                    (filename, inv_no))
                cur.execute(
                    "UPDATE credit_memos SET entered_in_qb = TRUE "
                    "WHERE email_filename = %s AND credit_memo_number = %s",
                    (filename, inv_no))
        conn.commit()
    finally:
        conn.close()
    return len(marked)


def main():
    print("Creating tables (if missing)...")
    db.init_db()

    results = _read(OUTPUT_DIR / "results.xlsx")
    if results:
        touched = {r.get("Filename") for r in results if r.get("Filename")}
        _, _ = db.save_results(results, touched)  # (protected, recorded) — unused in migration
        print(f"  migrated {len(results)} result row(s)")

    history = _read(OUTPUT_DIR / "invoice_history.xlsx")
    if history:
        db.save_history(history)
        print(f"  migrated {len(history)} invoice_history row(s)")

    queue = _read(OUTPUT_DIR / "review_queue.xlsx")
    if queue:
        db.save_review_queue(queue)
        print(f"  migrated {len(queue)} review_queue row(s)")

    learnings = _read(OUTPUT_DIR / "learnings.xlsx")
    if learnings:
        db.save_learnings(learnings)
        print(f"  migrated {len(learnings)} learning(s)")

    qb = _read(OUTPUT_DIR / "qb_ready.xlsx")
    if qb:
        n = _apply_entered_in_qb(qb)
        print(f"  applied 'Entered in QB' to {n} record(s)")

    print("Migration complete.")


if __name__ == "__main__":
    main()
